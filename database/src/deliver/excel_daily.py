"""매일 엑셀 `data/deliver/daily/model_scores_<YYYYMMDD>_<basis>.xlsx` (MODEL_EXCEL_SPEC A).

시트: 점수 → 점수 원자료 → 지표(표시용) → 실적 → 업종 → 모델 비교 → 메타. 값만(수식 없음).
행 = 주 모델(run.primary_spec) 모집단 — 순위 종목(순위순) 뒤에 제외 종목(코드순).
지표 열은 레지스트리 `role` 이 가른다: score → 점수 원자료, display → 지표(표시용).
"""
from __future__ import annotations

import os
import statistics
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import Workbook

from .common import (
    PREV_NEGATIVE_FLAGS,
    REVISION_KEYS,
    analyst_note,
    bucket_label,
    change_pct,
    coverage_label,
    exclude_label,
    ind_meta,
    market_label,
    missing,
    split_flags,
    winsorize,
    yoy,
)
from .qpack import (
    C_HDR,
    FIRST_DATA_ROW,
    Col,
    Group,
    Title,
    add_sparklines,
    font,
    new_workbook,
    write_meta,
    write_table,
)
from .reader import (
    DeliverError,
    deployed_rev,
    fi_run_meta,
    iso,
    load_run,
    previous_run,
    ql_home,
    read_fi,
    read_scores,
    ymd,
)
from .trend import WINDOWS, Trend, load_trend, spark_groups
from .trend import write_sheet as write_trend_sheet
from .view import DayView, composite_of, load_day, rank_of, ticker_of

Row = dict[str, object]
Dictionary = list[tuple[str, str, str, str]]

MODEL_LABELS = {"scope@1.0": "scope_v1.0", "v4_rank@0.1": "v4 기본", "v4_rank@0.2": "v4 동일가중",
                "v3_zscore": "v3 원본", "v2_percentrank": "v2 원본"}
FOOTNOTES = (
    "① 유니버스 컷: 주 모델 모집단은 메타 시트 '유니버스 규칙' 줄을 따른다. 탈락 종목은 점수 없이 "
    "제외 사유만 적는다.",
    "② 창: 가격 지표는 수정종가 세션 수, 재무는 D 이전 공시(PIT), 추정치는 당해 결산기 컨센서스.",
    "③ 색: 3색 백분위 10/50/90 — 빨강(낮음)·노랑·초록(높음), 초록 = 좋음. "
    "순위 열은 반전(1위 = 초록). Δ순위 1M 은 가운데 = 0(하락 빨강 · 상승 초록). "
    "레벨 값은 무색. 형광 노랑 = 신규 진입.",
    "④ 정렬 키: 점수·원자료·지표·실적·모델 비교 = 주 모델 순위 오름차순, 제외 종목은 뒤에 코드순.",
    "⑤ 흐름: 셀 안 꺾은선(엑셀 스파크라인) = 1M 동안 판마다의 순위(원순위), "
    "위로 갈수록 순위 상승. 선 색 = 그 기간 Δ순위(초록 상승 · 빨강 하락 · 회색 같음·모름) — "
    "선의 처음 → 끝 방향과 같다. 1일·1W Δ순위는 싣지 않는다.",
)


# 비교 열에서 빼는 모델(model_id) — v4 는 결함 수정 전까지 뺀다(N-27 §8-17). 판 계산·저장은 그대로
HIDDEN_COMPARE_MODELS = frozenset({"v4_rank"})


def model_label(spec_id: str) -> str:
    if spec_id in MODEL_LABELS:
        return MODEL_LABELS[spec_id]
    return MODEL_LABELS.get(spec_id.split("@")[0], spec_id)


def compared_specs(view: DayView) -> list[str]:
    """엑셀에 싣는 비교 모델 — 주 모델 밖 spec 중 `HIDDEN_COMPARE_MODELS` 를 뺀 것(제목·점수 시트
    '다른 모델 순위'·모델 비교 시트·메타가 이 한 곳을 따른다)."""
    return [s for s in sorted(view.run.specs)
            if s != view.spec_id and s.split("@")[0] not in HIDDEN_COMPARE_MODELS]


def left_out_specs(view: DayView) -> list[str]:
    """판에는 있지만 비교 열에서 뺀 모델(N-27)."""
    shown = set(compared_specs(view))
    return [s for s in sorted(view.run.specs) if s != view.spec_id and s not in shown]


def _num(v: object) -> float | None:
    if isinstance(v, int | float) and not isinstance(v, bool) and v == v:
        return float(v)
    return None


def period_add(period: str, months: int) -> str:
    """'YYYY/MM' + 개월."""
    y, m = int(period[:4]), int(period[5:7])
    k = y * 12 + (m - 1) + months
    return f"{k // 12:04d}/{k % 12 + 1:02d}"


# ── fi 원자료 ────────────────────────────────────────────────────────────────
@dataclass
class FiData:
    """엑셀이 읽는 factor_inputs 원자료(판 id 고정)."""

    close: dict[str, float | None] = field(default_factory=dict)
    fins: dict[str, list[Row]] = field(default_factory=dict)
    cons: dict[str, list[Row]] = field(default_factory=dict)
    flow_last: dict[str, str | None] = field(default_factory=dict)

    def annual(self, t: str, need: Sequence[str] = ("ni",)) -> Row | None:
        """need 열 중 하나라도 값이 있는 가장 최근 연간 확정 행."""
        rows = [r for r in self.fins.get(t, ()) if r.get("period_type") == "annual"
                and any(r.get(c) is not None for c in need)]
        return max(rows, key=lambda r: str(r["period"]), default=None)

    def annual_at(self, t: str, period: str) -> Row | None:
        return next((r for r in self.fins.get(t, ()) if r.get("period_type") == "annual"
                     and str(r["period"]) == period), None)

    def quarters(self, t: str) -> dict[str, Row]:
        return {str(r["period"]): r for r in self.fins.get(t, ())
                if r.get("period_type") == "quarter"}

    def cons_at(self, t: str, period: str, horizon: str = "cur") -> Row | None:
        return next((r for r in self.cons.get(t, ()) if str(r["target_period"]) == period
                     and r.get("horizon") == horizon), None)

    def fy0(self, t: str, d: str) -> str | None:
        """당해 결산기 = D 의 'YYYY/MM' 이상 중 최소(엔진 `_current_period` 와 같은 규칙)."""
        asof = f"{d[:4]}/{d[5:7]}"
        periods = [str(r["target_period"]) for r in self.cons.get(t, ())
                   if r.get("target_period") is not None and str(r["target_period"]) >= asof]
        return min(periods, default=None)

    def fy_month(self, t: str) -> str:
        """결산월 'MM' — 컨센서스 결산기(가장 늦은 것)의 월, 컨센서스가 없으면 확정 연간 행의 월,
        둘 다 없으면 '12'. 실적 시트 연간 칸 = 결산기가 끝나는 연도(사용자 결정 10-06).
        U22 전까지 fi 연간 행은 12월 결산만이라 두 번째 대안은 늘 '12' — 컨센서스 없는
        비12월 결산사는 12월로 떨어진다(물결 4)."""
        p = max((str(r["target_period"]) for r in self.cons.get(t, ())
                 if r.get("target_period") is not None), default=None)
        if p is None:
            a = self.annual(t, ("revenue", "op", "ni"))
            p = None if a is None else str(a["period"])
        return "12" if p is None else p[5:7]


def _group(rows: Sequence[Row]) -> dict[str, list[Row]]:
    out: dict[str, list[Row]] = {}
    for r in rows:
        out.setdefault(str(r["ticker"]), []).append(r)
    return out


def load_fi(fi_root: Path, view: DayView) -> FiData:
    d = date.fromisoformat(view.date).isoformat()          # SQL 에 넣기 전 형식 검증
    bid = view.run.fi_build_id
    close = {str(r["ticker"]): _num(r["close"]) for r in read_fi(
        fi_root, "fi_prices", bid, columns=("ticker", "close"), where=f"date = DATE '{d}'")}
    flow_last = {str(r["ticker"]): iso(r["last_date"]) for r in read_fi(
        fi_root, "fi_flows", bid, select="ticker, max(date) AS last_date", tail="GROUP BY ticker")}
    return FiData(close=close, fins=_group(read_fi(fi_root, "fi_fin_summary", bid)),
                  cons=_group(read_fi(fi_root, "fi_consensus", bid)), flow_last=flow_last)


# ── 공통 열 ──────────────────────────────────────────────────────────────────
def id_cols(full: bool = False) -> tuple[Col, ...]:
    base = (Col("code", "코드", "id", None, 7.0, definition="종목코드 6자리(문자)"),
            Col("name", "이름", "id", None, 14.0, definition="종목명(fi_universe)"))
    l1 = Col("l1", "대분류", "id", None, 10.0, definition="WICS 대분류(fi_universe 기준 스냅샷)")
    if not full:
        return (*base, l1)
    return (*base,
            Col("market", "시장", "id", None, 5.0, definition="KS = KOSPI · KQ = KOSDAQ"),
            l1,
            Col("l2", "중분류", "id", None, 14.0, definition="WICS 중분류"),
            Col("mcap", "시총(억)", "num", "#,##0", 9.0,
                definition="D 시가총액, 억원(fi_universe)"),
            Col("adv20", "거래대금\n20일(억)", "num", "#,##0", 9.0,
                definition="최근 20세션 평균 거래대금, 억원(D-13 적격성 재료)"),
            Col("price", "주가", "num", "#,##0", 9.0, definition="D 종가, 원(fi_prices)"))


def id_values(view: DayView, t: str, fi: FiData | None = None) -> Row:
    u = view.uni.get(t, {})
    row = view.by_ticker.get(t, {})
    return {"code": t, "name": u.get("name") or "",
            "market": market_label(u.get("market")),
            "l1": u.get("sector_l1_name") or row.get("sector_l1") or "",
            "l2": u.get("sector_l2_name") or row.get("sector_l2") or "",
            "mcap": _num(u.get("market_cap")), "adv20": _num(u.get("adv20")),
            "price": None if fi is None else fi.close.get(t)}


def title_for(view: DayView, sheet: str, note: str = "") -> Title:
    others = compared_specs(view)
    model_line = f"주 모델 {view.spec_id}({model_label(view.spec_id)})"
    if others:
        model_line += " · 비교 " + ", ".join(f"{s}({model_label(s)})" for s in others)
    source = (f"{view.date} 기준({view.run.basis}) · Source: quant-ledger model·factor_inputs "
              f"· 판 {view.run.build_id} · fi {view.run.fi_build_id}")
    return Title(sheet, model_line, source, note)


def bucket_weight_label(view: DayView, b: str) -> str:
    if view.spec is None:
        return bucket_label(b)
    return f"{bucket_label(b)} {view.spec.buckets[b] * 100:.0f}%"


def score_indicators(view: DayView) -> list[tuple[str, str, str]]:
    """(key, bucket, definition) — role=score, 레지스트리 순서. spec 이 없으면 데이터에서."""
    return _indicators(view, "score")


def display_indicators(view: DayView) -> list[tuple[str, str, str]]:
    return _indicators(view, "display")


def _indicators(view: DayView, role: str) -> list[tuple[str, str, str]]:
    if view.spec is not None:
        return [(i.key, i.bucket, i.definition) for i in view.spec.indicators if i.role == role]
    seen: dict[str, tuple[str, str, str]] = {}
    for per in view.ind.values():
        for k, r in per.items():
            if r.get("role") == role and k not in seen:
                seen[k] = (k, str(r.get("bucket")), "")
    return list(seen.values())


# ── scope·v3 점수 표 원값(E-10, N-25 Q7) ─────────────────────────────────────
RATIO, PCT, TIMES, FLOW = "비율", "%", "배", "순매수/시총"
UNIT_FMT = {RATIO: "0.0000", PCT: "#,##0.00", TIMES: "#,##0.00", FLOW: "0.000"}
UNIT_TEXT = {RATIO: "비율(0.05 = 5%)", PCT: "%", TIMES: "배",
             FLOW: "순매수/시총(순매수 백만원 ÷ 시총 억원, 1 = 시총의 1%)"}


@dataclass(frozen=True)
class RawCol:
    """점수 표의 원값·표식 열 하나. unit 이 빈 문자열이면 표식(글자) 열."""

    key: str
    bucket: str
    label: str
    unit: str
    definition: str


def _rev_cols(metric: str, name: str) -> tuple[RawCol, ...]:
    out: list[RawCol] = []
    for p, ago in (("1w", "1주 전"), ("1m", "1개월 전"), ("3m", "3개월 전")):
        out += [RawCol(f"{metric}_change_{p}", "revision", f"{name} 추정 {p.upper()}", RATIO,
                       f"당해 결산기 {name} 컨센서스 (cur − {ago}) ÷ |{ago}|. 비교값이 없거나 "
                       "0 이면 0(v3 규약), 컨센서스 쌍이 없으면 빈칸"),
                RawCol(f"{metric}_{p}_flag", "revision", f"{name} {p.upper()}\n표식", "",
                       f"{name} {p.upper()} 부호 전환 — 흑전 · 적전 · 적확 · 적축(v3 규약, 변화율 "
                       "칸은 그대로)")]
    return tuple(out)


def _flow_cols(prefix: str, name: str, subject: str) -> tuple[RawCol, ...]:
    return tuple(RawCol(f"flow_{prefix}_{n}d", "flow", f"{name} {n}일", FLOW,
                        f"{subject} 순매수 — D 이하 최근 수급 행 {n}개 합(백만원) ÷ 시총(억원). "
                        "1 = 시총의 1%")
                 for n in (5, 20))


# scope·v3_zscore 점수 표(contracts.V3_SCORE_COLUMNS) 원값 27열 + 표식 6열 — 엔진 값을 배율 없이
# 그대로 싣는다. 순서 = 버킷 순서(momentum·revision·flow·quality·valuation) · 버킷 안 엔진 순서
V3_RAW: tuple[RawCol, ...] = (
    *(RawCol(f"r{m}", "momentum", f"{m.upper()} 수익률", RATIO,
             f"{n}세션 전 대비 수익률(수정종가, 없으면 종가). 가격 행이 {n}개 이하면 빈칸")
      for m, n in (("1m", 20), ("3m", 60), ("6m", 120), ("9m", 180), ("12m", 240))),
    *_rev_cols("op", "영업이익"), *_rev_cols("ni", "순이익"),
    *_flow_cols("inst", "기관", "기관 합계"), *_flow_cols("for", "외국인", "외국인"),
    *_flow_cols("pe", "사모", "사모펀드"),
    RawCol("qual_gpa", "quality", "GP/A", RATIO, "매출총이익 ÷ 자산총계(최근 연간 확정)"),
    RawCol("qual_roa", "quality", "ROA", PCT, "최근 연간 확정 ROA(DART 순이익 ÷ 자산총계 × 100)"),
    RawCol("qual_fcf_assets", "quality", "FCF/자산", RATIO,
           "FCF(영업현금흐름 − |capex|) ÷ 자산총계(최근 연간 확정)"),
    RawCol("qual_debt_ratio", "quality", "부채비율", PCT,
           "부채총계 ÷ 자본총계 × 100(최근 연간 확정, 낮을수록 좋다)"),
    RawCol("qual_gpa_change", "quality", "GP/A 변화", RATIO,
           "(GP/A − 전기 GP/A) ÷ |전기 GP/A|(연간 2기)"),
    RawCol("qual_std_20d", "quality", "20일 변동성", RATIO,
           "최근 21행 일간 수익률 표본 표준편차(낮을수록 좋다)"),
    RawCol("val_per", "valuation", "PER", TIMES, "최근 연간 확정 PER(WISE, 양수만 — 적자는 빈칸)"),
    RawCol("val_pbr", "valuation", "PBR", TIMES, "최근 연간 확정 PBR(WISE, 양수만)"),
    RawCol("val_ev_ebitda", "valuation", "EV/EBITDA", TIMES,
           "최근 연간 EV/EBITDA(WISE, 양수만). scope 는 가중에서 뺐다(점수에 안 씀, 10-02 결정)"),
    RawCol("val_dividend_yield", "valuation", "배당수익률", PCT,
           "최근 연간 배당수익률(WISE 투자지표)"),
)


def score_raw_columns(view: DayView) -> tuple[RawCol, ...]:
    """주 모델 점수 표에 v3 원값 열이 다 있으면(scope·v3_zscore) 그 열들, 아니면 빈 튜플."""
    keys = view.rows[0].keys() if view.rows else ()
    return V3_RAW if all(c.key in keys for c in V3_RAW) else ()


def raw_units_text(cols: Sequence[RawCol]) -> str:
    """단위표 한 줄 — 단위마다 열 이름."""
    by_unit: dict[str, list[str]] = {}
    for c in cols:
        if c.unit:
            by_unit.setdefault(c.unit, []).append(c.key)
    return " / ".join(f"{UNIT_TEXT[u]}: {' · '.join(keys)}" for u, keys in by_unit.items())


def bucket_missing_reason(view: DayView, t: str, b: str) -> str:
    """버킷이 빈 이유 — 버킷결측(NN%) 표식이 있으면 그것, 아니면 첫 지표의 결측 사유."""
    per = view.ind.get(t, {})
    keys = [k for k, bk, _ in score_indicators(view) if bk == b]
    flags = [f for k in keys for f in split_flags(per.get(k, {}).get("flag"))]
    for f in flags:
        if f.startswith("버킷결측"):
            return f
    for k in keys:
        r = per.get(k, {})
        if r.get("raw") is None:
            fl = split_flags(r.get("flag"))
            if fl:
                return fl[0]
    return "원천없음"


# ── 시트: 점수 ───────────────────────────────────────────────────────────────
def sheet_scores(wb: Workbook, view: DayView, fi: FiData,
                 prev_rank: Mapping[str, int | None], trend: Trend) -> Dictionary:
    others = compared_specs(view)
    groups = [
        Group("종목", id_cols(full=True)),
        Group("종합", (
            Col("rank", "순위", "rank", "#,##0", 7.0, key_col=True, color=False,
                definition="주 모델 순위(1 = 최고). 제외 종목은 빈칸"),
            Col("composite", "종합\n점수", "num", "#,##0.00", definition=(
                "있는 버킷 점수의 가중평균(0~100, 엔진 값). 제외 종목도 점수는 남는다")),
            Col("prev_rank", "전일\n순위", "num", "#,##0", color=False,
                definition="직전 성공 판(같은 basis)의 주 모델 순위"),
            Col("d1m", "Δ순위\n1M", "chg", "#,##0", zero_mid=True, definition=(
                "1M 비교 판(D−1개월 이하 마지막 판) 순위 − 오늘 순위(양수 = 상승). "
                "색 가운데 = 0 — 상승 초록 · 하락 빨강")),
            Col("t1m", "1M\n흐름", "spark", None, 11.0, definition=(
                "1M 동안 판마다의 순위 꺾은선(원순위, 위 = 상승). 선 색 = Δ순위 1M 부호 = 선의 "
                "처음 → 끝 방향")),
            # 너비 상한 26 — 가장 긴 v4 라벨 '고점근접+반전 게이트(pull_gate)'가 qpack fit 25.05
            # (N-26 4.5 — 상한 16 이면 v4 라벨 8개 중 6개가 잘렸다)
            Col("excl", "제외 사유", "txt", None, 26.0, definition=(
                "엔진 exclude_reason — D-13 적격성(관리·정지·감사·지연·거래대금) · 데이터 부족 · "
                "버킷 게이트(고점근접+반전 하위 30%)")),
            Col("cov", "커버리지", "txt", None, 8.0, definition=(
                "추정치 신선도(T2.11): 신선 · 유예 D+n(마지막 신선일부터 거래일 n) · 소멸")),
            # 너비 상한 26 — 가장 긴 조합 '최근 3개월 의견 없음 · 11월 결산'이 qpack fit 25.75
            Col("note", "비고", "txt", None, 26.0, definition=(
                "추정기관수(WISE 최근 3개월 투자의견을 낸 증권사 수): 0 = 최근 3개월 의견 없음 · "
                "1~3 = 애널리스트 3명 이하 · 모름 = 애널리스트 수 미상 · 4 이상은 빈칸. "
                "결산월이 12월이 아니면 'N월 결산'을 ' · ' 로 덧붙인다(실적 시트 연간 칸 = "
                "결산기가 끝나는 연도)")),
        ), core=True),
    ]
    for b in view.buckets:
        name = bucket_label(b)
        groups.append(Group(bucket_weight_label(view, b), (
            Col(f"{b}_u", f"{name}\n유니버스", "pct", "#,##0.0", definition=(
                f"{name} 버킷 점수(엔진)의 유니버스 백분위 0~100 — 높을수록 좋다")),
            Col(f"{b}_s", f"{name}\n업종", "pct", "#,##0.0", definition=(
                f"{name} 버킷 점수의 WICS 대분류 안 백분위(표본 < 5 면 유니버스)")),
        )))
    groups.append(Group("결측", (Col("miss", "결측 축", "txt", None, 14.0, definition=(
        "버킷 점수가 빈 축(점수 대상 종목만). 사유는 축 칸의 '결측(사유)'")),)))
    if others:
        groups.append(Group("다른 모델 순위", (
            *(Col(f"r:{s}", model_label(s), "num", "#,##0", color=False,
                  definition=f"{s} 순위(같은 판)") for s in others),
            Col("max_diff", "최대 차이", "num", "#,##0",
                definition="주 모델 포함 모델 순위의 최댓값 − 최솟값(순위 둘 이상일 때)"),
        )))

    rows: list[Row] = []
    for row in view.rows:
        t = ticker_of(row)
        u = view.uni.get(t, {})
        r = rank_of(row)
        pr = prev_rank.get(t)
        mm = fi.fy_month(t)
        notes = (analyst_note(u.get("n_analysts")), None if mm == "12" else f"{int(mm)}월 결산")
        out: Row = {**id_values(view, t, fi), "rank": r, "composite": composite_of(row),
                    "prev_rank": pr, "d1m": trend.delta("1M", t, r),
                    "excl": exclude_label(row.get("exclude_reason")),
                    # v3·v2 점수 행엔 신선도 열이 없다 → 그 판 fi_universe 값(v4 는 같은 값을 행에 싣는다)
                    "cov": coverage_label(row.get("coverage_state", u.get("coverage_state")),
                                          u.get("coverage_age_days")),
                    "note": " · ".join(n for n in notes if n) or None}
        scored = view.scored(t)
        miss = []
        for b in view.buckets:
            if view.bucket_score(t, b) is not None:
                out[f"{b}_u"] = view.upct[b].get(t)
                out[f"{b}_s"] = view.spct[b].get(t)
            elif scored:
                out[f"{b}_u"] = out[f"{b}_s"] = missing(bucket_missing_reason(view, t, b))
                miss.append(bucket_label(b))
        out["miss"] = "·".join(miss)
        ranks = [x for x in (r, *(view.other_ranks[s].get(t) for s in others)) if x is not None]
        for s in others:
            out[f"r:{s}"] = view.other_ranks[s].get(t)
        out["max_diff"] = max(ranks) - min(ranks) if len(ranks) >= 2 else None
        rows.append(out)
    note = ("축 = 버킷 점수의 백분위(높을수록 좋다, 초록 = 상위). 제외 종목은 점수만 남고 순위가 "
            "없다. 결측은 '결측(사유)'.")
    return write_table(wb, title_for(view, "점수", note), groups, rows, sort_key="rank")


# ── 시트: 점수 원자료 ─────────────────────────────────────────────────────────
def ind_cells(view: DayView, t: str, key: str) -> tuple[object, object, str]:
    """(표시값, 엔진 백분위, 표식) — 결측은 '결측(사유)', 리비전 이전값<0 은 표식만(E)."""
    r = view.ind.get(t, {}).get(key)
    if r is None:
        return None, None, ""
    flags = split_flags(r.get("flag"))
    raw = _num(r.get("raw"))
    pct = _num(r.get("pct"))
    if raw is None:
        return missing(flags[0] if flags else None), pct, ";".join(flags[1:])
    if key in REVISION_KEYS and PREV_NEGATIVE_FLAGS.intersection(flags):
        return None, pct, ";".join(flags)
    if key in ("EP", "EP_FWD") and raw < 0:
        flags = ["적자", *flags]
    return raw * ind_meta(key).scale, pct, ";".join(flags)


def sheet_raw(wb: Workbook, view: DayView, fi: FiData) -> Dictionary:
    inds = score_indicators(view)
    groups = [Group("종목", id_cols())]
    for b in dict.fromkeys(bk for _, bk, _ in inds):
        cols: list[Col] = []
        for key, bk, definition in inds:
            if bk != b:
                continue
            m = ind_meta(key)
            extra = " · 표시 = 원값 × 100(%)" if m.scale == 100 else ""
            if key in REVISION_KEYS:
                extra += " · 이전값 < 0(흑전·적확·적축)이면 변화율 대신 표식만(E)"
            cols += [
                Col(f"v:{key}", m.label, "num", m.fmt, color=False,
                    definition=f"{key}: {definition}{extra}"),
                Col(f"p:{key}", f"{key}\n백분위", "num", "#,##0.0", color=False, definition=(
                    f"{key} 엔진 백분위 0~100(방향 적용, 동률 평균순위; 밸류·퀄리티는 대분류 안)")),
                Col(f"f:{key}", f"{key}\n표식", "txt", None, 9.0, definition=(
                    "엔진 flag — 부호 전환(흑전·적전·적확·적축) · 적자 · 무배당 · 업종 되돌림 · "
                    "버킷결측(NN%) · 적격미확인(표식)")),
            ]
        groups.append(Group(bucket_label(b), tuple(cols)))
    # scope·v3 는 지표 긴 표가 0행이고 원값이 점수 표 열에 있다 — 그 열을 그대로 싣는다(E-10)
    raw_cols = score_raw_columns(view)
    for b in dict.fromkeys(c.bucket for c in raw_cols):
        groups.append(Group(bucket_label(b), tuple(
            Col(f"s:{c.key}", f"{c.label}\n({c.unit})", "num", UNIT_FMT[c.unit], color=False,
                definition=f"{c.key}: {c.definition}. 단위 {UNIT_TEXT[c.unit]}") if c.unit
            else Col(f"s:{c.key}", c.label, "txt", None, 9.0, definition=f"{c.key}: {c.definition}")
            for c in raw_cols if c.bucket == b)))
    groups.append(Group("기준", (
        Col("fy", "연간\n기준기", "txt", None, 8.0,
            definition="순이익이 있는 최근 연간 확정 결산기(E/P·배당의 후보 기, 730일 창은 엔진)"),
        Col("fy_avail", "재무\n접수일", "txt", None, 10.0,
            definition="그 연간 행의 available_date(공시·수집으로 알게 된 날)"),
        Col("fq", "분기\n기준", "txt", None, 8.0,
            definition="영업이익이 있는 최근 분기(TTM 끝 분기)"),
        Col("obs", "추정\n관측일", "txt", None, 10.0,
            definition="당해 결산기 컨센서스(cur) 관측일 = 추정 base_date"),
        Col("n_an", "애널리스트\n수", "num", "#,##0", color=False,
            definition=("추정기관 수(fi_universe) = WISE 최근 3개월 투자의견을 낸 증권사 수. "
                        "0 = 의견 없음, 빈칸 = 모름")),
        Col("flow_d", "수급\n최신일", "txt", None, 10.0, definition="fi_flows 에 있는 최신 수급일"),
    )))
    rows: list[Row] = []
    for row in view.rows:
        t = ticker_of(row)
        out: Row = id_values(view, t)
        for key, _, _ in inds:
            out[f"v:{key}"], out[f"p:{key}"], out[f"f:{key}"] = ind_cells(view, t, key)
        for c in raw_cols:
            out[f"s:{c.key}"] = row.get(c.key)
        a = fi.annual(t, ("ni",))
        qs = [p for p, q in fi.quarters(t).items() if q.get("op") is not None]
        fy0 = fi.fy0(t, view.date)
        cur = None if fy0 is None else fi.cons_at(t, fy0)
        out.update(fy=None if a is None else a["period"],
                   fy_avail=None if a is None else iso(a.get("available_date")),
                   fq=max(qs, default=None), obs=None if cur is None else iso(cur.get("obs_date")),
                   n_an=view.uni.get(t, {}).get("n_analysts"), flow_d=fi.flow_last.get(t))
        rows.append(out)
    note = "점수에 들어간 하위 지표 원값 · 엔진 백분위 · 표식. 5행 = 정의(전문은 메타). 색 없음."
    if raw_cols:
        note = ("주 모델 점수 표(scores.parquet) 원값 그대로(배율 없음 · 빈칸 = 엔진 값 없음, "
                "v3 는 사유를 남기지 않는다). 단위는 머리글 괄호 — 비율(0.05 = 5%) · % · 배 · "
                "순매수/시총(1 = 시총의 1%), 전체는 메타 '점수 원자료 단위'. "
                "5행 = 정의(전문은 메타). 색 없음.")
    return write_table(wb, title_for(view, "점수 원자료", note), groups, rows, defs_row=True)


# ── 시트: 지표(표시용) ────────────────────────────────────────────────────────
def sheet_display(wb: Workbook, view: DayView, fi: FiData) -> Dictionary:
    inds = display_indicators(view)
    groups = [Group("종목", id_cols())]
    for b in dict.fromkeys(bk for _, bk, _ in inds):
        cols: list[Col] = []
        for key, bk, definition in inds:
            if bk != b:
                continue
            m = ind_meta(key)
            cols += [Col(f"v:{key}", m.label, "chg" if m.change else "num", m.fmt,
                         definition=f"{key}: {definition} · 표시 = 1/99 윈저라이즈"),
                     Col(f"f:{key}", f"{key}\n표식", "txt", None, 9.0,
                         definition="엔진 flag(결측 사유는 값 칸의 '결측(사유)')")]
        groups.append(Group(f"{bucket_label(b)}(표시)", tuple(cols)))
    groups.append(Group("재무(최근 연간)", (
        Col("roe", "ROE(%)", definition=(
            "최근 연간 확정 ROE(DART 사업보고서 — 순이익 ÷ 자본총계 × 100) · 1/99 윈저라이즈")),
        Col("roa", "ROA(%)", definition=(
            "최근 연간 확정 ROA(DART 사업보고서 — 순이익 ÷ 자산총계 × 100) · 1/99 윈저라이즈")),
        Col("debt", "부채비율(%)", definition=(
            "최근 연간 확정 부채비율(DART 사업보고서 — 부채총계 ÷ 자본총계 × 100) · "
            "1/99 윈저라이즈")),
    )))
    groups.append(Group("컨센서스(당해 결산기)", (
        Col("per", "선행 PER(배)", definition=(
            "당해 결산기 컨센서스 PER(cur). 컨센서스 순이익 < 0 이면 '적자' · 1/99 윈저라이즈")),
        Col("pbr", "선행 PBR(배)", fmt="#,##0.00",
            definition="당해 결산기 컨센서스 PBR(cur) · 1/99 윈저라이즈"),
        Col("eps1m", "선행 EPS\n1M 변화(%)", "chg", definition=(
            "(EPS cur − EPS 1개월 전) ÷ |1개월 전| × 100. 이전값 < 0 은 표식만(E) · 윈저라이즈")),
        Col("eps1m_f", "EPS 1M\n표식", "txt", None, 8.0, definition="흑전·적전·적확·적축(v3 규약)"),
    )))
    rows: list[Row] = []
    for row in view.rows:
        t = ticker_of(row)
        out: Row = id_values(view, t)
        for key, _, _ in inds:
            v, _p, f = ind_cells(view, t, key)
            out[f"v:{key}"], out[f"f:{key}"] = v, f
        a = fi.annual(t, ("roe", "roa", "debt_ratio"))
        if a is not None:
            out.update(roe=_num(a.get("roe")), roa=_num(a.get("roa")),
                       debt=_num(a.get("debt_ratio")))
        fy0 = fi.fy0(t, view.date)
        cur = None if fy0 is None else fi.cons_at(t, fy0)
        prev = None if fy0 is None else fi.cons_at(t, fy0, "1m")
        if cur is not None:
            ni = _num(cur.get("ni"))
            out["per"] = "적자" if ni is not None and ni < 0 else _num(cur.get("per"))
            out["pbr"] = _num(cur.get("pbr"))
            out["eps1m"], out["eps1m_f"] = change_pct(
                _num(cur.get("eps")), None if prev is None else _num(prev.get("eps")))
        rows.append(out)
    numeric = [f"v:{k}" for k, _, _ in inds] + ["roe", "roa", "debt", "per", "pbr", "eps1m"]
    for key in numeric:
        vals = winsorize([r.get(key) for r in rows])
        for r, v in zip(rows, vals, strict=True):
            r[key] = v
    note = ("점수에 넣지 않는 참고 지표(레지스트리 role=display) + 원자료. "
            "숫자는 1/99 윈저라이즈한 표시값. 색은 변화 열만.")
    return write_table(wb, title_for(view, "지표(표시용)", note), groups, rows, defs_row=True)


# ── 시트: 실적 ───────────────────────────────────────────────────────────────
METRICS = (("revenue", "매출"), ("op", "영업이익"), ("ni", "순이익"))
N_QUARTERS = 5          # fi_fin_summary 는 분기 5기만 싣는다(계약 창)


def _actual(fi: FiData, t: str, period: str, key: str) -> float | None:
    row = fi.annual_at(t, period)
    return None if row is None else _num(row.get(key))


def _estimate(fi: FiData, t: str, period: str, key: str) -> float | None:
    row = fi.cons_at(t, period)
    return None if row is None else _num(row.get(key))


def _either(fi: FiData, t: str, period: str, key: str) -> float | None:
    """y-y 비교 기준 — 확정 우선, 없으면 추정."""
    a = _actual(fi, t, period, key)
    return a if a is not None else _estimate(fi, t, period, key)


def sheet_earnings(wb: Workbook, view: DayView, fi: FiData) -> Dictionary:
    # 연간 머리글 = 판 날짜의 연도 Y 로 Y−1 · YE · (Y+1)E — 해가 바뀌면 저절로 넘어간다(N-19).
    # 칸 = 결산기가 끝나는 연도(9월 결산이면 2026E = 2026/09 결산기, 증권사 FY 표기 관례)
    year = int(view.date[:4])
    years = (("a", year - 1, f"{year - 1}"), ("e0", year, f"{year}E"),
             ("e1", year + 1, f"{year + 1}E"))
    groups = [Group("종목", id_cols())]
    for key, label in METRICS:
        cols: list[Col] = []
        for slot, yr, head in years:
            what = (f"{yr}년에 끝나는 결산기(12월 결산 = {yr}/12, 9월 결산이면 {yr}/09) "
                    "확정치(사업보고서). 확정치가 없으면 같은 결산기 추정치(컨센서스 cur)를 회색 'E' "
                    "서식으로, 둘 다 없으면 빈칸" if slot == "a"
                    else f"{yr}년에 끝나는 결산기 추정치(컨센서스 cur)")
            cols += [
                Col(f"{key}_{slot}", head, "num", "#,##0", definition=f"{label}(억원) {what}"),
                Col(f"{key}_{slot}_y", f"{head}\ny-y(%)", "chg", "#,##0.0", definition=(
                    f"{label} 전년 대비(%) — 한 해 앞 같은 결산월 결산기(확정 우선, 없으면 추정)와 "
                    "비교. 부호가 바뀌면 숫자 대신 흑자전환(−→+) · 적자전환(+→−) · 적자지속(−→−)"
                    + (". 옆 칸이 추정치(회색 'E')면 이 칸도 회색" if slot == "a" else ""))),
            ]
        groups.append(Group(f"{label}(억원)", tuple(cols)))
    # 분기 창 = 판 전 종목 중 영업이익이 있는 가장 최근 달력 분기를 끝으로 5칸(오래된 → 최근).
    # 머리글 'YY.nQ' — 'YYYY/MM' 의 03→1Q · 06→2Q · 09→3Q · 12→4Q. 기말이 분기말이 아닌 결산월
    # (예 11월 결산의 2026/08)은 창 끝을 정하지 않고 칸에도 들어가지 않는다(최근 분기 y-y 는 그대로)
    q_last = max((p for row in view.rows for p, q in fi.quarters(ticker_of(row)).items()
                  if q.get("op") is not None and p[5:7] in ("03", "06", "09", "12")),
                 default=None)
    q_periods = ([] if q_last is None
                 else [period_add(q_last, -3 * k) for k in range(N_QUARTERS - 1, -1, -1)])
    qcols = [Col(f"q:{p}", f"{p[2:4]}.{int(p[5:7]) // 3}Q", "num", "#,##0",
                 definition=f"{p} 분기 영업이익(3개월 값), 억원 — 그 분기 행이 없으면 빈칸")
             for p in q_periods]
    qcols.append(Col("q_y", "최근 분기\ny-y(%)", "chg", "#,##0.0", definition=(
        "종목마다 영업이익이 있는 최근 분기의 전년 동기 대비(%). 부호가 바뀌면 숫자 대신 "
        "흑자전환·적자전환·적자지속")))
    groups.append(Group("분기 영업이익(억원)", tuple(qcols)))

    rows: list[Row] = []
    est: list[tuple[int, str]] = []         # Y−1 칸에 추정치를 넣은 (행 순번, 지표)
    for i, row in enumerate(view.rows):
        t = ticker_of(row)
        out: Row = id_values(view, t)
        mm = fi.fy_month(t)                 # 결산월 — 비12월이면 점수 시트 비고에 'N월 결산'
        for key, _ in METRICS:
            for slot, yr, _head in years:
                p = f"{yr}/{mm}"
                v = _either(fi, t, p, key) if slot == "a" else _estimate(fi, t, p, key)
                if slot == "a" and v is not None and _actual(fi, t, p, key) is None:
                    est.append((i, key))
                y, f = yoy(v, _either(fi, t, f"{yr - 1}/{mm}", key))
                out[f"{key}_{slot}"], out[f"{key}_{slot}_y"] = v, f or y   # 부호 전환이면 글자
        qs = fi.quarters(t)
        for p in q_periods:
            q = qs.get(p)
            out[f"q:{p}"] = None if q is None else _num(q.get("op"))
        with_op = [p for p, q in qs.items() if q.get("op") is not None]
        if with_op:
            q0 = max(with_op)
            prev = qs.get(period_add(q0, -12))
            y, f = yoy(_num(qs[q0].get("op")), None if prev is None else _num(prev.get("op")))
            out["q_y"] = f or y
        rows.append(out)
    note = ("연간: Y−1 = 확정치(사업보고서) — 확정치가 없으면 Y−1 추정치를 회색 'E' 로"
            "(옆 y-y 도 회색) · "
            f"YE·(Y+1)E = 컨센서스 추정. 머리글 연도는 판 날짜 기준(Y = {year})이라 해가 바뀌면 "
            f"자동으로 넘어간다. 연간 칸 = 결산기가 끝나는 연도(9월 결산이면 {year}E = {year}/09 "
            "결산기, 비12월 결산은 점수 시트 비고에 'N월 결산'). "
            "분기 = 판 전 종목의 최근 달력 분기(기말 03·06·09·12월)까지 5분기 "
            "영업이익 — 기말이 다른 결산월의 분기는 칸이 비고 최근 분기 y-y 만 있다. 부호가 바뀌면 "
            "y-y 칸에 숫자 대신 흑자전환·적자전환·적자지속. "
            "어닝시즌(예정일·잠정치·서프라이즈)은 원천이 없어 비운다.")
    dictionary = write_table(wb, title_for(view, "실적", note), groups, rows)
    # qpack 은 서식을 열 단위로 건다 — Y−1 칸 중 추정치를 넣은 칸만 'E' 서식·회색 글자로 고친다.
    # 바로 옆 y-y 칸(같은 그룹의 다음 열)도 추정치 기준 증가율이라 회색 글자(N-25 Q10)
    ws = wb["실적"]
    y1_cols = dict(zip((k for k, _ in METRICS),
                       (c for c in range(1, ws.max_column + 1)
                        if ws.cell(7, c).value == f"{year - 1}"), strict=True))
    for i, key in est:
        cell = ws.cell(FIRST_DATA_ROW + i, y1_cols[key])
        cell.number_format, cell.font = '#,##0"E"', font(color=C_HDR)
        ws.cell(FIRST_DATA_ROW + i, y1_cols[key] + 1).font = font(color=C_HDR)
    return dictionary


# ── 시트: 업종 ───────────────────────────────────────────────────────────────
def _median(xs: Sequence[float]) -> float | None:
    return statistics.median(xs) if xs else None


def sheet_sectors(wb: Workbook, view: DayView) -> Dictionary:
    n_ranked = len(view.ranked)
    top_n = view.output.top_n
    cands = set(view.candidates)
    # 업종 지표 재료 — v4 는 지표 긴 표(EP·REV_OP_1M·R1M), scope·v3 는 점수 표 원값
    # (val_per·op_change_1m·r1m — 지표 긴 표가 0행이라 옛 코드는 세 열이 빈칸, E-10)
    v3 = bool(score_raw_columns(view))
    if v3:
        has = set(view.rows[0])
        lvl_key, rev_key, ret_key = "val_per", "op_change_1m", "r1m"
        metric_cols = (
            Col("per_med", "PER 중앙값\n(배)", "num", "#,##0.00",
                definition="최근 연간 확정 PER(엔진 val_per, 양수만) 중앙값"),
            Col("rev_up", "리비전 상향\n비율(%)", "num", "#,##0.0", definition=(
                "영업이익 추정 1M 변화(엔진 op_change_1m) > 0 종목 비율(값 있는 종목 중 — 비교값이 "
                "없거나 0 이면 엔진 값이 0 이라 상향이 아니다, v3 규약)")),
            Col("r1m", "1M 수익률\n(%)", "chg", "#,##0.0",
                definition="1M 수익률(엔진 r1m × 100) 시총가중 평균"))
    else:
        has = {k for per in view.ind.values() for k in per}
        lvl_key, rev_key, ret_key = "EP", "REV_OP_1M", "R1M"
        metric_cols = (
            Col("ep_med", "E/P 중앙값\n(%)", "num", "#,##0.00", definition="후행 E/P(EP) 중앙값"),
            Col("rev_up", "리비전 상향\n비율(%)", "num", "#,##0.0",
                definition="영업이익 추정 1M 변화(REV_OP_1M) > 0 종목 비율(값 있는 종목 중)"),
            Col("r1m", "1M 수익률\n(%)", "chg", "#,##0.0",
                definition="1M 수익률(R1M) 시총가중 평균"))
    groups = [
        Group("업종", (
            Col("level", "구분", "txt", None, 6.0, definition="대분류 · 중분류(WICS)"),
            Col("code", "코드", "id", None, 7.0, definition="WICS 코드"),
            Col("sname", "업종명", "id", None, 16.0, item=True, definition="WICS 업종명"),
            Col("parent", "대분류", "id", None, 12.0, definition="중분류의 상위 대분류"))),
        Group("규모", (
            Col("n", "종목 수", "num", "#,##0", definition="주 모델 모집단 종목 수(제외 포함)"),
            Col("n_ranked", "순위\n종목 수", "num", "#,##0", definition="순위가 있는 종목 수"),
            Col("mcap", "시총 합(억)", "num", "#,##0", definition="모집단 시가총액 합, 억원"),
            Col("w", "시총 비중(%)", "num", "#,##0.0", definition="모집단 시총 합 대비 비중"))),
        Group("쏠림(%)", (
            Col("t30", f"상위 {top_n}\n비중", "num", "#,##0.0",
                definition=f"순위 ≤ {top_n} 종목 중 이 업종 비중(업종 상한 적용 전)"),
            Col("t100", "상위 100\n비중", "num", "#,##0.0", definition="순위 ≤ 100 종목 중 비중"),
            Col("n_cand", "후보 수", "num", "#,##0",
                definition=f"업종 상한({view.output.max_per_sector}) 적용 후보 {top_n} 중 수"))),
        Group("종합", (Col("comp_med", "종합 점수\n중앙값", "num", "#,##0.00",
                           definition="순위 종목 종합 점수 중앙값"),), core=True),
        Group("축별 평균 백분위", tuple(
            Col(f"{b}_u", bucket_label(b), "pct", "#,##0.0",
                definition=f"{bucket_label(b)} 유니버스 백분위의 업종 평균")
            for b in view.buckets)),
        Group("업종 지표", metric_cols),
        Group("상위 종목", (Col("top", "상위 3(순위)", "txt", None, 34.0,
                              definition="업종 안 순위 상위 3 종목 '이름(순위)'"),)),
    ]
    total_cap = sum(_num(view.uni.get(ticker_of(r), {}).get("market_cap")) or 0.0
                    for r in view.rows)

    def aggregate(level: str, code: str, members: list[str]) -> Row:
        first = view.uni.get(members[0], {})
        name_key = "sector_l1_name" if level == "대분류" else "sector_l2_name"
        ranked = sorted((t for t in members if view.rank(t) is not None),
                        key=lambda t: view.rank(t) or 0)
        caps = {t: _num(view.uni.get(t, {}).get("market_cap")) for t in members}
        cap_sum = sum(v for v in caps.values() if v is not None)
        out: Row = {"level": level, "code": code, "sname": first.get(name_key) or code,
                    "parent": first.get("sector_l1_name") if level == "중분류" else "",
                    "n": len(members), "n_ranked": len(ranked), "mcap": cap_sum,
                    "w": cap_sum / total_cap * 100 if total_cap else None,
                    "t30": (sum(1 for t in ranked if (view.rank(t) or 0) <= top_n)
                            / min(top_n, n_ranked) * 100 if n_ranked else None),
                    "t100": (sum(1 for t in ranked if (view.rank(t) or 0) <= 100)
                             / min(100, n_ranked) * 100 if n_ranked else None),
                    "n_cand": sum(1 for t in members if t in cands),
                    "comp_med": _median([c for t in ranked
                                         if (c := composite_of(view.by_ticker[t]))
                                         is not None]),
                    "top": ", ".join(f"{view.name(t) or t}({view.rank(t)})" for t in ranked[:3])}
        for b in view.buckets:
            vals = [view.upct[b][t] for t in members if t in view.upct[b]]
            out[f"{b}_u"] = statistics.fmean(vals) if vals else None

        def raws(key: str) -> dict[str, float]:
            if v3:
                return {t: v for t in members
                        if (v := _num(view.by_ticker[t].get(key))) is not None}
            return {t: v for t in members
                    if (v := _num(view.ind.get(t, {}).get(key, {}).get("raw"))) is not None}
        if lvl_key in has:
            lv = raws(lvl_key)
            med = statistics.median(lv.values()) if lv else None
            if v3:
                out["per_med"] = med                                  # 배 그대로
            else:
                out["ep_med"] = None if med is None else med * 100    # 비율 → %
        if rev_key in has:
            rv = raws(rev_key)
            out["rev_up"] = sum(1 for v in rv.values() if v > 0) / len(rv) * 100 if rv else None
        if ret_key in has:
            r1 = {t: v for t, v in raws(ret_key).items() if caps.get(t)}
            w = sum(caps[t] or 0.0 for t in r1)
            out["r1m"] = sum(v * (caps[t] or 0.0) for t, v in r1.items()) / w * 100 if w else None
        return out

    rows: list[Row] = []
    for level, col in (("대분류", "sector_l1"), ("중분류", "sector_l2")):
        by: dict[str, list[str]] = {}
        for r in view.rows:
            code = r.get(col)
            if code is not None:
                by.setdefault(str(code), []).append(ticker_of(r))
        rows += [aggregate(level, code, by[code]) for code in sorted(by)]
    note = ("주 모델 모집단의 WICS 대분류·중분류 집계. 축 평균은 유니버스 백분위 평균, 1M 수익률은 "
            "시총가중.")
    return write_table(wb, title_for(view, "업종", note), groups, rows, style="sector")


# ── 시트: 모델 비교 ──────────────────────────────────────────────────────────
def sheet_models(wb: Workbook, view: DayView) -> Dictionary:
    specs = [view.spec_id, *compared_specs(view)]
    ranks = {view.spec_id: {ticker_of(r): rank_of(r) for r in view.rows}, **view.other_ranks}
    top_n = view.output.top_n
    groups = [
        Group("종목", id_cols()),
        Group("순위", tuple(Col(f"r:{s}", model_label(s), "rank", "#,##0",
                              key_col=s == view.spec_id, definition=f"{s} 순위(1 = 최고)")
                          for s in specs), core=True),
        Group(f"상위 {top_n}", tuple(Col(f"t:{s}", model_label(s), "txt", None, 8.0,
                                         definition=f"{s} 순위 ≤ {top_n} 이면 ●")
                                     for s in specs)),
        Group("차이", (Col("max_diff", "최대 차이", "num", "#,##0",
                         definition="모델 순위 최댓값 − 최솟값(순위 둘 이상)"),)),
    ]
    rows: list[Row] = []
    for row in view.rows:
        t = ticker_of(row)
        out: Row = id_values(view, t)
        got = []
        for s in specs:
            r = ranks[s].get(t)
            out[f"r:{s}"] = r
            out[f"t:{s}"] = "●" if r is not None and r <= top_n else ""
            if r is not None:
                got.append(r)
        out["max_diff"] = max(got) - min(got) if len(got) >= 2 else None
        rows.append(out)
    note = "주 모델 모집단 종목의 모델별 순위(같은 판). 모델마다 유니버스가 달라 빈칸이 있다."
    if left_out_specs(view):
        note += " v4 비교 열은 결함 수정 전까지 뺐다(N-27, 메타 '비교에서 뺀 모델')."
    return write_table(wb, title_for(view, "모델 비교", note), groups, rows,
                       sort_key=f"r:{view.spec_id}")


# ── 시트: 메타 ───────────────────────────────────────────────────────────────
def gates_text(g: object) -> str:
    """판 manifest 의 게이트 요약(dict · list 모두)."""
    if isinstance(g, Mapping):
        items = [(str(k), v) for k, v in g.items()]
    elif isinstance(g, list):
        items = [(str(x.get("name", "?")), x) for x in g if isinstance(x, Mapping)]
    else:
        return "" if g is None else str(g)
    return " · ".join(f"{k} {v.get('status', '?') if isinstance(v, Mapping) else v}"
                      for k, v in items)


def _builds_text(v: object) -> str:
    if isinstance(v, Mapping):
        ids = sorted({str(x.get("build_id", x)) if isinstance(x, Mapping) else str(x)
                      for x in v.values()})
        return ", ".join(ids)
    return "" if v is None else str(v)


def universe_text(view: DayView) -> str:
    if view.spec is None:
        return "레지스트리에 없는 spec — 규칙 미상"
    u = view.spec.universe
    g = u.coverage_grace_days
    parts = ["추정치 보유" if u.require_estimates else "추정치 무관",
             "/".join(u.sec_types), "/".join(u.markets),
             f"추정치 유예 {g} 거래일" if g else "추정치 유예 없음"]
    if u.min_analysts is not None:
        parts.append(f"추정기관수 ≥ {u.min_analysts}")
    if u.min_market_cap is not None:
        parts.append(f"시총 ≥ {u.min_market_cap:,.0f}억")
    if u.min_adv20 is not None:
        parts.append(f"20세션 거래대금 ≥ {u.min_adv20:,.0f}억")
    if u.exclude:
        parts.append("제외 " + "·".join(u.exclude))
    return " · ".join(parts)


def meta_pairs(view: DayView, fi_meta: Mapping[str, object], fi: FiData | None) -> list[
        tuple[str, object]]:
    run = view.run
    counts: dict[str, int] = {}
    for r in view.rows:
        if r.get("exclude_reason"):
            k = exclude_label(r["exclude_reason"])
            counts[k] = counts.get(k, 0) + 1
    scored = [t for t in view.by_ticker if view.scored(t)]
    miss = {b: sum(1 for t in scored if view.bucket_score(t, b) is None) for b in view.buckets}
    pairs: list[tuple[str, object]] = [
        ("기준일", view.date), ("basis", run.basis),
        ("주 모델", f"{view.spec_id}({model_label(view.spec_id)})"),
        ("비교 모델", ", ".join(f"{s}({model_label(s)})" for s in compared_specs(view))),
    ]
    left_out = left_out_specs(view)
    if left_out:
        pairs.append(("비교에서 뺀 모델",
                      ", ".join(f"{s}({model_label(s)})" for s in left_out)
                      + " — v4 비교 열은 결함 수정 전까지 뺌(N-27). 판 계산·저장은 그대로"
                        "(아래 판 게이트 줄)"))
    pairs += [
        ("model 판 id", run.build_id), ("model 생성 시각", run.generated_at),
        ("factor_inputs 판 id", run.fi_build_id),
        ("equity 판 id", _builds_text(fi_meta.get("equity_builds")) or "(fi manifest 없음)"),
        ("유니버스 규칙", universe_text(view)),
        ("종목 수", f"모집단 {len(view.rows)} · 순위 {len(view.ranked)} · "
                   f"제외 {len(view.rows) - len(view.ranked)}"),
        ("제외 사유별", " · ".join(f"{k} {v}" for k, v in sorted(counts.items())) or "없음"),
        ("축별 결측 수", " · ".join(f"{bucket_label(b)} {n}" for b, n in miss.items())),
    ]
    raw_cols = score_raw_columns(view)
    if raw_cols:
        pairs.append(("점수 원자료 단위", raw_units_text(raw_cols)))
    if fi is not None:
        fin = [a for rows in fi.fins.values() for r in rows if (a := iso(r.get("available_date")))]
        obs = [o for rows in fi.cons.values() for r in rows
               if r.get("horizon") == "cur" and (o := iso(r.get("obs_date")))]
        flows = [x for x in fi.flow_last.values() if x]
        pairs.append(("데이터 기준일", f"가격 {view.date} · "
                                  f"재무 최신 접수 {max(fin, default='-')} · "
                                  f"추정 관측 최신 {max(obs, default='-')} · "
                                  f"수급 최신 {max(flows, default='-')}"))
    if view.spec is not None:
        pairs += [
            ("모델 버전", f"{view.spec.model_id} {view.spec.version} · 엔진 {view.spec.engine}"),
            ("가중치", " · ".join(f"{bucket_label(b)} {w * 100:.1f}%"
                               for b, w in view.spec.buckets.items())),
            ("제외 게이트", " · ".join(f"{g.key} {g.rule} {g.value}" for g in view.spec.gates)
             or "없음"),
            ("후보 규칙", f"상위 {view.output.top_n} · {view.output.sector_level} 업종당 최대 "
                      f"{view.output.max_per_sector}"),
        ]
    for sid in sorted(run.specs):
        s = run.specs[sid]
        pairs.append((f"판 게이트 {sid}",
                      f"점수 {s.get('n_scores')} · 순위 {s.get('n_ranked')} · "
                      f"제외 {s.get('n_excluded')} · {gates_text(s.get('gates'))}"))
    excluded = run.meta.get("excluded_specs")
    if isinstance(excluded, Mapping):
        for sid, s in sorted(excluded.items()):
            if isinstance(s, Mapping) and "error" in s:      # 실행 예외로 뺀 비교 모델(D-01)
                pairs.append((f"제외된 비교 모델 {sid}",
                              f"실행 오류로 이번 판에서 뺐다 · {s['error']}"))
                continue
            gates_of = s.get("gates") if isinstance(s, Mapping) else None
            pairs.append((f"제외된 비교 모델 {sid}",
                          f"게이트 실패로 이번 판에서 뺐다 · {gates_text(gates_of)}"))
    pairs += [(f"각주 {i}", text) for i, text in enumerate(FOOTNOTES, start=1)]
    return pairs


def insert_after(pairs: list[tuple[str, object]], key: str,
                 new: Sequence[tuple[str, object]]) -> None:
    """메타 줄을 키 이름 뒤에 끼운다 — 고정 위치 숫자는 줄 수가 바뀌면 다른 줄 사이로 옮겨 갔다
    (E-09: 1W 를 뺀 뒤 '순위 흐름 판'이 factor_inputs·equity 판 id 사이에 끼었다)."""
    keys = [k for k, _ in pairs]
    if key not in keys:
        raise ValueError(f"메타 줄 {key!r} 가 없어 {[k for k, _ in new]} 를 끼울 자리를 모른다 — "
                         f"있는 줄: {keys}")
    at = keys.index(key) + 1
    pairs[at:at] = list(new)


# ── 조립 ─────────────────────────────────────────────────────────────────────
@dataclass(frozen=True)
class DailyResult:
    path: Path
    date: str
    basis: str
    spec_id: str
    top: list[tuple[int, str, str]]     # (순위, 코드, 이름) 상위 5
    n_rows: int
    n_ranked: int
    sheets: list[str]
    build_id: str                       # model 판 id — 발송 장부·정정 캡션용(N-25 Q9)
    generated_at: str                   # model 판 생성 시각(판 JSON generated_at)


def daily_path(out_root: Path, d: str, basis: str) -> Path:
    return Path(out_root) / "daily" / f"model_scores_{ymd(d)}_{basis}.xlsx"


def save_atomic(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp")
    wb.save(tmp)
    os.replace(tmp, path)


def build_daily(d: str | date, basis: str, *, model_root: Path, fi_root: Path, out_root: Path,
                config_dir: Path | None = None) -> DailyResult:
    run = load_run(model_root, d, basis)
    view = load_day(model_root, fi_root, run, config_dir=config_dir)
    fi = load_fi(fi_root, view)
    prev = previous_run(model_root, run.date, basis)
    prev_rank: dict[str, int | None] = {}
    if prev is not None:
        sid = view.spec_id if view.spec_id in prev.specs else prev.primary_spec
        try:
            prev_rank = {ticker_of(r): rank_of(r) for r in read_scores(model_root, prev, sid)}
        except DeliverError:
            prev_rank = {}

    trend = load_trend(model_root, run, view.spec_id, basis)

    wb = new_workbook()
    wb.remove(wb.worksheets[0])
    dictionary: Dictionary = []
    dictionary += sheet_scores(wb, view, fi, prev_rank, trend)
    dictionary += sheet_raw(wb, view, fi)
    dictionary += sheet_display(wb, view, fi)
    dictionary += sheet_earnings(wb, view, fi)
    dictionary += sheet_sectors(wb, view)
    dictionary += sheet_models(wb, view)
    pairs = meta_pairs(view, fi_run_meta(fi_root, run.date, basis, run.fi_build_id), fi)
    insert_after(pairs, "model 생성 시각", [
        ("전일 비교 판", "없음" if prev is None else f"{prev.date} {prev.build_id}"),
        *((f"{w} 비교 판", trend.base[w] or "없음(그 전 판이 없다)") for w in WINDOWS),
        ("순위 흐름 판", f"{trend.dates[0]} ~ {trend.dates[-1]} · {len(trend.dates)}개"
                     if trend.dates else "없음")])
    # 같은 model 판으로 코드만 바뀐 정정판을 파일로 구별한다(E-08)
    insert_after(pairs, "equity 판 id", [
        ("엑셀 생성 시각", datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")),
        ("코드 rev", deployed_rev(ql_home()))])
    write_meta(wb, title_for(view, "메타"), pairs, dictionary)
    tickers = [ticker_of(r) for r in view.rows]
    write_trend_sheet(wb, tickers, trend)
    path = daily_path(out_root, run.date, basis)
    save_atomic(wb, path)
    score_ws = wb["점수"]
    heads = {str(score_ws.cell(7, c).value): score_ws.cell(7, c).column_letter
             for c in range(1, score_ws.max_column + 1)}
    cols = {w: heads[f"{w}\n흐름"] for w in WINDOWS if f"{w}\n흐름" in heads}
    deltas = {w: {ticker_of(r): trend.delta(w, ticker_of(r), rank_of(r)) for r in view.rows}
              for w in WINDOWS}
    add_sparklines(path, "점수", spark_groups(tickers, trend, date.fromisoformat(run.date),
                                              cols, deltas))
    top = [(rank_of(r) or 0, ticker_of(r), view.name(ticker_of(r))) for r in view.ranked[:5]]
    return DailyResult(path, run.date, basis, view.spec_id, top, len(view.rows),
                       len(view.ranked),
                       [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"],
                       build_id=run.build_id, generated_at=run.generated_at)


__all__ = ["DailyResult", "FiData", "build_daily", "daily_path", "id_cols", "id_values",
           "ind_cells", "load_fi", "model_label", "period_add", "save_atomic", "title_for"]
