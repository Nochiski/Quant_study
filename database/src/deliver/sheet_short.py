"""Short 시트 — 공매도 · 대차 · 신용(DECISIONS §7 #11, 확정 샘플 30).

원천: equity `short_daily`(키움 공매도 · 대차잔고) · `price_daily`(거래대금 · 시총) · `credit_daily`
(KIS 신용잔고 — 원천이 2거래일쯤 늦다). 기준일 = 키움 공매도와 KRX 거래대금이 둘 다 있는 최신일.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date

import duckdb
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from .extra_common import (
    F_FL,
    F_PCT,
    F_PP,
    F_WON,
    G1,
    HR,
    Calendar,
    Universe,
    base_text,
    krx_calendar,
    latest,
    md,
)
from .qpack import CHART_BLUE, XRefs, XSheet, date_frame, line_series, x_note
from .sheet_market import HIST_DAYS, common_markets
from .sources import SourceMissing, Sources, day_of, lit, num, num0, text

NAME = "Short"
WIDTHS = {"A": 7.5, "B": 18, "C": 14} | dict.fromkeys("DEFGHIJKL", 8)
WEEKS = 26
CREDIT_1M = 20              # 신용 1M 전 = 신용 날짜 20개 전


@dataclass
class MarketDay:
    """그날 보통주 합 — 공매도 거래대금 · 거래대금 · 대차잔고(원)."""

    day: date
    sv: float
    tv: float
    lb: float


@dataclass
class StockDay:
    """종목 하루 — 공매도 거래대금 · 대차잔고 · 거래대금 · 시총(원, 빈칸 None)."""

    sv: float | None
    lb: float | None
    tv: float | None
    mc: float | None


@dataclass
class Credit:
    amt: float                                  # 신용융자 잔고(원)
    rate: float | None                          # 신용잔고율(%)


@dataclass
class ShortData:
    cal: Calendar
    mkt_short: list[MarketDay]                  # 날짜순
    tk_short: dict[str, dict[date, StockDay]]
    credit_dates: tuple[date, date]             # 최신 · 1M 전
    mkt_credit: dict[date, float]               # 날짜 → 신용융자 잔고 합(원)
    tk_credit: dict[str, dict[date, Credit]]


def load(con: duckdb.DuckDBPyConnection, src: Sources, upto: date, uni: Universe) -> ShortData:
    sh, pv = src.equity("short_daily"), src.equity("price_daily")
    day = min(latest(con, sh, "date", upto, "short_value_kiwoom_krw IS NOT NULL", "키움 공매도"),
              latest(con, pv, "date", upto, "value_krw IS NOT NULL", "KRX 종목 일별 거래대금"))
    cal = krx_calendar(con, src, day)
    d250, d25 = cal.back(HIST_DAYS), cal.back(25)
    common_markets(con, src, day)
    con.execute("CREATE OR REPLACE TEMP TABLE uc (ticker VARCHAR)")
    con.executemany("INSERT INTO uc VALUES (?)", [(c,) for c in uni])
    mkt = [MarketDay(day_of(dd), num0(sv), num0(tv), num0(lb)) for dd, sv, tv, lb in con.execute(f"""
        SELECT s.date, sum(coalesce(s.short_value_kiwoom_krw, 0)), sum(coalesce(p.value_krw, 0)),
               sum(coalesce(s.lending_balance_kiwoom_krw, 0))
        FROM {sh} s JOIN mk USING (ticker)
             LEFT JOIN {pv} p ON p.ticker = s.ticker AND p.date = s.date
        WHERE s.date > {lit(d250)} AND s.date <= {lit(day)} GROUP BY 1 ORDER BY 1""").fetchall()]
    tk: dict[str, dict[date, StockDay]] = defaultdict(dict)
    for t, dd, sv, lb, tv, mc in con.execute(f"""
            SELECT s.ticker, s.date, s.short_value_kiwoom_krw, s.lending_balance_kiwoom_krw,
                   p.value_krw, p.mktcap_krw
            FROM {sh} s JOIN uc USING (ticker)
                 LEFT JOIN {pv} p ON p.ticker = s.ticker AND p.date = s.date
            WHERE s.date > {lit(d25)} AND s.date <= {lit(day)}""").fetchall():
        tk[text(t)][day_of(dd)] = StockDay(num(sv), num(lb), num(tv), num(mc))
    cr = src.equity("credit_daily")
    cdays = [day_of(r[0]) for r in con.execute(
        f"SELECT DISTINCT date FROM {cr} WHERE date > {lit(d250)} AND date <= {lit(day)} "
        f"AND whol_loan_rmnd_amt IS NOT NULL ORDER BY 1").fetchall()]
    if len(cdays) <= CREDIT_1M:
        raise SourceMissing(f"KIS 신용잔고 날짜가 {len(cdays)}개뿐이다(1M 비교 {CREDIT_1M + 1}개 필요)")
    cd, cd1 = cdays[-1], cdays[-1 - CREDIT_1M]
    # 금액 단위 만원 → 원(Σ잔고주식수 × 종가 ÷ Σ금액 ≈ 1.09만)
    mkt_cr = {day_of(dd): num0(amt) * 1e4 for dd, amt in con.execute(
        f"SELECT c.date, sum(c.whol_loan_rmnd_amt) FROM {cr} c JOIN mk USING (ticker) "
        f"WHERE c.date > {lit(d250)} AND c.date <= {lit(day)} "
        f"AND c.whol_loan_rmnd_amt IS NOT NULL GROUP BY 1").fetchall()}
    tk_cr: dict[str, dict[date, Credit]] = defaultdict(dict)
    for t, dd, amt, rate in con.execute(
            f"SELECT c.ticker, c.date, c.whol_loan_rmnd_amt, c.whol_loan_rmnd_rate_pct "
            f"FROM {cr} c JOIN uc USING (ticker) WHERE c.date IN ({lit(cd)}, {lit(cd1)}) "
            f"AND c.whol_loan_rmnd_amt IS NOT NULL").fetchall():
        tk_cr[text(t)][day_of(dd)] = Credit(num0(amt) * 1e4, num(rate))
    return ShortData(cal, mkt, tk, (cd, cd1), mkt_cr, tk_cr)


def write(wb: Workbook, s: ShortData, uni: Universe, refs: XRefs, d: date,
          *, index: int | None = None) -> list[tuple[str, str, str]]:
    cal, day = s.cal, s.cal.day
    cd, cd1 = s.credit_dates
    lag = sum(1 for x in cal.days if x > cd)
    xs = XSheet(wb, NAME, WIDTHS, refs, index=index)
    ws = xs.ws
    xs.texts({"A1": "정렬: 블록마다(주 최근 순 · 비율 높은 순)", "A2": "Short: 공매도 · 대차 · 신용",
              "A3": f"자료: 키움 공매도 · 대차잔고 · KIS 신용잔고(보통주) · 모델 유니버스 {len(uni)}종목",
              "A4": base_text(day, d), "F4": "출처: quant-ledger 원장 · 시험 양식",
              "A5": (f"신용잔고는 원천이 {lag}거래일 늦어 최신 {md(cd)} · 1M 전 {md(cd1)}" if lag
                     else f"신용잔고 최신 {md(cd)} · 1M 전 {md(cd1)}")})
    put, header, close, cs = xs.put, xs.header, xs.close, xs.cs
    ms, mc = s.mkt_short, s.mkt_credit

    def wk(x: date) -> tuple[int, int]:
        return x.isocalendar()[:2]
    weeks: dict[tuple[int, int], list[MarketDay]] = defaultdict(list)
    for m in ms:
        weeks[wk(m.day)].append(m)
    r0 = 6
    header(r0, [("A", f"시장 추이: 최근 {WEEKS}주, 주 마지막 거래일"), ("C", "공매도(주간 합)"),
                ("F", "대차잔고(조원)"), ("H", "신용융자 잔고(조원)")],
           ["", "주", "공매도 거래대금\n(억원)", "거래대금\n(조원)", "공매도 비중\n(%)", "잔고", "주간 변화",
            "잔고", "주간 변화"], ab=True)
    r = r0 + 2
    rows_w: list[tuple[date, float, float, float, float | None, date | None]] = []
    for w in sorted(weeks)[-(WEEKS + 1):]:
        g = weeks[w]
        crd = [x for x in sorted(mc) if wk(x) == w]
        rows_w.append((g[-1].day, sum(x.sv for x in g), sum(x.tv for x in g), g[-1].lb,
                       mc[crd[-1]] if crd else None, crd[-1] if crd else None))
    for i in range(len(rows_w) - 1, 0, -1):
        last_d, sv, tv, lb, crv, crd_ = rows_w[i]
        lb0, cr0 = rows_w[i - 1][3], rows_w[i - 1][4]
        xs.lab(r, md(last_d) + (f" (신용 {md(crd_)})" if crd_ and crd_ != last_d else ""))
        put(r, 3, sv / 1e8, F_FL)
        put(r, 4, tv / 1e12, "#,##0.0")
        put(r, 5, sv / tv * 100 if tv else None, "0.00")
        put(r, 6, lb / 1e12, "#,##0.0")
        put(r, 7, (lb - lb0) / 1e12, "+#,##0.0;-#,##0.0;0.0")
        put(r, 8, crv / 1e12 if crv else None, "#,##0.0")
        put(r, 9, (crv - cr0) / 1e12 if crv and cr0 else None, "+#,##0.0;-#,##0.0;0.0")
        r += 1
    s1_end = r - 1
    close(s1_end, 9)
    cs(f"E{r0 + 2}:E{s1_end}")
    # 그래프 4개(1년) — 원자료 숨긴 AD~AH 열
    xs.hide(G1, 5, 10)
    for k, h in enumerate(("날짜", "공매도 비중 20일", "대차잔고", "신용융자 잔고", "공매도 거래대금 20일")):
        ws.cell(HR, G1 + k).value = h
    rr = HR + 1
    sv20: list[float] = []
    tv20: list[float] = []
    for m in ms:
        sv20.append(m.sv)
        tv20.append(m.tv)
        ws.cell(rr, G1).value = m.day
        ws.cell(rr, G1).number_format = "yyyy-mm-dd"
        ws.cell(rr, G1 + 1).value = (round(sum(sv20[-20:]) / sum(tv20[-20:]) * 100, 3)
                                     if sum(tv20[-20:]) else None)
        ws.cell(rr, G1 + 2).value = round(m.lb / 1e12, 3)
        ws.cell(rr, G1 + 3).value = round(mc[m.day] / 1e12, 3) if m.day in mc else None
        ws.cell(rr, G1 + 4).value = round(sum(sv20[-20:]) / len(sv20[-20:]) / 1e8, 1)
        rr += 1
    gend = rr - 1
    cs0 = s1_end + 2
    t1, t2 = cs0, cs0 + 16
    for cell, ttl, col, anchor, yf in (
            (f"A{t1}", "공매도 비중(%): 20일 평균, 1년", 1, f"A{t1 + 1}", "0.0"),
            (f"H{t1}", "대차잔고(조원): 1년", 2, f"H{t1 + 1}", "#,##0"),
            (f"A{t2}", "신용융자 잔고(조원): 1년", 3, f"A{t2 + 1}", "#,##0"),
            (f"H{t2}", "공매도 거래대금(억원): 20일 평균, 1년", 4, f"H{t2 + 1}", "#,##0")):
        xs.title(cell, ttl)
        ch = LineChart()
        ch.add_data(Reference(ws, min_col=G1 + col, min_row=HR, max_row=gend), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=G1, min_row=HR + 1, max_row=gend))
        ch.legend = None
        line_series(ch.series[0], CHART_BLUE)
        date_frame(ch, yf, "yy/mm", "days", 2, "months")
        ch.display_blanks = "gap"
        ws.add_chart(ch, anchor)
    xs.tight(cs0, cs0 + 32)
    d1m = cal.win["1M"]

    def sm(code: str) -> dict[str, float | None]:
        x = s.tk_short.get(code, {})
        ds = sorted(x)

        def ratio(n: int) -> float | None:
            sv = sum(x[q].sv or 0 for q in ds[-n:])
            tv = sum(x[q].tv or 0 for q in ds[-n:])
            return sv / tv * 100 if tv else None

        def lbr(q: date) -> float | None:
            y = x.get(q)
            return y.lb / y.mc * 100 if y and y.lb is not None and y.mc else None
        c = s.tk_credit.get(code, {})
        cr = c[cd].rate if cd in c else None
        cr0 = c[cd1].rate if cd1 in c else None
        l1, l0 = lbr(day), lbr(d1m)
        return {"r5": ratio(5), "r20": ratio(20), "lb": l1,
                "dlb": l1 - l0 if l1 is not None and l0 is not None else None, "cr": cr,
                "dcr": cr - cr0 if cr is not None and cr0 is not None else None}
    smv = {c: sm(c) for c in uni}
    s_end = cs0 + 32
    for title, key in (("공매도 비중 상위 20: 최근 5거래일", "r5"), ("대차잔고 ÷ 시총 상위 20", "lb"),
                       ("대차잔고 ÷ 시총 1M 상승 상위 20", "dlb"), ("신용잔고율 상위 20", "cr")):
        t0 = s_end + 2
        header(t0, [("A", title), ("C", "분류"), ("E", "규모"), ("F", "공매도 비중(%)"),
                    ("H", "대차잔고 ÷ 시총"), ("J", "신용잔고율"), ("L", "수익률")],
               ["코드", "이름", "대분류", "종합\n순위", "시총\n(억원)", "5일", "20일", f"{md(day)}\n(%)",
                "1M 변화\n(%p)", f"{md(cd)}\n(%)", "1M 변화\n(%p)", "1M\n(%)"])
        vals = {c: v for c in uni if (v := smv[c][key]) is not None}
        cand = sorted(vals, key=lambda c: -vals[c])[:20]
        r = t0 + 2
        for c in cand:
            v, u = smv[c], uni[c]
            xs.stock(r, c, u.name, u.big, u.rank, [
                (u.mcap, F_WON, False), (v["r5"], "0.00", False), (v["r20"], "0.00", False),
                (v["lb"], "0.00", False), (v["dlb"], F_PP, False), (v["cr"], "0.00", False),
                (v["dcr"], F_PP, False), (u.r1m, F_PCT, False)])
            r += 1
        s_end = r - 1
        close(s_end, 12)
        cs(f"L{t0 + 2}:L{s_end}")
    ws["E7"].comment = x_note(
        "공매도 비중(%): 공매도 거래대금 ÷ 전체 거래대금(보통주, 주간 합)")
    xs.page(f"A1:L{s_end}")
    return [("Short", "시장 추이 · 그래프",
             "공매도 거래대금 · 비중(키움, 보통주 합), 대차잔고(키움, 금액), 신용융자 잔고(KIS, 금액 단위 "
             f"만원을 원으로 — 원천 {lag}거래일 늦음). 그래프 원자료는 숨긴 AD~AH 열"),
            ("Short", "종목 목록",
             f"공매도 비중 = 기간 공매도 거래대금 ÷ 거래대금, 대차잔고 ÷ 시총({md(day)} · 1M 변화 %p), "
             "신용잔고율 = KIS 신용 잔고 주식 ÷ 상장주식(%), 1M 수익률 = Raw 모멘텀 1M")]


__all__ = ["NAME", "Credit", "MarketDay", "ShortData", "StockDay", "load", "write"]
