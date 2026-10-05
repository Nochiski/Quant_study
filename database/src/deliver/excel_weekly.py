"""주간 엑셀 `data/deliver/weekly/<YYYY-Www>/weekly_<YYYY-Www>.xlsx` (MODEL_EXCEL_SPEC B).

시트: 주간 후보 → 팩터 카드 → 주간 추이 → 이탈·진입 → 지난주 성적 → 메타. 값만(수식 없음).
재료 = 그 ISO 주 월~금의 아침 확정판(`_runs/<YYYYMMDD>_morning.json`). 기준일 = 판이 있는 마지막
거래일(보통 금요일). 판이 없는 날(휴장·미빌드)은 ✕ 이고 평균 순위는 순위가 있는 날만으로 낸다
(사용자 09-25). 후보 = 기준일 주 모델 순위순으로 업종(대분류)당 최대 9, 30 종목(OutputRule).
"""
from __future__ import annotations

import statistics
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from . import qpack
from .common import bucket_label, coverage_label, ind_meta, missing
from .excel_daily import (
    FiData,
    bucket_missing_reason,
    id_cols,
    id_values,
    ind_cells,
    load_fi,
    model_label,
    save_atomic,
    score_indicators,
    title_for,
    universe_text,
)
from .qpack import C_NEW, Col, Group, Title, write_meta, write_table
from .reader import DeliverError, find_run, iso, read_fi
from .view import DayView, composite_of, load_day

Row = dict[str, object]
Dictionary = list[tuple[str, str, str, str]]
BASIS = "morning"
DOW = ("월", "화", "수", "목", "금")
MARK_TOP, MARK_RANKED, MARK_OUT = "●", "○", "✕"


def _num(v: object) -> float | None:
    if isinstance(v, int | float) and not isinstance(v, bool) and v == v:
        return float(v)
    return None


def parse_week(week: str) -> tuple[int, int]:
    """'YYYY-Www' → (ISO 연도, 주)."""
    try:
        y, w = week.split("-W")
        year, num = int(y), int(w)
        date.fromisocalendar(year, num, 1)
    except (ValueError, TypeError) as e:
        raise DeliverError(f"주 형식은 YYYY-Www(예 2026-W39): {week!r}") from e
    return year, num


def week_days(week: str) -> list[str]:
    """그 ISO 주의 월~금(YYYY-MM-DD)."""
    y, w = parse_week(week)
    return [date.fromisocalendar(y, w, k).isoformat() for k in range(1, 6)]


def prev_week(week: str) -> str:
    y, w = parse_week(week)
    iy, iw, _ = (date.fromisocalendar(y, w, 1) - timedelta(days=7)).isocalendar()
    return f"{iy}-W{iw:02d}"


@dataclass
class Week:
    """한 주의 판들. `days` = 월~금, `views` = 판이 있는 날만(가벼운 적재)."""

    week: str
    days: list[str]
    views: dict[str, DayView]

    @property
    def available(self) -> list[str]:
        return [d for d in self.days if d in self.views]

    @property
    def base(self) -> DayView | None:
        avail = self.available
        return self.views[avail[-1]] if avail else None

    def mark(self, d: str, t: str) -> str:
        v = self.views.get(d)
        if v is None or v.rank(t) is None:
            return MARK_OUT
        return MARK_TOP if t in v.candidates else MARK_RANKED

    def ranks(self, t: str) -> list[int]:
        return [r for d in self.available if (r := self.views[d].rank(t)) is not None]


def load_week(model_root: Path, week: str, spec_id: str | None, *,
              config_dir: Path | None = None, fi_root: Path | None = None) -> Week:
    """판이 있는 날만 점수 표를 읽는다. spec_id 가 None 이면 마지막 판의 primary_spec.
    fi_root 는 업종 열이 없는 엔진(v3·v2)의 업종 상한 후보에 쓴다(`view.load_day`)."""
    days = week_days(week)
    runs = {d: r for d in days if (r := find_run(model_root, d, BASIS)) is not None}
    if spec_id is None and runs:
        spec_id = runs[max(runs)].primary_spec
    views: dict[str, DayView] = {}
    for d, run in runs.items():
        if spec_id not in run.specs:
            continue                      # 그날 판에 주 모델이 없으면 ✕ 로 둔다
        views[d] = load_day(model_root, fi_root, run, config_dir=config_dir, with_fi=False,
                            others=False, spec_id=spec_id)
    return Week(week, days, views)


def _dow(d: str) -> str:
    return DOW[date.fromisoformat(d).weekday()] if date.fromisoformat(d).weekday() < 5 else d


def _title(base: DayView, week: Week, sheet: str, note: str = "") -> Title:
    t = title_for(base, sheet, note)
    source = (f"{week.week} · 기준일 {base.date}({_dow(base.date)}) · 거래일 판 "
              f"{len(week.available)}/5 · Source: quant-ledger model·factor_inputs · "
              f"판 {base.run.build_id}")
    return Title(sheet, t.model_line, source, note)


# ── 시트: 주간 후보 ───────────────────────────────────────────────────────────
def _v3_spec(base: DayView) -> str | None:
    return next((s for s in sorted(base.other_ranks) if s.startswith("v3_zscore")), None)


def sheet_candidates(wb: Workbook, base: DayView, week: Week, last: Week | None) -> Dictionary:
    last_base = None if last is None else last.base
    v3 = _v3_spec(base)
    dow = _dow(base.date)
    groups = [
        Group("후보", (Col("pos", "순번", "num", "#,##0", 5.0, color=False,
                          definition="업종 상한 적용 뒤 순서(1~30)"),
                      *id_cols(full=True)[:5],
                      Col("mcap", "시총(억)", "num", "#,##0", 9.0,
                          definition="기준일 시가총액, 억원"))),
        Group("순위", (
            Col("rank", f"{dow}요일 순위\n(원본·상한 전)", "rank", "#,##0", 9.0, key_col=True,
                color=False, definition="기준일 주 모델 순위(엔진 rank, 업종 상한 전)"),
            Col("mean", "주간 평균\n순위", "rank", "#,##0.0", definition=(
                "이번 주 판이 있는 거래일 중 순위가 있는 날의 평균(휴장·미빌드 날은 뺀다)")),
            *(Col(f"m{k}", f"{DOW[k]}\n{d[5:]}", "txt", None, 5.0, definition=(
                "● 그날 후보(상위 30·업종 상한) · ○ 순위는 있으나 후보 밖 · ✕ 제외·판 없음(휴장)"))
              for k, d in enumerate(week.days)),
            Col("n_top", "등장\n(●)", "num", "#,##0", 6.0, color=False,
                definition="이번 주 ● 횟수"),
            Col("last_rank", "지난주\n순위", "num", "#,##0", color=False,
                definition="지난주 기준일 주 모델 순위"),
            Col("new", "신규/\n유지", "txt", None, 6.0,
                definition="지난주 후보에 없으면 신규(형광), 있으면 유지. 지난주 판이 없으면 빈칸"),
            *((Col("v3", "v3 순위\n(참고)", "num", "#,##0", color=False,
                   definition=f"{v3} 기준일 순위(참고)"),) if v3 else ()),
        ), core=True),
        Group("점수", (Col("composite", "종합\n점수", "num", "#,##0.00",
                         definition="기준일 종합 점수(엔진)"),
                     *(Col(f"{b}_u", bucket_label(b), "pct", "#,##0.0",
                           definition=f"{bucket_label(b)} 유니버스 백분위(버킷 점수 재순위)")
                       for b in base.buckets))),
    ]
    rows: list[Row] = []
    for pos, t in enumerate(base.candidates, start=1):
        ranks = week.ranks(t)
        out: Row = {**id_values(base, t), "pos": pos, "rank": base.rank(t),
                    "mean": statistics.fmean(ranks) if ranks else None,
                    "last_rank": None if last_base is None else last_base.rank(t),
                    "composite": composite_of(base.by_ticker.get(t))}
        marks = [week.mark(d, t) for d in week.days]
        for k, m in enumerate(marks):
            out[f"m{k}"] = m
        out["n_top"] = marks.count(MARK_TOP)
        if last_base is not None:
            is_new = t not in last_base.candidates
            out["new"] = "신규" if is_new else "유지"
            if is_new:
                out["_fills"] = {"new": C_NEW}
        if v3:
            out["v3"] = base.other_ranks[v3].get(t)
        for b in base.buckets:
            out[f"{b}_u"] = base.upct[b].get(t)
        rows.append(out)
    note = (f"기준일 {base.date} 주 모델 순위순, 대분류당 최대 {base.output.max_per_sector}. "
            "● 후보 · ○ 순위 있음 · ✕ 제외/판 없음. 평균 순위는 순위가 있는 날만.")
    return write_table(wb, _title(base, week, "주간 후보", note), groups, rows, sort_key="pos")


# ── 시트: 팩터 카드 ───────────────────────────────────────────────────────────
CARD_HEADERS = ("항목", "구분", "값", "유니버스 백분위", "업종 백분위", "업종 중앙값", "대비",
                "표식·결측")
CARD_WIDTHS = (22.0, 16.0, 13.0, 13.0, 13.0, 13.0, 13.0, 30.0)


def _sector_members(base: DayView, t: str) -> list[str]:
    l1 = base.by_ticker.get(t, {}).get("sector_l1")
    return [x for x, r in base.by_ticker.items() if r.get("sector_l1") == l1]


def sheet_cards(wb: Workbook, base: DayView, week: Week, fi: FiData) -> Dictionary:
    ws = wb.create_sheet("팩터 카드")
    note = ("후보 30 × 세로 블록: 축(엔진 버킷 점수) · 점수 지표(원값·엔진 백분위) · "
            "최근 5일 순위 · "
            "기준일. 업종 = WICS 대분류.")
    qpack.title_block(ws, _title(base, week, "팩터 카드", note), len(CARD_HEADERS), None)
    for k, w in enumerate(CARD_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    neutral = set()
    if base.spec is not None:
        nb = base.spec.params.get("sector_neutral_buckets") or ()
        neutral = {str(x) for x in nb} if isinstance(nb, list | tuple) else set()
    inds = score_indicators(base)
    r = 6
    for pos, t in enumerate(base.candidates, start=1):
        row = base.by_ticker[t]
        u = base.uni.get(t, {})
        ranks = week.ranks(t)
        head = (f"{pos}. {t} {base.name(t)} · {u.get('sector_l1_name') or ''} / "
                f"{u.get('sector_l2_name') or ''} · 순위 {base.rank(t)} · 주간 평균 "
                f"{statistics.fmean(ranks):.1f}" if ranks else f"{pos}. {t} {base.name(t)}")
        comp = composite_of(row)
        if comp is not None:
            head += f" · 종합 {comp:.2f}"
        for k in range(1, len(CARD_HEADERS) + 1):
            ws.cell(r, k).fill = qpack.fill(qpack.C_BAND)
        qpack.put(ws, r, 1, head).font = qpack.font(bold=True, color="FFFFFF")
        r += 1
        for k, text in enumerate(CARD_HEADERS, start=1):
            cell = qpack.put(ws, r, k, text)
            cell.fill = qpack.fill(qpack.C_HDR)
            cell.font = qpack.font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="top")
        r += 1
        members = _sector_members(base, t)
        lines: list[tuple[object, ...]] = []
        for b in base.buckets:
            score = base.bucket_score(t, b)
            peers = [v for x in members if (v := base.bucket_score(x, b)) is not None]
            med = statistics.median(peers) if peers else None
            w = "" if base.spec is None else f"축 {base.spec.buckets[b] * 100:.0f}%"
            lines.append((bucket_label(b), w, score, base.upct[b].get(t), base.spct[b].get(t),
                          med, None if score is None or med is None else score - med,
                          "" if score is not None else missing(bucket_missing_reason(base, t, b))))
        for key, b, _ in inds:
            v, pct, flag = ind_cells(base, t, key)
            scale = ind_meta(key).scale
            peers = [x * scale for m in members
                     if (x := _num(base.ind.get(m, {}).get(key, {}).get("raw"))) is not None]
            med = statistics.median(peers) if peers else None
            vnum = _num(v)
            upct, spct = (None, pct) if b in neutral else (pct, None)
            lines.append((ind_meta(key).label, f"{bucket_label(b)} · {key}", v, upct, spct, med,
                          None if vnum is None or med is None else vnum - med, flag))
        day_marks = tuple(
            f"{DOW[k]} {week.views[d].rank(t) or '✕'}" if d in week.views else f"{DOW[k]} ✕"
            for k, d in enumerate(week.days))
        lines.append(("최근 5일 순위", "월~금", *day_marks, ""))
        a = fi.annual(t, ("ni",))
        fy0 = fi.fy0(t, base.date)
        cur = None if fy0 is None else fi.cons_at(t, fy0)
        lines.append(("기준일", "재무·추정·수급",
                      f"재무 {'-' if a is None else iso(a.get('available_date'))}",
                      f"추정 {'-' if cur is None else iso(cur.get('obs_date'))}",
                      f"수급 {fi.flow_last.get(t) or '-'}", "", "",
                      "커버리지 " + coverage_label(row.get("coverage_state", u.get("coverage_state")),
                                                u.get("coverage_age_days"))))
        for line in lines:
            for k, v in enumerate(line, start=1):
                cell = qpack.put(ws, r, k, v)
                cell.font = qpack.font(bold=k == 1)
                cell.alignment = Alignment(vertical="top",
                                           horizontal="left" if isinstance(v, str) or k <= 2
                                           else "right")
                if isinstance(cell.value, int | float):
                    cell.number_format = "#,##0.00" if k in (3, 6, 7) else "#,##0.0"
            r += 1
        r += 1
    if r > 8:
        ws.conditional_formatting.add(f"D8:E{r}", qpack.scale_high_red())
    ws.freeze_panes = ws.cell(6, 1)
    qpack.tighten_rows(ws)
    return [("팩터 카드", "카드", h, d) for h, d in (
        ("항목", "축 이름 · 점수 지표 라벨 · 최근 5일 순위 · 기준일"),
        ("값", "축 = 엔진 버킷 점수(0~100) · 지표 = 원값(% 표시, 리비전 이전값<0 은 표식만)"),
        ("유니버스 백분위",
         "축 = 버킷 점수 유니버스 백분위 · 지표 = 엔진 백분위(유니버스 기준 버킷)"),
        ("업종 백분위", "축 = 버킷 점수 대분류 안 백분위 · 지표 = 엔진 백분위(대분류 기준 버킷)"),
        ("업종 중앙값", "같은 WICS 대분류 모집단의 중앙값(값 있는 종목)"),
        ("대비", "값 − 업종 중앙값"),
        ("표식·결측", "엔진 flag · 결측(사유)"))]


# ── 시트: 주간 추이 ───────────────────────────────────────────────────────────
def sheet_trend(wb: Workbook, base: DayView, week: Week, last: Week | None) -> Dictionary:
    last_c = [] if last is None or last.base is None else last.base.candidates
    order = list(base.candidates) + [t for t in last_c if t not in base.candidates]
    groups = [
        Group("종목", (*id_cols(), Col("status", "구분", "txt", None, 6.0, definition=(
            "유지(이번 주·지난주 후보) · 신규(이번 주만) · 이탈(지난주만)")))),
        Group("일별 순위", tuple(Col(f"r{k}", f"{DOW[k]}\n{d[5:]}", "rank", "#,##0", definition=(
            "그날 주 모델 순위. ✕ = 판 없음(휴장·미빌드), 제외 = 그날 순위 없음"))
            for k, d in enumerate(week.days)), core=True),
        Group("일별 종합 점수", tuple(Col(f"c{k}", f"{DOW[k]}\n{d[5:]}", "num", "#,##0.00",
                                        definition="그날 종합 점수(엔진)")
                                    for k, d in enumerate(week.days))),
    ]
    rows: list[Row] = []
    for t in order:
        out: Row = id_values(base, t)
        in_now, in_last = t in base.candidates, t in last_c
        out["status"] = ("유지" if in_now and in_last else "신규" if in_now and last_c
                         else "" if in_now else "이탈")
        for k, d in enumerate(week.days):
            v = week.views.get(d)
            if v is None:
                out[f"r{k}"], out[f"c{k}"] = MARK_OUT, None
                continue
            out[f"r{k}"] = v.rank(t) if v.rank(t) is not None else "제외"
            out[f"c{k}"] = composite_of(v.by_ticker.get(t))
        rows.append(out)
    note = "이번 주 후보 + 지난주 후보(이탈 포함)의 월~금 순위·종합 점수."
    return write_table(wb, _title(base, week, "주간 추이", note), groups, rows)


# ── 시트: 이탈·진입 ───────────────────────────────────────────────────────────
def _biggest_move(now: DayView, before: DayView, t: str) -> tuple[str, float | None]:
    best: tuple[str, float | None] = ("", None)
    for b in now.buckets:
        a, p = now.bucket_score(t, b), before.bucket_score(t, b)
        if a is None or p is None:
            continue
        if best[1] is None or abs(a - p) > abs(best[1]):
            best = (bucket_label(b), a - p)
    return best


def sheet_moves(wb: Workbook, base: DayView, week: Week, last: Week | None) -> Dictionary:
    lb = None if last is None else last.base
    groups = [
        Group("종목", (Col("dir", "방향", "txt", None, 6.0, definition="진입 · 이탈"),
                     *id_cols())),
        Group("순위", (
            Col("last_rank", "지난주\n순위", "rank", "#,##0", definition="지난주 기준일 순위"),
            Col("rank", "이번주\n순위", "rank", "#,##0", definition="이번 주 기준일 순위"),
            Col("d_rank", "Δ순위", "chg", "#,##0", definition="지난주 − 이번주(양수 = 상승)"),
            Col("state", "이번주 상태", "txt", None, 16.0, definition=(
                "이탈 사유 — 제외 사유(엔진) · 순위 밖(업종 상한 포함)")),
        ), core=True),
        Group("축 변화", (
            Col("axis", "가장 크게\n변한 축", "txt", None, 10.0,
                definition="두 기준일 모두 점수가 있는 축 중 |변화| 최대"),
            Col("move", "변화량(점)", "chg", "#,##0.0", definition="이번 주 − 지난주 버킷 점수"),
        )),
    ]
    rows: list[Row] = []
    if lb is not None:
        entered = [t for t in base.candidates if t not in lb.candidates]
        left = [t for t in lb.candidates if t not in base.candidates]
        for direction, tickers in (("진입", entered), ("이탈", left)):
            for t in tickers:
                now_r, last_r = base.rank(t), lb.rank(t)
                excl = base.by_ticker.get(t, {}).get("exclude_reason")
                state = "" if direction == "진입" else (
                    f"제외({excl})" if excl else ("순위 밖" if now_r is not None
                                                  else "모집단 밖"))
                axis, move = _biggest_move(base, lb, t)
                out: Row = {**id_values(base, t), "dir": direction, "last_rank": last_r,
                            "rank": now_r, "d_rank": None if now_r is None or last_r is None
                            else last_r - now_r, "state": state, "axis": axis, "move": move}
                if direction == "진입":
                    out["_fills"] = {"dir": C_NEW}
                rows.append(out)
    note = ("지난주 후보 대비 진입·이탈." if lb is not None
            else "지난주 판이 없어 비교하지 않는다.")
    return write_table(wb, _title(base, week, "이탈·진입", note), groups, rows)


# ── 시트: 지난주 성적 ─────────────────────────────────────────────────────────
@dataclass(frozen=True)
class Scorecard:
    mean: float | None
    median: float | None
    uni_mean: float | None
    uni_median: float | None
    win_rate: float | None
    n: int

    @property
    def excess(self) -> float | None:
        return None if self.mean is None or self.uni_mean is None else self.mean - self.uni_mean

    def text(self) -> str:
        def f(x: float | None, unit: str = "%") -> str:
            return "-" if x is None else f"{x:+.2f}{unit}"
        wr = "-" if self.win_rate is None else f"{self.win_rate:.0f}%"
        return (f"후보 평균 {f(self.mean)} · 중앙값 {f(self.median)} · 유니버스 평균 "
                f"{f(self.uni_mean)} · 중앙값 {f(self.uni_median)} · 초과(평균) "
                f"{f(self.excess, '%p')}"
                f" · 승률 {wr} (n={self.n})")


def _returns(fi_root: Path, base: DayView, start: str, end: str) -> dict[str, float | str]:
    """수정종가 1주 수익률(%). 두 날 adj_ok 가 다르면(미해결 사건을 넘음)
    '결측(수정주가미해결)'."""
    a, b = date.fromisoformat(start).isoformat(), date.fromisoformat(end).isoformat()
    rows = read_fi(fi_root, "fi_adj_prices", base.run.fi_build_id,
                   columns=("ticker", "date", "adj_close", "adj_ok"),
                   where=f"date IN (DATE '{a}', DATE '{b}')")
    px: dict[str, dict[str, tuple[float | None, object]]] = {}
    for r in rows:
        px.setdefault(str(r["ticker"]), {})[iso(r["date"]) or ""] = (_num(r["adj_close"]),
                                                                        r["adj_ok"])
    out: dict[str, float | str] = {}
    for t, by in px.items():
        s, e = by.get(a), by.get(b)
        if s is None or e is None or not s[0] or not e[0]:
            out[t] = missing("가격없음")
        elif s[1] is not None and e[1] is not None and bool(s[1]) != bool(e[1]):
            out[t] = missing("수정주가미해결")
        else:
            out[t] = (e[0] / s[0] - 1.0) * 100.0
    return out


def sheet_scorecard(wb: Workbook, fi_root: Path, base: DayView, week: Week,
                    last: Week | None) -> tuple[Dictionary, Scorecard | None]:
    lb = None if last is None else last.base
    groups = [
        Group("지난주 후보", (Col("pos", "순번", "num", "#,##0", 5.0, color=False,
                              definition="지난주 후보 순서"), *id_cols(),
                          Col("last_rank", "지난주\n순위", "num", "#,##0", color=False,
                              definition="지난주 기준일 순위"))),
        Group("1주 성과(%)", (
            Col("ret", "1주\n수익률", "chg", "#,##0.00",
                definition="지난주 기준일 → 이번 주 기준일 수정종가 수익률"),
            Col("uni", "유니버스\n평균", "num", "#,##0.00",
                definition="지난주 순위 종목 전체의 같은 기간 수익률 평균"),
            Col("excess", "초과\n(%p)", "chg", "#,##0.00", definition="1주 수익률 − 유니버스 평균"),
            Col("win", "승", "txt", None, 4.0, definition="초과 > 0 이면 ●(자기채점)"),
        ), core=True),
    ]
    rows: list[Row] = []
    card: Scorecard | None = None
    if lb is not None and lb.date != base.date:
        rets = _returns(fi_root, base, lb.date, base.date)
        uni = [v for t in (x for x in lb.by_ticker if lb.rank(x) is not None)
               if isinstance(v := rets.get(t), float)]
        uni_mean = statistics.fmean(uni) if uni else None
        cand = []
        for pos, t in enumerate(lb.candidates, start=1):
            ret = rets.get(t, missing("가격없음"))
            ex = ret - uni_mean if isinstance(ret, float) and uni_mean is not None else None
            if isinstance(ret, float):
                cand.append(ret)
            rows.append({**id_values(base, t), "name": base.name(t) or lb.name(t),
                         "pos": pos, "last_rank": lb.rank(t), "ret": ret, "uni": uni_mean,
                         "excess": ex, "win": "" if ex is None else ("●" if ex > 0 else "")})
        wins = [r for r in rows if r["excess"] is not None]
        card = Scorecard(
            statistics.fmean(cand) if cand else None,
            statistics.median(cand) if cand else None, uni_mean,
            statistics.median(uni) if uni else None,
            sum(1 for r in wins if r["win"] == "●") / len(wins) * 100 if wins else None,
            len(cand))
    note = (card.text() if card is not None else "지난주 판이 없어 성적을 매기지 않는다.")
    return write_table(wb, _title(base, week, "지난주 성적", note), groups, rows), card


# ── 조립 ─────────────────────────────────────────────────────────────────────
@dataclass(frozen=True)
class WeeklyResult:
    path: Path
    week: str
    base_date: str
    spec_id: str
    top: list[tuple[int, str, str]]      # (순번, 코드, 이름) 상위 5
    n_candidates: int
    n_new: int
    n_out: int
    n_days: int
    sheets: list[str]


def weekly_path(out_root: Path, week: str) -> Path:
    return Path(out_root) / "weekly" / week / f"weekly_{week}.xlsx"


def _meta_pairs(base: DayView, week: Week, last: Week | None,
                card: Scorecard | None) -> list[tuple[str, object]]:
    pairs: list[tuple[str, object]] = [
        ("주", week.week), ("기준일", f"{base.date}({_dow(base.date)}) — 판이 있는 마지막 거래일"),
        ("basis", BASIS), ("주 모델", f"{base.spec_id}({model_label(base.spec_id)})")]
    for k, d in enumerate(week.days):
        v = week.views.get(d)
        pairs.append((f"{DOW[k]} {d}", "판 없음(휴장·미빌드) — ✕" if v is None else
                      f"model {v.run.build_id} · fi {v.run.fi_build_id} · 순위 {len(v.ranked)} · "
                      f"모집단 {len(v.rows)}"))
    lb = None if last is None else last.base
    pairs.append(("지난주 기준일", "없음" if lb is None else f"{lb.date} · {lb.run.build_id}"))
    pairs += [("유니버스 규칙", universe_text(base)),
              ("후보 규칙", f"상위 {base.output.top_n} · {base.output.sector_level} 업종당 최대 "
                        f"{base.output.max_per_sector} · 기준일 순위순")]
    if base.spec is not None:
        pairs.append(("가중치", " · ".join(f"{bucket_label(b)} {w * 100:.1f}%"
                                        for b, w in base.spec.buckets.items())))
    miss = sum(1 for t in base.candidates
               if any(base.bucket_score(t, b) is None for b in base.buckets))
    pairs += [("결측 요약",
               f"거래일 판 {len(week.available)}/5 · 후보 중 결측 축 있는 종목 {miss}"),
              ("지난주 성적", "없음" if card is None else card.text()),
              ("표식", "● 그날 후보 · ○ 순위 있으나 후보 밖 · ✕ 제외 또는 판 없음(휴장). "
                     "평균 순위는 "
                     "순위가 있는 날만"),
              ("확신 표식(SUE×ΔP/E)", "미구현 — SUE·서프라이즈 원천이 factor_inputs 에 없다")]
    return pairs


def build_weekly(week: str, *, model_root: Path, fi_root: Path, out_root: Path,
                 config_dir: Path | None = None) -> WeeklyResult:
    now = load_week(model_root, week, None, config_dir=config_dir, fi_root=fi_root)
    if now.base is None:
        runs_dir = Path(model_root) / "_runs"
        raise DeliverError(f"{week} 에 {BASIS} model 판이 하나도 없다: {runs_dir}")
    spec_id = now.base.spec_id
    base = load_day(model_root, fi_root, now.base.run, config_dir=config_dir, spec_id=spec_id)
    now.views[base.date] = base
    last = load_week(model_root, prev_week(week), spec_id, config_dir=config_dir,
                     fi_root=fi_root)
    last_or_none = last if last.base is not None else None
    fi = load_fi(fi_root, base)

    wb = Workbook()
    wb.remove(wb.worksheets[0])
    dictionary: Dictionary = []
    dictionary += sheet_candidates(wb, base, now, last_or_none)
    dictionary += sheet_cards(wb, base, now, fi)
    dictionary += sheet_trend(wb, base, now, last_or_none)
    dictionary += sheet_moves(wb, base, now, last_or_none)
    d, card = sheet_scorecard(wb, fi_root, base, now, last_or_none)
    dictionary += d
    write_meta(wb, _title(base, now, "메타"), _meta_pairs(base, now, last_or_none, card),
               dictionary)
    path = weekly_path(out_root, week)
    save_atomic(wb, path)
    lb = None if last_or_none is None else last_or_none.base
    n_new = 0 if lb is None else sum(1 for t in base.candidates if t not in lb.candidates)
    n_out = 0 if lb is None else sum(1 for t in lb.candidates if t not in base.candidates)
    top = [(i, t, base.name(t)) for i, t in enumerate(base.candidates[:5], start=1)]
    return WeeklyResult(path, week, base.date, spec_id, top, len(base.candidates), n_new, n_out,
                        len(now.available), wb.sheetnames)


__all__ = ["BASIS", "Week", "WeeklyResult", "build_weekly", "load_week", "parse_week",
           "prev_week", "week_days", "weekly_path"]

