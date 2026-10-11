"""ETF 시트 — 유형별 자금 흐름과 수익률(DECISIONS §7 #10, 확정 샘플 28 `make_etf6.py`).

위에서 아래로: 유형별 요약(+12주 유입 셀 안 막대) · 그래프 4개(월간 순유입 · 26주 누적 순유입 ·
연말 순자산 · 거래대금 비중) · 주간(26주) · 월별(12개월) · 연도별(자료 첫해~) 표 · 운용사별 ·
1M 유입 상위 · 유출 상위 15 · 전체 목록(순자산 순). 계산은 `etf_data`.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from .etf_data import MONTHS_TABLE, WEEKS_SPARK, WEEKS_TABLE, Brand, EtfData, EtfRec, iso_week
from .extra_common import base_text
from .qpack import (
    CHART_BLACK,
    CHART_BLUE,
    XRefs,
    XSheet,
    bar_series,
    date_frame,
    line_series,
    x_note,
)

NAME = "ETF"
NC = 17                                                    # A..Q(제목 띠 폭)
WIDTHS = {"A": 7.5, "B": 26, "C": 11, "D": 13, "E": 10, "F": 8.5, "G": 6.4, "H": 6.4, "I": 6.4,
          "J": 6.4, "K": 6.4, "L": 8.5, "M": 8.5, "N": 8.5, "O": 9, "P": 12, "Q": 24}
HELP0 = 20                                                 # T 열부터 12주 흐름 원자료(숨김)
G0 = 33                                                    # AG 열부터 그래프 원자료(숨김)
F_AUM, F_R, F_FL, F_P, F_T = "#,##0", "#,##0.0;-#,##0.0", "#,##0;-#,##0", "0.00;-0.00", "#,##0.0"
CATS = ["국내주식", "해외주식", "채권", "단기자금·금리", "레버리지(롱)", "인버스", "혼합·자산배분",
        "원자재·통화", "리츠·부동산"]
DOM = ["시장대표", "반도체·IT", "전략(배당·커버드콜·밸류업)", "조선·방산·원전·전력", "2차전지", "그룹주",
       "금융", "자동차", "바이오·헬스", "인터넷·게임·콘텐츠", "기타 섹터·테마"]
GRP = [("국내주식", ("국내주식", "레버리지(롱)", "인버스")), ("해외주식", ("해외주식",)),
       ("채권·금리", ("채권", "단기자금·금리")), ("기타", ("혼합·자산배분", "원자재·통화", "리츠·부동산"))]
ROWS_PER_CHART = 16
CH = 15                                                    # 그래프 한 줄 = 제목 1 + 그래프 14행


@dataclass
class Written:
    """시트 뒤 손질(저장 뒤 셀 안 막대)과 Meta 열 사전."""

    defs: list[tuple[str, str, str]] = field(default_factory=list)
    sparks: list[tuple[str, str]] = field(default_factory=list)


def week_span(days: list[date]) -> str:
    """이번 주 거래일 '10-06~08' · '09-30~10-02' · '10-06'."""
    a, b = days[0], days[-1]
    if a == b:
        return f"{a:%m-%d}"
    return f"{a:%m-%d}~{b:%d}" if a.month == b.month else f"{a:%m-%d}~{b:%m-%d}"


def write(wb: Workbook, e: EtfData, refs: XRefs, d: date, *, index: int | None = None) -> Written:
    xs = XSheet(wb, NAME, WIDTHS, refs, ncol=NC, index=index)
    ws = xs.ws
    xs.hide(HELP0, WEEKS_SPARK, 6)
    day, recs = e.day, e.recs
    md = f"{day:%m-%d}"
    xs.texts({
        "A1": "정렬: 블록마다(요약 유형 순 · 시계열 최근 순 · 목록 순자산 순)",
        "A2": "ETF: 유형별 자금 흐름과 수익률",
        "A3": f"자료: KRX ETF 일별 {e.first_day:%Y-%m}~{day.isoformat()}(종가 · NAV · 상장좌수 · "
              f"거래대금) · 지금 {len(recs):,}종목",
        "A4": base_text(day, d),
        "F4": "출처: quant-ledger 원장(KRX ETF 일별) · 시험 양식",
        "A5": "유형: 이름 · 기초지수로 자동 분류(시험)",
        "G5": "수익률: 종가, 분배금 미반영",
        "L5": "자금 유입: 상장좌수 변화 × NAV(설정 − 환매 추정)"})
    cs, put, lab, close, header = xs.cs, xs.put, xs.lab, xs.close, xs.header

    weeks = sorted(e.weekly)
    w12, w26 = weeks[-WEEKS_SPARK:], weeks[-WEEKS_TABLE:]
    last_week = [x for x in e.days if iso_week(x) == weeks[-1]]

    def wseries(kind: str, key: str | None) -> list[float]:
        if kind == "all":
            return [sum(e.weekly[w].values()) for w in w12]
        if kind == "big":
            return [e.weekly[w].get(str(key), 0.0) for w in w12]
        return [e.weekly_sub.get(w, {}).get(str(key), 0.0) for w in w12]

    # ── 유형별 요약 ──
    r0 = 6
    header(r0, [("A", "유형별 요약"), ("C", "규모"), ("E", "순자산 · 거래"),
                ("G", "수익률(%, 순자산가중)"), ("L", "자금 유입(억원)"), ("P", "흐름")],
           ["", "유형", "종목 수", "순자산\n비중(%)", "순자산\n(억원)", "거래대금\n20일(억원)", "1D", "1W",
            "1M", "3M", "YTD", "1W", "1M", "3M", "YTD", "12주 유입\n(주간)"], ab=True)
    # (이름, 유형, 세부, 들여쓰기, 12주 흐름 키) — 유형 · 세부가 None 이면 거르지 않는다
    rows_: list[tuple[str, str | None, str | None, int, tuple[str, str | None]]] = [
        ("전체 ETF", None, None, 0, ("all", None)), ("국내주식", "국내주식", None, 1, ("big", "국내주식")),
        *((s, "국내주식", s, 2, ("sub", f"국내주식|{s}")) for s in DOM),
        ("해외주식", "해외주식", None, 1, ("big", "해외주식")),
        *((s, "해외주식", s, 2, ("sub", f"해외주식|{s}")) for s in ("미국", "중국·홍콩", "기타 해외")),
        *((b, b, None, 1, ("big", b)) for b in ("레버리지(롱)", "인버스", "채권", "단기자금·금리",
                                                 "혼합·자산배분", "원자재·통화", "리츠·부동산"))]
    tot = sum(x.aum for x in recs)
    r = r0 + 2
    out = Written()
    for label, big, sub, lvl, wk in rows_:
        g = [x for x in recs if (big is None or x.big == big) and (sub is None or x.sub == sub)]
        a = sum(x.aum for x in g)

        def wavg(k: str, g: list[EtfRec] = g) -> float | None:
            pts = [(v, x.aum) for x in g if (v := x.ret[k]) is not None]
            w = sum(p[1] for p in pts)
            return sum(v * w_ for v, w_ in pts) / w if w else None
        lab(r, label, bold=(lvl <= 1), indent=lvl)
        put(r, 3, len(g), "#,##0")
        put(r, 4, a / tot * 100 if tot else None, "0.0")
        put(r, 5, a, F_AUM)
        put(r, 6, sum((x.tv20 or 0) for x in g), F_AUM)
        for c, k in zip(range(7, 12), ("1D", "1W", "1M", "3M", "YTD"), strict=True):
            put(r, c, wavg(k), F_R)
        for c, k in zip(range(12, 16), ("1W", "1M", "3M", "YTD"), strict=True):
            put(r, c, sum(x.flow[k] for x in g), F_FL)
        put(r, 16, None)
        for i, v in enumerate(wseries(*wk)):
            ws.cell(r, HELP0 + i).value = round(v, 1)
        out.sparks.append((f"ETF!{ws.cell(r, HELP0).column_letter}{r}:"
                           f"{ws.cell(r, HELP0 + WEEKS_SPARK - 1).column_letter}{r}", f"P{r}"))
        r += 1
    a_end = r - 1
    close(a_end, 16)
    cs(f"G{r0 + 2}:K{a_end}")
    cs(f"L{r0 + 2}:O{a_end}")
    ws["A7"].comment = x_note(
        "유형: ETF 이름과 기초지수 이름으로 자동 분류(시험). 순서: 레버리지 · 인버스 → 단기자금 · 금리 → "
        "혼합 → 채권 → 원자재 · 통화 → 리츠 → 해외주식(미국 · 중국·홍콩 · 기타) → 국내주식(시장대표 · "
        "반도체·IT · 전략 · 업종·테마). 이름에 코리아 · 한국 · K- 가 있으면 국내")
    ws["E7"].comment = x_note(
        "순자산(억원) = NAV × 상장좌수(KRX 순자산 칸과 같은 값"
        + (f" — {md} 은 그 칸이 비어 있어 이 식으로 계산)" if e.netasset_empty else ")"))
    ws["G7"].comment = x_note(
        "수익률(%) = 종가 ÷ 기간 첫날 종가 − 1, 분배금 미반영. 1D · 1W · 1M · 3M = 1 · 5 · 20 · "
        f"60거래일 전, YTD = {e.win['YTD'].isoformat()} 대비. 요약 표는 유형 안 순자산가중")
    ws["L7"].comment = x_note(
        "자금 유입(억원) = 기간 안 날마다 (상장좌수 − 전날 좌수) × 그날 NAV 의 합. 설정(유입) − "
        "환매(유출) 추정이라 가격 변동은 들어가지 않는다. 신규 상장은 첫날 좌수 전부를 유입, 상장폐지는 "
        "다음 거래일 마지막 순자산만큼 유출(시계열 표만 — 위 요약 · 목록은 지금 상장된 ETF 기준). "
        f"액면분할(좌수 배수와 NAV 가 정확히 맞는 날)은 유입에서 빼는데 {e.first_day.year}년 이후 "
        f"{e.splits}건")
    ws["P7"].comment = x_note(
        f"12주 유입(주간): 최근 {WEEKS_SPARK}주 주별 자금 유입 막대(위 = 유입 초록, 아래 = 유출 빨강). "
        f"마지막 막대는 이번 주({week_span(last_week)}, {len(last_week)}거래일)")

    # ── 그래프 자리 · 주간 26주 ──
    cs0 = a_end + 2
    t_w = cs0 + 2 * ROWS_PER_CHART + 1
    header(t_w, [("A", f"주간 자금 유입(억원) — 최근 {WEEKS_TABLE}주"), ("C", "유형별"),
                 ("M", f"{WEEKS_TABLE}주 누적")],
           ["", "주(마지막 거래일)", *CATS, "합계", "국내주식", "해외주식"], ab=True)
    r = t_w + 2
    cum_d = cum_f = 0.0
    for w in w26:
        lab(r, f"{e.weekly_end[w]:%m-%d}" + (f"({len(last_week)}일)" if w == weeks[-1] else ""))
        vals = [e.weekly[w].get(c, 0.0) for c in CATS]
        for i, v in enumerate(vals):
            put(r, 3 + i, v, F_FL)
        put(r, 3 + len(CATS), sum(vals), F_FL)
        cum_d += vals[0]
        cum_f += vals[1]
        put(r, 13, cum_d, F_FL)
        put(r, 14, cum_f, F_FL)
        for c in range(15, NC + 1):
            put(r, c, None)
        r += 1
    w_end = r - 1
    close(w_end, 14)
    cs(f"C{t_w + 2}:K{w_end}")
    cs(f"L{t_w + 2}:L{w_end}")

    # ── 월별 12개월 ──
    t_m = w_end + 2
    months = sorted(e.monthly)[-MONTHS_TABLE:]
    header(t_m, [("A", f"월별 자금 유입(억원) — 최근 {MONTHS_TABLE}개월"), ("C", "유형별"), ("M", "규모")],
           ["", "월", *CATS, "합계", "월말 순자산\n(조원)", "ETF 거래대금\n÷ 주식(%)"], ab=True)
    r = t_m + 2
    for m in months:
        lab(r, f"{m[:4]}-{m[4:]}" + (f" ({day.day}일까지)" if m == months[-1] else ""))
        vals = [e.monthly[m].get(c, 0.0) for c in CATS]
        for i, v in enumerate(vals):
            put(r, 3 + i, v, F_FL)
        put(r, 12, sum(vals), F_FL)
        put(r, 13, sum(e.aum_month_end[m].values()) / 1e4, F_T)
        put(r, 14, e.tv_share_month.get(m), F_T)
        for c in range(15, NC + 1):
            put(r, c, None)
        r += 1
    m_end = r - 1
    close(m_end, 14)
    cs(f"C{t_m + 2}:K{m_end}")
    cs(f"L{t_m + 2}:L{m_end}")

    # ── 연도별 ──
    t_y = m_end + 2
    years = sorted(e.yearly)
    header(t_y, [("A", "연도별 자금 유입(억원) · 연말 순자산(조원) — "
                       f"{e.first_day.year}~"), ("C", "자금 유입(억원)"), ("M", "연말 순자산(조원)")],
           ["", "연도", *CATS, "합계", "전체", *(g for g, _ in GRP)], ab=True)
    r = t_y + 2
    for y in years:
        lab(r, y + (f" ({md}까지)" if y == years[-1] else ""))
        vals = [e.yearly[y].get(c, 0.0) for c in CATS]
        for i, v in enumerate(vals):
            put(r, 3 + i, v, F_FL)
        put(r, 12, sum(vals), F_FL)
        ae = e.aum_year_end[y]
        put(r, 13, sum(ae.values()) / 1e4, F_T)
        for i, (_, cl) in enumerate(GRP):
            put(r, 14 + i, sum(ae.get(c, 0.0) for c in cl) / 1e4, F_T)
        r += 1
    y_end = r - 1
    close(y_end, 17)
    cs(f"C{t_y + 2}:K{y_end}")
    cs(f"L{t_y + 2}:L{y_end}")

    # ── 운용사 ──
    t_b = y_end + 2
    brands = sorted(e.brand.items(), key=lambda kv: -kv[1].aum)
    top_b, rest = brands[:15], brands[15:]
    header(t_b, [("A", "운용사(브랜드)별 — 순자산 상위 15"), ("C", "규모"), ("F", "자금 유입(억원)")],
           ["", "브랜드", "종목 수", "점유율\n(%)", "순자산\n(억원)", "1M", "YTD"], ab=True)
    r = t_b + 2
    btot = sum(v.aum for _, v in brands)
    others = Brand(aum=sum(x.aum for _, x in rest), n=sum(x.n for _, x in rest),
                   f1m=sum(x.f1m for _, x in rest), fytd=sum(x.fytd for _, x in rest))
    for name, v in [*top_b, (f"그 밖 {len(rest)}곳", others)]:
        lab(r, name)
        put(r, 3, v.n, "#,##0")
        put(r, 4, v.aum / btot * 100 if btot else None, "0.0")
        put(r, 5, v.aum, F_AUM)
        put(r, 6, v.f1m, F_FL)
        put(r, 7, v.fytd, F_FL)
        r += 1
    b_end = r - 1
    close(b_end, 7)
    cs(f"F{t_b + 2}:F{b_end}")
    cs(f"G{t_b + 2}:G{b_end}")

    # ── 목록 ──
    heads = ["코드", "이름", "유형", "세부", "순자산\n(억원)", "거래대금\n20일(억원)", "1D", "1W", "1M",
             "3M", "YTD", "1W", "1M", "3M", "YTD", "괴리율\n(%)", "기초지수"]
    gl = [("C", "분류"), ("E", "규모"), ("G", "수익률(%)"), ("L", "자금 유입(억원)"), ("P", "괴리 · 기초지수")]

    def list_row(r: int, x: EtfRec) -> None:
        vals: list[tuple[object, str | None, bool]] = [
            (x.code, None, True), (x.name, None, True), (x.big, None, True),
            (x.sub or None, None, True), (x.aum, F_AUM, False), (x.tv20, F_AUM, False),
            *((x.ret[k], F_R, False) for k in ("1D", "1W", "1M", "3M", "YTD")),
            *((x.flow[k], F_FL, False) for k in ("1W", "1M", "3M", "YTD")),
            (x.prem, F_P, False), (x.idx, None, True)]
        for c, (v, fmt, text) in enumerate(vals, 1):
            put(r, c, v, fmt, text)

    def flow1m(x: EtfRec) -> float:
        return x.flow["1M"]
    t_t = b_end + 2
    header(t_t, [("A", "1M 자금 유입 상위 15 · 유출 상위 15"), *gl], heads)
    r = t_t + 2
    for x in sorted(recs, key=lambda x: -flow1m(x))[:15]:
        list_row(r, x)
        r += 1
    close(r - 1, 17)
    for x in sorted(recs, key=flow1m)[:15]:
        list_row(r, x)
        r += 1
    t_end = r - 1
    close(t_end, 17)
    cs(f"G{t_t + 2}:K{t_end}")
    cs(f"L{t_t + 2}:O{t_end}")
    t_c = t_end + 2
    header(t_c, [("A", "전체 ETF(순자산 순)"), *gl], heads)
    r = t_c + 2
    for x in sorted(recs, key=lambda x: -x.aum):
        list_row(r, x)
        r += 1
    c_end = r - 1
    close(c_end, 17)
    cs(f"G{t_c + 2}:K{c_end}")
    cs(f"L{t_c + 2}:O{c_end}")
    ws.auto_filter.ref = f"A{t_c + 1}:Q{c_end}"
    ws[f"P{t_t + 1}"].comment = x_note("괴리율(%) = 종가 ÷ NAV − 1. 양수 = NAV 보다 비싸게 거래")

    # ── 그래프 4개 — 원자료는 숨긴 AG~AN 열(날짜 · 값 사본; 표의 긴 이름 '2026-10 (8일까지)' 대신
    # 날짜). 주간 막대는 날짜 축에서 가늘어져 월간 막대(12개월)로 그린다
    xs.hide(G0, 8, 10)

    def gput(r: int, c: int, v: object, fmt: str | None = None) -> None:
        ws.cell(r, c).value = v
        if fmt:
            ws.cell(r, c).number_format = fmt
    gput(t_w + 1, G0, "주")
    gput(t_w + 1, G0 + 1, "국내주식")
    gput(t_w + 1, G0 + 2, "해외주식")
    for k, w in enumerate(w26):
        rr = t_w + 2 + k
        gput(rr, G0, e.weekly_end[w], "yyyy-mm-dd")
        gput(rr, G0 + 1, ws.cell(rr, 13).value)
        gput(rr, G0 + 2, ws.cell(rr, 14).value)
    gput(t_m + 1, G0 + 3, "월")
    gput(t_m + 1, G0 + 4, "전체 ETF")
    gput(t_m + 1, G0 + 5, "거래대금 비중")
    for k, m in enumerate(months):
        rr = t_m + 2 + k
        gput(rr, G0 + 3, date(int(m[:4]), int(m[4:]), 1), "yyyy-mm-dd")
        gput(rr, G0 + 4, ws.cell(rr, 12).value)
        gput(rr, G0 + 5, ws.cell(rr, 14).value)
    gput(t_y + 1, G0 + 6, "연도")
    gput(t_y + 1, G0 + 7, "순자산")
    for k, y in enumerate(years):
        rr = t_y + 2 + k
        gput(rr, G0 + 6, date(int(y), 12, 31), "yyyy-mm-dd")
        gput(rr, G0 + 7, ws.cell(rr, 13).value)

    def gref(c: int, r0: int, r1: int) -> Reference:
        return Reference(ws, min_col=c, min_row=r0, max_row=r1)
    t1, t2 = cs0, cs0 + CH + 1
    xs.title(f"A{t1}", f"월간 순유입(억원): 최근 {MONTHS_TABLE}개월, 전체 ETF, "
                       f"{day.month}월은 {day.day}일까지")
    c1 = BarChart()
    c1.type = "col"
    c1.gapWidth = 60
    c1.add_data(gref(G0 + 4, t_m + 1, m_end), titles_from_data=True)
    c1.set_categories(gref(G0 + 3, t_m + 2, m_end))
    c1.legend = None
    bar_series(c1.series[0], CHART_BLUE)
    date_frame(c1, "#,##0", "yy/mm", "months", 2, "months")
    ws.add_chart(c1, f"A{t1 + 1}")
    xs.title(f"I{t1}", f"{WEEKS_TABLE}주 누적 순유입(억원): 국내주식 · 해외주식")
    c2 = LineChart()
    c2.add_data(Reference(ws, min_col=G0 + 1, max_col=G0 + 2, min_row=t_w + 1, max_row=w_end),
                titles_from_data=True)
    c2.set_categories(gref(G0, t_w + 2, w_end))
    line_series(c2.series[0], CHART_BLUE)
    line_series(c2.series[1], CHART_BLACK)
    if c2.legend is not None:
        c2.legend.position = "t"
    date_frame(c2, "#,##0", "yy/mm", "days", 1, "months")
    ws.add_chart(c2, f"I{t1 + 1}")
    xs.title(f"A{t2}", f"ETF 순자산(조원): 연말, {day.year} 은 {md}")
    c3 = LineChart()
    c3.add_data(gref(G0 + 7, t_y + 1, y_end), titles_from_data=True)
    c3.set_categories(gref(G0 + 6, t_y + 2, y_end))
    c3.legend = None
    line_series(c3.series[0], CHART_BLUE)
    date_frame(c3, "#,##0", "yyyy", "years", 2, "years")
    ws.add_chart(c3, f"A{t2 + 1}")
    xs.title(f"I{t2}", f"ETF 거래대금 ÷ 주식 거래대금(%): 최근 {MONTHS_TABLE}개월")
    c4 = LineChart()
    c4.add_data(gref(G0 + 5, t_m + 1, m_end), titles_from_data=True)
    c4.set_categories(gref(G0 + 3, t_m + 2, m_end))
    c4.legend = None
    line_series(c4.series[0], CHART_BLUE)
    date_frame(c4, "#,##0", "yy/mm", "months", 2, "months")
    ws.add_chart(c4, f"I{t2 + 1}")
    xs.tight(cs0, t_w - 1)
    xs.page(f"A1:Q{c_end}")

    out.defs = [
        ("ETF", "유형별 요약 · 유형 · 종목 수 · 순자산 비중",
         "유형 = 이름 · 기초지수 자동 분류(시험), 레버리지(롱) · 인버스는 따로. 비중 = 유형 순자산 ÷ 전체"),
        ("ETF", f"유형별 요약 · {WEEKS_SPARK}주 유입(주간)",
         f"최근 {WEEKS_SPARK}주 주별 자금 유입 막대(셀 안 그래프 — 원자료는 숨긴 T~AE 열)"),
        ("ETF", "주간 · 월별 · 연도별 표",
         "유형별 자금 유입 합(억원). 월별 표에 월말 순자산(조원) · ETF 거래대금 ÷ 코스피 + 코스닥 주식 "
         "거래대금(%), 연도별 표에 연말 순자산(조원, 4묶음)"),
        ("ETF", "그래프 4개",
         f"월간 순유입(전체) · {WEEKS_TABLE}주 누적 순유입(국내주식 · 해외주식, 주간 표 "
         f"'{WEEKS_TABLE}주 누적' 열) · 연말 순자산 · 거래대금 비중. 원자료는 숨긴 AG~AN 열(날짜 · 값 사본)"),
        ("ETF", "운용사(브랜드)",
         "이름 첫 낱말(KODEX · TIGER · RISE …) 기준 순자산 · 점유율 · 1M · YTD 자금 유입"),
        ("ETF", "분류 · 코드 / 이름 / 유형 / 세부",
         "KRX ETF 일별(krx_etf_bydd_trd) 종목코드 · 종목명, 유형 · 세부는 자동 분류"),
        ("ETF", "규모 · 순자산 (억원)", "NAV × 상장좌수 ÷ 1억(KRX 순자산 칸과 같음)"),
        ("ETF", "규모 · 거래대금 20일 (억원)", "최근 20거래일 거래대금 평균"),
        ("ETF", "수익률 · 1D / 1W / 1M / 3M / YTD (%)",
         f"종가 기준, 1 · 5 · 20 · 60거래일 전 · {e.win['YTD'].isoformat()} 대비, 분배금 미반영. "
         "요약 표는 순자산가중"),
        ("ETF", "자금 유입 · 1W / 1M / 3M / YTD (억원)",
         "날마다 상장좌수 변화 × 그날 NAV 의 합(설정 − 환매 추정). 신규 상장 = 첫날 좌수 전부 유입, "
         "상장폐지 = 다음 거래일 유출(시계열 표만), 액면분할(배수 맞는 날) 보정 — "
         f"{e.first_day.year}년 이후 {e.splits}건"),
        ("ETF", "괴리 · 괴리율 (%) / 기초지수", "종가 ÷ NAV − 1 · KRX 기초지수 이름")]
    return out


__all__ = ["NAME", "Written", "week_span", "write"]
