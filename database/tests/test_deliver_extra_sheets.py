"""매일 엑셀 새 시트 여섯(Analyst · Events · Market · ETF · Flow · Short — DECISIONS §7 #10 · #11) —
합성 equity · stage 원천으로 계산 · 창 · 날짜 글 · '데이터 없음'을 검증한다.

합성 세계: 거래일 = 2025-01-02 ~ 2026-09-25 평일(2026-09-23 휴장). 모델 · fi 판은 매일 엑셀 테스트
(`test_deliver_excel`)의 합성 판을 그대로 쓴다(D = 2026-09-25, 45종목 — tick(0~14) IT · tick(15~24)
에너지 · tick(25~34) 산업재 …). 원천 종목: tick(0) 오름 · tick(1) 내림 · tick(2) 보합 · tick(15) 오름 ·
tick(25) 오름(2026-06-01 2:1 분할 — 수정종가로만 오름) · 900001(유니버스 밖 코스닥 보통주) 오름.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
import zipfile
from collections.abc import Iterator
from datetime import date, timedelta
from pathlib import Path

import duckdb
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from deliver import (
    etf_data,
    excel_daily,
    qpack,
    sheet_analyst,
    sheet_etf,
    sheet_events,
    sheet_flow,
    sheet_market,
    sheet_short,
    sheets_extra,
)
from deliver.excel_daily import build_daily
from deliver.extra_common import (
    Stock,
    base_text,
    common_markets,
    krx_calendar,
    latest_complete,
    span,
)
from deliver.qpack import XRefs, new_workbook
from deliver.sources import SourceMissing, Sources, connect
from openpyxl import Workbook, load_workbook
from stage import manifest
from test_deliver_excel import FRI, HOLIDAY, W38, W39, tick, write_fi_day, write_model_day

D = date.fromisoformat(FRI)
DAYS = [date(2025, 1, 2) + timedelta(days=k) for k in range((D - date(2025, 1, 2)).days + 1)]
DAYS = [x for x in DAYS if x.weekday() < 5 and x != date.fromisoformat(HOLIDAY)]
SPLIT_DAY = date(2026, 6, 1)
OTHER = "900001"
MARKET = {tick(0): "KOSPI", tick(1): "KOSPI", tick(2): "KOSDAQ", tick(15): "KOSDAQ",
          tick(25): "KOSPI", OTHER: "KOSDAQ"}
PRE_SPLIT = 1000.0
SATURDAY = date(2026, 9, 26)
MONDAY = date(2026, 9, 28)


def close_of(t: str, i: int) -> float:
    """원주가 — tick(25) 는 분할일부터 반값(수정종가로는 계속 오름)."""
    if t == tick(1):
        return 2000.0 - i
    if t == tick(2):
        return 500.0
    c = PRE_SPLIT + i
    return c / 2 if t == tick(25) and DAYS[i] >= SPLIT_DAY else c


# ── 원천 쓰기 ──────────────────────────────────────────────────────────────────
def write_table(root: Path, table: str, cols: dict[str, tuple[pa.DataType, list[object]]],
                build_id: str = "b_test_0001") -> None:
    """`<root>/<table>/v=<build>/part0.parquet` + MANIFEST(current_build) — equity · stage 규약."""
    vdir = root / table / f"v={build_id}"
    vdir.mkdir(parents=True, exist_ok=True)
    tbl = pa.table({k: pa.array(v, type=t) for k, (t, v) in cols.items()})
    pq.write_table(tbl, vdir / "part0.parquet")
    manifest.commit(root / table, manifest.BuildRecord(
        build_id=build_id, snapshot_id="", rules_version="v_test",
        built_at_utc="2026-09-26T00:00:00+00:00", n_rows=tbl.num_rows, content_hash="test",
        partitions=[{"path": f"v={build_id}", "n_rows": tbl.num_rows}]))


def cols_of(rows: list[dict[str, object]], types: dict[str, pa.DataType]
            ) -> dict[str, tuple[pa.DataType, list[object]]]:
    return {k: (t, [r.get(k) for r in rows]) for k, t in types.items()}


S, F, DT, B = pa.string(), pa.float64(), pa.date32(), pa.bool_()


def build_equity(root: Path, last: date = D, *, flow_last: date | None = None,
                 flow_partial: bool = False, holder_last: date | None = None,
                 saturday_fetch: bool = False, index_last: date | None = None) -> None:
    """합성 equity. flow_last = 키움 수급 마지막 날(지수보다 늦게 끝나게), flow_partial = 마지막 날
    수급이 한 종목뿐(일부 수집), holder_last = 지분 보고 마지막 날, saturday_fetch = 토요일 의견 수집,
    index_last = KRX 지수만 `last` 보다 늦게 끝나게(D+1 아침 체인 뒤 생성 — 지수는 current 판,
    나머지는 fi 판에 고정된 하루 늦은 판)."""
    days = [x for x in DAYS if x <= last]
    idx_rows, price, uni, flow, short, credit = [], [], [], [], [], []
    for i, dd in enumerate(x for x in DAYS if x <= max(last, index_last or last)):
        for cls, nm, c in (("KOSPI", "코스피", 2000.0 + i), ("KOSDAQ", "코스닥", 800.0 + i / 2),
                           ("KOSPI", "코스피 200", 300.0 + i), ("KOSPI", "전기전자", 100.0 + i),
                           ("KOSDAQ", "제약", 100.0 - i / 10), ("KOSPI", "제조", 50.0)):
            idx_rows.append({"index_class": cls, "index_name": nm, "date": dd, "close_idx": c,
                             "value_krw": 10e12 + i * 1e9, "mktcap_krw": 2000e12})
    for i, dd in enumerate(days):
        for t, mkt in MARKET.items():
            c = close_of(t, i)
            price.append({"ticker": t, "date": dd, "open": c, "high": c, "low": c, "close": c,
                          "volume_shr": 1000.0, "price_kind": "krx", "basis": "krx",
                          "corp_action_pending": False, "value_krw": 1e9, "mktcap_krw": 1e12})
            uni.append({"date": dd, "ticker": t, "sec_type": "common", "market": mkt})
            keep = ((flow_last is None or dd <= flow_last)
                    and not (flow_partial and dd == days[-1] and t != tick(0)))
            if keep:
                flow.append({"date": dd, "ticker": t, "src": "kiwoom", "frgnr_invsr_krw": 1e9,
                             "natfor_krw": 1e8, "orgn_krw": -5e8, "fnnc_invt_krw": -3e8,
                             "penfnd_etc_krw": -2e8, "ind_invsr_krw": -6e8, "etc_corp_krw": 0.0,
                             "insrnc_krw": 0.0, "invtrt_krw": 0.0, "etc_fnnc_krw": 0.0,
                             "bank_krw": 0.0, "samo_fund_krw": 0.0, "natn_krw": 0.0,
                             "foreign_wght_pct": 10.0 + i / 100})
            short.append({"date": dd, "ticker": t, "short_value_kiwoom_krw": 1e8,
                          "lending_balance_kiwoom_krw": 1e10 + i * 1e6})
            if i < len(days) - 2:            # 신용은 원천이 2거래일 늦다
                credit.append({"date": dd, "ticker": t, "whol_loan_rmnd_amt": 1e5,
                               "whol_loan_rmnd_rate_pct": 1.0 + i / 1000})
    write_table(root, "index_daily", cols_of(idx_rows, {
        "index_class": S, "index_name": S, "date": DT, "close_idx": F, "value_krw": F,
        "mktcap_krw": F}))
    write_table(root, "price_daily", cols_of(price, {
        "ticker": S, "date": DT, "open": F, "high": F, "low": F, "close": F, "volume_shr": F,
        "price_kind": S, "basis": S, "corp_action_pending": B, "value_krw": F,
        "mktcap_krw": F}))
    write_table(root, "universe_daily", cols_of(uni, {"date": DT, "ticker": S, "sec_type": S,
                                                      "market": S}))
    write_table(root, "flow_daily", cols_of(flow, {
        "date": DT, "ticker": S, "src": S, **dict.fromkeys(
            ("frgnr_invsr_krw", "natfor_krw", "orgn_krw", "fnnc_invt_krw", "penfnd_etc_krw",
             "ind_invsr_krw", "etc_corp_krw", "insrnc_krw", "invtrt_krw", "etc_fnnc_krw",
             "bank_krw", "samo_fund_krw", "natn_krw", "foreign_wght_pct"), F)}))
    write_table(root, "short_daily", cols_of(short, {
        "date": DT, "ticker": S, "short_value_kiwoom_krw": F, "lending_balance_kiwoom_krw": F}))
    write_table(root, "credit_daily", cols_of(credit, {
        "date": DT, "ticker": S, "whol_loan_rmnd_amt": F, "whol_loan_rmnd_rate_pct": F}))
    write_table(root, "adj_factor", cols_of([{
        "ticker": tick(25), "apply_date": SPLIT_DAY, "available_date": SPLIT_DAY,
        "price_factor": 0.5, "share_factor": 2.0, "factor_ok": True, "price_only_factor": None,
        "price_resolution": "split"}], {
        "ticker": S, "apply_date": DT, "available_date": DT, "price_factor": F,
        "share_factor": F, "factor_ok": B, "price_only_factor": F, "price_resolution": S}))
    write_table(root, "trading_calendar", cols_of([
        {"date": x, "prev_td": days[k - 1] if k else None,
         "next_td": days[k + 1] if k + 1 < len(days) else None} for k, x in enumerate(days)],
        {"date": DT, "prev_td": DT, "next_td": DT}))
    d1, d2, d3 = days[-2], days[-3], days[-4]
    ops = [  # (종목, 수집일, 증권사, 보고서 날짜, 목표가, 직전 목표가, 변화, 의견, 등급, 직전 등급)
        (tick(0), d2, "A증권", d2, 1100.0, 1000.0, 10.0, "BUY", "buy", "buy"),
        (tick(0), d1, "A증권", d2, 1200.0, 1000.0, 20.0, "BUY", "buy", "buy"),   # 나중 수집이 이긴다
        (tick(1), d1, "B증권", d3, 900.0, 1000.0, -10.0, "HOLD", "hold", "buy"),
        (tick(2), d1, "C증권", d1, 500.0, None, None, "BUY", "buy", None),       # 새 커버
        (tick(15), last, "D증권", d2, 1000.0, 1000.0, 0.0, "BUY", "buy", "buy"),  # 유지(세지 않음)
        # 직전 목표가 없이 등급만 오름 — 신규 커버가 아니라 의견 상향(10-08 006800 한국투자 같은 경우)
        (tick(25), d1, "E증권", d1, 800.0, None, None, "BUY", "buy", "hold")]
    if saturday_fetch:
        ops.append((tick(0), SATURDAY, "F증권", last, 1300.0, 1200.0, 8.33, "BUY", "buy", "buy"))
    write_table(root, "opinion_broker_daily", {
        "ticker": (S, [o[0] for o in ops]), "fetched_date": (DT, [o[1] for o in ops]),
        "broker": (S, [o[2] for o in ops]), "opinion_date": (DT, [o[3] for o in ops]),
        "target_price_krw": (F, [o[4] for o in ops]),
        "prev_target_price_krw": (F, [o[5] for o in ops]), "change_pct": (F, [o[6] for o in ops]),
        "opinion": (S, [o[7] for o in ops]), "opinion_class": (S, [o[8] for o in ops]),
        "prev_opinion": (S, [None if o[9] is None else ("BUY" if o[9] == "buy" else "HOLD")
                             for o in ops]),
        "prev_opinion_class": (S, [o[9] for o in ops]),
        "prev_opinion_date": (DT, [None for _ in ops])})
    cons = [(tick(0), 2 * close_of(tick(0), len(days) - 1), 5, "wise"), (tick(1), 800.0, 2, "wise"),
            (tick(2), 600.0, 3, "wise"), (tick(2), 9999.0, 9, "v3")]   # v3 판 겹침 — 쓰지 않는다
    write_table(root, "opinion_daily", {
        "ticker": (S, [c[0] for c in cons]), "obs_date": (DT, [last] * len(cons)),
        "src": (S, [c[3] for c in cons]), "target_price_krw": (F, [c[1] for c in cons]),
        "analyst_count": (F, [c[2] for c in cons])})
    hl = holder_last or d1
    hold = [  # (보고번호, src, 보고자, corp, 접수일, 직위, 등기, 사유, 주식 수 변동, 지분, 지분 변화)
        ("H1", "elestock", "홍길동", "C2", hl, "이사", "등기임원", None, 1000.0, 0.1, 0.0),
        ("H2", "elestock", "큰손", "C2", hl, "-", "-", None, 99999.0, 12.0, 1.0),       # 임원 아님
        ("H5", "elestock", "최대주주씨", "C2", hl, "최대주주", "-", None, 5000.0, 30.0, 0.0),  # 임원 아님
        ("H6", "elestock", "의장님", "C0", hl, "이사회 의장", "등기임원", None, -2000.0, 15.0, 0.0),
        ("H3", "majorstock", "가나자산", "C1", hl, None, None, "단순  투자", None, 6.0, 1.5),
        ("H4", "majorstock", "다라자산", "C1", hl, None, None, "담보 변경", None, 5.5, 0.0)]  # 변화 0
    write_table(root, "holder_daily", {
        "rcept_no": (S, [h[0] for h in hold]), "src": (S, [h[1] for h in hold]),
        "repror": (S, [h[2] for h in hold]), "corp_code": (S, [h[3] for h in hold]),
        "rcept_dt": (DT, [h[4] for h in hold]), "ofcps": (S, [h[5] for h in hold]),
        "exec_rgist": (S, [h[6] for h in hold]), "report_resn": (S, [h[7] for h in hold]),
        "qty_change_shr": (F, [h[8] for h in hold]), "rate_pct": (F, [h[9] for h in hold]),
        "rate_change_pct": (F, [h[10] for h in hold])})


ETFS = {"A00001": ("KODEX 200", "코스피 200"), "A00002": ("TIGER 미국S&P500", "S&P 500"),
        "A00003": ("KODEX 레버리지", "코스피 200 레버리지"), "A00004": ("ARIRANG 테스트", "코스피 200")}
DELIST = date(2026, 3, 31)


def build_stage(root: Path, last: date = D) -> None:
    days = [x for x in DAYS if x <= last]
    etf = []
    for i, dd in enumerate(days):
        for t, (nm, ix) in ETFS.items():
            if t == "A00004" and dd > DELIST:
                continue
            nav, sh = 10000.0 + i, 1_000_000.0
            if t == "A00001" and dd == days[-1]:
                sh += 100_000.0                     # 마지막 날 10만 좌 설정
            if t == "A00002" and dd >= date(2026, 7, 1):
                nav, sh = nav / 2, sh * 2            # 액면분할 2:1 — 유입 아님
            etf.append({"ticker": t, "date": dd, "name": nm, "obj_index_name": ix,
                        "close_krw": nav + 5, "nav_krw": nav, "list_shrs": sh, "value_krw": 1e9,
                        "netasset_krw": 0.0})
    write_table(root, "stg_etf_price_daily", cols_of(etf, {
        "ticker": S, "date": DT, "name": S, "obj_index_name": S, "close_krw": F, "nav_krw": F,
        "list_shrs": F, "value_krw": F, "netasset_krw": F}))
    d1 = days[-2]
    no = f"{d1:%Y%m%d}"
    disc = [  # (접수번호, 접수일, 종목, 제목, 정정본?, 뒤에 정정됨?)
        (f"{no}000001", d1, tick(0), "단일판매ᆞ공급계약체결", False, False),         # stage NFKC 'ᆞ'
        (f"{no}000002", d1, tick(0), "[기재정정]단일판매ᆞ공급계약체결", True, False),  # 같은 날 — 한 건
        (f"{no}000003", d1, tick(1), "주식등의대량보유상황보고서", False, False),
        (f"{no}000004", d1, tick(2), "기업설명회(IR)개최", False, False),
        (f"{no}000005", d1, tick(15), "[기재정정]주요사항보고서(자기주식취득결정)", True, False),
        (f"{no}000010", d1, tick(0), "주요사항보고서(자기주식취득결정)", False, True),   # 정정됨 — 뺀다
        (f"{no}000013", d1, tick(0), "[기재정정]주요사항보고서(자기주식취득결정)", True, False)]
    write_table(root, "stg_disclosure", {
        "rcept_no": (S, [x[0] for x in disc]), "rcept_dt": (DT, [x[1] for x in disc]),
        "ticker": (S, [x[2] for x in disc]), "report_nm": (S, [x[3] for x in disc]),
        "is_correction": (B, [x[4] for x in disc]),
        "rm_corrected_later": (B, [x[5] for x in disc])})
    write_table(root, "stg_corp_map", {"corp_code": (S, ["C0", "C1", "C2"]),
                                       "ticker": (S, [tick(0), tick(1), tick(2)])})
    write_table(root, "stg_event_tsstk_aq", {
        "rcept_no": (S, [f"{no}000010", f"{no}000013"]), "corp_code": (S, ["C0", "C0"]),
        "aq_mth": (S, ["장내매수", "장내매수"]), "aq_pp": (S, ["주주가치\n제고", "주주가치\n제고"]),
        "aqpln_prc_ostk": (F, [5e9, 1e10]), "aqexpd_bgd": (DT, [D, D]),
        "aqexpd_edd": (DT, [None, None])})
    write_table(root, "stg_event_tsstk_dp", {
        "rcept_no": (S, []), "corp_code": (S, []), "dp_pp": (S, []), "dppln_prc_ostk": (F, []),
        "dpprpd_bgd": (DT, []), "dpprpd_edd": (DT, [])})
    # 같은 날 같은 금액의 유상증자 결정 두 건(접수번호가 다르면 다른 결정 — 10-06 알테오젠 같은 경우)
    write_table(root, "stg_event_piic", {
        "rcept_no": (S, [f"{no}000011", f"{no}000014"]), "corp_code": (S, ["C1", "C1"]),
        "ic_mthn": (S, ["제3자배정", "제3자배정"]), "nstk_ostk_cnt": (F, [1234567.0, None]),
        **{k: (F, [5e9 if k == "fdpp_op" else None] * 2) for k in sheet_events.FUNDS}})
    write_table(root, "stg_event_cvbd_is", {
        "rcept_no": (S, [f"{no}000012"]), "corp_code": (S, ["C2"]), "bd_knd": (S, ["전환사채"]),
        "bd_fta": (F, [3e8]), "cv_prc": (F, [None])})


UNI = {tick(0): Stock(tick(0), "종목0", "IT", 1, 100.0, 5.0),
       tick(1): Stock(tick(1), "종목1", "IT", 2, 200.0, -5.0),
       tick(2): Stock(tick(2), "종목2", "IT", 3, 300.0, None),
       tick(15): Stock(tick(15), "종목15", "에너지", 4, 400.0, 1.0),
       tick(25): Stock(tick(25), "종목25", "산업재", None, 500.0, 2.0)}


@pytest.fixture(scope="module")
def roots(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path]:
    eq, st = tmp_path_factory.mktemp("equity"), tmp_path_factory.mktemp("stage")
    build_equity(eq)
    build_stage(st)
    return eq, st


@pytest.fixture(scope="module")
def src(roots: tuple[Path, Path]) -> Sources:
    return Sources(roots[0], roots[1], {})


@pytest.fixture
def con() -> Iterator[duckdb.DuckDBPyConnection]:
    c = connect()
    yield c
    c.close()


def book_with_refs() -> tuple[Workbook, XRefs]:
    wb = new_workbook()
    return wb, XRefs.of(wb.worksheets[0])


def values(ws, col: int = 1) -> list[object]:
    return [ws.cell(r, col).value for r in range(1, ws.max_row + 1)]


def row_of(ws, text: str, col: int = 1) -> int:
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value == text)


def col_values(ws, r0: int, col: int, n: int) -> list[object]:
    return [ws.cell(r0 + k, col).value for k in range(n)]


# ── 공용 ──────────────────────────────────────────────────────────────────────
def test_base_text_and_window_text() -> None:
    assert base_text(D, D) == "기준일: 2026-09-25"
    assert base_text(D - timedelta(days=1), D) == "기준일: 2026-09-24(원천 최신 · 엑셀 09-25)"
    assert span(date(2026, 9, 18), D) == "09-18 다음 날~09-25"
    assert sheet_etf.week_span([date(2026, 9, 21), D]) == "09-21~25"
    assert sheet_etf.week_span([date(2026, 9, 29), date(2026, 10, 2)]) == "09-29~10-02"


def test_non_trading_source_day_maps_to_last_trading_day(src: Sources, con) -> None:
    """원천 날짜가 거래일이 아니면(토요일 수집 · 12-31 공시) 창은 그 앞 마지막 거래일에서 센다 —
    엑셀 D(월요일)는 지수에 아직 없어도 거래일이다."""
    cal = krx_calendar(con, src, SATURDAY, MONDAY)
    assert (cal.anchor, cal.day, cal.win["1W"]) == (SATURDAY, D, DAYS[-6])
    assert krx_calendar(con, src, MONDAY, MONDAY).day == MONDAY


def test_latest_complete_drops_a_partial_collection_day(tmp_path: Path, con) -> None:
    eq = tmp_path / "equity"
    build_equity(eq, flow_partial=True)
    fl = Sources(eq, None, {}).equity("flow_daily")
    lt = latest_complete(con, fl, "date", D, "src = 'kiwoom'", "수급")
    assert (lt.day, lt.partial) == (DAYS[-2], D)
    assert lt.note() == " · 09-25 은 일부만 수집돼 뺌"


# ── ETF ──────────────────────────────────────────────────────────────────────
def test_etf_flows_split_delist_and_windows(src: Sources, con) -> None:
    e = etf_data.load(con, src, D)
    assert e.day == D and e.first_day == DAYS[0]
    assert e.win["1D"] == DAYS[-2] and e.win["1W"] == DAYS[-6] and e.win["YTD"] == date(2025, 12, 31)
    rec = {r.code: r for r in e.recs}
    assert set(rec) == {"A00001", "A00002", "A00003"}           # 폐지 ETF 는 지금 목록에 없다
    nav_d = 10000.0 + len(DAYS) - 1
    assert rec["A00001"].flow["1D"] == pytest.approx(100_000 * nav_d / 1e8)
    assert rec["A00001"].flow["1M"] == rec["A00001"].flow["1D"]
    assert rec["A00002"].flow["3M"] == 0.0 and e.splits == 1     # 분할은 유입이 아니다
    assert (rec["A00003"].big, rec["A00002"].big, rec["A00002"].sub) == ("레버리지(롱)", "해외주식",
                                                                          "미국")
    assert rec["A00001"].ret["1D"] == pytest.approx(((nav_d + 5) / (nav_d + 4) - 1) * 100)
    assert rec["A00001"].aum == pytest.approx(nav_d * 1_100_000 / 1e8)
    # 시계열: 첫날 상장은 유입 아님, 폐지 다음 거래일에 마지막 순자산만큼 유출
    nxt = DAYS[DAYS.index(DELIST) + 1]
    w = f"{nxt.isocalendar()[0]}-W{nxt.isocalendar()[1]:02d}"
    i_last = DAYS.index(DELIST)
    assert e.weekly[w]["국내주식"] == pytest.approx(-(10000.0 + i_last) * 1_000_000 / 1e8)
    assert e.brand["KODEX"].n == 2 and e.brand["ARIRANG"].n == 0
    assert e.netasset_empty
    # 공용 연결은 행 순서 보존을 되돌려 놓고 임시 표를 남기지 않는다
    assert con.execute("SELECT current_setting('preserve_insertion_order')").fetchone()[0] is True
    assert con.execute("SELECT count(*) FROM duckdb_tables() WHERE temporary").fetchone()[0] == 0


def test_etf_sheet_texts_charts_and_sparklines(src: Sources, con) -> None:
    wb, refs = book_with_refs()
    w = sheet_etf.write(wb, etf_data.load(con, src, D), refs, D)
    ws = wb["ETF"]
    assert ws["A3"].value == ("자료: KRX ETF 일별 2025-01~2026-09-25(종가 · NAV · 상장좌수 · 거래대금) · "
                              "지금 3종목")
    assert ws["A4"].value == "기준일: 2026-09-25"
    assert "이번 주(09-21~25, 4거래일)" in ws["P7"].comment.text
    assert "2025-12-31 대비" in ws["G7"].comment.text.replace("\n", " ")
    assert "2025년 이후 1건" in ws["L7"].comment.text.replace("\n", " ")
    col_a = values(ws)
    assert "09-25(4일)" in col_a and "2026-09 (25일까지)" in col_a and "2026 (09-25까지)" in col_a
    assert ws.freeze_panes == "C6" and len(ws._charts) == 4
    # 요약 24줄(전체 · 국내 · 국내 세부 11 · 해외 · 해외 세부 3 · 그 밖 유형 7)마다 12주 막대 하나
    assert len(w.sparks) == 24 and w.sparks[0] == ("ETF!T8:AE8", "P8")
    assert [d[1] for d in w.defs][:2] == ["유형별 요약 · 유형 · 종목 수 · 순자산 비중",
                                          "유형별 요약 · 12주 유입(주간)"]


def test_etf_uses_latest_source_day_before_d(tmp_path: Path, roots: tuple[Path, Path],
                                              con) -> None:
    """장 마감 엑셀(사용자 결정 Q-E2): 원천이 D 전날까지면 그날로 만들고 4행에 적는다."""
    st = tmp_path / "stage"
    build_stage(st, last=DAYS[-2])
    wb, refs = book_with_refs()
    e = etf_data.load(con, Sources(roots[0], st, {}), D)
    sheet_etf.write(wb, e, refs, D)
    assert e.day == DAYS[-2]
    assert wb["ETF"]["A4"].value == "기준일: 2026-09-24(원천 최신 · 엑셀 09-25)"


# ── Market ───────────────────────────────────────────────────────────────────
def test_market_index_breadth_and_investors(src: Sources, con) -> None:
    m = sheet_market.load(con, src, D)
    ix = {(i.cls, i.name): i for i in m.index}
    n = len(DAYS) - 1
    assert ix[("KOSPI", "코스피")].ret["1D"] == pytest.approx(((2000 + n) / (1999 + n) - 1) * 100)
    b = m.breadth[-1]
    assert b.day == D and len(m.breadth) == 20
    # 코스피: tick0 · tick25(수정종가) 오름, tick1 내림 / 코스닥: tick15 · 900001 오름, tick2 보합.
    # 52주 신고가 = 오른 넷 + 보합(최고 = 최저), 신저가 = 내림 + 보합
    assert (b.kospi, b.kosdaq, b.hi, b.lo) == ((2, 1, 0), (2, 0, 1), 5, 2)
    f = m.mkt_flow[("KOSPI", "1W")]
    assert (f["frgnr_invsr_krw"] or 0) + (f["natfor_krw"] or 0) == pytest.approx(3 * 5 * 11e8)


def test_market_sheet_layout(src: Sources, con) -> None:
    wb, refs = book_with_refs()
    defs = sheet_market.write(wb, sheet_market.load(con, src, D), refs, D)
    ws = wb["Market"]
    assert ws["A3"].value.endswith("(보통주) 2025~")              # 지수 자료 첫해
    assert ws["A4"].value == "기준일: 2026-09-25"
    assert ws["A5"].value.endswith("YTD = 2025-12-31 대비")
    r = row_of(ws, "외국인")
    assert ws.cell(r, 4).value == pytest.approx(3 * 5 * 11)        # 코스피 1W(억원)
    r = row_of(ws, "코스피 업종 지수: 1M 높은 순")
    assert ws.cell(r + 2, 1).value == "전기전자"                   # 제조 · 규모 지수는 뺀다
    assert "제조" not in values(ws)
    assert len(ws._charts) == 4 and [d[1] for d in defs][0] == "대표 지수 · 업종 지수"


def test_market_flow_lagging_index_gets_its_own_date_and_blanks(tmp_path: Path, con) -> None:
    """키움 수급이 지수보다 하루 늦으면 수급 표는 수급 날짜로 세고(1D = 그 전날 하루) 4행 · 표 제목에
    날짜를 적으며, 수급이 없는 지수 날의 누적 점은 0 이 아니라 빈칸이다."""
    eq = tmp_path / "equity"
    build_equity(eq, flow_last=DAYS[-2])
    wb, refs = book_with_refs()
    m = sheet_market.load(con, Sources(eq, None, {}), D)
    sheet_market.write(wb, m, refs, D)
    ws = wb["Market"]
    assert (m.cal.day, m.flow.day) == (D, DAYS[-2])
    assert ws["A4"].value == "기준일: 2026-09-25 · 수급 09-24"
    r = row_of(ws, "시장 투자자별 순매수(억원): 보통주, 09-24까지") + 2
    assert ws.cell(r, 1).value == "개인"
    r = row_of(ws, "외국인")
    assert ws.cell(r, 3).value == pytest.approx(3 * 11)            # 1D 코스피 = 09-24 하루(0 아님)
    last = max(rr for rr in range(8, 400) if ws.cell(rr, 30).value is not None)
    assert ws.cell(last, 30).value == D                             # 그래프 원자료 끝 = 지수 날
    assert ws.cell(last, 33).value is None and ws.cell(last - 1, 33).value is not None


def sheet_digest(ws) -> tuple[int, str]:
    """시트 전 칸(좌표 · 값 · 숫자 서식 · 메모)의 지문 — 실수는 유효숫자 12자리로 적는다(합계 순서에
    따른 끝자리 흔들림을 지문에 넣지 않는다)."""
    cells = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is None and c.comment is None:
                continue
            v = f"{c.value:.12g}" if isinstance(c.value, float) else repr(c.value)
            memo = "" if c.comment is None else c.comment.text
            cells.append(f"{c.coordinate}|{v}|{c.number_format}|{memo}")
    return len(cells), hashlib.sha256("\n".join(cells).encode("utf-8")).hexdigest()


# (칸 수, sha256) — 78c557d6(결함 1 수정 전) 코드에서 아래 테스트와 같은 합성 원천(build_equity(last=
# DAYS[-2]))으로 Market · Short 시트를 만들어 `sheet_digest` 로 잰 값. 정상 운영 날짜 정렬에서는 칸이 하나도
# 바뀌지 않아야 한다. 의도한 표시 변경(문구 · 서식 · 열)을 넣으면 그 커밋에서 다시 잰다 — 이 테스트의 본문
# 그대로 시트를 만들고 `print(sheet_digest(wb["Market"]), sheet_digest(wb["Short"]))` 값을 옮기고, 바뀐
# 칸이 의도한 것뿐인지 옛 · 새 시트 칸 비교로 먼저 확인한다
ALIGNED_MARKET = (1772, "9e0213b75f218f675233351c6e58455d3e2410fba82de61abb9664c95cbcf039")
ALIGNED_SHORT = (1841, "281a885c9dcb2a92fb00ebbd7ea5d99461a79db0b801e9ac4847621a7b041a7f")


def test_market_and_short_cells_unchanged_when_all_sources_end_on_the_same_day(
        tmp_path: Path, con) -> None:
    """정상 운영(D 15:41 생성 — 지수 · 유니버스 · 가격 · 수급 원천이 모두 D−1 에서 끝난다)의 Market ·
    Short 시트는 결함 1 수정(시장 구분 as-of · 시장 폭 블록 기준일) 전과 칸 단위로 같다."""
    eq = tmp_path / "equity"
    build_equity(eq, last=DAYS[-2])
    src = Sources(eq, None, {})
    wb, refs = book_with_refs()
    sheet_market.write(wb, sheet_market.load(con, src, D), refs, D)
    sheet_short.write(wb, sheet_short.load(con, src, D, UNI), UNI, refs, D)
    assert sheet_digest(wb["Market"]) == ALIGNED_MARKET
    assert sheet_digest(wb["Short"]) == ALIGNED_SHORT


def test_market_built_after_the_next_morning_chain_uses_the_pinned_day(tmp_path: Path,
                                                                      con) -> None:
    """결함 1 — D+1 아침 체인 뒤에 만들면(늦은 재발송 · 수동 재생성) 지수(current 판)는 D 까지,
    fi 판에 고정된 유니버스 · 가격 · 수급은 D−1 까지다. 시장 구분은 D 이하 최신 유니버스 날짜로 만들어
    시장 폭 · 시장 투자자별 순매수 · 외국인 누적이 D−1 값으로 차고, 시장 폭 · 거래대금 블록은 가격
    최신일(D−1)까지만 행을 내며 제목에 그 날짜를 적는다(데이터 없는 D 행을 빈칸으로 내지 않는다)."""
    eq = tmp_path / "equity"
    build_equity(eq, last=DAYS[-2], index_last=D)
    m = sheet_market.load(con, Sources(eq, None, {}), D)
    assert (m.cal.day, m.flow.day) == (D, DAYS[-2])
    b = m.breadth[-1]
    assert b.day == DAYS[-2] and len(m.breadth) == 20
    assert (b.kospi, b.kosdaq, b.hi, b.lo) == ((2, 1, 0), (2, 0, 1), 5, 2)
    f = m.mkt_flow[("KOSPI", "1W")]
    assert (f["frgnr_invsr_krw"] or 0) + (f["natfor_krw"] or 0) == pytest.approx(3 * 5 * 11e8)
    wb, refs = book_with_refs()
    sheet_market.write(wb, m, refs, D)
    ws = wb["Market"]
    assert ws["A4"].value == "기준일: 2026-09-25 · 수급 09-24"
    row_of(ws, "시장 투자자별 순매수(억원): 보통주, 09-24까지")
    r = row_of(ws, "시장 폭: 최근 20거래일, 보통주, 09-24까지") + 2
    assert [ws.cell(r, c).value for c in (1, 3, 4, 5)] == ["09-24", 2, 1, 0]
    assert ws.cell(r, 11).value == pytest.approx(3e9 / 1e12)       # 코스피 거래대금(조원)
    assert "09-25" not in values(ws)                               # 데이터 없는 날 행은 없다
    r = row_of(ws, "외국인")
    assert ws.cell(r, 3).value == pytest.approx(3 * 11)            # 1D 코스피(억원)
    last = max(rr for rr in range(8, 400) if ws.cell(rr, 30).value is not None)
    assert ws.cell(last, 30).value == D                            # 그래프 원자료 끝 = 지수 날
    assert ws.cell(last, 33).value is None and ws.cell(last - 1, 33).value is not None


def test_common_markets_is_as_of_and_an_empty_universe_is_no_data(src: Sources, con) -> None:
    """시장 구분(mk)은 기준일 이하 최신 유니버스 날짜의 보통주 코스피 · 코스닥이고 그 날짜를
    돌려준다. 0행이면 조용한 빈칸 대신 '데이터 없음'(SourceMissing)."""
    assert common_markets(con, src, SATURDAY) == D
    assert con.execute("SELECT count(*) FROM mk").fetchone()[0] == len(MARKET)
    with pytest.raises(SourceMissing, match="유니버스"):
        common_markets(con, src, DAYS[0] - timedelta(days=1))


# ── Flow ─────────────────────────────────────────────────────────────────────
def test_flow_sector_sums_and_lists(src: Sources, con) -> None:
    wb, refs = book_with_refs()
    f = sheet_flow.load(con, src, D, UNI)
    sheet_flow.write(wb, f, UNI, refs, D)
    ws = wb["Flow"]
    r = row_of(ws, "IT")
    assert ws.cell(r, 3).value == 3 and ws.cell(r, 4).value == pytest.approx(3 * 5 * 11)
    assert f.fw[tick(0)]["now"] == pytest.approx(10 + (len(DAYS) - 1) / 100)
    r = row_of(ws, "외국인 지분율 변화: 1M 상승 상위 15 · 하락 상위 15")
    assert ws.cell(r + 2, 11).value == pytest.approx(0.20)          # 20거래일 × 0.01%p


def test_flow_skips_a_partial_latest_day(tmp_path: Path, con) -> None:
    eq = tmp_path / "equity"
    build_equity(eq, flow_partial=True)
    wb, refs = book_with_refs()
    f = sheet_flow.load(con, Sources(eq, None, {}), D, UNI)
    sheet_flow.write(wb, f, UNI, refs, D)
    assert f.cal.day == DAYS[-2]
    assert wb["Flow"]["A4"].value == ("기준일: 2026-09-24(원천 최신 · 엑셀 09-25) · 09-25 은 일부만 "
                                      "수집돼 뺌")


# ── Short ────────────────────────────────────────────────────────────────────
def test_short_credit_lag_and_weeks(src: Sources, con) -> None:
    wb, refs = book_with_refs()
    s = sheet_short.load(con, src, D, UNI)
    sheet_short.write(wb, s, UNI, refs, D)
    ws = wb["Short"]
    assert s.credit_dates[0] == DAYS[-3]
    assert ws["A5"].value == f"신용잔고는 원천이 2거래일 늦어 최신 09-22 · 1M 전 {DAYS[-23]:%m-%d}"
    assert ws["A8"].value == "09-25 (신용 09-22)"                   # 이번 주 끝 · 신용 날짜
    assert ws["E8"].value == pytest.approx(10.0)                   # 공매도 ÷ 거래대금 = 1e8 ÷ 1e9
    assert len(ws._charts) == 4


# ── Analyst ──────────────────────────────────────────────────────────────────
def test_analyst_counts_dedupe_and_upside(src: Sources, con) -> None:
    wb, refs = book_with_refs()
    a = sheet_analyst.load(con, src, D, UNI)
    acts = {(x.code, x.broker): x for x in a.acts}
    assert acts[(tick(0), "A증권")].chg == 20.0                    # 같은 날 같은 증권사 — 나중 수집
    assert a.cal.day == D and a.cons_day == D
    assert a.cons[tick(2)].tp == 600.0                             # src 'wise' 만(v3 판 9999 아님)
    sheet_analyst.write(wb, a, UNI, refs, D)
    ws = wb["Analyst"]
    r = row_of(ws, "IT")
    # 1W 상향 1 · 하향 1 · 신규 1 · 상향 비율 50%
    assert [ws.cell(r, c).value for c in range(3, 8)] == [3, 1, 1, 1, 50.0]
    r = row_of(ws, "산업재")                                      # 등급만 오른 보고서는 신규가 아니다
    assert [ws.cell(r, c).value for c in range(4, 7)] == [0, 0, 0]
    r = row_of(ws, "목표가 여력 상위 20: 증권사 3곳 이상")
    assert col_values(ws, r + 2, 1, 2) == [tick(0), tick(2)]
    assert ws.cell(r + 2, 10).value == pytest.approx(100.0)          # 목표가 = 종가 × 2
    # 등급 변경(1M)이 먼저, 신규 커버(1W)는 직전 의견 · 목표가가 모두 없는 보고서만 — 한 번씩
    r = row_of(ws, "투자의견 변경 · 신규 커버: 1M")
    rows = [(ws.cell(rr, 1).value, ws.cell(rr, 12).value) for rr in range(r + 2, r + 6)]
    assert rows == [(tick(25), "의견 상향"), (tick(1), "의견 하향"), (tick(2), "신규 커버"),
                    (None, None)]
    # 일별 그래프 원자료(AD 열)는 기준일 전날까지 — WISE 가 그날 보고서를 다음 날 올린다
    days = [ws.cell(rr, 30).value for rr in range(8, 40) if ws.cell(rr, 30).value]
    assert max(days) == DAYS[-2] and len(days) == 19


def test_analyst_saturday_fetch_counts_windows_from_friday(tmp_path: Path, con) -> None:
    eq = tmp_path / "equity"
    build_equity(eq, saturday_fetch=True)
    wb, refs = book_with_refs()
    a = sheet_analyst.load(con, Sources(eq, None, {}), MONDAY, UNI)
    sheet_analyst.write(wb, a, UNI, refs, MONDAY)
    assert (a.cal.anchor, a.cal.day, a.cal.win["1W"]) == (SATURDAY, D, DAYS[-6])
    assert wb["Analyst"]["A4"].value == "기준일: 2026-09-26(원천 최신 · 엑셀 09-28)"
    assert (tick(0), "F증권") in {(x.code, x.broker) for x in a.acts}


def test_analyst_bar_chart_without_values_is_skipped(src: Sources, con) -> None:
    """그릴 막대가 없으면 빈 범위 그래프(가짜 막대) 대신 제목 칸에 '데이터 없음'."""
    wb, refs = book_with_refs()
    a = sheet_analyst.load(con, src, D, UNI)
    a.cons, a.acts = {}, []
    sheet_analyst.write(wb, a, UNI, refs, D)
    ws = wb["Analyst"]
    titles = [c.value for row in ws.iter_rows(min_col=1, max_col=8) for c in row
              if isinstance(c.value, str) and c.value.endswith("데이터 없음")]
    assert titles == ["대분류별 1M 목표가 상향 비율(%): 데이터 없음",
                      "대분류별 목표가 여력 중앙값(%): 09-25 컨센서스 목표가 ÷ 종가 − 1: 데이터 없음"]
    assert len(ws._charts) == 2


# ── Events ───────────────────────────────────────────────────────────────────
def test_events_types_titles_holders_and_decisions(src: Sources, con) -> None:
    wb, refs = book_with_refs()
    sheet_events.write(wb, sheet_events.load(con, src, D, UNI), UNI, refs, D)
    ws = wb["Events"]
    r = row_of(ws, "공급계약")
    assert [ws.cell(r, c).value for c in range(3, 7)] == [1, 1, 1, 1]   # 정정본은 한 건
    r = row_of(ws, "주요 공시: 1W, 5% · 임원 소유 보고와 그 밖 유형 뺌")
    assert col_values(ws, r + 2, 7, 4) == [   # 유형 순 · 같은 유형은 종합 순위 순 · 원문 'ㆍ'
        "주요사항보고서(자기주식취득결정)", "[정정] 주요사항보고서(자기주식취득결정)",
        "단일판매ㆍ공급계약체결", None]
    r = row_of(ws, "임원 보유 순변동: 1M 순증 상위 15 · 순감 상위 15")
    # 임원(등기 · 비등기)만 — 최대주주 · 10% 주주 보고는 빼고, 임원을 겸한 주요주주(의장)는 남는다
    assert col_values(ws, r + 2, 1, 3) == [tick(2), tick(0), None]
    assert ws.cell(r + 2, 7).value == "홍길동(이사) +1,000주 · 1건"
    assert ws.cell(r + 3, 7).value == "의장님(이사회 의장) -2,000주 · 1건"
    r = row_of(ws, "5% 대량보유 보고: 1W, 지분 변화 큰 순 30")
    assert ws.cell(r + 2, 7).value == "가나자산 · 단순 투자 · 지분 6.00%"
    assert ws.cell(r + 3, 1).value != tick(1)                          # 변화 0 보고는 뺐다
    r = row_of(ws, "자사주 · 증자 · 사채 결정: 1M")
    # 뒤에 정정된 판본(…010)은 빼고 정정본(…013)은 [정정], 같은 금액의 다른 결정 두 건은 둘 다
    assert col_values(ws, r + 2, 7, 5) == [
        "유상증자 · 제3자배정", "[정정] 자사주 취득 · 장내매수 · 주주가치 제고 · 2026년 09월 25일~-",
        "전환사채 · 전환가 -원", "유상증자 · 제3자배정 · 새 주식 1,234,567주", None]
    assert ws.cell(r + 3, 6).value == pytest.approx(100.0)            # 1e10 원 = 100 억원


def test_events_holder_source_lag_is_stated_in_titles(tmp_path: Path, roots: tuple[Path, Path],
                                                       con) -> None:
    eq = tmp_path / "equity"
    build_equity(eq, holder_last=DAYS[-3])
    wb, refs = book_with_refs()
    e = sheet_events.load(con, Sources(eq, roots[1], {}), D, UNI)
    sheet_events.write(wb, e, UNI, refs, D)
    ws = wb["Events"]
    assert (e.cal.anchor, e.hcal.anchor) == (DAYS[-2], DAYS[-3])
    assert ws["A4"].value == "기준일: 2026-09-24(원천 최신 · 엑셀 09-25)"
    row_of(ws, "임원 보유 순변동: 1M 순증 상위 15 · 순감 상위 15, 09-22까지")
    row_of(ws, "5% 대량보유 보고: 1W, 지분 변화 큰 순 30, 09-22까지")
    row_of(ws, "자사주 · 증자 · 사채 결정: 1M")                      # 결정 원천은 공시와 같은 날


# ── 매일 엑셀에 붙이기 · '데이터 없음' · 저장 뒤 손질 ─────────────────────────────
@pytest.fixture(scope="module")
def fi_world(tmp_path_factory: pytest.TempPathFactory, roots: tuple[Path, Path]
             ) -> dict[str, Path]:
    root = tmp_path_factory.mktemp("deliver_extra")
    model_root, fi_root = root / "model", root / "fi"
    for day in [*W38, *[x for x in W39 if x != HOLIDAY]]:
        write_model_day(model_root, day, write_fi_day(fi_root, day))
    runs = fi_root / "_runs"
    runs.mkdir(parents=True, exist_ok=True)
    (runs / "20260925_morning.json").write_text(json.dumps({
        "build_id": "m_20260925T010000Z", "equity_root": str(roots[0]),
        "stage_root": str(roots[1]), "equity_builds": {"price_daily": "b_test_0001"}}),
        encoding="utf-8")
    return {"model": model_root, "fi": fi_root, "out": root / "out"}


def test_daily_excel_has_new_sheets_in_order_with_meta(fi_world: dict[str, Path]) -> None:
    res = build_daily(FRI, "morning", model_root=fi_world["model"], fi_root=fi_world["fi"],
                      out_root=fi_world["out"])
    wb = load_workbook(res.path)
    assert wb.sheetnames == ["Score", "Sector", "Earnings", "Analyst", "Events", "Market", "ETF",
                             "Flow", "Short", "Compare", "Raw", "Meta", "trend"]
    meta = wb["Meta"]
    pairs = {meta.cell(r, 1).value: meta.cell(r, 2).value for r in range(8, meta.max_row + 1)}
    used = pairs["새 시트 데이터 id"]
    assert "price_daily=b_test_0001" in used and "stg_etf_price_daily=b_test_0001" in used
    sheets = [meta.cell(r, 1).value for r in range(8, meta.max_row + 1)]
    assert sheets[-1] == "Short" and sheets.index("ETF") < sheets.index("Analyst")
    assert not (fi_world["out"] / "tmp").exists()                     # DuckDB 임시 폴더는 지웠다
    with zipfile.ZipFile(res.path) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        assert len(charts) == 16
        xml = z.read(charts[0]).decode("utf-8")
        assert '<autoTitleDeleted val="1" />' in xml and "<numCache>" in xml
        assert any('type="column"' in z.read(n).decode("utf-8") and 'displayHidden="1"'
                   in z.read(n).decode("utf-8") for n in z.namelist()
                   if n.startswith("xl/worksheets/sheet"))


def test_line_chart_date_axis_sits_at_bottom(fi_world: dict[str, Path]) -> None:
    """꺾은선 날짜 축은 값 축 최솟값에서 교차한다(min) — autoZero 면 음수 계열(외국인 누적 순매수 등)에서
    축 선이 0 높이로 그래프를 가로지른다(10-16 사용자 '가로선이 길게'). 막대는 0 에서 자라야 해서 autoZero."""
    res = build_daily(FRI, "morning", model_root=fi_world["model"], fi_root=fi_world["fi"],
                      out_root=fi_world["out"])
    with zipfile.ZipFile(res.path) as z:
        xmls = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/charts/chart")]
    dated = [x for x in xmls if "<dateAx>" in x]
    lines = [x for x in dated if "<lineChart>" in x]
    bars = [x for x in dated if "<barChart>" in x]
    assert lines and bars
    for x in lines:
        assert '<crosses val="min" /><auto val="1" />' in x and "autoZero" not in x.split("<dateAx>")[1].split("</dateAx>")[0]
    for x in bars:
        assert '<crosses val="autoZero" /><auto val="1" />' in x


def test_missing_source_gives_one_line_sheet_and_warning(fi_world: dict[str, Path],
                                                        tmp_path: Path,
                                                        caplog: pytest.LogCaptureFixture) -> None:
    """원천 하나가 없으면 그 시트만 '데이터 없음: 사유'(warn), 나머지 시트 · 엑셀은 그대로."""
    st = tmp_path / "stage"
    build_stage(st)
    (st / "stg_etf_price_daily" / "MANIFEST.json").unlink()
    eq = Path(json.loads((fi_world["fi"] / "_runs" / "20260925_morning.json").read_text(
        encoding="utf-8"))["equity_root"])
    with caplog.at_level(logging.WARNING, logger="deliver"):
        res = build_daily(FRI, "morning", model_root=fi_world["model"], fi_root=fi_world["fi"],
                          out_root=tmp_path / "out", equity_root=eq, stage_root=st)
    wb = load_workbook(res.path)
    assert wb["ETF"]["A6"].value == "데이터 없음: stage stg_etf_price_daily 데이터를 못 찾았다"
    assert wb["ETF"]["A2"].value == "ETF: 데이터 없음"
    assert wb["Market"]["A2"].value == "Market: 지수 · 시장 수급 · 시장 폭"
    assert any("ETF" in r.getMessage() and "데이터 없음" in r.getMessage() for r in caplog.records)
    assert not re.search(r"/(home|Users|private)/", wb["ETF"]["A6"].value)
    meta = wb["Meta"]
    assert "ETF" not in [meta.cell(r, 1).value for r in range(8, meta.max_row + 1)]


def test_unexpected_error_in_one_sheet_is_logged_with_traceback(
        fi_world: dict[str, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
        caplog: pytest.LogCaptureFixture) -> None:
    def boom(*_args: object) -> list[tuple[str, str, str]]:
        raise RuntimeError("계산 중 터짐")
    monkeypatch.setitem(sheets_extra.BUILDERS, "Flow", boom)
    with caplog.at_level(logging.WARNING, logger="deliver"):
        res = build_daily(FRI, "morning", model_root=fi_world["model"], fi_root=fi_world["fi"],
                          out_root=tmp_path / "out")
    wb = load_workbook(res.path)
    assert wb["Flow"]["A6"].value == "데이터 없음: 계산 오류(RuntimeError) — 로그 참고"
    assert wb["Short"]["A2"].value == "Short: 공매도 · 대차 · 신용"
    rec = next(r for r in caplog.records if "Flow" in r.getMessage())
    assert rec.exc_info is not None and "계산 중 터짐" in str(rec.exc_info[1])


def test_post_save_failure_keeps_the_saved_workbook(fi_world: dict[str, Path], tmp_path: Path,
                                                    monkeypatch: pytest.MonkeyPatch,
                                                    caplog: pytest.LogCaptureFixture) -> None:
    """저장 뒤 손질(그래프 값 사본)이 실패해도 손질 전 엑셀을 내고 경고만 — 임시 파일은 남지 않는다."""
    def broken(path: Path, wb: Workbook) -> int:
        raise ValueError("그래프 XML 모양이 예상과 다르다")
    monkeypatch.setattr(excel_daily, "finish_charts", broken)
    out = tmp_path / "out"
    with caplog.at_level(logging.WARNING, logger="deliver"):
        res = build_daily(FRI, "morning", model_root=fi_world["model"], fi_root=fi_world["fi"],
                          out_root=out)
    assert load_workbook(res.path).sheetnames[3] == "Analyst"
    assert any("저장 뒤 손질" in r.getMessage() for r in caplog.records)
    assert not [p for p in res.path.parent.iterdir() if p.name.endswith(".tmp")]


def test_finish_charts_failure_removes_its_temp_file(fi_world: dict[str, Path],
                                                     tmp_path: Path) -> None:
    """이미 손질한 파일을 다시 손질하면 모양 검사에서 멈춘다 — 원본은 그대로, 임시 파일은 없다."""
    res = build_daily(FRI, "morning", model_root=fi_world["model"], fi_root=fi_world["fi"],
                      out_root=tmp_path / "out")
    before = res.path.read_bytes()
    with pytest.raises(ValueError, match="그래프 XML 모양"):
        qpack.finish_charts(res.path, load_workbook(res.path))
    assert res.path.read_bytes() == before
    assert not [p for p in res.path.parent.iterdir() if p.name.endswith(".tmp")]
