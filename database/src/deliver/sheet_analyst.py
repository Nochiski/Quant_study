"""Analyst 시트 — 증권사 목표가 · 투자의견 변화(DECISIONS §7 #11, 확정 샘플 30).

원천: equity `opinion_broker_daily`(WISE 종목 화면 '제공처별 투자의견'을 매일 수집, 09-01~) ·
`opinion_daily`(WISE 컨센서스, src 'wise' — 09-01 까지는 v3 판과 겹친다) · `price_daily`(종가).
기준일 = 증권사 의견 최신 수집일(토요일 수집이면 창은 그 앞 마지막 거래일에서 센다).
같은 날 같은 증권사 보고서는 한 건(나중 수집이 이긴다). 일별 그래프는 기준일(최신 수집일) 앞
마지막 거래일까지 — WISE 가 그날 보고서를 다음 날 올려 수집일은 늘 비어 보인다.
WISE 원천의 수집일 상한은 fi 판 JSON 의 wise_asof(`Sources.wise_upto`) — 아침판 D, 장 마감 판 T(N-53),
주말 보강판은 휴일 수집일 W 라 D 분 보고서가 들어오고 일별 그래프가 D 까지 간다(N-45). 그때 글의
날짜는 WISE 기준일 D 다 — 4행 '기준일: D(WISE 휴일 아침 재수집 W)', 보고서 날짜 · 컨센서스 날짜도 D.
"""
from __future__ import annotations

import bisect
import statistics
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date

import duckdb
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from .extra_common import (
    BIGS,
    F_CNT,
    F_PCT,
    F_WON,
    G1,
    HR,
    Calendar,
    Universe,
    base_text,
    close_day,
    closes,
    krx_calendar,
    latest,
    md,
    span,
    tail_cap,
)
from .qpack import (
    CHART_BLACK,
    CHART_BLUE,
    XRefs,
    XSheet,
    bar_series,
    cat_frame,
    chart_text,
    date_frame,
    line_series,
    x_note,
)
from .sources import Sources, day_of, lit, num, text

NAME = "Analyst"
WIDTHS = {"A": 7.5, "B": 18, "C": 14, "D": 7, "E": 7.5, "F": 7.5, "G": 8, "H": 10, "I": 9,
          "J": 7.5, "K": 7.5, "L": 46}
ORD = {"sell": 0, "hold": 1, "buy": 2}
MIN_BROKERS = 3             # 여력 상위 표 — 의견 낸 증권사 3곳 이상


@dataclass
class Act:
    """증권사 보고서 하나(같은 날 같은 증권사 1건)."""

    code: str
    broker: str
    day: date                                   # 보고서 날짜
    tp: float | None                            # 목표가
    ptp: float | None                           # 직전 목표가
    chg: float | None                           # 목표가 변화율(%)
    op: str | None                              # 투자의견 · 등급(sell · hold · buy)
    oc: str | None
    pop: str | None                             # 직전 투자의견 · 등급
    poc: str | None

    @property
    def kind(self) -> str:
        """목표가 변화 — 상향 · 하향 · 유지(직전 목표가 대비), 신규(직전 의견 · 목표가가 모두 없고
        지금 목표가가 있다 = 새 커버), 비교 없음(직전 목표가가 없는 그 밖 보고서 — 직전 의견만 있는
        등급 변경 등. 세지 않는다)."""
        if self.ptp is None:
            new = not self.poc and not (self.pop or "").strip() and self.tp is not None
            return "신규" if new else "비교 없음"
        c = self.chg or 0
        return "상향" if c > 0 else "하향" if c < 0 else "유지"

    @property
    def grade_change(self) -> bool:
        """직전 등급이 있고 지금 등급과 다르다(의견 상향 · 하향)."""
        return bool(self.poc and self.oc and self.poc != self.oc)


@dataclass
class Consensus:
    tp: float | None                            # 컨센서스 목표가
    n: int | None                               # 최근 3개월 의견 낸 증권사 수


@dataclass
class PerStock:
    """종목의 1W 보고서 집계 — 상향 · 하향 · 신규 건수, 변화율, '증권사 변화' 글."""

    up: int = 0
    dn: int = 0
    nw: int = 0
    chg: list[float] = field(default_factory=list)
    who: list[str] = field(default_factory=list)


@dataclass
class AnalystData:
    cal: Calendar
    first: date                                 # 증권사 의견 첫 수집일
    acts: list[Act]
    cons_day: date | None                       # 컨센서스 날짜(원천 최신)
    cons: dict[str, Consensus]
    close_day: date                             # 여력 계산 종가 날짜(원천 최신)
    close: dict[str, float | None]


def _opt(x: object) -> str | None:
    return None if x is None else str(x)


def load(con: duckdb.DuckDBPyConnection, src: Sources, upto: date, uni: Universe) -> AnalystData:
    ob = src.equity("opinion_broker_daily")
    # WISE 원천은 수집일 ≤ wise_asof — 아침판 D · 장 마감 판 T(N-53) · 주말 보강판 W(N-45). 종가는 D
    wise = src.wise_upto(upto)
    day = latest(con, ob, "fetched_date", wise, "true", "WISE 증권사별 투자의견")
    cal = krx_calendar(con, src, day, upto)
    d1m = cal.win["1M"]
    got = con.execute(f"SELECT min(fetched_date) FROM {ob}").fetchone()
    first = day_of(got[0]) if got and got[0] else day
    # 보고서 날짜는 엑셀 D 까지 — 주말 보강판도 WISE 기준일 D 의 보고서까지만 센다(휴일 수집일 W 날짜 보고서 X)
    op = con.execute(f"""
        SELECT ticker, fetched_date, broker, opinion_date, target_price_krw,
               prev_target_price_krw, change_pct, opinion, opinion_class, prev_opinion,
               prev_opinion_class
        FROM {ob} WHERE opinion_date > {lit(d1m)} AND opinion_date <= {lit(min(day, upto))}
          AND fetched_date <= {lit(day)}""").fetchall()
    # 나중 수집이 이긴다(파이썬 정렬은 안정 — 같은 수집일은 원장 행 순서)
    seen: dict[tuple[str, str, date], Act] = {}
    for t, _, br, od, tp, ptp, chg, opn, oc, pop, poc in sorted(op, key=lambda r: day_of(r[1])):
        if t in uni:
            seen[(t, br, od)] = Act(text(t), text(br), day_of(od), num(tp), num(ptp), num(chg),
                                    _opt(opn), _opt(oc), _opt(pop), _opt(poc))
    od_ = src.equity("opinion_daily")
    # 컨센서스 · 종가는 그 원천의 최신 날짜(≤ 엑셀 D) — 의견 수집일과 따로 늦을 수 있다
    got = con.execute(f"SELECT max(obs_date) FROM {od_} WHERE src = 'wise' "
                      f"AND obs_date <= {lit(wise)}").fetchone()
    cons_day = day_of(got[0]) if got and got[0] else None
    cons: dict[str, Consensus] = {}
    if cons_day is not None:
        for t, tp, n in con.execute(f"SELECT ticker, target_price_krw, analyst_count FROM {od_} "
                                    f"WHERE src = 'wise' AND obs_date = {lit(cons_day)}").fetchall():
            if t in uni:
                cnt = num(n)
                cons[text(t)] = Consensus(num(tp), None if cnt is None else int(cnt))
    cd = close_day(con, src, upto)
    return AnalystData(cal, first, list(seen.values()), cons_day, cons, cd,
                       closes(con, src, cd, list(uni)))


def write(wb: Workbook, a: AnalystData, uni: Universe, refs: XRefs, d: date,
          *, index: int | None = None) -> list[tuple[str, str, str]]:
    cal, day = a.cal, a.cal.anchor       # day = 의견 최신 수집일(거래일이 아닐 수 있다)
    win = cal.win
    # 주말 보강판 — WISE 를 엑셀 D 뒤 휴일 아침에 다시 받았다(N-45). 글의 날짜는 WISE 기준일 D 이고
    # 수집일(W)은 4행 괄호 · 5행 안내에만 적는다
    holiday = day > d
    base = d if holiday else day
    cday = md(min(a.cons_day, d)) if a.cons_day else md(base)
    pday = md(a.close_day)
    xs = XSheet(wb, NAME, WIDTHS, refs, index=index)
    ws = xs.ws
    xs.texts({"A1": "정렬: 블록마다(요약 대분류 순 · 목록 건수 · 날짜 순)",
              "A2": "Analyst: 증권사 목표가 · 투자의견 변화",
              "A3": f"자료: WISE 증권사별 투자의견({md(a.first)}부터 매일 수집) · 컨센서스 목표가 · "
                    f"모델 유니버스 {len(uni)}종목",
              "A4": (f"기준일: {d.isoformat()}(WISE 휴일 아침 재수집 {md(day)})" if holiday
                     else base_text(day, d)),
              "F4": "출처: quant-ledger 원장 · 시험 양식",
              "A5": f"1W: 보고서 날짜 {span(win['1W'], base)} · 1M: {span(win['1M'], base)}",
              "H5": "상향 · 하향: 직전 목표가 대비, 신규: 직전 의견 · 목표가 없음 · " + (
                  f"{md(d)} 보고서는 휴일 아침 WISE 재수집({md(day)})으로 포함" if holiday
                  else f"{md(day)} 보고서는 WISE 가 다음 날 올림")})
    put, header, close, cs = xs.put, xs.header, xs.close, xs.cs

    oa = [x for x in a.acts if x.code in uni]

    def in_1w(x: Act) -> bool:
        return x.day > win["1W"]
    cons, close_now = a.cons, a.close
    r0 = 6
    header(r0, [("A", "대분류별 요약: 건수"), ("C", "커버"), ("D", "1W"), ("H", "1M")],
           ["", "대분류", "종목 수", "상향", "하향", "신규", "상향 비율\n(%)", "상향", "하향", "신규",
            "상향 비율\n(%)"], ab=True)
    r = r0 + 2
    for big in [*BIGS, "전체"]:
        g = [x for x in oa if big == "전체" or uni[x.code].big == big]
        cov = sum(1 for c, u in uni.items() if (big == "전체" or u.big == big)
                  and c in cons and cons[c].n)
        xs.lab(r, big, bold=(big == "전체"))
        put(r, 3, cov, F_CNT)
        for base, sub in ((4, [x for x in g if in_1w(x)]), (8, g)):
            up, dn, nw = (sum(1 for x in sub if x.kind == k) for k in ("상향", "하향", "신규"))
            put(r, base, up, F_CNT)
            put(r, base + 1, dn, F_CNT)
            put(r, base + 2, nw, F_CNT)
            put(r, base + 3, up / (up + dn) * 100 if up + dn else None, F_PCT)
        r += 1
    a_end = r - 1
    close(a_end, 11)
    cs(f"G{r0 + 2}:G{a_end}")
    cs(f"K{r0 + 2}:K{a_end}")
    ws["G7"].comment = x_note("상향 비율(%): 상향 ÷ (상향 + 하향). 같은 날 같은 증권사 보고서는 한 건. "
                              "목표가를 그대로 둔 보고서(유지)는 세지 않는다")
    ws["C7"].comment = x_note(f"커버 종목 수: {cday} 컨센서스에 최근 3개월 의견을 낸 증권사가 1곳 이상인 "
                              "종목")

    ps: dict[str, PerStock] = defaultdict(PerStock)
    for x in oa:
        if not in_1w(x):
            continue
        s, k = ps[x.code], x.kind
        if k == "상향":
            s.up += 1
        if k == "하향":
            s.dn += 1
        if k == "신규":
            s.nw += 1
        if k in ("상향", "하향") and x.chg is not None:
            s.chg.append(x.chg)
        if k in ("상향", "하향", "신규"):
            s.who.append(f"{x.broker} " + ("신규" if k == "신규" else format(x.chg, "+.1f") + "%"))

    def upside(code: str) -> float | None:
        c, p = cons.get(code), close_now.get(code)
        return (c.tp / p - 1) * 100 if c and c.tp and p else None

    def avg(xs_: list[float]) -> float | None:
        return sum(xs_) / len(xs_) if xs_ else None
    lhead = ["코드", "이름", "대분류", "종합\n순위", "1W 상향\n(건)", "1W 하향\n(건)", "평균 변화\n(%)",
             "컨센서스\n목표가(원)", "현재가\n(원)", "여력\n(%)", "증권사\n(3개월)", "1W 증권사 · 변화"]
    lgrp = [("C", "분류"), ("E", "1W 목표가"), ("H", "컨센서스")]

    def list_rows(r: int, codes: list[str]) -> int:
        for code in codes:
            s = ps.get(code, PerStock())
            c = cons.get(code, Consensus(None, None))
            u = uni[code]
            xs.stock(r, code, u.name, u.big, u.rank, [
                (s.up, F_CNT, False), (s.dn, F_CNT, False), (avg(s.chg), F_PCT, False),
                (c.tp, F_WON, False), (close_now.get(code), F_WON, False),
                (upside(code), F_PCT, False), (c.n, F_CNT, False),
                (tail_cap(" · ".join(s.who)) or None, None, True)])
            r += 1
        return r
    # 그래프 4개 — 원자료는 숨긴 AD~AJ 열
    xs.hide(G1, 7, 12)
    # 일별 그래프 끝 = WISE 자료가 담을 수 있는 마지막 보고서 날짜 = 최신 수집일 앞 마지막 거래일
    # (WISE 는 그날 보고서를 다음 날 올린다). 아침 · 장 마감 = 수집일 전날, 주말 보강판은 휴일 수집이
    # D 분 보고서를 담아 D 까지(N-45)
    td = [x for x in cal.days if win["1M"] < x < day]
    cnt = {x: [0, 0] for x in td}
    for x in oa:
        i = bisect.bisect_left(td, x.day)                               # 휴일 날짜 보고서는 다음 거래일로
        if i < len(td) and x.kind in ("상향", "하향"):
            cnt[td[i]][0 if x.kind == "상향" else 1] += 1
    for k, h in enumerate(("날짜", "상향", "하향", "상향 비율 5일")):
        ws.cell(HR, G1 + k).value = h
    rr = HR + 1
    for i, x in enumerate(td):
        ws.cell(rr, G1).value = x
        ws.cell(rr, G1).number_format = "yyyy-mm-dd"
        ws.cell(rr, G1 + 1).value, ws.cell(rr, G1 + 2).value = cnt[x]
        up5 = sum(cnt[q][0] for q in td[max(0, i - 4): i + 1])
        dn5 = sum(cnt[q][1] for q in td[max(0, i - 4): i + 1])
        ws.cell(rr, G1 + 3).value = round(up5 / (up5 + dn5) * 100, 1) if up5 + dn5 else None
        rr += 1
    gend = rr - 1
    sec = []
    for big in BIGS:
        g = [x for x in oa if uni[x.code].big == big]
        up = sum(1 for x in g if x.kind == "상향")
        dn = sum(1 for x in g if x.kind == "하향")
        ups = [v for c, u in uni.items() if u.big == big and (v := upside(c)) is not None]
        sec.append((big, round(up / (up + dn) * 100, 1) if up + dn else None,
                    round(statistics.median(ups), 1) if ups else None))

    def bar_chart(head: int, col: int, label: str, cell: str, ttl: str, anchor: str) -> int:
        """대분류 가로 막대 — 아래부터 그려지므로 작은 값부터 적어 큰 값이 위로 오게"""
        ws.cell(head, G1 + 4).value, ws.cell(head, G1 + 5).value = "대분류", label
        vals = [(x[0], v) for x in sec if (v := x[col]) is not None]
        vals.sort(key=lambda x: x[1])
        for j, (big, v) in enumerate(vals):
            ws.cell(head + 1 + j, G1 + 4).value = big
            ws.cell(head + 1 + j, G1 + 5).value = v
        end = head + len(vals)
        if not vals:                    # 그릴 막대가 없으면 빈 범위로 가짜 막대가 서지 않게 그래프를 뺀다
            xs.title(cell, f"{ttl}: 데이터 없음")
            return end
        xs.title(cell, ttl)
        ch = BarChart()
        ch.type = "bar"
        ch.gapWidth = 60
        ch.add_data(Reference(ws, min_col=G1 + 5, min_row=head, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=G1 + 4, min_row=head + 1, max_row=end))
        bar_series(ch.series[0], CHART_BLUE)
        cat_frame(ch, "0")
        ws.add_chart(ch, anchor)
        return end
    cs0 = a_end + 2
    t1, t2 = cs0, cs0 + 16
    ttl1 = (f"일별 목표가 상향 · 하향 건수: 최근 1M({md(td[-1]) if td else md(day)}까지), "
            "상향 파랑 · 하향 검정")
    ttl2 = "목표가 상향 비율(%): 최근 5거래일 이동, 상향 ÷ (상향 + 하향)"
    if not td:                          # 1M 안 거래일이 없으면 빈 범위 그래프를 만들지 않는다
        xs.title(f"A{t1}", f"{ttl1}: 데이터 없음")
        xs.title(f"H{t1}", f"{ttl2}: 데이터 없음")
    else:
        xs.title(f"A{t1}", ttl1)
        ch = LineChart()
        ch.add_data(Reference(ws, min_col=G1 + 1, max_col=G1 + 2, min_row=HR, max_row=gend),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=G1, min_row=HR + 1, max_row=gend))
        line_series(ch.series[0], CHART_BLUE)
        line_series(ch.series[1], CHART_BLACK)
        date_frame(ch, "#,##0", "mm/dd", "days", 7, "days")
        if ch.legend is not None:
            ch.legend.position = "t"
            ch.legend.txPr = chart_text()
        ws.add_chart(ch, f"A{t1 + 1}")
        xs.title(f"H{t1}", ttl2)
        ch = LineChart()
        ch.add_data(Reference(ws, min_col=G1 + 3, min_row=HR, max_row=gend), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=G1, min_row=HR + 1, max_row=gend))
        ch.legend = None
        line_series(ch.series[0], CHART_BLUE)
        date_frame(ch, "0", "mm/dd", "days", 7, "days")
        ws.add_chart(ch, f"H{t1 + 1}")
    # 대분류 막대 2개: 1M 상향 비율 · 여력 중앙값(각자 정렬해 따로 적는다)
    e1 = bar_chart(gend + 3, 1, "1M 상향 비율", f"A{t2}", "대분류별 1M 목표가 상향 비율(%)",
                   f"A{t2 + 1}")
    bar_chart(e1 + 3, 2, "여력 중앙값", f"H{t2}",
              f"대분류별 목표가 여력 중앙값(%): {cday} 컨센서스 목표가 ÷ 종가 − 1", f"H{t2 + 1}")
    xs.tight(cs0, cs0 + 32)
    r = cs0 + 32
    tabs = []
    for title, key, rev in (("목표가 상향 종목: 1W 상향 건수 순 상위 20", "up", False),
                            ("목표가 하향 종목: 1W 하향 건수 순 상위 20", "dn", True)):
        t0 = r + 1
        header(t0, [("A", title), *lgrp], lhead)
        n_of = {c: s.up if key == "up" else s.dn for c, s in ps.items()}
        cand = [c for c in ps if n_of[c] > 0]
        cand.sort(key=lambda c: (-n_of[c], (avg(ps[c].chg) or 0) * (1 if rev else -1)))
        r = list_rows(t0 + 2, cand[:20])
        close(r - 1, 12)
        cs(f"G{t0 + 2}:G{r - 1}")
        cs(f"J{t0 + 2}:J{r - 1}")
        tabs.append(t0)
    ws[f"G{tabs[0] + 1}"].comment = x_note(
        "평균 변화(%): 1W 안 상향 · 하향 보고서의 목표가 변화율 평균(직전 목표가 대비)")
    ws[f"J{tabs[0] + 1}"].comment = x_note(f"여력(%): {cday} 컨센서스 목표가 ÷ {pday} 종가 − 1")
    t0 = r + 1
    header(t0, [("A", f"목표가 여력 상위 20: 증권사 {MIN_BROKERS}곳 이상"), *lgrp], lhead)
    ups = {c: v for c in uni if (v := upside(c)) is not None and (cons[c].n or 0) >= MIN_BROKERS}
    cand = sorted(ups, key=lambda c: -ups[c])[:20]
    r = list_rows(t0 + 2, cand)
    close(r - 1, 12)
    cs(f"G{t0 + 2}:G{r - 1}")
    cs(f"J{t0 + 2}:J{r - 1}")
    t0 = r + 1
    header(t0, [("A", "투자의견 변경 · 신규 커버: 1M"), ("C", "분류"), ("E", "보고서"), ("G", "투자의견"),
                ("I", "목표가(원)")],
           ["코드", "이름", "대분류", "종합\n순위", "날짜", "증권사", "이전", "현재", "이전", "현재", "변화\n(%)",
            "구분"])
    # 등급 변경(1M, 직전 등급이 있고 다르다)을 먼저 가르고, 신규 커버(1W)는 직전 의견 · 목표가가 모두
    # 없는 보고서만 — 둘은 겹치지 않아 한 보고서는 한 번만 나온다
    grade = sorted([x for x in oa if x.grade_change], key=lambda x: x.day, reverse=True)
    news = sorted([x for x in oa if x.kind == "신규" and in_1w(x)], key=lambda x: x.day,
                  reverse=True)
    r = t0 + 2
    for x in [*grade, *news]:
        kd = ("신규 커버" if not x.grade_change else
              ("의견 상향" if ORD.get(str(x.oc), 1) > ORD.get(str(x.poc), 1) else "의견 하향"))
        u = uni[x.code]
        xs.stock(r, x.code, u.name, u.big, u.rank, [
            (md(x.day), None, True), (x.broker, None, True), (x.pop, None, True),
            (x.op, None, True), (x.ptp, F_WON, False), (x.tp, F_WON, False),
            (x.chg if x.ptp else None, F_PCT, False), (kd, None, True)])
        r += 1
    ch_end = r - 1
    close(ch_end, 12)
    ws[f"L{t0 + 1}"].comment = x_note(
        "구분: 의견 상향 · 하향은 1M 안 투자의견 등급(매도 < 중립 < 매수)이 바뀐 보고서, 신규 커버는 "
        "1W 안 직전 의견 · 목표가가 모두 없는 보고서")
    xs.page(f"A1:L{ch_end}")
    return [("Analyst", "대분류별 요약 · 상향 / 하향 / 신규",
             "WISE 증권사별 투자의견(종목 화면 '제공처별 투자의견')을 매일 받아 보고서 날짜 기준으로 셈. "
             "상향 · 하향 = 직전 목표가 대비, 신규 = 직전 의견 · 목표가 없음, 같은 날 같은 증권사는 1건"),
            ("Analyst", "컨센서스 목표가 · 여력 · 증권사(3개월)",
             f"{cday} WISE 컨센서스(최근 3개월 의견) 목표가 · 그 목표가 ÷ {pday} 종가 − 1 · 의견 낸 "
             "증권사 수"),
            ("Analyst", "그래프 4개",
             "일별 목표가 상향 · 하향 건수(휴일 보고서는 다음 거래일, "
             + ("휴일 수집이 기준일 보고서까지 담아 기준일까지" if day > d
                else "기준일 보고서는 WISE 가 다음 날 올려 전날까지")
             + ") · 상향 비율 5거래일 이동 · 대분류별 1M 상향 비율 · 대분류별 목표가 여력 중앙값. "
             "원자료는 숨긴 AD~AJ 열"),
            ("Analyst", "이력 한계",
             f"증권사별 의견 수집은 {md(a.first)} 시작"
             + (" — 3M 창은 아직 없음" if a.first > win["3M"] else ""))]


__all__ = ["NAME", "Act", "AnalystData", "Consensus", "PerStock", "load", "write"]
