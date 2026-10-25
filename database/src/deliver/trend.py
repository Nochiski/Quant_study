"""주 모델 순위 흐름 — Δ순위 1W·1M 과 셀 안 꺾은선(엑셀 스파크라인) 원자료.

2026-10-06 사용자 결정 N-16: 1일 Δ순위는 노이즈라 빼고, 1M 순위 변화를 숫자와 셀 안 꺾은선으로
보인다(1W 는 같은 날 '없애도 되겠다'로 뺐다 — 계산 함수는 '1W' 도 받는다).
그래프는 모양만 알아보면 된다.

  비교 판   1W = D−7일, 1M = D−1개월(달력. 같은 날이 없으면 그달 말일) **이하**의 마지막 성공 판
            (휴장이면 직전 거래일 판 — Q.Pack 의 −1W 표기와 같은 달력 기준).
  Δ순위     비교 판 순위 − 오늘 순위(양수 = 상승). 어느 한쪽에 순위가 없으면 빈칸.
  흐름      1M 비교 판(없으면 D−1개월 다음 날)부터 D 까지 성공 판의 일별 **원순위 × −1**
            (그래프 위 = 상승 — 스파크라인은 세로축을 뒤집지 못해 부호를 바꿔 싣는다).
            선의 처음(비교 판) → 끝(D) 차이가 곧 Δ순위라 선 방향·선 색(Δ 부호)·Δ 숫자가 늘
            같은 말을 한다.
            10-07 까지는 백분위 100·(N−순위)/(N−1)였다 — 유니버스가 커지는 날 순위가 밀려도 백분위가
            올라 '선은 오르는데 빨강'이 생겼다(E-06, N-25 Q8 '원순위 선').
  판 파일이 지워졌거나(keep) 그 판에 주 모델이 없으면 그날은 흐름에서 빠진다. 단 주 모델이 없고
            앞 주 모델(`contracts.PRIMARY_PREDECESSOR` — scope@2.0 의 앞은 scope@1.0)이 있으면
            그 순위로 잇고(N-47 · Q-E1, 바뀐 날부터 약 4주) `Trend.older` 에 날짜를 남긴다.
  T-50(PR-6b) 장 마감 판(evening)의 1M 비교 판이 evening 이력에 없거나 주 모델 순위를
            못 읽으면, 대체 루트(연구 루트 data/model)가 주어졌을 때만 그 루트 아침
            확정판(morning)의 같은 규칙 비교 판을 쓴다. 흐름 첫 점도 그 판이라 위 '처음 → 끝 = Δ'
            가 그대로다. evening 이력이 D−1개월까지 닿으면 저절로 evening 판만 쓴다.
  비교 이력 basis  주말 보강판(supplement)은 휴장 앞 마지막 거래일에만 서는 드문 판이라 자기 이력이 없다.
            컷오버 뒤 사용자가 매일 받는 엑셀이 장 마감 판이라 장 마감 판처럼 견준다(N-45 · 10-12 결정):
            전일 순위 · 비교 판 · 흐름 = 장 마감 판 이력(`history_root` — data/model_db/model), 그 판이
            없는 날(10-14 전 · 빠진 날)은 날짜마다 아침 확정판(보강판과 같은 연구 루트)으로 메운다 —
            T-50 대체와 같은 뜻을 날짜마다 적용(`HISTORY_BASIS` · `history_sources` · `latest_in`).
            흐름 끝 점(D)만 보강판 자신이다.
"""
from __future__ import annotations

import calendar
import logging
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import duckdb
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from . import qpack
from .reader import (
    DeliverError,
    ModelRun,
    history_spec,
    previous_run,
    read_scores,
    runs_between,
)
from .view import rank_of, ticker_of

log = logging.getLogger(__name__)
SHEET = "trend"              # 그래프 원자료(숨김 시트). 행 번호를 Score 시트와 맞춘다
WINDOWS = ("1M",)             # 엑셀에 싣는 창(1W 는 10-06 사용자 결정으로 뺐다)
LINE_UP, LINE_DOWN, LINE_FLAT = "1E8C45", "D0312D", "8C8C8C"   # 상승 초록 · 하락 빨강 · 같음 회색
# Meta '1M 비교 데이터' 값 뒤에 붙인다
FALLBACK_NOTE = "(아침 확정 데이터 — 장 마감 데이터 비교 데이터 없음, T-50)"
# 비교 이력을 다른 basis 에서 읽는 판 — 주말 보강판(N-45)은 장 마감 판처럼 견준다(컷오버 뒤 사용자가
# 받는 엑셀 = 장 마감 판, 10-12 결정). 그 이력이 없는 날은 `FILL_BASIS`(아침 확정판)로 메운다 — 장 마감 판
# 1M 비교 판의 T-50 대체와 같은 뜻. 표에 없는 basis(morning · evening)는 자기 이력이다
HISTORY_BASIS: dict[str, str] = {"supplement": "evening"}
FILL_BASIS = "morning"
Source = tuple[Path, str]           # (model 루트, basis)


def history_sources(model_root: Path, basis: str, history_root: Path | None = None) -> list[Source]:
    """전일 순위 · 비교 판 · 흐름을 읽을 (루트, basis) — 앞이 우선. 주말 보강판은 장 마감 판
    루트(`history_root`, 없으면 건너뜀) 다음 자기 루트(연구 루트)의 아침 확정판."""
    hist = HISTORY_BASIS.get(basis)
    if hist is None:
        return [(model_root, basis)]
    return [*([] if history_root is None else [(history_root, hist)]), (model_root, FILL_BASIS)]


def latest_in(sources: Sequence[Source], before: str | date) -> tuple[Path, ModelRun] | None:
    """`before` 보다 앞선 가장 최근 성공 판과 그 루트 — 원천 여럿이면 날짜가 늦은 것, 같은 날이면
    앞 원천(장 마감 판이 빠진 날만 아침판이 메운다)."""
    best: tuple[Path, ModelRun] | None = None
    for root, b in sources:
        r = previous_run(root, before, b)
        if r is not None and (best is None or r.date > best[1].date):
            best = (root, r)
    return best


def month_back(d: date) -> date:
    """한 달 전 같은 날(없으면 그달 말일)."""
    y, m = (d.year, d.month - 1) if d.month > 1 else (d.year - 1, 12)
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def window_start(d: date, which: str) -> date:
    return d - timedelta(days=7) if which == "1W" else month_back(d)


def _on_or_before(model_root: Path, d: date, basis: str) -> ModelRun | None:
    return previous_run(model_root, d + timedelta(days=1), basis)


def source_word(r: ModelRun) -> str:
    """주말 보강판 Meta 비교 표시 — 장 마감 판이면 '장 마감 데이터', 메운 아침판이면 '아침 확정
    데이터(대체)'."""
    return "장 마감 데이터" if r.basis == "evening" else "아침 확정 데이터(대체)"


@dataclass(frozen=True)
class Trend:
    dates: tuple[str, ...]                        # 흐름에 든 판 날짜(오름차순, 마지막 = D)
    line: Mapping[str, tuple[int | None, ...]]    # 종목 → dates 순 −순위(위 = 상승)
    base: Mapping[str, str | None]                # '1W'·'1M' → 비교 판 날짜(없으면 None)
    base_rank: Mapping[str, Mapping[str, int | None]]   # '1W'·'1M' → 비교 판 순위
    fallback: frozenset[str] = frozenset()        # 비교 판을 대체 루트 아침판에서 쓴 창(T-50)
    # 날짜 → 그날 순위를 읽은 앞 주 모델 spec(주 모델이 그 판에 없어 앞 주 모델로 이은 날만, N-47)
    older: Mapping[str, str] = field(default_factory=dict)
    # 날짜 → 흐름 · 비교 판으로 쓴 model 판(그날 fi 판을 다시 읽는 곳 — 1M 유예 일수)
    runs: Mapping[str, ModelRun] = field(default_factory=dict)

    def delta(self, which: str, ticker: str, rank_now: int | None) -> int | None:
        then = self.base_rank[which].get(ticker)
        return None if then is None or rank_now is None else then - rank_now

    def first_index(self, which: str, d: date) -> int:
        """그 창의 흐름이 시작하는 dates 위치 — 비교 판, 없으면 창 시작 다음 날 이후 첫 판."""
        base = self.base[which]
        after = window_start(d, which) + timedelta(days=1)
        start = base if base is not None else after.isoformat()
        return next((i for i, x in enumerate(self.dates) if x >= start), len(self.dates))


def _ranks(model_root: Path, run: ModelRun,
           spec_id: str) -> tuple[str, dict[str, int | None]] | None:
    """(읽은 spec, 종목 → 순위) — 그 판에 주 모델이 없으면 앞 주 모델(`history_spec`). 둘 다
    없거나 점수 표를 못 읽으면 None."""
    sid = history_spec(run, spec_id)
    if sid is None:
        return None
    try:
        return sid, {ticker_of(r): rank_of(r) for r in read_scores(model_root, run, sid)}
    except DeliverError:
        return None


def load_trend(model_root: Path, run: ModelRun, spec_id: str, basis: str,
               fallback_root: Path | None = None, history_root: Path | None = None) -> Trend:
    """`history_root` = 주말 보강판의 장 마감 판 model 루트(`history_sources`, 그 밖 basis 는 무시)."""
    d = date.fromisoformat(run.date)
    srcs = history_sources(model_root, basis, history_root)
    bases = {w: latest_in(srcs, window_start(d, w) + timedelta(days=1)) for w in WINDOWS}
    b1m = bases["1M"]
    start = b1m[1].date if b1m is not None else (month_back(d) + timedelta(days=1)).isoformat()
    by_date: dict[str, dict[str, int | None]] = {}
    older: dict[str, str] = {}
    runs: dict[str, ModelRun] = {}

    def take(r: ModelRun, got: tuple[str, dict[str, int | None]]) -> dict[str, int | None]:
        sid, ranks = got
        by_date[r.date], runs[r.date] = ranks, r
        if sid != spec_id:
            older[r.date] = sid
        return ranks
    # 날짜마다 앞 원천의 판, 없거나 순위를 못 읽으면 다음 원천(보강판: 장 마감 판 → 아침 확정판).
    # 흐름 끝 점(D)은 이 판 자신이다
    cands: dict[str, list[tuple[Path, ModelRun]]] = {}
    for root, b in srcs:
        for x in runs_between(root, start, run.date, b):
            if x.date != run.date:
                cands.setdefault(x.date, []).append((root, x))
    cands[run.date] = [(model_root, run)]
    for day in sorted(cands):
        for root, r in cands[day]:
            if (got := _ranks(root, r, spec_id)) is not None:
                take(r, got)
                break
    base: dict[str, str | None] = {}
    base_rank: dict[str, Mapping[str, int | None]] = {}
    for w, found in bases.items():
        ranks = None
        if found is not None:
            root, b = found
            ranks = by_date.get(b.date)
            if ranks is None and (got := _ranks(root, b, spec_id)) is not None:
                ranks = take(b, got)
        base[w] = None if ranks is None or found is None else found[1].date
        base_rank[w] = ranks or {}
    fallback: frozenset[str] = frozenset()
    if basis == "evening" and fallback_root is not None and base["1M"] is None:
        # T-50 — 대체 루트 아침 확정판의 같은 규칙 판(D−1개월 이하 마지막 성공 판). 그 날짜는
        # evening 흐름 판들(D−1개월 다음 날부터)보다 앞이라 흐름 첫 점이 된다
        cut, why = month_back(d), ""
        try:
            fb = _on_or_before(fallback_root, cut, "morning")
            got = None if fb is None else _ranks(fallback_root, fb, spec_id)
        except (DeliverError, OSError, ValueError, duckdb.Error) as e:
            # 보조 입력이라 연구 루트가 깨져도 장 마감 판 엑셀은 멈추지 않는다 — 대체 없이 빈칸(P1)
            fb, got, why = None, None, f"읽기 실패({type(e).__name__}: {e})"
        if fb is not None and got is not None:
            base["1M"], base_rank["1M"], fallback = fb.date, take(fb, got), frozenset({"1M"})
        else:
            # 대체를 시도했는데 못 쓴 경우 — evening 1M 판이 없는 이행 기간에만 오는 자리다
            if not why:
                why = ("판 없음" if fb is None
                       else f"주 모델 없음({fb.date} 판 specs 에 {spec_id} 없음)"
                       if history_spec(fb, spec_id) is None
                       else f"점수 표 없음({fb.date} 판 {history_spec(fb, spec_id)})")
            log.warning("T-50 1M 비교 판 대체 안 씀(Δ순위 1M 빈칸) — 대체 루트 %s · 기준일 %s"
                        "(D−1개월) 이하 · 사유 %s", fallback_root, cut, why)
    dates = tuple(sorted(by_date))
    line: dict[str, list[int | None]] = {}
    for i, day in enumerate(dates):
        for t, rk in by_date[day].items():
            row = line.setdefault(t, [None] * len(dates))
            row[i] = None if rk is None else -rk
    return Trend(dates, {t: tuple(v) for t, v in line.items()}, base, base_rank, fallback,
                 older, runs)


def sheet_ref(name: str) -> str:
    """수식 안 시트 이름 — 영숫자 · 밑줄만이면 그대로(trend), 그 밖은 작은따옴표로 감싼다."""
    return name if name.replace("_", "").isalnum() and name.isascii() else f"'{name}'"


def line_color(delta: int | None) -> str:
    if delta is None or delta == 0:
        return LINE_FLAT
    return LINE_UP if delta > 0 else LINE_DOWN


def write_sheet(wb: Workbook, tickers: Sequence[str], trend: Trend) -> None:
    """숨김 원자료 시트 — 7행 머리(코드·날짜), 8행부터 Score 시트와 같은 행 순서. 값 = −순위."""
    ws = wb.create_sheet(SHEET)
    qpack.put(ws, 7, 1, "코드").font = qpack.font(bold=True)
    for j, day in enumerate(trend.dates, start=2):
        qpack.put(ws, 7, j, day[5:]).font = qpack.font(bold=True)
    for i, t in enumerate(tickers, start=qpack.FIRST_DATA_ROW):
        qpack.put(ws, i, 1, t).font = qpack.font()
        for j, v in enumerate(trend.line.get(t, ()), start=2):
            if v is not None:
                cell = qpack.put(ws, i, j, v)
                cell.font = qpack.font()
                cell.alignment = Alignment(horizontal="right")
    qpack.tighten_rows(ws)
    ws.sheet_state = "hidden"


def spark_groups(tickers: Sequence[str], trend: Trend, d: date,
                 cols: Mapping[str, str], deltas: Mapping[str, Mapping[str, int | None]],
                 ) -> list[qpack.SparkGroup]:
    """창(1W·1M)마다 선 색(Δ 부호)별 묶음. 흐름 판이 2개 미만이면 그 창은 그리지 않는다.

    `cols` = 창 → Score 시트의 그래프 열 글자, `deltas` = 창 → 종목 → Δ순위."""
    out: list[qpack.SparkGroup] = []
    last = len(trend.dates) + 1                    # 원자료 시트 마지막 날짜 열(1 = 코드)
    for w in WINDOWS:
        first = trend.first_index(w, d) + 2
        if last - first + 1 < 2 or w not in cols:
            continue
        ref_cols = f"{get_column_letter(first)}{{r}}:{get_column_letter(last)}{{r}}"
        by_color: dict[str, list[tuple[str, str]]] = {}
        for i, t in enumerate(tickers, start=qpack.FIRST_DATA_ROW):
            if t not in trend.line:
                continue
            ref = f"{sheet_ref(SHEET)}!" + ref_cols.format(r=i)
            by_color.setdefault(line_color(deltas[w].get(t)), []).append((ref, f"{cols[w]}{i}"))
        out += [qpack.SparkGroup(c, tuple(cells)) for c, cells in sorted(by_color.items())]
    return out


__all__ = ["FALLBACK_NOTE", "FILL_BASIS", "HISTORY_BASIS", "SHEET", "Source", "Trend",
           "history_sources", "latest_in", "line_color", "load_trend", "month_back", "source_word",
           "spark_groups", "window_start", "write_sheet"]
