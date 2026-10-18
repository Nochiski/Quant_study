"""Market 시트 — 지수 · 시장 수급 · 시장 폭(DECISIONS §7 #11, 확정 샘플 30 `x5_sheets.py`).

원천: equity `index_daily`(KRX 지수) · `price_daily` · `universe_daily` · `flow_daily`(키움, 보통주) ·
수정종가(카탈로그 매크로 `v_adj_price` 와 같은 정의 — `equity.views` 템플릿을 판 고정 원천에 올린다).
기준일: 지수 = KRX 지수(코스피) 최신일, 시장 폭 · 거래대금 = 그날과 종목 가격 최신일 중 이른 날,
투자자 수급 = 키움 수급 최신 완전한 날(블록마다 따로 — 지수보다 늦으면 4행 · 표 제목에 그 날짜를 적고,
없는 날은 0 이 아니라 빈칸). 시장 구분 = 지수 기준일 이하 최신 유니버스 날짜(`common_markets`).
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date

import duckdb
from equity import views
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from .extra_common import (
    F_CNT,
    F_FL,
    F_PCT,
    G1,
    HIST_DAYS,
    HR,
    INV,
    Calendar,
    Latest,
    base_text,
    close_day,
    common_markets,
    krx_calendar,
    latest,
    latest_complete,
    md,
    until,
)
from .qpack import CHART_BLUE, XRefs, XSheet, date_frame, line_series, x_note
from .sources import Sources, day_of, lit, num, text

NAME = "Market"
DAILY: tuple[str, ...] = ("frgnr_invsr_krw", "natfor_krw", "orgn_krw", "ind_invsr_krw", "penfnd_etc_krw")
# (이름, 열, 들여쓰기) — 국가는 원천 값이 늘 0 이라 행을 뺀다
INVS = [("개인", ["ind_invsr_krw"], 0), ("외국인", ["frgnr_invsr_krw", "natfor_krw"], 0),
        ("기관계", ["orgn_krw"], 0), ("금융투자", ["fnnc_invt_krw"], 1), ("투신", ["invtrt_krw"], 1),
        ("사모", ["samo_fund_krw"], 1), ("연기금", ["penfnd_etc_krw"], 1), ("보험", ["insrnc_krw"], 1),
        ("은행", ["bank_krw"], 1), ("기타금융", ["etc_fnnc_krw"], 1), ("기타법인", ["etc_corp_krw"], 0)]
MAIN = [("KOSPI", "코스피"), ("KOSDAQ", "코스닥"), ("KOSPI", "코스피 200"), ("KOSDAQ", "코스닥 150"),
        ("KOSPI", "코스피 대형주"), ("KOSPI", "코스피 중형주"), ("KOSPI", "코스피 소형주"),
        ("KOSDAQ", "코스닥 대형주"), ("KOSDAQ", "코스닥 중형주"), ("KOSDAQ", "코스닥 소형주")]
WIDTHS = {"A": 7.5, "B": 20, "C": 9, "D": 7.5, "E": 7.5, "F": 7.5, "G": 7.5, "H": 7.5, "I": 9,
          "J": 9, "K": 7.5, "L": 7.5}
HIGH_LOW_DAYS = 250         # 52주 = 250거래일
HIGH_LOW_MIN = 200          # 이력 200거래일 이상만 52주 판정
BREADTH_DAYS = 20
ADJ_DAYS = 320              # 수정종가를 읽는 범위(거래일) — 시장 폭 20일 × 52주 창 250일 + 여유
NO_RETURN = -999            # 업종 정렬에서 1M 수익률이 없는 지수를 맨 뒤로


@dataclass
class IndexRec:
    """지수 한 줄 — 기준일 종가 · 시총, 창별 수익률(%), 거래대금 20일 평균(원)."""

    cls: str
    name: str
    close: float | None
    mcap: float | None
    ret: dict[str, float | None]
    tv20: float | None


@dataclass
class Breadth:
    """그날 보통주 상승 · 하락 · 보합 종목 수(시장별)와 52주 신고가 · 신저가 종목 수 — 그날 비교할
    가격 행이 없으면 모두 None(빈칸)."""

    day: date
    kospi: tuple[int | None, int | None, int | None]
    kosdaq: tuple[int | None, int | None, int | None]
    hi: int | None
    lo: int | None


@dataclass
class MarketData:
    cal: Calendar                                       # 지수 · 시장 폭 기준
    flow: Latest                                        # 투자자 수급 기준일(일부 수집일 뺌)
    flow_cal: Calendar
    first_year: int                                     # 지수 자료 첫해
    index: list[IndexRec]
    series: dict[date, dict[str, float | None]]         # 날짜 → 코스피 · 코스닥 종가(1년)
    breadth: list[Breadth]                              # 최근 20거래일, 날짜순
    tv: dict[date, dict[str, float]]                    # 날짜 → 시장 → 보통주 거래대금(원)
    mkt_flow: dict[tuple[str, str], dict[str, float | None]]   # (시장, 창) → 투자자 열 → 순매수(원)
    daily_flow: dict[date, dict[str, float | None]]     # 날짜 → 열 → 순매수(원)
    breadth_day: date                                   # 시장 폭 · 거래대금 블록 기준일


def adj_price_macro(con: duckdb.DuckDBPyConnection, src: Sources, since: date) -> None:
    """수정종가 매크로 `v_adj_price(as_of)` 를 이 연결에 올린다 — 카탈로그와 같은 템플릿
    (`equity.views`)을 판 고정 원천으로 채운다(카탈로그 파일은 서버 절대 경로를 담아 쓰지 않는다).
    price_daily 는 `since` 이후 행만 — 그 행들의 값은 앞 행과 무관해(계수 누적은 adj_factor 에서)
    같다."""
    sources = {"price_daily": f"(SELECT * FROM {src.equity('price_daily')} "
                              f"WHERE date >= {lit(since)})",
               "adj_factor": src.equity("adj_factor"),
               "trading_calendar": src.equity("trading_calendar")}
    views.install_temp_macros(con, sources, names=("v_adj_price",))


def load(con: duckdb.DuckDBPyConnection, src: Sources, upto: date) -> MarketData:
    ix = src.equity("index_daily")
    day = latest(con, ix, "date", upto, "index_name = '코스피'", "KRX 지수 일별")
    cal = krx_calendar(con, src, day, upto)
    win = cal.win
    d250 = cal.back(HIST_DAYS)
    # ── 지수 ──
    byix: dict[tuple[str, str], dict[date, tuple[float | None, float | None, float | None]]] = \
        defaultdict(dict)
    for cls, nm, dd, c, v, mc in con.execute(
            f"SELECT index_class, index_name, date, close_idx, value_krw, mktcap_krw FROM {ix} "
            f"WHERE date >= {lit(min(win['YTD'], cal.back(60)))} AND date <= {lit(day)}"
    ).fetchall():
        byix[(text(cls), text(nm))][day_of(dd)] = (num(c), num(v), num(mc))
    index = []
    for (cls, nm), s in byix.items():
        if day not in s:
            continue
        c, _, mcap = s[day]
        ret: dict[str, float | None] = {}
        for k, d0 in win.items():
            p0 = s[d0][0] if d0 in s else None
            ret[k] = (c / p0 - 1) * 100 if p0 and c is not None else None
        tv20 = [v for x in cal.days[-20:] if x in s and (v := s[x][1]) is not None]
        index.append(IndexRec(cls, nm, c, mcap, ret, sum(tv20) / len(tv20) if tv20 else None))
    series: dict[date, dict[str, float | None]] = defaultdict(dict)
    for nm, dd, c in con.execute(
            f"SELECT index_name, date, close_idx FROM {ix} WHERE index_name IN ('코스피', '코스닥') "
            f"AND date >= {lit(d250)} AND date <= {lit(day)} ORDER BY date").fetchall():
        series[day_of(dd)][text(nm)] = num(c)
    # ── 시장 폭(보통주, 수정종가) — 오늘 오른 · 내린 종목, 52주(250거래일) 최고 · 최저와 같은 종목 ──
    common_markets(con, src, day)
    # 블록 기준일 = 지수 기준일과 종목 가격(fi 판 고정) 최신일 중 이른 날 — 지수는 current 판이라 D+1
    # 아침 체인 뒤에 만들면 하루 앞선다. 그 뒤 날짜 행은 내지 않는다(N-49 ④ 블록마다 원천 날짜)
    bday = close_day(con, src, day)
    bdays = [x for x in cal.days if x <= bday]
    adj_price_macro(con, src, cal.back(ADJ_DAYS))
    con.execute("CREATE OR REPLACE TEMP TABLE dd (date DATE, i INTEGER)")
    con.executemany("INSERT INTO dd VALUES (?, ?)", [(d, i) for i, d in enumerate(cal.days)])
    last20 = len(bdays) - BREADTH_DAYS
    got = {day_of(r[0]): r for r in con.execute(f"""
        WITH ap AS (
            SELECT a.ticker, mk.market, dd.i, dd.date, a.adj_close AS c
            FROM v_adj_price({lit(bday)}, NULL) a JOIN mk USING (ticker) JOIN dd USING (date)
            WHERE a.date >= {lit(cal.back(ADJ_DAYS))} AND a.adj_close IS NOT NULL),
        w AS (
            SELECT ticker, market, i, date, c, lag(c) OVER p AS pc, lag(i) OVER p AS pi,
                   max(c) OVER h AS hi, min(c) OVER h AS lo, count(c) OVER h AS n
            FROM ap
            WINDOW p AS (PARTITION BY ticker ORDER BY i),
                   h AS (PARTITION BY ticker ORDER BY i
                         RANGE BETWEEN {HIGH_LOW_DAYS - 1} PRECEDING AND CURRENT ROW))
        SELECT date,
               count(*) FILTER (WHERE market = 'KOSPI' AND c > pc),
               count(*) FILTER (WHERE market = 'KOSPI' AND c < pc),
               count(*) FILTER (WHERE market = 'KOSPI' AND c = pc),
               count(*) FILTER (WHERE market = 'KOSDAQ' AND c > pc),
               count(*) FILTER (WHERE market = 'KOSDAQ' AND c < pc),
               count(*) FILTER (WHERE market = 'KOSDAQ' AND c = pc),
               count(*) FILTER (WHERE n >= {HIGH_LOW_MIN} AND c >= hi),
               count(*) FILTER (WHERE n >= {HIGH_LOW_MIN} AND c <= lo)
        FROM w WHERE i >= {last20} AND pi = i - 1 GROUP BY date""").fetchall()}
    # 비교할 행이 하나도 없는 날은 빈칸 줄(0 이 아니다)
    breadth = [Breadth(x, (r[1], r[2], r[3]), (r[4], r[5], r[6]), r[7], r[8]) if r is not None
               else Breadth(x, (None, None, None), (None, None, None), None, None)
               for x in bdays[last20:] for r in (got.get(x),)]
    tv: dict[date, dict[str, float]] = defaultdict(dict)
    for dd, mkt, v in con.execute(f"""
            SELECT p.date, u.market, sum(p.value_krw)
            FROM {src.equity('price_daily')} p JOIN {src.equity('universe_daily')} u
                 ON u.ticker = p.ticker AND u.date = p.date
            WHERE p.date >= {lit(d250)} AND p.date <= {lit(day)} AND u.sec_type = 'common'
              AND u.market IN ('KOSPI', 'KOSDAQ') GROUP BY 1, 2""").fetchall():
        if (x := num(v)) is not None:
            tv[day_of(dd)][text(mkt)] = x
    first = con.execute(f"SELECT min(date) FROM {ix} WHERE index_name = '코스피'").fetchone()
    # ── 시장 투자자별 순매수(키움, 보통주 — 시장은 위 `mk`). 수급 기준일 · 창은 따로 ──
    fl = src.equity("flow_daily")
    flow = latest_complete(con, fl, "date", upto, "src = 'kiwoom' AND frgnr_invsr_krw IS NOT NULL",
                           "키움 투자자별 매매")
    fcal = krx_calendar(con, src, flow.day, upto)
    fwin = fcal.win
    con.execute("CREATE OR REPLACE TEMP TABLE w (k VARCHAR, d0 DATE)")
    con.executemany("INSERT INTO w VALUES (?, ?)", list(fwin.items()))
    sums = ", ".join(f"sum({c}) AS {c}" for c in INV)
    mkt_flow = {(text(r[0]), text(r[1])): dict(zip(INV, map(num, r[2:]), strict=True))
                for r in con.execute(f"""
            SELECT mk.market, w.k, {sums} FROM {fl} f JOIN mk USING (ticker) JOIN w ON f.date > w.d0
            WHERE f.src = 'kiwoom' AND f.date <= {lit(flow.day)}
              AND f.date > {lit(min(fwin['YTD'], fcal.back(HIST_DAYS)))} GROUP BY 1, 2""").fetchall()}
    daily = ", ".join(f"sum({c}) AS {c}" for c in DAILY)
    daily_flow = {day_of(r[0]): dict(zip(DAILY, map(num, r[1:]), strict=True))
                  for r in con.execute(f"""
        SELECT f.date, {daily} FROM {fl} f JOIN mk USING (ticker)
        WHERE f.src = 'kiwoom' AND f.date > {lit(d250)} AND f.date <= {lit(flow.day)}
        GROUP BY 1""").fetchall()}
    first_day = day_of(first[0]) if first and first[0] else day
    return MarketData(cal, flow, fcal, first_day.year, index, series, breadth, tv, mkt_flow,
                      daily_flow, bday)


def write(wb: Workbook, m: MarketData, refs: XRefs, d: date,
          *, index: int | None = None) -> list[tuple[str, str, str]]:
    cal, day = m.cal, m.cal.day
    fday = m.flow.day
    xs = XSheet(wb, NAME, WIDTHS, refs, index=index)
    ws = xs.ws
    flow_note = ("" if fday == day and m.flow.partial is None
                 else f" · 수급 {md(fday)}{m.flow.note()}")
    xs.texts({"A1": "정렬: 블록마다(대표 지수 고정 · 업종 1M 높은 순 · 날짜 최근 순)",
              "A2": "Market: 지수 · 시장 수급 · 시장 폭",
              "A3": f"자료: KRX 지수 · 종목 일별 · 키움 투자자별 매매(보통주) {m.first_year}~",
              "A4": base_text(day, d) + flow_note, "F4": "출처: quant-ledger 원장 · 시험 양식",
              "A5": "수익률: 지수 종가, 1D · 1W · 1M · 3M = 1 · 5 · 20 · 60거래일 전, "
                    f"YTD = {cal.ytd0.isoformat()} 대비"})
    put, header, close, cs = xs.put, xs.header, xs.close, xs.cs
    ixd = {(i.cls, i.name): i for i in m.index}
    ihead = ["", "지수", "종가", "1D", "1W", "1M", "3M", "YTD", "거래대금\n20일(조원)", "시총\n(조원)"]

    def idx_rows(r: int, keys: list[tuple[str, str]]) -> int:
        for k in keys:
            i = ixd[k]
            xs.lab(r, i.name)
            put(r, 3, i.close, "#,##0.00")
            for c, w in zip(range(4, 9), ("1D", "1W", "1M", "3M", "YTD"), strict=True):
                put(r, c, i.ret[w], F_PCT)
            put(r, 9, i.tv20 / 1e12 if i.tv20 else None, "#,##0.0")
            put(r, 10, i.mcap / 1e12 if i.mcap else None, "#,##0")
            r += 1
        return r
    r0 = 6
    header(r0, [("A", "대표 지수"), ("D", "수익률(%)"), ("I", "규모")], ihead, ab=True)
    r = idx_rows(r0 + 2, [k for k in MAIN if k in ixd])
    m1_end = r - 1
    close(m1_end, 10)
    cs(f"D{r0 + 2}:H{m1_end}")
    # 그래프 4개(1년) — 원자료는 숨긴 AD~AH 열
    xs.hide(G1, 5, 10)
    ser = m.series
    tvs = {x: sum(v.values()) for x, v in m.tv.items() if len(v) == 2}   # 두 시장이 다 있는 날만
    for k, h in enumerate(("날짜", "코스피", "코스닥", "외국인 누적", "거래대금 20일")):
        ws.cell(HR, G1 + k).value = h
    rr = HR + 1
    cum = 0.0
    tvw: list[float] = []
    dates = sorted(ser)
    for dd in dates:
        # 누적은 첫 날(1년 전)을 0 으로 두고 다음 날부터 더한다. 수급 · 거래대금이 없는 날(수급이 늦은
        # 지수 날 등)은 0 이 아니라 빈 점
        f = m.daily_flow.get(dd, {})
        fr, nf = f.get("frgnr_invsr_krw"), f.get("natfor_krw")
        cum_v = round(cum, 3) if dd == dates[0] else None
        if fr is not None and nf is not None:
            cum += (fr + nf) / 1e12
            cum_v = round(cum, 3)
        a20 = None
        if dd in tvs:
            tvw.append(tvs[dd] / 1e12)
            a20 = round(sum(tvw[-20:]) / len(tvw[-20:]), 3)
        ws.cell(rr, G1).value = dd
        ws.cell(rr, G1).number_format = "yyyy-mm-dd"
        ws.cell(rr, G1 + 1).value = ser[dd].get("코스피")
        ws.cell(rr, G1 + 2).value = ser[dd].get("코스닥")
        ws.cell(rr, G1 + 3).value = cum_v
        ws.cell(rr, G1 + 4).value = a20
        rr += 1
    gend = rr - 1
    cs0 = m1_end + 2
    t1, t2 = cs0, cs0 + 16
    for cell, ttl, col, anchor in (
            (f"A{t1}", "코스피: 1년", 1, f"A{t1 + 1}"), (f"H{t1}", "코스닥: 1년", 2, f"H{t1 + 1}"),
            (f"A{t2}", "외국인 누적 순매수(조원): 코스피 + 코스닥, 1년", 3, f"A{t2 + 1}"),
            (f"H{t2}", "시장 거래대금 20일 평균(조원): 코스피 + 코스닥 보통주, 1년", 4, f"H{t2 + 1}")):
        xs.title(cell, ttl)
        ch = LineChart()
        ch.add_data(Reference(ws, min_col=G1 + col, min_row=HR, max_row=gend), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=G1, min_row=HR + 1, max_row=gend))
        ch.legend = None
        line_series(ch.series[0], CHART_BLUE)
        date_frame(ch, "#,##0" if col != 3 else "#,##0.0", "yy/mm", "days", 2, "months")
        ws.add_chart(ch, anchor)
    xs.tight(cs0, cs0 + 32)
    # 시장 투자자별 순매수
    t0 = cs0 + 33
    header(t0, [("A", f"시장 투자자별 순매수(억원): 보통주{until(fday, day)}"), ("C", "코스피"),
                ("H", "코스닥")],
           ["", "투자자", "1D", "1W", "1M", "3M", "YTD", "1D", "1W", "1M", "3M", "YTD"], ab=True)
    r = t0 + 2
    for nm, cols, ind in INVS:
        xs.lab(r, nm, bold=(ind == 0), indent=ind)
        for base, mkt in ((3, "KOSPI"), (8, "KOSDAQ")):
            for j, w in enumerate(("1D", "1W", "1M", "3M", "YTD")):
                f = m.mkt_flow.get((mkt, w), {})
                vals = [f.get(c) for c in cols]          # 자료가 없는 창은 빈칸(0 이 아니다)
                put(r, base + j, None if any(v is None for v in vals)
                    else sum(v for v in vals if v is not None) / 1e8, F_FL)
        r += 1
    m2_end = r - 1
    close(m2_end, 12)
    cs(f"C{t0 + 2}:G{m2_end}")
    cs(f"H{t0 + 2}:L{m2_end}")
    ws[f"A{t0 + 1}"].comment = x_note(
        "투자자: 키움 투자자별 매매(보통주 합). 외국인 = 외국인 + 기타외국인. 기관계 = 금융투자 · 투신 · "
        "사모 · 연기금 · 보험 · 은행 · 기타금융 · 국가의 합(국가는 원천 값이 늘 0 이라 행을 뺌). "
        "개인 + 외국인 + 기관계 + 기타법인 = 0")
    # 시장 폭
    t0 = m2_end + 2
    header(t0, [("A", f"시장 폭: 최근 {BREADTH_DAYS}거래일, 보통주{until(m.breadth_day, day)}"),
                ("C", "코스피(종목 수)"),
                ("F", "코스닥(종목 수)"), ("I", "52주(두 시장)"), ("K", "거래대금(조원)")],
           ["", "날짜", "상승", "하락", "보합", "상승", "하락", "보합", "신고가", "신저가", "코스피",
            "코스닥"], ab=True)
    r = t0 + 2
    for b in reversed(m.breadth):
        xs.lab(r, md(b.day))
        for j, v in enumerate((*b.kospi, *b.kosdaq)):
            put(r, 3 + j, v, F_CNT)
        put(r, 9, b.hi, F_CNT)
        put(r, 10, b.lo, F_CNT)
        tvd = m.tv.get(b.day, {})
        kp, kq = tvd.get("KOSPI"), tvd.get("KOSDAQ")
        put(r, 11, None if kp is None else kp / 1e12, "#,##0.0")
        put(r, 12, None if kq is None else kq / 1e12, "#,##0.0")
        r += 1
    m3_end = r - 1
    close(m3_end, 12)
    ws[f"I{t0 + 1}"].comment = x_note(
        f"52주 신고가 · 신저가: 수정종가가 최근 {HIGH_LOW_DAYS}거래일 최고 · 최저와 같은 종목 수"
        f"(이력 {HIGH_LOW_MIN}거래일 이상만)")
    # 업종 지수(제조 · 규모 지수 뺌)
    sect = {"KOSPI": [k for k in ixd if k[0] == "KOSPI" and not k[1].startswith("코스피")
                      and k[1] != "제조"],
            "KOSDAQ": [k for k in ixd if k[0] == "KOSDAQ" and not k[1].startswith("코스닥")
                       and k[1] != "제조"]}
    r = m3_end + 1
    for mkt, title in (("KOSPI", "코스피 업종 지수: 1M 높은 순"), ("KOSDAQ", "코스닥 업종 지수: 1M 높은 순")):
        t0 = r + 1
        header(t0, [("A", title), ("D", "수익률(%)"), ("I", "규모")], ihead, ab=True)
        keys = sorted(sect[mkt], key=lambda k: -(v if (v := ixd[k].ret["1M"]) is not None
                                                  else NO_RETURN))
        r = idx_rows(t0 + 2, keys)
        close(r - 1, 10)
        cs(f"D{t0 + 2}:H{r - 1}")
    xs.page(f"A1:L{r - 1}")
    return [("Market", "대표 지수 · 업종 지수",
             "KRX 지수 일별 종가 수익률(%) · 거래대금 20일 평균(조원) · 시총(조원). 업종 = KRX 업종 지수"
             "(제조 · 규모 지수 뺌)"),
            ("Market", "시장 투자자별 순매수",
             "키움 투자자별 매매를 보통주 전체로 더함(억원). 외국인 = 외국인 + 기타외국인"),
            ("Market", "시장 폭 · 그래프",
             "보통주 상승 · 하락 · 보합 종목 수, 52주 신고가 · 신저가(수정종가) · 거래대금. 그래프 원자료는 "
             "숨긴 AD~AH 열")]


__all__ = ["NAME", "Breadth", "IndexRec", "MarketData", "adj_price_macro", "load", "write"]
