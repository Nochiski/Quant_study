"""삼성증권 Q.Pack 문법 시트 작성기 (MODEL_EXCEL_SPEC C·D).

  1행   정렬 마커(연한 파랑 BDD7EE, 가운데) — 'sort' · 정렬 키 열 '▲' · 나머지 's'
  2~4행 제목 띠(파랑 4F81BD, 흰 글씨) — 시트명 / 모델·버전 / 기준일·Source·판 id
  5행   각주(회색 기울임). `defs_row` 면 열마다 정의 문자열
  6~7행 2단 헤더 — 6행 그룹(파랑, 종합·순위 그룹은 금색 FFC000), 7행 열 이름(회색 7F7F7F),
        핵심 열은 보라 7030A0, 얇은 테두리로 열마다 상자, 7행 높이 33.75·줄바꿈
  8행~  데이터(채움 없음 · 위 정렬 · 맑은 고딕 8 · 기본 높이 11.25)
  열 너비 — 값·헤더에 맞춘 최소 폭(Q.Pack 실측: 보이는 숫자 열 4.4~6.4, 종목명 13~15). 10-05 전엔
        숫자 열을 13 으로 고정해 Q.Pack 의 두 배 넘게 넓었다(사용자 지적).
  그룹 사이 빈 구분 열(너비 5 · 숨김 · 아웃라인 1) · 식별 열 끝에서 틀고정 · 7행 자동필터
  색 스케일 3색 백분위 10/50/90: 빨강 F8696B → 노랑 FFEB84 → 초록 63BE7B(높음 = 초록 = 좋음,
  2026-10-06 사용자 결정 N-15 — Q.Pack 의 '높음 = 빨강'과 반대). 순위 열은 반전(1위 = 초록).
  Δ순위 열(`zero_mid`)은 가운데 = 0, 끝점 ±|값| 90 백분위(N-25 Q8). 레벨 값은 무색.

**값만 쓴다.** '=' 로 시작하는 문자열도 수식이 되지 않게 문자열 형으로 고정한다.

매일 엑셀은 위 문법을 고친 확정 양식(샘플 22 — `write_confirmed` · `write_meta_confirmed`)을 쓴다.
주간 엑셀은 위 문법(`write_table` · `write_meta`) 그대로.
"""
from __future__ import annotations

import math
import os
import re
import unicodedata
import zipfile
from collections.abc import Callable, Mapping, Sequence
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree
from xml.sax.saxutils import escape

from openpyxl.chart import BarChart, LineChart
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.drawing.text import Font as DFont
from openpyxl.formatting.rule import ColorScaleRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook import Workbook
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from .stats import quantile

FONT_NAME = "맑은 고딕"
FONT_SIZE = 8
ROW_HEIGHT = 11.25
HEADER_HEIGHT = 33.75
DEFS_HEIGHT = ROW_HEIGHT    # 정의 행도 한 줄(Q.Pack 5행 각주와 같은 높이) — 전문은 메타 시트
SEP_WIDTH = 5.0
MIN_WIDTH = 4.4             # Q.Pack 실측 보이는 열 최소(Sector 4.43)
NUM_MAX_WIDTH = 12.0        # 숫자 열 상한 — 넘는 값은 없다(시총 억원 9자리도 9 안팎)
TXT_MAX_WIDTH = 9.0         # 너비를 지정하지 않은 문자 열 상한(지정하면 그 값이 상한)
FIRST_DATA_ROW = 8
# 표준(Normal) 스타일 글꼴 — 큐팩 원본 styles.xml 과 같은 Arial 10
# (글꼴 0번·'표준' 스타일). 엑셀은 행 번호·열 머리를 이 글꼴로 그린다.
# openpyxl 기본 Calibri 11 은 줄 높이(약 18px)가 11.25pt 행(15px)보다 커서
# 행 번호가 잘린다(10-06 사용자 지적).
NORMAL_FONT = Font(name="Arial", sz=10, family=2)

C_STRIP = "BDD7EE"      # 1행 — accent1 60% 밝게
C_BAND = "4F81BD"       # 2~4행 제목 띠 · 6행 그룹 헤더 — accent1
C_HDR = "7F7F7F"        # 7행 열 이름 — 흰색 50% 어둡게
C_KEY = "7030A0"        # 핵심 열(보라)
C_GOLD = "FFC000"       # 종합·순위 그룹 — accent4
C_STYLE = "0070C0"      # 업종 시트 헤더(Q.Pack Style 시트)
C_NOTE = "595959"
C_NEW = "FFFF00"        # 형광 — '신규 진입' 한 종류만(D)
C_WHITE = "FFFFFF"
GREEN, YELLOW, RED = "63BE7B", "FFEB84", "F8696B"

# spark = 셀 안 꺾은선(엑셀 스파크라인) 자리 — 값 없이 지정 너비 그대로.
# 그림은 저장 뒤 add_sparklines 가 넣는다.
KINDS = ("id", "txt", "num", "pct", "chg", "rank", "spark")
_THIN = Side(style="thin", color="808080")
_MEDIUM = Side(style="medium", color="1F3864")


def new_workbook() -> Workbook:
    """표준 글꼴을 `NORMAL_FONT` 로 둔 새 통합문서."""
    wb = Workbook()
    # openpyxl 에는 기본 글꼴을 바꾸는 공개 API 가 없다. 셀 스타일이 하나도
    # 없는 지금 글꼴 0번을 갈아 끼우고 표준 스타일을 다시 계산하게 한다.
    wb._fonts = IndexedList([copy(NORMAL_FONT)])
    wb._named_styles["Normal"].font = copy(NORMAL_FONT)
    return wb


def font(**kw: object) -> Font:
    return Font(name=FONT_NAME, size=FONT_SIZE, **kw)  # type: ignore[arg-type]


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def scale_high_good() -> Rule:
    """백분위·변화율 열 — 높음 = 초록(좋음), 낮음 = 빨강(나쁨)."""
    return ColorScaleRule(start_type="percentile", start_value=10, start_color=RED,
                          mid_type="percentile", mid_value=50, mid_color=YELLOW,
                          end_type="percentile", end_value=90, end_color=GREEN)


def scale_fixed_zero(span: float) -> Rule:
    """가운데 = 0(노랑), 끝점 ±span 고정 — 음수 빨강 · 양수 초록(높음 = 좋음, N-15)."""
    return ColorScaleRule(start_type="num", start_value=-span, start_color=RED,
                          mid_type="num", mid_value=0, mid_color=YELLOW,
                          end_type="num", end_value=span, end_color=GREEN)


def scale_zero_mid(values: Sequence[object]) -> Rule:
    """Δ순위 열 — 가운데 = 0(노랑), 음수 빨강 · 양수 초록. 끝점 = |값|의 90 백분위(선형 보간)를
    ±M 으로 대칭에 둔다 — 백분위 50 가운데는 유니버스가 바뀌는 날 중앙값이 0 에서 벗어나 소폭 하락이
    연두로 보였다(E-06, Q8). 값이 없거나 전부 0 이면 M = 1."""
    mags = sorted(abs(float(v)) for v in values
                  if isinstance(v, int | float) and not isinstance(v, bool) and math.isfinite(v))
    m = quantile(mags, 0.9) if mags else 0.0
    return scale_fixed_zero(m if m > 0 else 1.0)


def scale_zero_anchor(values: Sequence[object]) -> Rule:
    """부호가 뜻을 갖는 증감 열(Earnings YoY) — 가운데 = 0(노랑), 음수 끝 = −lo · 양수 끝 = +hi.
    lo · hi = 음수 · 양수 |값|의 90% 분위(정렬한 값의 int(0.9·(n−1)) 번째)를 유효숫자 2자리로
    반올림. 한쪽 부호가 없으면 다른 쪽 값, 둘 다 없으면 1. 글자 칸은 뺀다."""
    nums = [float(v) for v in values
            if isinstance(v, int | float) and not isinstance(v, bool) and math.isfinite(v)]

    def p90(xs: list[float]) -> float | None:
        return sorted(xs)[int(0.9 * (len(xs) - 1))] if xs else None
    pos, neg = p90([v for v in nums if v > 0]), p90([-v for v in nums if v < 0])
    hi = float(f"{(pos or neg or 1):.2g}")
    lo = float(f"{(neg or pos or 1):.2g}")
    return ColorScaleRule(start_type="num", start_value=-lo, start_color=RED,
                          mid_type="num", mid_value=0, mid_color=YELLOW,
                          end_type="num", end_value=hi, end_color=GREEN)


def scale_rank() -> Rule:
    """순위 열 — 1위(작은 값) = 초록(좋음)."""
    return ColorScaleRule(start_type="percentile", start_value=10, start_color=GREEN,
                          mid_type="percentile", mid_value=50, mid_color=YELLOW,
                          end_type="percentile", end_value=90, end_color=RED)


def clean(v: object) -> object:
    """셀 값 정리 — NaN·inf 는 빈칸, bool 은 'Y'/''(엑셀 TRUE/FALSE 대신)."""
    if isinstance(v, bool):
        return "Y" if v else ""
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def put(ws: Worksheet, row: int, col: int, v: object) -> Any:
    """값만 쓴다 — '=' 로 시작하는 문자열도 수식이 아니라 문자열로 둔다. 반환 = openpyxl 셀."""
    cell: Any = ws.cell(row, col)
    val: Any = clean(v)
    cell.value = val
    if isinstance(val, str) and val.startswith("="):
        cell.data_type = "s"
    return cell


@dataclass(frozen=True)
class Col:
    """열 하나. kind: id(식별) · txt(문자) · num(레벨, 무색) · pct(백분위) · chg(변화) ·
    rank(순위)."""

    key: str
    label: str
    kind: str = "num"
    fmt: str | None = "#,##0.0"
    width: float | None = None
    key_col: bool = False       # 보라 헤더
    definition: str = ""
    item: bool = False          # 업종 시트 항목명(보라 글씨)
    color: bool | None = None   # None = kind 기본(pct·chg·rank 만 색)
    zero_mid: bool = False      # 색 가운데 = 0(Δ순위 — 부호와 색이 같은 쪽, E-06)
    span: float | None = None   # 색 가운데 0 · 끝점 ±span 고정(z 칸 — −3·0·+3)
    # 아래는 확정 양식(write_confirmed)만 쓴다
    fit_label: str | None = None  # 폭을 잴 때 쓸 머리 글(없으면 label) — 머리를 줄여도 폭은 그대로
    fmt_all: str | None = None  # 값과 무관하게 데이터 칸 전부에 거는 숫자 서식(빈칸·글자 칸 포함)
    fixed: float | None = None  # 값에 맞추지 않는 고정 폭
    core: bool = False          # 6·7행 핵심 블록 색(8064A2)
    fit_lines: bool = False     # 머리 폭을 낱말이 아니라 줄 단위로 잰다(줄 안에서 접지 않는 머리)
    dict_label: str | None = None   # Meta 열 사전 이름(없으면 머리 글의 줄바꿈을 공백으로)

    @property
    def colored(self) -> bool:
        return self.kind in ("pct", "chg", "rank") if self.color is None else self.color


@dataclass(frozen=True)
class Group:
    name: str
    cols: tuple[Col, ...]
    core: bool = False          # 종합·순위 그룹(금색 — Q.Pack 문법 write_table 만)
    gap: bool = True            # 앞 그룹과 사이에 빈 구분 열을 둔다(False = 바로 붙인 하위 블록)


@dataclass(frozen=True)
class Title:
    sheet: str
    model_line: str
    source_line: str
    note: str = ""


def _text_width(s: str, wide: float, narrow: float) -> float:
    """맑은 고딕 8 표시 폭(엑셀 열 너비 단위) 근사 — 한글 등 전각은 wide, 그 밖은 narrow."""
    return sum(wide if unicodedata.east_asian_width(ch) in ("W", "F") else narrow for ch in s)


def _shown(v: object, fmt: str | None) -> str:
    """셀에 보이는 문자열 근사 — 숫자는 서식(#,##0.0 등)대로, 그 밖은 str."""
    if isinstance(v, bool) or not isinstance(v, int | float):
        return "" if v is None else str(v)
    if isinstance(v, float) and not math.isfinite(v):
        return ""
    if not fmt or fmt == "General":
        return str(v)
    pct = fmt.endswith("%")
    dec = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
    x = v * 100 if pct else v
    s = f"{x:,.{dec}f}" if "," in fmt else f"{x:.{dec}f}"
    return s + ("%" if pct else "")


def fit_width(col: Col, rows: Sequence[Mapping[str, object]]) -> float:
    """열 너비 = 값·헤더가 들어가는 최소 폭(Q.Pack 처럼 촘촘히). 숫자는 '####' 가 안 나게 넉넉히 재고,
    헤더는 낱말(공백·줄바꿈으로 나뉨)이 한 줄에 온전히 들어가게 잰다 — 7행 3줄 안에서 낱말 사이로만
    접혀 '유니버/스' 처럼 가운데서 끊기지 않는다(10-05 렌더 확인). 상한 = 지정 너비 또는 종류별 기본."""
    if col.fixed is not None:
        return col.fixed
    if col.kind == "spark":
        return col.width or MIN_WIDTH
    numeric = col.kind not in ("id", "txt")
    wide, narrow, pad = (1.6, 0.85, 1.0) if numeric else (1.55, 0.8, 0.7)
    label = col.label if col.fit_label is None else col.fit_label
    words = label.split("\n") if col.fit_lines else label.split()
    need_hdr = max((_text_width(w, 1.55, 0.8) for w in words), default=0.0) + 0.7
    need_val = max((_text_width(_shown(r.get(col.key), col.fmt), wide, narrow) for r in rows),
                   default=0.0) + pad
    cap = col.width or (NUM_MAX_WIDTH if numeric else TXT_MAX_WIDTH)
    return round(min(cap, max(MIN_WIDTH, need_hdr, need_val)), 2)


def _layout(groups: Sequence[Group]) -> list[tuple[int, Group | None, Col | None]]:
    out: list[tuple[int, Group | None, Col | None]] = []
    c = 1
    for gi, g in enumerate(groups):
        if gi > 0 and g.gap:
            out.append((c, None, None))
            c += 1
        for col in g.cols:
            out.append((c, g, col))
            c += 1
    return out


def tighten_rows(ws: Worksheet) -> None:
    """높이를 정하지 않은 행에 11.25 를 박는다(이미 정한 7행 33.75 등은 그대로).

    시트 기본 높이(sheetFormatPr)만 두면 엑셀·LibreOffice 가 통합 문서 기본 글꼴(Calibri 11)로 행을
    다시 잡아 Q.Pack 보다 1.4배 높게 그린다 — 10-05 렌더: 한 쪽 29행 vs Q.Pack 37행, 행마다 박으면 37행.
    """
    for r in range(1, ws.max_row + 1):
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = ROW_HEIGHT


def title_block(ws: Worksheet, title: Title, ncol: int, sort_col: int | None) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = ROW_HEIGHT
    ws.sheet_format.customHeight = True
    for i in range(1, ncol + 1):
        marker = "sort" if i == 1 else ("▲" if i == sort_col else "s")
        cell = put(ws, 1, i, marker)
        cell.font = font()
        cell.fill = fill(C_STRIP)
        cell.alignment = Alignment(horizontal="center")
        for r in (2, 3, 4):
            ws.cell(r, i).fill = fill(C_BAND)
    for r, text, bold in ((2, title.sheet, True), (3, title.model_line, False),
                          (4, title.source_line, False)):
        cell = put(ws, r, 1, text)
        cell.font = font(bold=bold, color=C_WHITE)
        cell.alignment = Alignment(horizontal="left")
    if title.note:
        put(ws, 5, 1, title.note).font = font(italic=True, color=C_NOTE)


def write_table(wb: Workbook, title: Title, groups: Sequence[Group],
                rows: Sequence[Mapping[str, object]], *, sort_key: str | None = None,
                defs_row: bool = False, style: str = "qpack",
                n_freeze_groups: int = 1) -> list[tuple[str, str, str, str]]:
    """표 시트 하나. 반환 = 열 사전 [(시트, 그룹, 열, 정의)] — 메타 시트가 모은다.

    행 dict 의 `_fills`({열 key: 색})로 셀 채움을 줄 수 있다(신규 진입 형광).
    """
    ws = wb.create_sheet(title.sheet)
    layout = _layout(groups)
    ncol = layout[-1][0] if layout else 1
    sort_col = next((c for c, _, col in layout if col is not None and col.key == sort_key), None)
    title_block(ws, title, ncol, sort_col)

    hdr_side = _MEDIUM if style == "sector" else _THIN
    ident = {id(g) for g in groups[:n_freeze_groups]}     # 식별 그룹 = 틀고정 그룹(회색 헤더)
    first_in_group: set[int] = set()
    seen: set[str] = set()
    for c, g, col in layout:
        letter = get_column_letter(c)
        if g is None or col is None:
            dim = ws.column_dimensions[letter]
            dim.width, dim.hidden, dim.outlineLevel = SEP_WIDTH, True, 1
            continue
        if g.name not in seen:
            seen.add(g.name)
            first_in_group.add(c)
        ws.column_dimensions[letter].width = fit_width(col, rows)
        is_id = col.kind == "id" or id(g) in ident
        if style == "sector":
            g_fill, l_fill = C_STYLE, C_STYLE
        else:
            g_fill = C_KEY if col.key_col else (C_HDR if is_id else (C_GOLD if g.core else C_BAND))
            l_fill = C_KEY if col.key_col else C_HDR
        on_gold = g_fill == C_GOLD
        h6 = put(ws, 6, c, g.name if c in first_in_group else None)
        h6.fill = fill(g_fill)
        h6.font = font(bold=True, color="000000" if on_gold else C_WHITE)
        h6.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)
        h6.border = Border(left=hdr_side, right=hdr_side, top=hdr_side)
        h7 = put(ws, 7, c, col.label)
        h7.fill = fill(l_fill)
        h7.font = font(bold=col.key_col or style == "sector", color=C_WHITE)
        h7.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        h7.border = Border(left=hdr_side, right=hdr_side, bottom=hdr_side)
        if defs_row and col.definition and not is_id:
            d5 = put(ws, 5, c, col.definition)
            d5.font = font(italic=True, color=C_NOTE)
            d5.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[7].height = HEADER_HEIGHT
    if defs_row:
        ws.row_dimensions[5].height = DEFS_HEIGHT

    for ri, row in enumerate(rows, start=FIRST_DATA_ROW):
        fills = row.get("_fills") or {}
        assert isinstance(fills, Mapping)
        for c, _g, col in layout:
            if col is None:
                continue
            v = row.get(col.key)
            cell = put(ws, ri, c, v)
            cell.font = font(color=C_KEY) if col.item else font()
            text_like = col.kind in ("id", "txt") or isinstance(cell.value, str)
            cell.alignment = Alignment(vertical="top",
                                       horizontal="left" if text_like else "right")
            if col.fmt and col.kind not in ("id", "txt") and isinstance(cell.value, int | float):
                cell.number_format = col.fmt
            if col.key in fills:
                cell.fill = fill(str(fills[col.key]))

    last = FIRST_DATA_ROW + len(rows) - 1
    if rows:
        for c, _g, col in layout:
            if col is None or not col.colored:
                continue
            letter = get_column_letter(c)
            if col.span is not None:
                rule = scale_fixed_zero(col.span)
            elif col.zero_mid:
                rule = scale_zero_mid([r.get(col.key) for r in rows])
            else:
                rule = scale_rank() if col.kind == "rank" else scale_high_good()
            ws.conditional_formatting.add(f"{letter}{FIRST_DATA_ROW}:{letter}{last}", rule)
    n_freeze = sum(len(g.cols) for g in groups[:n_freeze_groups]) + max(0, n_freeze_groups - 1)
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, n_freeze + 1)
    ws.auto_filter.ref = f"A7:{get_column_letter(ncol)}{max(7, last)}"
    tighten_rows(ws)
    return [(title.sheet, g.name, col.label.replace("\n", " "), col.definition)
            for _c, g, col in layout if g is not None and col is not None]


def write_meta(wb: Workbook, title: Title, pairs: Sequence[tuple[str, object]],
               dictionary: Sequence[tuple[str, str, str, str]]) -> None:
    """메타 시트 — 키·값 표(8행~) 뒤에 열 사전(시트 · 그룹·열 · 정의)."""
    ws = wb.create_sheet(title.sheet)
    title_block(ws, title, 3, None)
    def_width = 110.0
    for letter, width in (("A", 18.0), ("B", 34.0), ("C", def_width)):
        ws.column_dimensions[letter].width = width

    def header(r: int, labels: Sequence[str]) -> None:
        for i, text in enumerate(labels, start=1):
            cell = put(ws, r, i, text)
            cell.fill = fill(C_HDR)
            cell.font = font(bold=True, color=C_WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="top")
            cell.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

    g6 = put(ws, 6, 1, "메타")
    g6.fill, g6.font = fill(C_BAND), font(bold=True, color=C_WHITE)
    for i in (2, 3):
        ws.cell(6, i).fill = fill(C_BAND)
    header(7, ("항목", "값", ""))
    r = FIRST_DATA_ROW
    for key, value in pairs:
        put(ws, r, 1, key).font = font(bold=True)
        cell = put(ws, r, 2, value)
        cell.font = font()
        cell.alignment = Alignment(vertical="top", horizontal="left")
        ws.cell(r, 1).alignment = Alignment(vertical="top")
        r += 1
    r += 1
    header(r, ("시트", "그룹 · 열", "정의"))
    r += 1
    for sheet, group, label, definition in dictionary:
        put(ws, r, 1, sheet).font = font()
        put(ws, r, 2, f"{group} · {label}").font = font()
        cell = put(ws, r, 3, definition)
        cell.font = font()
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        # 정의가 접히는 줄 수만큼만 높인다(자동 높이는 기본 글꼴 기준이라 헐겁다 — tighten_rows)
        lines = math.ceil(_text_width(str(definition or ""), 1.55, 0.8) / (def_width - 1))
        ws.row_dimensions[r].height = ROW_HEIGHT * max(1, lines)
        r += 1
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, 1)
    tighten_rows(ws)


# ── 확정 양식(샘플 22 — DECISIONS §7 #6·#7) ─────────────────────────────────
# 매일 엑셀만 이 양식을 쓴다(주간 엑셀은 위 write_table · write_meta 그대로).
#   1행   A1 '정렬: …'(왼쪽) · 정렬 열 위에만 '▲'
#   2행   'Sheet: 설명'(파랑 띠) · 3·4행 연한 파랑 DCE6F1 + 진남색 글자 — 틀 고정 폭에 다 들어가지
#         않으면 ' · ' 뒷부분을 첫 스크롤 열로 나눈다
#   5행   A5 읽는 법 + 블록 첫 열 위에 그 블록 설명
#   6·7행 그룹 595959 · 열 7F7F7F · 핵심 블록 8064A2, 그룹 이름은 칠이 같은 여러 칸이면 가운데 맞춤
#   선    블록 첫 열 왼쪽 · 표 오른쪽 끝 medium 595959, 블록 안 thin D9D9D9, 표 끝 = 다음 행 위쪽
#   인쇄  A4 · 1쪽 너비 · 2~7행 반복 · 바닥글 '&A' · '&P / &N'
C_GROUP = "FF595959"        # 6행 그룹 머리
C_COLHDR = "FF7F7F7F"       # 7행 열 머리
C_CORE = "FF8064A2"         # 핵심 블록
C_BAND34 = "FFDCE6F1"       # 3·4행 띠
C_BAND34_TXT = "FF1F497D"
C_POS_TXT, C_NEG_TXT = "FF0070C0", "FFC00000"     # 흑자 파랑 · 적자 빨강 글자
SIDE_BLOCK = Side(style="medium", color="FF595959")
SIDE_INNER = Side(style="thin", color="FFD9D9D9")
SIDE_PAIR = Side(style="thin", color="FFA6A6A6")    # Earnings 연도 묶음 사이 · 분기 YoY 앞
SIDE_HDR = Side(style="thin", color="FF808080")
_THIN_SECTOR = Side(style="thin", color="1F3864")   # Sector 머리 — 옛 medium 과 같은 색의 thin
SIGN_COLORS = {"흑자전환": C_POS_TXT, "흑전": C_POS_TXT, "적자전환": C_NEG_TXT,
               "적자지속": C_NEG_TXT, "적전": C_NEG_TXT, "적확": C_NEG_TXT, "적축": C_NEG_TXT,
               "적자": C_NEG_TXT}
PRINT_MARGINS = {"left": 0.25, "right": 0.25, "top": 0.35, "bottom": 0.45, "header": 0.15,
                 "footer": 0.2}


def text_width(s: str) -> float:
    """맑은 고딕 8 글 폭(열 너비 단위) — 전각 1.55 · 그 밖 0.8."""
    return _text_width(s, 1.55, 0.8)


def text_lines(text: str, cap: float) -> int:
    """공백에서 끊었을 때 폭 cap 안에 드는 줄 수(Meta 여러 줄 칸 높이)."""
    n, cur = 1, 0.0
    for word in text.split(" "):
        w = text_width(word)
        if cur and cur + 0.8 + w > cap:
            n, cur = n + 1, w
        else:
            cur = cur + 0.8 + w if cur else w
    return n


# 셀 메모 — 엑셀 메모 상자는 글에 맞춰 자라지 않는다. 줄을 300px 로 끊고 상자 크기를 줄 수 ·
# 글 폭으로 정한다(한글 14px · 그 밖 8px · 줄 21px · 여백 30px — 메모 기본 글꼴에 한글 대체 글꼴이
# 쓰여 넉넉히 잡는다).
NOTE_CJK, NOTE_OTHER, NOTE_LINE_H, NOTE_LINE_W, NOTE_PAD = 14, 8, 21, 300, 30
NOTE_GLUE = frozenset({"=", "<", ">", "≥", "≤", "−", "+", "×", "÷", "~"})
NBSP = "\xa0"


def note_px(s: str) -> int:
    return sum(NOTE_CJK if ord(ch) > 0x2E80 else NOTE_OTHER for ch in s)


def _wrap_px(text: str) -> list[str]:
    """공백에서 끊어 한 줄 폭 NOTE_LINE_W 이하로."""
    out: list[str] = []
    cur = ""
    for word in text.split(" "):
        cand = f"{cur} {word}" if cur else word
        if cur and note_px(cand) > NOTE_LINE_W:
            out.append(cur)
            cur = word
        else:
            cur = cand
    out.append(cur)
    return out


def _wrap_glued(text: str) -> list[str]:
    """기호(= < ≥ − × ÷ ~ …)는 앞뒤 낱말과, '·' 는 앞 낱말과 붙여(줄 안 공백 = NBSP) 줄 머리·끝에
    홀로 남지 않게 끊는다."""
    units: list[str] = []
    glue_next = False
    for word in text.split(" "):
        if units and (glue_next or word in NOTE_GLUE or word == "·"):
            units[-1] += " " + word
        else:
            units.append(word)
        glue_next = word in NOTE_GLUE
    return _wrap_px(" ".join(u.replace(" ", NBSP) for u in units))


def note_lines(text: str, *, glue: bool = True, wrap: bool = True) -> list[str]:
    """메모 줄 — 문단('\\n')마다 끊는다. wrap=False 면 문단 그대로(줄을 직접 정한 범례)."""
    paras = text.split("\n")
    if not wrap:
        return paras
    cut = _wrap_glued if glue else _wrap_px
    return [line for p in paras for line in cut(p)]


def note(text: str, author: str = "설명", *, glue: bool = True, wrap: bool = True) -> Comment:
    """셀 메모 + 상자 크기(너비 = 가장 긴 줄 + 여백, 높이 = 줄 수 × 21 + 여백)."""
    lines = note_lines(text, glue=glue, wrap=wrap)
    cm = Comment("\n".join(lines), author)
    cm.width = max(note_px(line) for line in lines) + NOTE_PAD
    cm.height = NOTE_LINE_H * len(lines) + NOTE_PAD
    return cm


@dataclass(frozen=True)
class Band:
    """확정 양식 1~5행. model·source = (앞, 뒤) — 틀 고정 폭에 다 들어가면 'A · B' 한 칸, 아니면
    뒤를 첫 스크롤 열로. notes = (열 key — None 이면 A 열, 글)."""

    sort_label: str | None
    title: str
    model: tuple[str, str]
    source: tuple[str, str]
    notes: tuple[tuple[str | None, str], ...] = ()


@dataclass(frozen=True)
class Page:
    """인쇄 — A4, fit_width 쪽 너비(높이는 fit_height, 0 = 제한 없음)."""

    portrait: bool = False
    fit_width: int = 1
    fit_height: int = 0
    over_then_down: bool = False
    title_cols: str | None = None


@dataclass
class Sheet:
    """write_confirmed 결과 — 표 뒤 손질(시트별 색 규칙·메모)에 쓴다."""

    ws: Worksheet
    cols: dict[str, int]                    # 열 key → 열 번호
    blocks: list[tuple[int, int, int]]      # 표마다 (6행 머리 행, 첫 데이터 행, 끝 데이터 행)
    last_col: int
    dictionary: list[tuple[str, str, str, str]]

    def letter(self, key: str) -> str:
        return get_column_letter(self.cols[key])


def _band_rows(ws: Worksheet, band: Band, ncol: int, sort_col: int | None, freeze_col: int,
               cols: Mapping[str, int]) -> None:
    for i in range(1, ncol + 1):
        cell = put(ws, 1, i, None)
        cell.font = font()
        cell.fill = fill(C_STRIP)
        cell.alignment = Alignment(horizontal="center")
        ws.cell(2, i).fill = fill(C_BAND)
        for r in (3, 4):
            c34 = ws.cell(r, i)
            c34.fill = fill(C_BAND34)
            c34.font = Font(name=NORMAL_FONT.name, sz=NORMAL_FONT.sz, family=NORMAL_FONT.family,
                            color=C_BAND34_TXT)
    if band.sort_label:
        a1 = put(ws, 1, 1, band.sort_label)
        a1.alignment = Alignment(horizontal="left")
    if sort_col is not None:
        put(ws, 1, sort_col, "▲")
    a2 = put(ws, 2, 1, band.title)
    a2.font = font(bold=True, color=C_WHITE)
    a2.alignment = Alignment(horizontal="left")

    def visible(c: int) -> bool:
        return not ws.column_dimensions[get_column_letter(c)].hidden
    room = (sum(ws.column_dimensions[get_column_letter(c)].width or 8.43
                for c in range(1, freeze_col) if visible(c)) if freeze_col > 1 else math.inf)
    scroll = freeze_col
    while scroll <= ncol and not visible(scroll):
        scroll += 1
    for r, (head, tail) in ((3, band.model), (4, band.source)):
        whole = f"{head} · {tail}"
        parts = ([(1, whole)] if text_width(whole) <= room - 1
                 else [(1, head), (scroll, tail)])
        for c, text in parts:
            cell = put(ws, r, c, text)
            cell.font = font(color=C_BAND34_TXT)
            cell.alignment = Alignment(horizontal="left")
    for key, text in band.notes:
        cell = put(ws, 5, 1 if key is None else cols[key], text)
        cell.font = font(italic=True, color=C_NOTE)
        cell.alignment = Alignment()


def _group_align(gi: int, g: Group) -> str:
    """그룹 이름 맞춤 — 칠이 같은 여러 칸 그룹은 '선택 영역의 가운데로', 첫(식별) 그룹 · 한 칸
    그룹 · 칠이 섞인 그룹(핵심 열만 보라)은 첫 칸 가운데."""
    uniform = len({c.core for c in g.cols}) == 1
    return "centerContinuous" if gi > 0 and len(g.cols) >= 2 and uniform else "center"


def _header_rows(ws: Worksheet, groups: Sequence[Group], layout: Sequence[tuple[int, Group | None,
                 Col | None]], r6: int, style: str) -> None:
    align = {id(g): _group_align(gi, g) for gi, g in enumerate(groups)}
    seen: set[int] = set()
    for c, g, col in layout:
        if g is None or col is None:
            continue
        first = id(g) not in seen
        seen.add(id(g))
        h6 = put(ws, r6, c, g.name if first else None)
        h6.fill = fill(C_CORE if col.core else C_GROUP)
        h6.font = font(bold=True, color=C_WHITE)
        h6.alignment = Alignment(horizontal=align[id(g)], vertical="top", wrap_text=False)
        h7 = put(ws, r6 + 1, c, col.label)
        h7.fill = fill(C_CORE if col.core else C_COLHDR)
        h7.font = font(bold=col.key_col or style == "sector", color=C_WHITE)
        h7.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.row_dimensions[r6 + 1].height = HEADER_HEIGHT


def _body_rows(ws: Worksheet, layout: Sequence[tuple[int, Group | None, Col | None]],
               rows: Sequence[Mapping[str, object]], r0: int, sign_colors: bool) -> None:
    for ri, row in enumerate(rows, start=r0):
        for c, _g, col in layout:
            if col is None:
                continue
            cell = put(ws, ri, c, row.get(col.key))
            color = C_KEY if col.item else None
            if sign_colors and isinstance(cell.value, str):
                color = SIGN_COLORS.get(cell.value, color)
            cell.font = font(color=color) if color else font()
            text_like = col.kind in ("id", "txt") or isinstance(cell.value, str)
            cell.alignment = Alignment(vertical="top",
                                       horizontal="left" if text_like else "right")
            if col.fmt_all:
                cell.number_format = col.fmt_all
            elif col.fmt and col.kind not in ("id", "txt") and isinstance(cell.value, int | float):
                cell.number_format = col.fmt


def _color_rules(ws: Worksheet, layout: Sequence[tuple[int, Group | None, Col | None]],
                 rows: Sequence[Mapping[str, object]], r0: int) -> None:
    if not rows:
        return
    r1 = r0 + len(rows) - 1
    for c, _g, col in layout:
        if col is None or not col.colored:
            continue
        letter = get_column_letter(c)
        if col.span is not None:
            rule = scale_fixed_zero(col.span)
        elif col.zero_mid:
            rule = scale_zero_mid([r.get(col.key) for r in rows])
        else:
            rule = scale_rank() if col.kind == "rank" else scale_high_good()
        ws.conditional_formatting.add(f"{letter}{r0}:{letter}{r1}", rule)


def _block_lines(ws: Worksheet, layout: Sequence[tuple[int, Group | None, Col | None]],
                 starts: set[int], last_col: int, r6: int, r1: int, hdr_side: Side,
                 inner_sides: Mapping[str, Side] | None) -> None:
    """블록 세로선 — 블록 첫 열 왼쪽 · 표 오른쪽 끝 medium, 블록 안은 데이터 행만 thin D9D9D9
    (`inner_sides` 를 주면 머리 행까지 그 열의 선). 머리 행의 나머지 선은 hdr_side 그대로."""
    r0 = r6 + 2
    for r in range(r6, r1 + 1):
        for c, _g, col in layout:
            if col is None:
                continue
            if c in starts:
                left = SIDE_BLOCK
            elif inner_sides is not None:
                left = inner_sides.get(col.key, SIDE_INNER)
            else:
                left = SIDE_INNER if r >= r0 else hdr_side
            right = SIDE_BLOCK if c == last_col else (hdr_side if r < r0 else None)
            ws.cell(r, c).border = Border(left=left, right=right,
                                          top=hdr_side if r == r6 else None,
                                          bottom=hdr_side if r == r6 + 1 else None)


def _close_line(ws: Worksheet, layout: Sequence[tuple[int, Group | None, Col | None]],
                r: int) -> None:
    """표 끝 선 — 마지막 행 다음 행의 위쪽(정렬 범위 밖이라 정렬해도 표 끝에 남는다)."""
    for c, _g, col in layout:
        if col is not None:
            ws.cell(r, c).border = Border(top=SIDE_BLOCK)


def page_setup(ws: Worksheet, page: Page, area: str) -> None:
    ws.sheet_view.zoomScale = 100
    ps = ws.page_setup
    ps.paperSize = 9
    ps.orientation = "portrait" if page.portrait else "landscape"
    ps.fitToWidth = page.fit_width
    ps.fitToHeight = page.fit_height
    if page.over_then_down:
        ps.pageOrder = "overThenDown"
    if page.title_cols:
        ws.print_title_cols = page.title_cols
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(**PRINT_MARGINS)
    ws.print_area = area
    ws.print_title_rows = "2:7"
    odd = ws.oddFooter
    if odd is not None:                 # openpyxl 형식상 Optional — 새 시트에는 늘 있다
        odd.left.text, odd.left.size = "&A", 8
        odd.center.text, odd.center.size = "&P / &N", 8


def write_confirmed(wb: Workbook, sheet: str, band: Band, groups: Sequence[Group],
                    blocks: Sequence[Sequence[Mapping[str, object]]], *,
                    sort_key: str | None = None, n_freeze_groups: int = 1, style: str = "table",
                    defs_notes: bool = False, inner_sides: Mapping[str, Side] | None = None,
                    sign_colors: bool = False, page: Page | None = None) -> Sheet:
    """확정 양식 표 시트. `blocks` = 세로로 쌓는 표(행 목록)들 — 둘째 표부터는 빈 줄 하나 뒤에
    6·7행 머리를 다시 쓴다(Sector 대분류 · 업종). 열 폭은 모든 표의 값으로 잰다.

    `defs_notes` = 열 정의를 7행 머리 칸 메모로(5행 정의 칸은 서식만 남긴다 — 확정 샘플 그대로),
    `inner_sides` = 블록 안 열 사이 선을 열마다 정한다(Earnings 3단 선 — 머리 행까지),
    `sign_colors` = 흑자 · 적자 글자 색(`SIGN_COLORS`)."""
    ws = wb.create_sheet(sheet)
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = ROW_HEIGHT
    ws.sheet_format.customHeight = True
    layout = _layout(groups)
    ncol = layout[-1][0] if layout else 1
    every = [r for b in blocks for r in b]
    cols = {col.key: c for c, _g, col in layout if col is not None}
    for c, _g, col in layout:
        dim = ws.column_dimensions[get_column_letter(c)]
        if col is None:
            dim.width, dim.hidden, dim.outlineLevel = SEP_WIDTH, True, 1
        else:
            dim.width = fit_width(col, every)
    freeze_col = 1 + max(c for c, g, _col in layout
                         if g is not None and any(g is x for x in groups[:n_freeze_groups]))
    visible = [c for c, _g, col in layout if col is not None]
    last_col = visible[-1]
    starts = {visible[0]}
    seen: set[int] = set()
    for c, g, _col in layout:
        if g is not None and id(g) not in seen:
            seen.add(id(g))
            starts.add(c)
    hdr_side = _THIN_SECTOR if style == "sector" else _THIN

    _header_rows(ws, groups, layout, 6, style)
    if defs_notes:
        ident = {id(g) for g in groups[:n_freeze_groups]}
        for c, g, col in layout:
            if g is None or col is None or not col.definition or col.kind == "id" \
                    or id(g) in ident:
                continue
            d5 = ws.cell(5, c)
            d5.font = font(italic=True, color=C_NOTE)
            d5.alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(7, c).comment = note(col.definition.replace(" — ", ": "), "정의", glue=False)
    sort_col = None if sort_key is None else cols.get(sort_key)
    _band_rows(ws, band, ncol, sort_col, freeze_col, cols)

    spans: list[tuple[int, int, int]] = []
    r6 = 6
    for bi, rows in enumerate(blocks):
        if bi > 0:
            _close_line(ws, layout, r6)            # 앞 표 끝 선 = 빈 줄 위쪽
            r6 += 1
            _header_rows(ws, groups, layout, r6, style)
        r0 = r6 + 2
        _body_rows(ws, layout, rows, r0, sign_colors)
        _color_rules(ws, layout, rows, r0)
        r1 = r0 + len(rows) - 1
        _block_lines(ws, layout, starts, last_col, r6, max(r1, r6 + 1), hdr_side, inner_sides)
        spans.append((r6, r0, r1))
        r6 = r1 + 1
    last = spans[-1][2] if every else FIRST_DATA_ROW - 1
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, freeze_col)
    h7 = spans[-1][0] + 1
    ws.auto_filter.ref = f"A{h7}:{get_column_letter(ncol)}{max(h7, spans[-1][2])}"
    tighten_rows(ws)
    if every:
        _close_line(ws, layout, last + 1)
    page_setup(ws, page or Page(), f"A1:{get_column_letter(ncol)}{max(last, 7)}")
    dictionary = [(sheet, g.name, col.dict_label or col.label.replace("\n", " "),
                   col.definition)
                  for _c, g, col in layout if g is not None and col is not None]
    return Sheet(ws, cols, spans, last_col, dictionary)


def write_meta_confirmed(wb: Workbook, band: Band, pairs: Sequence[tuple[str, object]],
                         dictionary: Sequence[tuple[str, str, str, str]],
                         tail: Sequence[tuple[str, str, str, str]] = (),
                         extra: Sequence[tuple[str, str, str]] = ()) -> Worksheet:
    """확정 양식 Meta — 키·값 표(8행~) 뒤 열 사전. `tail` = 사전 끝에 붙이는 줄(그룹 · 열 이름이
    길어 B 칸도 줄바꿈), `extra` = 그 뒤 새 시트 줄(시트, B 칸 이름, 정의 — 서식은 tail 과 같다).
    band.notes 의 '{hdr}' 는 열 사전 머리 행 번호로 바꾼다."""
    ws = wb.create_sheet("Meta")
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = ROW_HEIGHT
    ws.sheet_format.customHeight = True
    def_width = 110.0
    key_width = max((text_width(str(k)) + 1 for k, _ in pairs), default=0.0)
    for letter, width in (("A", max(18.0, key_width)), ("B", 34.0), ("C", def_width)):
        ws.column_dimensions[letter].width = width
    hdr = FIRST_DATA_ROW + len(pairs) + 1
    band = Band(band.sort_label, band.title, band.model, band.source,
                tuple((k, t.replace("{hdr}", str(hdr))) for k, t in band.notes))
    _band_rows(ws, band, 3, None, 1, {})

    def header(r: int, labels: Sequence[str]) -> None:
        for i, text in enumerate(labels, start=1):
            cell = put(ws, r, i, text)
            cell.fill = fill(C_COLHDR)
            cell.font = font(bold=True, color=C_WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="top")
            cell.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

    put(ws, 6, 1, "Meta").font = font(bold=True, color=C_WHITE)
    for i in (1, 2, 3):
        ws.cell(6, i).fill = fill(C_GROUP)
    header(7, ("항목", "값", ""))
    r = FIRST_DATA_ROW
    for key, value in pairs:
        put(ws, r, 1, key).font = font(bold=True)
        ws.cell(r, 1).alignment = Alignment(vertical="top")
        cell = put(ws, r, 2, value)
        cell.font = font()
        cell.alignment = Alignment(vertical="top", horizontal="left")
        r += 1
    header(hdr, ("시트", "그룹 · 열", "정의"))
    r = hdr + 1
    for sheet, group, label, definition in dictionary:
        put(ws, r, 1, sheet).font = font()
        put(ws, r, 2, f"{group} · {label}").font = font()
        cell = put(ws, r, 3, definition)
        cell.font = font()
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        # 정의가 접히는 줄 수만큼만 높인다(자동 높이는 기본 글꼴 기준이라 헐겁다 — tighten_rows)
        lines = math.ceil(text_width(str(definition or "")) / (def_width - 1))
        ws.row_dimensions[r].height = ROW_HEIGHT * max(1, lines)
        r += 1
    for sheet, name, definition in [*((s, f"{g} · {lb}", d) for s, g, lb, d in tail), *extra]:
        for c, v in ((1, sheet), (2, name), (3, definition)):
            put(ws, r, c, v).font = font()
        ws.cell(r, 1).alignment = Alignment(vertical="top")
        ws.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(r, 3).alignment = Alignment(vertical="top", wrap_text=True)
        # 폭 110 · 34 에서 실제로 끊기는 줄 수(여백 3 · 2)
        ws.row_dimensions[r].height = ROW_HEIGHT * max(text_lines(str(definition), def_width - 3),
                                                      text_lines(name, 34.0 - 2))
        r += 1
    # 여러 줄 행은 시트 · 열 이름도 위로 맞춘다(설명 첫 줄과 나란히)
    for row in range(hdr + 1, r):
        if (ws.row_dimensions[row].height or ROW_HEIGHT) > ROW_HEIGHT:
            for c in (1, 2):
                a = ws.cell(row, c).alignment
                ws.cell(row, c).alignment = Alignment(horizontal=a.horizontal, vertical="top",
                                                      wrap_text=a.wrap_text)
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, 1)
    tighten_rows(ws)
    page_setup(ws, Page(), f"A1:C{r - 1}")
    return ws


# ── 셀 안 꺾은선(엑셀 스파크라인) ───────────────────────────────────────────
# openpyxl 은 스파크라인을 읽지도 쓰지도 못한다. 저장한 xlsx 의 시트 XML 끝 extLst 에 x14
# sparklineGroups 를 넣는다(엑셀 2010+ 확장. LibreOffice 도 읽는다 — 10-06 왕복 확인).
_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
_SPARK_URI = "{05C60535-1F16-4fd2-B633-F4F36F0B64E0}"
# 선 두께(pt) — 엑셀이 고를 수 있는 가장 얇은 값. 1pt 는 11.25pt 행에서 굵어 보였다(10-06 지적).
SPARK_LINE_WEIGHT = 0.25


@dataclass(frozen=True)
class SparkGroup:
    """같은 선 색의 스파크라인 묶음. cells = (원자료 범위, 그릴 셀) — 예 ("'시트'!B8:W8", "M8")."""

    color: str
    cells: tuple[tuple[str, str], ...]


def sparkline_xml(groups: Sequence[SparkGroup]) -> str:
    body = []
    for g in groups:
        c = f'rgb="FF{g.color}"'
        colors = "".join(f"<x14:{k} {c}/>" for k in (
            "colorSeries", "colorNegative", "colorAxis", "colorMarkers", "colorFirst",
            "colorLast", "colorHigh", "colorLow"))
        lines = "".join(f"<x14:sparkline><xm:f>{escape(ref)}</xm:f><xm:sqref>{escape(cell)}"
                        "</xm:sqref></x14:sparkline>" for ref, cell in g.cells)
        body.append('<x14:sparklineGroup displayEmptyCellsAs="gap" '
                    f'lineWeight="{SPARK_LINE_WEIGHT}">'
                    f"{colors}<x14:sparklines>{lines}</x14:sparklines></x14:sparklineGroup>")
    return (f'<ext uri="{_SPARK_URI}" '
            'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
            '<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
            + "".join(body) + "</x14:sparklineGroups></ext>")


def _sheet_part(z: zipfile.ZipFile, title: str) -> str:
    wb = ElementTree.fromstring(z.read("xl/workbook.xml"))
    rid = next((s.get(f"{{{_NS_REL}}}id") for s in wb.iter(f"{{{_NS_MAIN}}}sheet")
                if s.get("name") == title), None)
    rels = ElementTree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = next((r.get("Target") for r in rels.iter(f"{{{_NS_PKG}}}Relationship")
                   if rid is not None and r.get("Id") == rid), None)
    if target is None:
        raise ValueError(f"시트 {title!r} 를 xlsx 에서 찾지 못했다")
    return target.lstrip("/") if target.startswith("/") else f"xl/{target}"


def add_sparklines(path: Path, sheet: str, groups: Sequence[SparkGroup]) -> int:
    """저장된 xlsx 의 `sheet` 에 스파크라인을 넣고 그린 셀 수를 돌려준다(파일은 원자적 교체)."""
    groups = [g for g in groups if g.cells]
    if not groups:
        return 0
    _add_ext(Path(path), sheet, sparkline_xml(groups))
    return sum(len(g.cells) for g in groups)


def add_column_sparklines(path: Path, sheet: str, cells: Sequence[tuple[str, str]]) -> int:
    """셀 안 막대(세로 막대 스파크라인) — 위 = 초록 · 아래 = 빨강, 원자료가 숨긴 열이라
    `displayHidden` 을 켠다(안 켜면 엑셀에서 빈 칸 — 샘플 24~25 결함). cells = (원자료, 그릴 칸)."""
    if not cells:
        return 0
    items = "".join(f"<x14:sparkline><xm:f>{escape(ref)}</xm:f><xm:sqref>{escape(cell)}"
                    "</xm:sqref></x14:sparkline>" for ref, cell in cells)
    dark = "FF1E8C45"
    ext = (f'<ext uri="{_SPARK_URI}" '
           'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
           '<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
           '<x14:sparklineGroup type="column" displayEmptyCellsAs="gap" negative="1" '
           'displayHidden="1">'
           f'<x14:colorSeries rgb="FF{GREEN}"/><x14:colorNegative rgb="FF{RED}"/>'
           '<x14:colorAxis rgb="FF000000"/>'
           + "".join(f'<x14:{k} rgb="{dark}"/>' for k in (
               "colorMarkers", "colorFirst", "colorLast", "colorHigh", "colorLow"))
           + f"<x14:sparklines>{items}</x14:sparklines></x14:sparklineGroup>"
           "</x14:sparklineGroups></ext>")
    _add_ext(Path(path), sheet, ext)
    return len(cells)


def _rewrite(path: Path, suffix: str, edit: Callable[[str, bytes], bytes]) -> None:
    """저장된 xlsx 의 파트마다 edit(이름, 내용)을 거쳐 원자적으로 바꾼다. 도중에 실패하면 임시 파일을
    지우고 원본은 그대로 둔다(손질 전 엑셀도 그대로 열린다)."""
    tmp = path.with_name(f".{path.name}.{suffix}.tmp")
    try:
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, edit(item.filename, zin.read(item.filename)))
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def _add_ext(path: Path, sheet: str, ext: str) -> None:
    """저장된 xlsx 의 `sheet` XML 끝 extLst 에 ext 하나를 넣는다(파일은 원자적 교체)."""
    with zipfile.ZipFile(path) as z:
        part = _sheet_part(z, sheet)

    def edit(name: str, data: bytes) -> bytes:
        if name != part:
            return data
        xml = data.decode("utf-8")
        if "</extLst>" in xml:
            xml = xml.replace("</extLst>", ext + "</extLst>", 1)
        else:
            xml = xml.replace("</worksheet>", f"<extLst>{ext}</extLst></worksheet>", 1)
        return xml.encode("utf-8")
    _rewrite(path, "ext", edit)


# ── 새 시트 문법(ETF · Analyst · Events · Market · Flow · Short — 확정 샘플 28 · 30) ─────────
# 1~5행 띠 · 6·7행 머리 · 본문 칸의 서식은 Sector 시트 칸에서 옮긴다(확정 샘플이 그렇게 만들었다 —
# 같은 엑셀 안에서 띠 · 머리 색이 시트마다 같다). 표마다 머리 색 · 테두리는 그 표의 마지막 데이터
# 열까지, 코드 없는 표는 A · B 를 이름 칸 하나로 합친다. 색은 빨강 · 노랑 · 초록 그라데이션
# (백분위 5 · 0 · 백분위 95). 틀 고정은 1~5행 띠와 B 열(C6) — 표를 위아래로 쌓은 시트라 표 하나의
# 머리를 고정하면 아래 표 열과 어긋난다(샘플 26 결함).
CHART_BLUE, CHART_BLACK = "0078C8", "000000"
CHART_W_CM, CHART_H_CM = 12.6, 5.6


@dataclass(frozen=True)
class XRefs:
    """새 시트가 옮겨 쓰는 서식(StyleArray) — 기준 시트(Sector)의 1~5행 칸 · A6 그룹 머리 ·
    A7 열 머리 · C8 글자 본문 · H8 숫자 본문."""

    band: Mapping[tuple[int, int], StyleArray]
    max_col: int
    group: StyleArray
    head: StyleArray
    text: StyleArray
    num: StyleArray

    @classmethod
    def of(cls, ws: Worksheet) -> XRefs:
        n = ws.max_column
        return cls({(r, c): copy(ws.cell(r, c)._style) for r in range(1, 6)
                    for c in range(1, n + 1)}, n, copy(ws["A6"]._style),
                   copy(ws["A7"]._style), copy(ws["C8"]._style), copy(ws["H8"]._style))


class XSheet:
    """새 시트 하나 — 1~5행 띠를 깔고 표 · 그래프를 얹는 도우미(샘플 28 · 30 의 st · header ·
    put · lab · close · cs · title_cell 그대로)."""

    def __init__(self, wb: Workbook, name: str, widths: Mapping[str, float], refs: XRefs,
                 *, ncol: int | None = None, index: int | None = None) -> None:
        self.ws = ws = wb.create_sheet(name, index)
        self.refs = refs
        ws.sheet_view.showGridLines = False
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        for r in range(1, 6):
            for c in range(1, (ncol or len(widths)) + 1):
                cell = ws.cell(r, c)
                cell._style = copy(refs.band[(r, min(c, refs.max_col))])
                cell.value = None
            ws.row_dimensions[r].height = ROW_HEIGHT

    def texts(self, texts: Mapping[str, str]) -> None:
        """1~5행 글 — A 열 밖 칸은 그 행 A 칸 서식."""
        for k, v in texts.items():
            self.ws[k] = v
            if k[0] != "A":
                self.ws[k]._style = copy(self.ws[f"A{k[1:]}"]._style)

    def hide(self, first: int, n: int, width: float) -> None:
        """그래프 · 셀 안 막대 원자료 열(숨김)."""
        for c in range(first, first + n):
            dim = self.ws.column_dimensions[get_column_letter(c)]
            dim.hidden, dim.width = True, width

    def header(self, r: int, groups: Sequence[tuple[str, str]], heads: Sequence[str],
               ab: bool = False) -> None:
        """r 행 그룹 · r+1 행 열 이름 — 칠은 이 표의 마지막 열까지. ab = A · B 를 이름 칸 하나로."""
        ws = self.ws
        for c in range(1, len(heads) + 1):
            ws.cell(r, c)._style = copy(self.refs.group)
            ws.cell(r, c).value = None
            ws.cell(r + 1, c)._style = copy(self.refs.head)
            ws.cell(r + 1, c).value = None
        for col, t in groups:
            ws[f"{col}{r}"] = t
        for i, t in enumerate(heads, 1):
            if t:
                ws.cell(r + 1, i).value = t
        if ab:
            ws.cell(r + 1, 1).value = heads[1]
            ws.cell(r + 1, 2).value = None
            for rr in (r, r + 1):
                ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
        ws.row_dimensions[r].height = ROW_HEIGHT
        ws.row_dimensions[r + 1].height = HEADER_HEIGHT

    def put(self, r: int, c: int, v: object, fmt: str | None = None, text: bool = False,
            bold: bool = False, indent: int = 0) -> None:
        cell = put(self.ws, r, c, v)                       # 값만('=' 글자도 수식이 아니게)
        cell._style = copy(self.refs.text if text else self.refs.num)
        if fmt:
            cell.number_format = fmt
        if bold or indent:
            cell.font = Font(name=cell.font.name, sz=cell.font.sz, b=bold)
            if indent:
                cell.alignment = Alignment(horizontal="left", indent=indent)

    def lab(self, r: int, v: object, bold: bool = False, indent: int = 0) -> None:
        """코드 없는 표의 이름 칸(A · B 합침)."""
        self.put(r, 1, v, text=True, bold=bold, indent=indent)
        self.ws.cell(r, 2)._style = copy(self.refs.text)
        self.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)

    def stock(self, r: int, code: str, name: object, big: object, rank: object,
              extra: Sequence[tuple[object, str | None, bool]]) -> None:
        """종목 열 A 코드 · B 이름 · C 대분류 · D 종합 순위 · E~ = extra[(값, 서식, 글자?)]."""
        self.put(r, 1, code, text=True)
        self.put(r, 2, name, text=True)
        self.put(r, 3, big, text=True)
        self.put(r, 4, rank, "#,##0")
        for i, (v, fmt, text) in enumerate(extra, 5):
            self.put(r, i, v, fmt, text)

    def close(self, r: int, last: int) -> None:
        """표 끝 선 — 마지막 행 아래쪽 medium."""
        for c in range(1, last + 1):
            b = copy(self.ws.cell(r, c).border)
            self.ws.cell(r, c).border = Border(left=b.left, right=b.right, top=b.top,
                                               bottom=SIDE_BLOCK)

    def cs(self, rng: str) -> None:
        """빨강 · 노랑 · 초록 그라데이션 — 백분위 5 · 0(가운데) · 백분위 95. 빈 범위는 건너뛴다."""
        a, b = rng.split(":")
        if int("".join(ch for ch in b if ch.isdigit())) < int("".join(ch for ch in a
                                                                      if ch.isdigit())):
            return
        self.ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="percentile", start_value=5, start_color="FF" + RED,
            mid_type="num", mid_value=0, mid_color="FF" + YELLOW,
            end_type="percentile", end_value=95, end_color="FF" + GREEN))

    def tight(self, first: int, stop: int) -> None:
        """그래프 자리 행 높이 11.25 — first ≤ 행 < stop."""
        for r in range(first, stop):
            self.ws.row_dimensions[r].height = ROW_HEIGHT

    def title(self, cell: str, text: str) -> None:
        """그래프 제목 칸(그래프 안 제목 대신 위 칸)."""
        self.ws[cell] = text
        self.ws[cell].font = Font(name=FONT_NAME, sz=8, b=True)

    def page(self, area: str) -> None:
        ws = self.ws
        ws.freeze_panes = "C6"
        ws.print_title_rows = "2:5"
        ws.print_area = area
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True


def x_note(text: str) -> Comment:
    """머리 칸 메모(새 시트) — 줄 300px · 기호 붙임 없이(샘플 그대로)."""
    return note(text, "설명", glue=False)


# ── 그래프(Q.Pack FICC 문법) — 계열 1~2개 · 파랑 0078C8 / 검정 1.5pt 실선 · 점 없음 · 제목은 위 칸 ·
# 격자선 · 테두리 · 채움 없음 · 축 가는 검정 · 글자 8pt 맑은 고딕(latin · ea · cs — 글자 깨짐 방지).
# 가로축은 날짜 축(눈금 간격 단위라 글자가 세로로 쪼개지거나 눕지 않는다), 대분류 이름 축은 글자
def chart_text() -> RichText:
    cp = CharacterProperties(sz=800, solidFill=CHART_BLACK, latin=DFont(typeface=FONT_NAME),
                             ea=DFont(typeface=FONT_NAME), cs=DFont(typeface=FONT_NAME))
    return RichText(bodyPr=RichTextProperties(rot=0, vert="horz"),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _noline() -> GraphicalProperties:
    return GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))


def _plain(ch: BarChart | LineChart) -> None:
    """채움 · 테두리 없음, 숨긴 원자료도 그림('숨긴 데이터 표시' — 안 켜면 엑셀에서 빈 그래프), 크기."""
    ch.roundedCorners = False
    ch.graphical_properties = _noline()
    ch.plot_area.graphicalProperties = _noline()
    ch.visible_cells_only = False
    ch.width = CHART_W_CM  # pyright: ignore[reportAttributeAccessIssue]  # reason: openpyxl stub 이 width 를 int 로 적었다(실제 cm 실수)
    ch.height = CHART_H_CM


def _axes(ch: BarChart | LineChart) -> None:
    for a in (ch.x_axis, ch.y_axis):
        a.delete = False
        a.majorGridlines = None
        a.majorTickMark = "out"
        a.minorTickMark = "none"
        a.txPr = chart_text()
        a.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=CHART_BLACK,
                                                                      w=3175))


TimeUnit = Literal["days", "months", "years"]


def date_frame(ch: BarChart | LineChart, yfmt: str, xfmt: str, base: TimeUnit, major: int,
               unit: TimeUnit) -> None:
    """날짜 가로축 그래프 틀 — 축 칸은 Q.Pack(엑셀 저장본) 날짜 축과 같게(아래 · autoZero · auto ·
    lblOffset 100). 원자료가 숨긴 열이라 '숨긴 데이터도 표시'."""
    ax = DateAxis(crossAx=100, axPos="b", crosses="autoZero", auto=True, lblOffset=100,
                  baseTimeUnit=base, majorUnit=major, majorTimeUnit=unit)
    ax.number_format = xfmt
    vars(ch)["x_axis"] = ax                 # openpyxl 은 막대그래프에 날짜 축 대입을 막아서 직접 넣는다
    ch.y_axis.crossAx = 500
    _axes(ch)
    ch.x_axis.tickLblPos = "low"
    ch.y_axis.number_format = yfmt
    if ch.legend is not None:
        ch.legend.txPr = chart_text()
    _plain(ch)


def cat_frame(ch: BarChart | LineChart, yfmt: str) -> None:
    """글자(대분류 이름) 가로축 그래프 틀 — 범례 없음."""
    _axes(ch)
    ch.x_axis.tickLblPos = "low"
    ch.y_axis.number_format = yfmt
    ch.legend = None
    _plain(ch)


def line_series(s: Series, color: str) -> None:
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 19050
    if s.marker is None:                    # 꺾은선 계열은 늘 표식이 있다 — 형 좁히기
        s.marker = Marker()
    s.marker.symbol = "none"
    s.smooth = False


def bar_series(s: Series, color: str) -> None:
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.solidFill = color


_REF = re.compile(r"'([^']+)'!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?")


def finish_charts(path: Path, wb: Workbook) -> int:
    """저장된 xlsx 의 그래프 XML 손질 — 제목 자동 생성 끔(autoTitleDeleted) · 빈 글 조각 제거 ·
    날짜 축 `auto` · 값 사본(cache: 엑셀 밖 미리보기는 셀을 다시 읽지 않고 이것만 그린다, 날짜는 엑셀
    일련번호 · 대분류 이름 축은 글자 strRef). 값은 `wb` 의 칸에서 읽는다. 손질한 그래프 수."""
    def vals(ref: str) -> list[object]:
        m = _REF.fullmatch(ref)
        if m is None:
            raise ValueError(f"그래프 범위를 읽지 못했다: {ref}")
        w = wb[m[1]]
        c0, r0 = column_index_from_string(m[2]), int(m[3])
        c1, r1 = (column_index_from_string(m[4]), int(m[5])) if m[4] else (c0, r0)
        return [w.cell(r, c).value for c in range(c0, c1 + 1) for r in range(r0, r1 + 1)]

    def cache(kind: str, ref: str) -> str:
        v = vals(ref)
        if kind == "date":
            v = [(x - date(1899, 12, 30)).days if isinstance(x, date) else x for x in v]
        pts = "".join(f'<pt idx="{i}"><v>{escape(str(x))}</v></pt>'
                      for i, x in enumerate(v) if x is not None)
        if kind == "str":
            return f'<strCache><ptCount val="{len(v)}" />{pts}</strCache>'
        fc = "yyyy\\-mm\\-dd" if kind == "date" else "General"
        return f'<numCache><formatCode>{fc}</formatCode><ptCount val="{len(v)}" />{pts}</numCache>'

    def cat(m: re.Match[str]) -> str:
        if any(isinstance(x, str) for x in vals(m[1])):
            return f"<cat><strRef><f>{m[1]}</f>{cache('str', m[1])}</strRef></cat>"
        return f"<cat><numRef><f>{m[1]}</f>{cache('date', m[1])}</numRef></cat>"

    n = 0

    def edit(name: str, data: bytes) -> bytes:
        nonlocal n
        if not (name.startswith("xl/charts/chart") and name.endswith(".xml")):
            return data
        s = data.decode("utf-8")        # openpyxl 은 차트를 기본 이름공간(접두사 없음)으로 쓴다
        if s.count("<chart><plotArea>") != 1 or "autoTitleDeleted" in s or "majorGridlines" in s:
            raise ValueError(f"그래프 XML 모양이 예상과 다르다(openpyxl 판 변경?): {name}")
        s = s.replace("<chart><plotArea>", '<chart><autoTitleDeleted val="1" /><plotArea>')
        s = s.replace("<a:r><a:t /></a:r>", "")
        s = s.replace('<crosses val="autoZero" /><lblOffset',
                      '<crosses val="autoZero" /><auto val="1" /><lblOffset')
        s = re.sub(r"<cat><numRef><f>(.*?)</f></numRef></cat>", cat, s)
        s = re.sub(r"<val><numRef><f>(.*?)</f></numRef></val>",
                   lambda m: f"<val><numRef><f>{m[1]}</f>{cache('num', m[1])}</numRef></val>", s)
        s = re.sub(r"<tx><strRef><f>(.*?)</f></strRef></tx>",
                   lambda m: f"<tx><strRef><f>{m[1]}</f>{cache('str', m[1])}</strRef></tx>", s)
        n += 1
        return s.encode("utf-8")
    _rewrite(Path(path), "chart", edit)
    return n


__all__ = [
    "C_BAND", "C_GOLD", "C_HDR", "C_KEY", "C_NEW", "C_STRIP", "C_STYLE", "FIRST_DATA_ROW",
    "FONT_NAME", "GREEN", "RED", "YELLOW", "Col", "Group", "Title", "clean", "fill", "font",
    "SparkGroup", "add_sparklines", "new_workbook", "put", "scale_high_good", "scale_rank",
    "scale_fixed_zero", "scale_zero_mid", "sparkline_xml", "title_block", "write_meta",
    "write_table",
    # 확정 양식
    "C_COLHDR", "C_CORE", "C_GROUP", "SIDE_BLOCK", "SIDE_HDR", "SIDE_INNER", "SIDE_PAIR",
    "SIGN_COLORS", "Band", "Page", "Sheet", "note", "note_lines", "note_px", "page_setup",
    "scale_zero_anchor", "text_lines", "text_width", "write_confirmed", "write_meta_confirmed",
    # 새 시트(ETF · Analyst · Events · Market · Flow · Short)
    "CHART_BLACK", "CHART_BLUE", "XRefs", "XSheet", "add_column_sparklines", "bar_series",
    "cat_frame", "chart_text", "date_frame", "finish_charts", "line_series", "x_note",
]
