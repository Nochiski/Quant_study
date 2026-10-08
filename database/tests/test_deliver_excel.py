"""전달층 엑셀(`src/deliver/`, 플랜 M2 W2 T2.5b · MODEL_EXCEL_SPEC) — 합성 판으로 검증.

합성 세계: ISO 2026-W39(월 09-21 ~ 금 09-25, 수 09-23 휴장 = 판 없음) + 지난주 목·금(09-17·18).
종목 45(대분류 G45 15 · G10 10 · G20 10 · G30 5 · G40 5). 금요일 순위는 G45 가 1~15위라
업종 상한 9 가 걸린다. model 판·factor_inputs 판은 고정 계약 경로에 parquet 로 쓴다.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import statistics
import zipfile
from collections.abc import Callable
from datetime import date
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from deliver import __main__ as cli
from deliver.common import cap_candidates, change_pct, quantile, winsorize, yoy
from deliver.excel_daily import build_daily, model_label
from deliver.excel_weekly import build_weekly, prev_week, week_days
from deliver.qpack import Col, fit_width
from deliver.reader import DeliverError, find_run
from deliver.view import load_day
from model import registry
from model.contracts import FI_TABLES, score_columns
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScale

W39 = ["2026-09-21", "2026-09-22", "2026-09-23", "2026-09-24", "2026-09-25"]
HOLIDAY = "2026-09-23"
W38 = ["2026-09-17", "2026-09-18"]
FRI, THU, MON, TUE = "2026-09-25", "2026-09-24", "2026-09-21", "2026-09-22"
PRIMARY = "v4_rank@0.1"
SPECS = ("v2_percentrank@1.0", "v3_zscore@1.0", "v4_rank@0.1", "v4_rank@0.2")
N = 45
SECTOR_NAMES = {"G45": "IT", "G10": "에너지", "G20": "산업재", "G30": "필수소비재", "G40": "금융"}
FORMULA_NAME = '=HYPERLINK("http://x")'
# fi_universe 추정기관수(최근 3개월 투자의견 증권사 수) — 비고 열 대조용. 나머지 종목은 5명.
ANALYSTS = {1: 0, 2: 2, 3: 3, 4: None}
TOKEN, CHAT = "123456:SECRET-TOKEN-abc", "-100987654321"
PA = {"VARCHAR": pa.string(), "DATE": pa.date32(), "BIGINT": pa.int64(),
      "INTEGER": pa.int32(), "DOUBLE": pa.float64(), "BOOLEAN": pa.bool_()}


def tick(i: int) -> str:
    return f"{i + 1:06d}"


def sector(i: int) -> str:
    return ("G45" if i < 15 else "G10" if i < 25 else "G20" if i < 35 else "G30" if i < 40
            else "G40")


BASE_ORDER = [*range(0, 40), 44]


def order_for(day: str) -> list[int]:
    o = list(BASE_ORDER)
    if day == MON:                      # 월: 38 이 1위(월만 ●)
        o.remove(38)
        o.insert(0, 38)
    if day == THU:                      # 목: 0·1 자리 바꿈 → 금 Δ순위 +1
        o[0], o[1] = o[1], o[0]
    if day in W38:                      # 지난주: 34·39 가 맨 앞
        o = [34, 39] + [x for x in o if x not in (34, 39)]
    return o


def excluded_for(day: str) -> dict[int, str]:
    ex = {40: "admin", 41: "adv20", 42: "pull_gate", 43: "insufficient_data"}
    if day == TUE:
        ex[44] = "pull_gate"            # 화: 44 제외 → ✕
    return ex


def ret_of(i: int) -> float:
    return (i % 7 - 3) / 100


# ── parquet 쓰기 ───────────────────────────────────────────────────────────────
def _write(path: Path, schema: pa.Schema, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = {f.name: [r.get(f.name) for r in rows] for f in schema}
    pq.write_table(pa.table(cols, schema=schema), path)


def _d(v: object) -> object:
    return date.fromisoformat(v) if isinstance(v, str) else v


def write_fi(root: Path, table: str, bid: str, rows: list[dict[str, object]]) -> None:
    c = FI_TABLES[table]
    schema = pa.schema([(col.name, PA[col.dtype]) for col in c.columns])
    fixed = [{k: (_d(v) if dict(zip(c.column_names, (x.dtype for x in c.columns),
                                    strict=True)).get(k) == "DATE" else v)
              for k, v in r.items()} for r in rows]
    _write(root / table / f"v={bid}" / "part0.parquet", schema, fixed)


def _score_schema(cols: tuple[str, ...]) -> pa.Schema:
    text = {"ticker", "stock_code", "score_date", "spec_id", "exclude_reason", "sector_l1",
            "sector_l2", "coverage_state"}
    fields = []
    for c in cols:
        if c in text or c.endswith("_flag"):
            fields.append((c, pa.string()))
        elif c in ("rank", "n_buckets_used"):
            fields.append((c, pa.int64()))
        elif c == "excluded":
            fields.append((c, pa.bool_()))
        else:
            fields.append((c, pa.float64()))
    return pa.schema(fields)


IND_SCHEMA = pa.schema([("ticker", pa.string()), ("score_date", pa.string()),
                        ("spec_id", pa.string()), ("key", pa.string()), ("bucket", pa.string()),
                        ("role", pa.string()), ("raw", pa.float64()), ("pct", pa.float64()),
                        ("flag", pa.string())])


# ── 합성 판 ────────────────────────────────────────────────────────────────────
def _raw(i: int, key: str) -> tuple[float | None, str | None]:
    """지표 원값·flag — 표식·결측 사례를 종목에 심는다."""
    if key.startswith("REV_") and i == 3:
        return None, "커버리지<3"
    if key == "REV_OP_1M" and i == 1:
        return 1.5, "흑전"
    if key == "REV_OP_1M" and i == 2:
        return -0.5, "적전"
    if key == "EP":
        return (i - 5) * 0.01, None
    if key == "DY0" and i % 4 == 0:
        return 0.0, "무배당"
    if key == "R1M" and i == 4:
        return 50.0, None               # 5000% — 윈저라이즈 대상
    if key == "OPM_TTM" and i == 6:
        return None, "원천없음"
    return 0.01 * ((i * 7 + len(key)) % 29) - 0.05, None


def write_model_day(model_root: Path, day: str, fi_bid: str, *, status: str = "ok") -> str:
    bid = f"m_{day.replace('-', '')}T000000Z"
    order, excl = order_for(day), excluded_for(day)
    ranked = [i for i in order if i not in excl]
    rank = {i: k + 1 for k, i in enumerate(ranked)}
    specs: dict[str, dict[str, object]] = {}
    for sid in SPECS:
        spec = registry.get(sid)
        cols = score_columns(spec)
        rows: list[dict[str, object]] = []
        inds: list[dict[str, object]] = []
        if spec.engine == "v4_rank":
            buckets = list(spec.buckets)
            local = dict(rank)
            if sid == "v4_rank@0.2":            # 동일가중판은 0·2 순위를 바꾼다
                local[0], local[2] = rank[2], rank[0]
            for i in range(N):
                reason = excl.get(i)
                d13 = reason in ("admin", "adv20")
                r = None if reason else local.get(i)
                comp = None if d13 else 100.0 - 2.0 * (local.get(i) or 44)
                bs: dict[str, float | None] = {}
                for k, b in enumerate(buckets):
                    bs[b] = None if d13 or comp is None else max(0.0, min(100.0, comp + k - 2.5))
                if i == 3:
                    bs["revision"] = None
                if i == 43:
                    bs = {b: (v if b in ("low_risk", "value") else None) for b, v in bs.items()}
                rows.append({"ticker": tick(i), "score_date": day, "spec_id": sid, "rank": r,
                             "composite": comp, "excluded": reason is not None,
                             "exclude_reason": reason, "sector_l1": sector(i),
                             "sector_l2": sector(i) + "10",
                             "coverage_state": "grace" if i == 5 else "fresh",
                             "n_buckets_used": sum(v is not None for v in bs.values()),
                             **{f"{b}_score": v for b, v in bs.items()}})
                if d13:
                    continue
                for ind in spec.indicators:
                    raw, flag = _raw(i, ind.key)
                    inds.append({"ticker": tick(i), "score_date": day, "spec_id": sid,
                                 "key": ind.key, "bucket": ind.bucket, "role": ind.role,
                                 "raw": raw, "pct": None, "flag": flag})
            for ind in spec.indicators:          # 점수 지표 백분위 = 원값 순위(합성)
                if ind.role != "score":
                    continue
                mine = sorted((r for r in inds if r["key"] == ind.key and r["raw"] is not None),
                              key=lambda r: (r["raw"], r["ticker"]))
                for k, r in enumerate(mine):
                    r["pct"] = 100.0 * k / (len(mine) - 1)
        else:
            key = "stock_code"
            comp_col = "composite_score" if spec.engine == "v3_zscore" else "total_score"
            universe = [*reversed(range(40)), 45] if spec.engine == "v3_zscore" else BASE_ORDER
            for k, i in enumerate(universe, start=1):
                rows.append({key: tick(i), "score_date": day, "rank": k,
                             comp_col: 100.0 - k})
        d = model_root / sid / f"v={bid}"
        _write(d / "scores.parquet", _score_schema(cols), rows)
        if inds:
            _write(d / "indicators.parquet", IND_SCHEMA, inds)
        n_ranked = sum(1 for r in rows if r.get("rank") is not None)
        specs[sid] = {"n_scores": len(rows), "n_ranked": n_ranked,
                      "n_excluded": len(rows) - n_ranked,
                      "gates": {"MG0": {"status": "pass"}, "MG1": "pass"}}
    meta = {"layer": "model", "status": status, "build_id": bid, "date": day,
            "basis": "morning", "fi_build_id": fi_bid, "generated_at": f"{day}T00:00:00Z",
            "specs": specs, "primary_spec": PRIMARY}
    runs = model_root / "_runs"
    runs.mkdir(parents=True, exist_ok=True)
    (runs / f"{day.replace('-', '')}_morning.json").write_text(
        json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    if status == "ok":
        (model_root / "latest_morning.json").write_text(json.dumps(meta), encoding="utf-8")
    return bid


def write_fi_day(fi_root: Path, day: str,
                 edit: Callable[[str, list[dict[str, object]]], None] | None = None) -> str:
    """합성 fi 판 하나. `edit(표 이름, 행 목록)` 으로 쓰기 직전에 표마다 행을 고칠 수 있다."""
    bid = f"m_{day.replace('-', '')}T010000Z"
    uni, prices, adj, flows, fins, cons = [], [], [], [], [], []
    for i in range(N + 1):
        t = tick(i)
        uni.append({"ticker": t, "date": day, "name": FORMULA_NAME if i == 7 else f"종목{i}",
                    "market": "KOSPI" if i % 2 else "KOSDAQ", "sec_type": "common",
                    "shares": 1_000_000, "market_cap": 1000.0 + 100 * i, "mktcap_basis": "krx",
                    "sector_l1": sector(i), "sector_l1_name": SECTOR_NAMES[sector(i)],
                    "sector_l2": sector(i) + "10",
                    "sector_l2_name": SECTOR_NAMES[sector(i)] + " 중분류",
                    "has_estimates": i != N, "coverage_state": "grace" if i == 5 else "fresh",
                    "coverage_age_days": 2 if i == 5 else 0, "n_analysts": ANALYSTS.get(i, 5), "adv20": 50.0,
                    "is_admin": i == 40, "is_halted": False, "audit_adverse": False,
                    "filing_late": False, "eligible": i != N, "exclude_reason": None})
        for dd, px, ok in ((W38[1], 100.0, True),
                           (day, 100.0 * (1 + ret_of(i)), i != 35)):
            prices.append({"ticker": t, "date": dd, "close": int(10_000 + i * 100),
                           "price_source": "krx"})
            adj.append({"ticker": t, "date": dd, "adj_close": px, "adj_factor": 1.0,
                        "adj_ok": ok})
        if i != 10:
            flows.append({"ticker": t, "date": day, "foreign_investor": 1.0})
        op24, op25, ni24, ni25 = 100.0 + i, 120.0 + i, 50.0, 60.0
        if i == 1:
            op24, op25 = -10.0, 20.0
        if i == 2:
            ni24, ni25 = 30.0, -5.0
        if i == 3:
            op24, op25 = -10.0, -20.0
        for per, rev, op, ni in (("2024/12", 1000.0 + 10 * i, op24, ni24),
                                 ("2025/12", 1100.0 + 10 * i, op25, ni25)):
            fins.append({"ticker": t, "period": per, "period_type": "annual", "revenue": rev,
                         "op": op, "ni": ni, "roe": 5.0 + i, "roa": 2.0, "debt_ratio": 80.0,
                         "available_date": "2026-03-15" if per == "2025/12" else "2025-03-15"})
        for k, per in enumerate(("2025/06", "2025/09", "2025/12", "2026/03", "2026/06")):
            op = 10.0 + k
            if i == 1:
                op = -5.0 if k == 0 else 8.0
            fins.append({"ticker": t, "period": per, "period_type": "quarter",
                         "revenue": 250.0, "op": op, "available_date": "2026-08-14"})
        for per, h, eps, ni in (("2026/12", "cur", 1200.0, 70.0), ("2026/12", "1m", 1000.0, 65.0),
                                ("2027/12", "cur", 1400.0, 80.0)):
            if i == 2 and h == "cur" and per == "2026/12":
                ni = -3.0
            if i == 4 and h == "1m":
                eps = -100.0
            cons.append({"ticker": t, "target_period": per, "horizon": h,
                         "revenue": 1200.0 + 10 * i, "op": 130.0 + i, "ni": ni, "eps": eps,
                         "per": 12.5, "pbr": 1.1, "n_analysts": 5, "obs_date": "2026-09-24",
                         "fetched_date": "2026-09-24"})
    for table, rows in (("fi_universe", uni), ("fi_prices", prices), ("fi_adj_prices", adj),
                        ("fi_flows", flows), ("fi_fin_summary", fins), ("fi_consensus", cons),
                        ("fi_credit", []), ("fi_consensus_annual", [])):
        if edit is not None:
            edit(table, rows)
        write_fi(fi_root, table, bid, rows)
    return bid


@pytest.fixture(scope="module", autouse=True)
def _pin_ql_home(tmp_path_factory: pytest.TempPathFactory):
    """QL_HOME 을 빈 임시 폴더로 고정 — 개발 머신·서버의 실제 QL_HOME 이 테스트에 새지 않게."""
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("QL_HOME", str(tmp_path_factory.mktemp("ql_home")))
        yield


@pytest.fixture(scope="module")
def world(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Path]:
    root = tmp_path_factory.mktemp("deliver")
    model_root, fi_root, out_root = root / "model", root / "fi", root / "deliver"
    for day in [*W38, *[d for d in W39 if d != HOLIDAY]]:
        write_model_day(model_root, day, write_fi_day(fi_root, day))
    return {"model": model_root, "fi": fi_root, "out": out_root}


@pytest.fixture(scope="module")
def daily(world: dict[str, Path]):
    res = build_daily(FRI, "morning", model_root=world["model"], fi_root=world["fi"],
                      out_root=world["out"])
    return res, load_workbook(res.path)


@pytest.fixture(scope="module")
def weekly(world: dict[str, Path]):
    res = build_weekly("2026-W39", model_root=world["model"], fi_root=world["fi"],
                       out_root=world["out"])
    return res, load_workbook(res.path)


def test_normal_style_font_is_arial_10_like_qpack(daily, weekly) -> None:
    """엑셀은 행 번호를 표준 스타일 글꼴로 그린다 — openpyxl 기본 Calibri 11 이면
    11.25pt 행에서 잘린다. 큐팩 원본처럼 글꼴 0번과 표준 스타일이 Arial 10 이어야
    한다(10-06 사용자 지적)."""
    for _, book in (daily, weekly):
        normal = book._named_styles["Normal"].font
        assert (normal.name, normal.sz) == ("Arial", 10.0)
        assert (book._fonts[0].name, book._fonts[0].sz) == ("Arial", 10.0)
        first = book["점수" if "점수" in book.sheetnames else book.sheetnames[0]]
        # 데이터 칸은 그대로 맑은 고딕 8
        assert {c.font.name for c in first[8] if c.value is not None} == {"맑은 고딕"}


def header(ws) -> dict[str, int]:
    """7행 열 이름 → 열 번호(줄바꿈은 공백으로)."""
    return {str(c.value).replace("\n", " "): c.column for c in ws[7] if c.value is not None}


def rows_by_code(ws) -> dict[str, int]:
    return {str(ws.cell(r, 1).value): r for r in range(8, ws.max_row + 1)}


def col_of(ws, label: str, occurrence: int = 0) -> int:
    hits = [c.column for c in ws[7] if str(c.value).replace("\n", " ") == label]
    return hits[occurrence]


# ── 순수 함수 ──────────────────────────────────────────────────────────────────
def test_yoy_sign_flags_replace_growth() -> None:
    assert yoy(20.0, -10.0) == (None, "흑자전환")
    assert yoy(-5.0, 30.0) == (None, "적자전환")
    assert yoy(-20.0, -10.0) == (None, "적자지속")
    assert yoy(5.0, 0.0) == (None, "흑자전환")              # 이전값 0 도 부호 전환
    assert yoy(-5.0, 0.0) == (None, "적자전환")
    assert yoy(110.0, 100.0) == (pytest.approx(10.0), None)  # 흑자지속은 숫자
    assert yoy(None, 1.0) == (None, None)


def test_change_pct_follows_v3_revision_convention() -> None:
    assert change_pct(10.0, -5.0) == (None, "흑전")          # 이전값<0 → 표식만
    assert change_pct(-20.0, -10.0) == (None, "적확")
    assert change_pct(-5.0, -10.0) == (None, "적축")
    v, f = change_pct(-5.0, 10.0)                            # 적전 = 변화율 + 표식
    assert f == "적전" and v == pytest.approx(-150.0)
    assert change_pct(12.0, 10.0) == (pytest.approx(20.0), None)


def test_winsorize_clips_to_1_99_and_keeps_text() -> None:
    vals: list[object] = [float(x) for x in range(100)] + [10_000.0, "결측(x)", None]
    out = winsorize(vals)
    nums = sorted(float(v) for v in vals if isinstance(v, float))
    hi = quantile(nums, 0.99)
    assert max(v for v in out if isinstance(v, float)) == pytest.approx(hi)
    assert min(v for v in out if isinstance(v, float)) == pytest.approx(quantile(nums, 0.01))
    assert out[-2:] == ["결측(x)", None]


def test_cap_candidates_limits_per_sector() -> None:
    rows = [{"t": i, "s": "A" if i < 12 else f"S{i % 4}"} for i in range(60)]
    got = cap_candidates(rows, 30, 9, lambda r: r["s"])
    assert len(got) == 30
    assert sum(1 for r in got if r["s"] == "A") == 9
    assert [r["t"] for r in got[:10]] == [*range(9), 12]         # 10~11번째 A 는 건너뛴다
    assert len(cap_candidates(rows[:12], 30, 9, lambda r: r["s"])) == 9


def test_week_helpers() -> None:
    assert week_days("2026-W39") == W39
    assert prev_week("2026-W39") == "2026-W38"
    assert prev_week("2027-W01") == "2026-W53"
    with pytest.raises(DeliverError):
        week_days("2026-39")


# ── 매일 엑셀 ─────────────────────────────────────────────────────────────────
def test_daily_sheets_and_row_counts(daily) -> None:
    res, wb = daily
    assert res.path.name == "model_scores_20260925_morning.xlsx"
    assert res.path.parent.name == "daily"
    assert wb.sheetnames == [
        "점수", "점수 원자료", "지표(표시용)", "실적", "업종", "모델 비교", "메타", "순위 흐름"]
    assert wb["순위 흐름"].sheet_state == "hidden"                 # 그래프 원자료(숨김)
    assert res.sheets == wb.sheetnames[:-1]                         # 보이는 시트만 알린다
    for name in ("점수", "점수 원자료", "지표(표시용)", "실적", "모델 비교"):
        assert wb[name].max_row == 7 + N, name              # 주 모델 모집단 45행
    assert wb["업종"].max_row == 7 + 10                       # 대분류 5 + 중분류 5
    assert res.n_rows == N and res.n_ranked == 41
    assert [t for _, t, _ in res.top] == [tick(0), tick(1), tick(2), tick(3), tick(4)]


def test_daily_scores_sheet_contents(daily) -> None:
    _, wb = daily
    ws = wb["점수"]
    h = header(ws)
    for label in ("코드", "이름", "시장", "대분류", "중분류", "시총(억)", "거래대금 20일(억)",
                  "주가",
                  "순위", "종합 점수", "전일 순위", "Δ순위 1M", "1M 흐름",
                  "제외 사유", "커버리지",
                  "저위험 유니버스", "저위험 업종", "리비전 유니버스", "결측 축", "최대 차이"):
        assert label in h, label
    rows = rows_by_code(ws)
    assert ws.cell(8, 1).value == tick(0) and ws.cell(8, h["순위"]).value == 1
    r0 = rows[tick(0)]
    assert ws.cell(r0, h["전일 순위"]).value == 2                      # 1일·1W Δ 없음(N-16)
    assert not {"Δ순위", "Δ순위 1W", "1W 흐름"} & set(h)
    assert ws.cell(r0, h["Δ순위 1M"]).value is None                 # 1M 비교 판이 없다
    assert ws.cell(r0, h["1M 흐름"]).value is None                  # 그래프 칸은 값이 없다
    assert ws.cell(rows[tick(40)], h["제외 사유"]).value == "관리종목(admin)"
    assert ws.cell(rows[tick(42)], h["제외 사유"]).value == "고점근접+반전 게이트(pull_gate)"
    assert ws.cell(rows[tick(40)], h["순위"]).value is None
    assert ws.cell(rows[tick(5)], h["커버리지"]).value == "유예 D+2"
    r3 = rows[tick(3)]
    assert ws.cell(r3, h["리비전 유니버스"]).value == "결측(커버리지<3)"
    assert ws.cell(r3, h["결측 축"]).value == "리비전"
    assert isinstance(ws.cell(r3, h["저위험 유니버스"]).value, float)
    # 제외 순서: 순위 종목 뒤에 코드순
    tail = [ws.cell(r, 1).value for r in range(8 + 41, 8 + N)]
    assert tail == sorted(tail) == [tick(i) for i in (40, 41, 42, 43)]
    # 다른 모델 순위 — v3 는 역순 유니버스
    v3 = col_of(ws, "v3 원본")                       # 헤더는 모델 이름만(spec id 는 정의·메타)
    assert ws.cell(r0, v3).value == 40


def test_daily_scores_sheet_note_flags_thin_coverage(daily) -> None:
    _, wb = daily
    ws = wb["점수"]
    h = header(ws)
    assert h["비고"] == h["커버리지"] + 1
    rows = rows_by_code(ws)
    note = {i: ws.cell(rows[tick(i)], h["비고"]).value for i in range(5)}
    assert note == {0: None, 1: "최근 3개월 의견 없음", 2: "애널리스트 3명 이하",
                    3: "애널리스트 3명 이하", 4: "애널리스트 수 미상"}


def test_daily_raw_sheet_flags_and_definitions(daily) -> None:
    _, wb = daily
    ws = wb["점수 원자료"]
    h = header(ws)
    rows = rows_by_code(ws)
    v, f = h["영업이익 추정 1M(%)"], h["REV_OP_1M 표식"]
    assert ws.cell(rows[tick(1)], v).value is None                      # 흑전 → 표식만(E)
    assert ws.cell(rows[tick(1)], f).value == "흑전"
    assert ws.cell(rows[tick(2)], v).value == pytest.approx(-50.0)      # 적전 → 값 + 표식
    assert ws.cell(rows[tick(2)], f).value == "적전"
    assert ws.cell(rows[tick(3)], v).value == "결측(커버리지<3)"
    assert ws.cell(rows[tick(0)], h["EP 표식"]).value == "적자"
    assert ws.cell(rows[tick(0)], h["DY0 표식"]).value == "무배당"
    assert ws.cell(rows[tick(0)], h["배당수익률(%)"]).value == 0          # 무배당은 0(결측 아님)
    assert isinstance(ws.cell(rows[tick(9)], h["VOL60 백분위"]).value, float)
    spec = registry.get(PRIMARY)
    vol_def = next(i.definition for i in spec.indicators if i.key == "VOL60")
    assert vol_def in str(ws.cell(5, h["60일 변동성(%)"]).value)
    assert ws.cell(rows[tick(0)], h["연간 기준기"]).value == "2025/12"
    assert ws.cell(rows[tick(0)], h["재무 접수일"]).value == "2026-03-15"
    assert ws.cell(rows[tick(10)], h["수급 최신일"]).value is None
    assert ws.cell(rows[tick(0)], h["수급 최신일"]).value == FRI
    # 표시 지표는 이 시트에 없다(중복 표기 금지)
    assert "1M 수익률(%)" not in h


def test_daily_display_sheet_is_winsorized(daily) -> None:
    _, wb = daily
    ws = wb["지표(표시용)"]
    h = header(ws)
    rows = rows_by_code(ws)
    c = h["1M 수익률(%)"]
    vals = [ws.cell(r, c).value for r in range(8, ws.max_row + 1)]
    nums = [v for v in vals if isinstance(v, float)]
    raw = sorted(_raw(i, "R1M")[0] * 100 for i in range(N) if i not in (40, 41))
    assert max(nums) == pytest.approx(quantile(raw, 0.99))
    assert max(nums) < 5000.0
    assert min(nums) == pytest.approx(quantile(raw, 0.01))
    assert ws.cell(rows[tick(2)], h["선행 PER(배)"]).value == "적자"
    assert ws.cell(rows[tick(4)], h["EPS 1M 표식"]).value == "흑전"
    assert ws.cell(rows[tick(4)], h["선행 EPS 1M 변화(%)"]).value is None
    assert "VOL60 백분위" not in h and "60일 변동성(%)" not in h


def test_daily_earnings_real_years_quarters_and_sign_words(daily) -> None:
    """N-19: 실적 시트 머리글은 실제 연도·분기다. 합성 판(09-25)은 10-02 판과 같은 구조다
    (Y = 2026, 확정 2025/12, 컨센서스 2026·2027, 최근 분기 2026/06).
    결산기·표식·최근 분기·Q0 공시일 열은 없고, 부호 전환은 y-y 칸 안의 글자다."""
    _, wb = daily
    ws = wb["실적"]
    labels = [str(c.value).replace("\n", " ") for c in ws[7] if c.value is not None]
    years = ["2025", "2025 y-y(%)", "2026E", "2026E y-y(%)", "2027E", "2027E y-y(%)"]
    assert labels == ["코드", "이름", "대분류", *years * 3,
                      "25.2Q", "25.3Q", "25.4Q", "26.1Q", "26.2Q", "최근 분기 y-y(%)"]
    groups = [str(c.value) for c in ws[6] if c.value is not None]
    assert groups == ["종목", "매출(억원)", "영업이익(억원)", "순이익(억원)", "분기 영업이익(억원)"]
    gone = ("표식", "공시일", "결산기", "FY-1", "FY0")
    assert not [s for s in labels + groups if any(w in s for w in gone)]
    rows = rows_by_code(ws)
    r10 = rows[tick(10)]
    rev = col_of(ws, "2025", 0)                       # 0 = 매출, 1 = 영업이익, 2 = 순이익
    assert ws.cell(r10, rev).value == 1200.0
    assert ws.cell(r10, rev).number_format == "#,##0"                   # 확정치 — 'E' 아님
    assert ws.cell(r10, rev + 1).value == pytest.approx((1200 / 1100 - 1) * 100)
    e0 = col_of(ws, "2026E", 0)
    assert ws.cell(r10, e0).value == 1300.0                              # 2026/12 컨센서스
    assert ws.cell(r10, e0 + 1).value == pytest.approx((1300 / 1200 - 1) * 100)
    op_y = col_of(ws, "2025 y-y(%)", 1)
    assert ws.cell(rows[tick(1)], op_y).value == "흑자전환"               # −10 → 20
    assert ws.cell(rows[tick(3)], op_y).value == "적자지속"               # −10 → −20
    assert ws.cell(rows[tick(2)], col_of(ws, "2025 y-y(%)", 2)).value == "적자전환"  # 순이익
    h = header(ws)
    assert [ws.cell(r10, h[q]).value for q in ("25.2Q", "25.3Q", "25.4Q", "26.1Q", "26.2Q")] == [
        10, 11, 12, 13, 14]
    assert ws.cell(r10, h["최근 분기 y-y(%)"]).value == pytest.approx((14 / 10 - 1) * 100)
    assert ws.cell(rows[tick(1)], h["최근 분기 y-y(%)"]).value == "흑자전환"   # −5 → 8


def _roll_to_2027(table: str, rows: list[dict[str, object]]) -> None:
    """2027-01-04 판 합성 — 2026 사업보고서는 0번 종목만 나왔고 나머지는 공시 전이다.
    컨센서스는 2026·2027·2028 결산기다(11번 종목은 2026 추정치도 없다).
    2026/09 분기는 0번 종목만 있다(분기 5기 창이라 0번의 2025/06 은 빠진다)."""
    if table == "fi_fin_summary":
        rows[:] = [r for r in rows if not (r["ticker"] == tick(0) and r["period"] == "2025/06")]
        rows += [{"ticker": tick(0), "period": "2026/12", "period_type": "annual",
                  "revenue": 1450.0, "op": 150.0, "ni": 75.0, "available_date": "2027-01-03"},
                 {"ticker": tick(0), "period": "2026/09", "period_type": "quarter",
                  "revenue": 250.0, "op": 20.0, "available_date": "2026-11-14"}]
    elif table == "fi_consensus":
        rows[:] = [r for r in rows
                   if not (r["ticker"] == tick(11) and r["target_period"] == "2026/12")]
        rows += [{**r, "target_period": "2028/12", "revenue": float(r["revenue"]) + 200}
                 for r in rows if r["target_period"] == "2027/12"]


def test_daily_earnings_headers_roll_over_with_the_board_year(tmp_path: Path) -> None:
    """N-19: 연간 머리글은 판 날짜의 연도 Y 로 `Y−1 · YE · (Y+1)E` 다 — 2027-01-04 판은
    `2026 · 2027E · 2028E`. 2026 확정치가 없으면 2026 추정치를 숫자 서식 `#,##0"E"`·회색 글자로
    그 칸만 표시한다(값은 숫자). 분기 창도 판 전체의 최근 분기(2026/09)를 따라 넘어간다."""
    day = "2027-01-04"
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, day, write_fi_day(fi_root, day, edit=_roll_to_2027))
    res = build_daily(day, "morning", model_root=model_root, fi_root=fi_root, out_root=tmp_path)
    ws = load_workbook(res.path)["실적"]
    labels = [str(c.value).replace("\n", " ") for c in ws[7] if c.value is not None]
    assert labels[3:9] == ["2026", "2026 y-y(%)", "2027E", "2027E y-y(%)", "2028E",
                           "2028E y-y(%)"]
    assert labels[-6:] == ["25.3Q", "25.4Q", "26.1Q", "26.2Q", "26.3Q", "최근 분기 y-y(%)"]
    assert "2025" not in labels
    rows = rows_by_code(ws)
    rev = col_of(ws, "2026", 0)
    est = ws.cell(rows[tick(10)], rev)                  # 2026 확정치 없음 → 2026 추정치
    assert est.value == 1300 and '"E"' in est.number_format
    assert est.font.color.rgb.endswith("7F7F7F")        # 회색 글자
    assert ws.cell(rows[tick(10)], rev + 1).value == pytest.approx((1300 / 1200 - 1) * 100)
    act = ws.cell(rows[tick(0)], rev)                   # 2026 확정치 있음 → 그 칸은 그대로
    assert act.value == 1450 and act.number_format == "#,##0"
    assert act.font.color is None or not act.font.color.rgb.endswith("7F7F7F")
    assert ws.cell(rows[tick(11)], rev).value is None   # 확정·추정 둘 다 없으면 빈칸
    # 열 전체 대조 — 'E' 칸은 0번(확정)·11번(빈칸)을 뺀 전 종목, 서식이 한 행 밀리면 깨진다
    e_rows = {code for code, r in rows.items() if '"E"' in ws.cell(r, rev).number_format}
    assert e_rows == set(rows) - {tick(0), tick(11)}
    assert ws.cell(rows[tick(10)], col_of(ws, "2028E", 0)).value == 1500
    h = header(ws)
    assert ws.cell(rows[tick(0)], h["26.3Q"]).value == 20
    assert ws.cell(rows[tick(10)], h["26.3Q"]).value is None          # 그 분기 행이 없다
    # 최근 분기 y-y 는 종목마다 자기 최근 분기 — 0번 2026/09 vs 2025/09, 10번 2026/06 vs 2025/06
    q_y = h["최근 분기 y-y(%)"]
    assert ws.cell(rows[tick(0)], q_y).value == pytest.approx((20 / 11 - 1) * 100)
    assert ws.cell(rows[tick(10)], q_y).value == pytest.approx((14 / 10 - 1) * 100)


def _off_calendar_quarters(table: str, rows: list[dict[str, object]]) -> None:
    """5번 종목을 11월 결산사처럼 — 분기 기말이 2025/08·11·2026/02·05·08(DART period_end 그대로)이고
    마지막 2026/08 은 12월 결산 종목들의 최근 분기(2026/06)보다 늦다."""
    if table == "fi_fin_summary":
        rows[:] = [r for r in rows
                   if not (r["ticker"] == tick(5) and r["period_type"] == "quarter")]
        rows += [{"ticker": tick(5), "period": per, "period_type": "quarter", "revenue": 250.0,
                  "op": 30.0 + k, "available_date": "2026-10-15"}
                 for k, per in enumerate(("2025/08", "2025/11", "2026/02", "2026/05", "2026/08"))]


def test_daily_earnings_quarter_window_ignores_off_calendar_quarters(tmp_path: Path) -> None:
    """분기 창 끝은 달력 분기(기말 03·06·09·12월)만 본다. 11월 결산 종목의 더 늦은 2026/08 분기가
    창을 끌고 가면 머리글이 '26.0Q' 로 깨지고 12월 결산 종목 칸이 전부 빈다(검토 재현).
    그 종목 분기는 달력 칸에 넣지 않고, 최근 분기 y-y 는 자기 최근 분기(2026/08 vs 2025/08)다."""
    day = "2026-10-20"
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, day, write_fi_day(fi_root, day, edit=_off_calendar_quarters))
    res = build_daily(day, "morning", model_root=model_root, fi_root=fi_root, out_root=tmp_path)
    ws = load_workbook(res.path)["실적"]
    quarters = ["25.2Q", "25.3Q", "25.4Q", "26.1Q", "26.2Q"]
    labels = [str(c.value).replace("\n", " ") for c in ws[7] if c.value is not None]
    assert labels[-6:] == [*quarters, "최근 분기 y-y(%)"]
    h, rows = header(ws), rows_by_code(ws)
    assert [ws.cell(rows[tick(10)], h[q]).value for q in quarters] == [10, 11, 12, 13, 14]
    assert [ws.cell(rows[tick(5)], h[q]).value for q in quarters] == [None] * 5
    assert ws.cell(rows[tick(5)], h["최근 분기 y-y(%)"]).value == pytest.approx((34 / 30 - 1) * 100)


def _september_fy(table: str, rows: list[dict[str, object]]) -> None:
    """2번 종목을 9월 결산사처럼 — 컨센서스 결산기 2026/09·2027/09·2028/09, 연간 확정 행 없음
    (fi 연간 행은 12월 결산만 싣는다, U22)."""
    if table == "fi_fin_summary":
        rows[:] = [r for r in rows
                   if not (r["ticker"] == tick(2) and r["period_type"] == "annual")]
    elif table == "fi_consensus":
        base = next(r for r in rows if r["ticker"] == tick(2) and r["horizon"] == "cur")
        rows[:] = [r for r in rows if r["ticker"] != tick(2)]
        rows += [{**base, "target_period": p, "revenue": rev, "op": rev / 10, "ni": rev / 20}
                 for p, rev in (("2026/09", 900.0), ("2027/09", 990.0), ("2028/09", 1100.0))]


def test_daily_earnings_non_december_fy_uses_the_year_the_fy_ends(tmp_path: Path) -> None:
    """비12월 결산(사용자 결정 10-06): 연간 칸 = 결산기가 끝나는 연도(증권사 FY 표기 관례) —
    9월 결산이면 2026E 칸 = 2026/09 결산기, y-y 는 한 해 앞 같은 결산월과 비교한다. 점수 시트
    비고엔 'N월 결산'을 기존 비고와 ' · ' 로 잇는다. 12월 결산 종목은 그대로다."""
    day = "2026-10-02"
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, day, write_fi_day(fi_root, day, edit=_september_fy))
    wb = load_workbook(build_daily(day, "morning", model_root=model_root, fi_root=fi_root,
                                   out_root=tmp_path).path)
    ws = wb["실적"]
    rows = rows_by_code(ws)
    r2, r10 = rows[tick(2)], rows[tick(10)]
    rev = {y: col_of(ws, y, 0) for y in ("2025", "2026E", "2027E")}
    assert ws.cell(r2, rev["2025"]).value is None                    # 2025/09 확정·추정 없음
    assert ws.cell(r2, rev["2026E"]).value == 900                    # 2026/09 결산기 추정
    assert ws.cell(r2, rev["2027E"]).value == 990                    # 2027/09 결산기 추정
    assert ws.cell(r2, rev["2027E"] + 1).value == pytest.approx((990 / 900 - 1) * 100)
    assert [ws.cell(r10, rev[y]).value for y in ("2025", "2026E", "2027E")] == [1200, 1300, 1300]
    sc = wb["점수"]
    note, srows = header(sc)["비고"], rows_by_code(sc)
    assert sc.cell(srows[tick(2)], note).value == "애널리스트 3명 이하 · 9월 결산"
    assert sc.cell(srows[tick(10)], note).value is None              # 12월 결산은 붙이지 않는다
    assert sc.cell(srows[tick(1)], note).value == "최근 3개월 의견 없음"


def _november_fy_no_opinion(table: str, rows: list[dict[str, object]]) -> None:
    """1번 종목(추정기관수 0 → '최근 3개월 의견 없음')을 11월 결산사로 — 비고가 가장 긴 조합이다."""
    if table == "fi_consensus":
        rows[:] = [{**r, "target_period": str(r["target_period"])[:5] + "11"}
                   if r["ticker"] == tick(1) else r for r in rows]


def test_daily_scores_note_column_fits_the_longest_note(tmp_path: Path) -> None:
    """비고 열은 가장 긴 조합('최근 3개월 의견 없음 · 11월 결산', qpack fit 25.75)도 잘리지 않는
    너비다 — 상한 16 이면 잘려 보인다(사용자는 잘린 표시에 민감하다)."""
    day = "2026-10-02"
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, day, write_fi_day(fi_root, day, edit=_november_fy_no_opinion))
    ws = load_workbook(build_daily(day, "morning", model_root=model_root, fi_root=fi_root,
                                   out_root=tmp_path).path)["점수"]
    col, longest = header(ws)["비고"], "최근 3개월 의견 없음 · 11월 결산"
    assert ws.cell(rows_by_code(ws)[tick(1)], col).value == longest
    need = fit_width(Col("note", "비고", "txt", None, 99.0), [{"note": longest}])
    assert ws.column_dimensions[ws.cell(7, col).column_letter].width >= need


def test_daily_sector_and_model_sheets(daily) -> None:
    _, wb = daily
    ws = wb["업종"]
    h = header(ws)
    first = {ws.cell(r, h["코드"]).value: r for r in range(8, ws.max_row + 1)}
    r45 = first["G45"]
    assert ws.cell(r45, h["구분"]).value == "대분류"
    assert ws.cell(r45, h["종목 수"]).value == 15
    assert ws.cell(r45, h["후보 수"]).value == 9
    assert ws.cell(r45, h["상위 30 비중"]).value == pytest.approx(15 / 30 * 100)
    assert str(ws.cell(r45, h["상위 3(순위)"]).value).startswith("종목0(1)")
    assert ws.cell(7, 1).fill.fgColor.rgb.endswith("0070C0")            # Style 시트 헤더
    assert ws.cell(r45, h["업종명"]).font.color.rgb.endswith("7030A0")   # 항목명 보라
    m = wb["모델 비교"]
    hm = header(m)
    for sid in SPECS:
        if sid == "v4_rank@0.2":               # v4 비교 열은 뺀다(N-27) — 주 모델 v4 기본은 남는다
            assert model_label(sid) not in hm
        else:
            assert model_label(sid) in hm, sid     # 헤더는 모델 이름(spec id 는 정의·메타)
    assert m.cell(8, col_of(m, "v4 기본")).value == 1


def test_daily_meta_sheet(daily) -> None:
    _, wb = daily
    ws = wb["메타"]
    pairs = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(8, ws.max_row + 1)}
    assert pairs["기준일"] == FRI and pairs["basis"] == "morning"
    assert pairs["model 판 id"] == "m_20260925T000000Z"
    assert pairs["factor_inputs 판 id"] == "m_20260925T010000Z"
    assert pairs["전일 비교 판"].startswith(THU)
    assert "1W 비교 판" not in pairs and pairs["1M 비교 판"].startswith("없음")
    assert pairs["순위 흐름 판"] == "2026-09-17 ~ 2026-09-25 · 6개"   # 09-23 휴장
    assert "MG0 pass" in pairs["판 게이트 v4_rank@0.1"]
    assert "저위험 25.0%" in pairs["가중치"]
    assert "20세션 거래대금 ≥ 10억" in pairs["유니버스 규칙"]
    assert all(f"각주 {k}" in pairs for k in (1, 2, 3, 4, 5))
    sheets_in_dict = {ws.cell(r, 1).value for r in range(8, ws.max_row + 1)}
    assert {"점수", "점수 원자료", "지표(표시용)", "실적", "업종", "모델 비교"} <= sheets_in_dict


def test_daily_values_only_no_formulas(daily) -> None:
    res, wb = daily
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate)
    with zipfile.ZipFile(res.path) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/sheet"):
                assert "<f>" not in z.read(name).decode("utf-8"), name
    ws = wb["점수"]
    r7 = rows_by_code(ws)[tick(7)]
    assert ws.cell(r7, 2).value == FORMULA_NAME                         # 문자열로 남는다


def test_daily_qpack_styles(daily) -> None:
    _, wb = daily
    ws = wb["점수"]
    h = header(ws)
    assert ws.sheet_view.showGridLines is False
    assert ws.freeze_panes == "I8"                                       # 식별 8열 뒤
    assert ws.auto_filter.ref == f"A7:{ws.cell(7, ws.max_column).column_letter}{7 + N}"
    a8 = ws.cell(8, 1)
    assert a8.font.name == "맑은 고딕" and a8.font.sz == 8
    assert ws.cell(1, 1).fill.fgColor.rgb.endswith("BDD7EE")
    assert ws.cell(1, 1).value == "sort" and ws.cell(1, h["순위"]).value == "▲"
    for r in (2, 3, 4):
        assert ws.cell(r, 1).fill.fgColor.rgb.endswith("4F81BD")
    assert ws.cell(2, 1).value == "점수" and ws.cell(2, 1).font.b
    assert ws.cell(7, 1).fill.fgColor.rgb.endswith("7F7F7F")
    assert ws.cell(7, h["순위"]).fill.fgColor.rgb.endswith("7030A0")      # 핵심 열 보라
    assert ws.cell(6, h["종합 점수"]).fill.fgColor.rgb.endswith("FFC000")  # 종합 그룹 금색
    assert ws.cell(6, h["저위험 유니버스"]).fill.fgColor.rgb.endswith("4F81BD")
    assert ws.row_dimensions[7].height == 33.75
    assert ws.cell(7, 1).border.bottom.style == "thin"
    assert ws.cell(8, h["시총(억)"]).number_format == "#,##0"
    sep = ws.column_dimensions[ws.cell(7, h["주가"] + 1).column_letter]
    assert sep.hidden and sep.outlineLevel == 1 and sep.width == 5
    # Q.Pack 실측(보이는 숫자 열 4.4~6.4) — 고정 13 이 아니라 값·헤더에 맞춘 최소 폭
    comp_w = ws.column_dimensions[ws.cell(7, h["종합 점수"]).column_letter].width
    assert 4.4 <= comp_w <= 6.5, comp_w
    shown = [ws.column_dimensions[ws.cell(7, c).column_letter] for c in range(1, ws.max_column + 1)]
    # 문자 열 상한 — 제외 사유 16 · 비고 26(가장 긴 조합 '최근 3개월 의견 없음 · 11월 결산'이
    # 25.75) · 그 밖은 14 이하
    assert max(d.width for d in shown if not d.hidden) <= 26, [d.width for d in shown]
    wide = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)
            if (d := ws.column_dimensions[ws.cell(7, c).column_letter]).width > 14 and not d.hidden]
    assert set(wide) <= {"제외 사유", "비고"}, wide
    assert wb["점수 원자료"].row_dimensions[5].height == 11.25        # 정의 행도 한 줄
    # 행 높이는 행마다 박는다 — 시트 기본값만 두면 기본 글꼴(Calibri 11)로 다시 잡혀 Q.Pack 보다
    # 1.4배 높게 그려진다(10-05 렌더: 한 쪽 29행 vs Q.Pack 37행)
    assert all(ws.row_dimensions[r].height == 11.25 for r in range(1, ws.max_row + 1) if r != 7)
    for s in wb.worksheets:
        assert all(s.row_dimensions[r].height for r in range(1, s.max_row + 1)), s.title
    rules = [(str(rng.sqref), rule) for rng in ws.conditional_formatting for rule in rng.rules]
    scales = [(ref, r.colorScale) for ref, r in rules if isinstance(r.colorScale, ColorScale)]
    colored = {ref.split(":")[0].rstrip("0123456789") for ref, _ in scales}
    letter = ws.cell(7, h["저위험 유니버스"]).column_letter
    assert letter in colored and ws.cell(7, h["Δ순위 1M"]).column_letter in colored
    assert ws.cell(7, h["1M 흐름"]).column_letter not in colored        # 그래프 칸은 색 없음
    assert ws.cell(7, h["순위"]).column_letter not in colored            # 점수 시트 순위 무색
    cs = next(c for ref, c in scales if ref.startswith(letter))
    assert [v.type for v in cs.cfvo] == ["percentile"] * 3
    assert [v.val for v in cs.cfvo] == [10, 50, 90]
    assert [c.rgb[-6:] for c in cs.color] == ["F8696B", "FFEB84", "63BE7B"]   # 높음 = 초록(N-15)


def test_daily_rank_trend_sheet_and_sparklines(daily) -> None:
    """숨김 '순위 흐름' 시트 = 판마다의 −순위(원순위, 위 = 상승 — E-06)(행은 점수 시트와 같은 순서),
    점수 시트에는 1M 흐름 칸마다 엑셀 스파크라인이 붙는다. 선 색 = Δ순위 1M 부호(N-16)."""
    res, wb = daily
    tr, sc = wb["순위 흐름"], wb["점수"]
    assert [tr.cell(7, c).value for c in range(1, tr.max_column + 1)] == [
        "코드", "09-17", "09-18", "09-21", "09-22", "09-24", "09-25"]
    assert [tr.cell(r, 1).value for r in range(8, tr.max_row + 1)] == [
        sc.cell(r, 1).value for r in range(8, sc.max_row + 1)]
    r0 = rows_by_code(tr)[tick(0)]
    assert tr.cell(r0, 2).value == -3                                       # 09-17 3위
    assert tr.cell(r0, 7).value == -1                                       # 09-25 1위
    assert tr.cell(rows_by_code(tr)[tick(40)], 7).value is None             # 제외 종목
    with zipfile.ZipFile(res.path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    h = header(sc)
    c1m = sc.cell(7, h["1M 흐름"]).column_letter
    assert xml.count("<x14:sparkline>") == N
    assert 'lineWeight="0.25"' in xml and 'lineWeight="1"' not in xml    # 가장 얇은 선
    assert f"<xm:f>'순위 흐름'!B{r0}:G{r0}</xm:f><xm:sqref>{c1m}{r0}</xm:sqref>" in xml
    groups = xml.split("<x14:sparklineGroup ")[1:]
    def color_of(cell: str) -> str:
        g = next(g for g in groups if f"<xm:sqref>{cell}</xm:sqref>" in g)
        return g.split('colorSeries rgb="FF')[1][:6]
    assert color_of(f"{c1m}{r0}") == "8C8C8C"                    # 1M 비교 판이 없어 모름 → 회색


def test_spark_groups_color_lines_by_delta_sign() -> None:
    """선 색 = Δ순위 1M 부호 — 상승 초록 · 하락 빨강 · 모름 회색.
    원자료 범위는 비교 판 열부터 D 까지."""
    from deliver.trend import Trend, spark_groups
    tr = Trend(dates=("2026-09-02", "2026-09-15", "2026-10-02"),
               line={"A": (-50, -20, -5), "B": (-5, -20, -45), "C": (None, -20, -20)},
               base={"1M": "2026-09-02"}, base_rank={"1M": {"A": 50, "B": 5}})
    deltas = {"1M": {"A": 45, "B": -40, "C": None}}
    groups = spark_groups(["A", "B", "C"], tr, date(2026, 10, 2), {"1M": "M"}, deltas)
    assert {g.color: [c for _, c in g.cells] for g in groups} == {
        "1E8C45": ["M8"], "D0312D": ["M9"], "8C8C8C": ["M10"]}
    assert {ref for g in groups for ref, _ in g.cells} == {
        "'순위 흐름'!B8:D8", "'순위 흐름'!B9:D9", "'순위 흐름'!B10:D10"}


def test_month_back_and_windows() -> None:
    from deliver.trend import month_back, window_start
    assert month_back(date(2026, 3, 31)) == date(2026, 2, 28)
    assert month_back(date(2026, 1, 15)) == date(2025, 12, 15)
    assert window_start(date(2026, 10, 2), "1W") == date(2026, 9, 25)
    assert window_start(date(2026, 10, 2), "1M") == date(2026, 9, 2)


def test_daily_rank_scale_is_reversed_on_model_sheet(daily) -> None:
    _, wb = daily
    ws = wb["모델 비교"]
    c = ws.cell(7, col_of(ws, "v4 기본")).column_letter
    cs = next(r.colorScale for rng in ws.conditional_formatting for r in rng.rules
              if str(rng.sqref).startswith(c))
    assert [x.rgb[-6:] for x in cs.color] == ["63BE7B", "FFEB84", "F8696B"]   # 1위 = 초록(N-15)


def test_daily_missing_or_failed_run_is_an_error(world: dict[str, Path], tmp_path: Path) -> None:
    with pytest.raises(DeliverError, match="없다"):
        build_daily("2026-09-26", "morning", model_root=world["model"], fi_root=world["fi"],
                    out_root=tmp_path)
    model_root = tmp_path / "model"
    write_model_day(model_root, FRI, write_fi_day(tmp_path / "fi", FRI), status="gate_failed")
    with pytest.raises(DeliverError, match="성공 판이 아니다"):
        build_daily(FRI, "morning", model_root=model_root, fi_root=tmp_path / "fi",
                    out_root=tmp_path)


# ── 주간 엑셀 ─────────────────────────────────────────────────────────────────
EXPECTED_CANDIDATES = [*range(0, 9), *range(15, 24), *range(25, 34), 35, 36, 37]


def test_weekly_sheets_and_candidates_with_sector_cap(weekly) -> None:
    res, wb = weekly
    assert res.path.parts[-2:] == ("2026-W39", "weekly_2026-W39.xlsx")
    assert wb.sheetnames == [
        "주간 후보", "팩터 카드", "주간 추이", "이탈·진입", "지난주 성적", "메타"]
    assert res.base_date == FRI and res.n_days == 4
    ws = wb["주간 후보"]
    h = header(ws)
    assert ws.max_row == 7 + 30
    codes = [ws.cell(r, h["코드"]).value for r in range(8, 38)]
    assert codes == [tick(i) for i in EXPECTED_CANDIDATES]
    per_sector: dict[str, int] = {}
    for r in range(8, 38):
        s = ws.cell(r, h["대분류"]).value
        per_sector[s] = per_sector.get(s, 0) + 1
    assert max(per_sector.values()) == 9 and per_sector["IT"] == 9
    assert [ws.cell(r, h["순번"]).value for r in range(8, 38)] == list(range(1, 31))
    assert ws.cell(8, h["금요일 순위 (원본·상한 전)"]).value == 1
    assert ws.cell(8 + 9, h["금요일 순위 (원본·상한 전)"]).value == 16   # 상한으로 10~15위 건너뜀
    assert res.n_new == 2 and res.n_out == 2


def test_weekly_rows_have_explicit_heights(weekly) -> None:
    _, wb = weekly
    for s in wb.worksheets:
        assert all(s.row_dimensions[r].height for r in range(1, s.max_row + 1)), s.title


def test_weekly_presence_marks_and_holiday(weekly) -> None:
    _, wb = weekly
    ws = wb["주간 후보"]
    h = header(ws)
    rows = {ws.cell(r, h["코드"]).value: r for r in range(8, 38)}
    marks = [f"{d} {day[5:]}" for d, day in zip("월화수목금", W39, strict=True)]
    r0 = rows[tick(0)]
    assert [ws.cell(r0, h[m]).value for m in marks] == ["●", "●", "✕", "●", "●"]
    assert ws.cell(r0, h["주간 평균 순위"]).value == pytest.approx(statistics.fmean([2, 1, 2, 1]))
    assert ws.cell(r0, h["등장 (●)"]).value == 4
    r37 = rows[tick(37)]                 # 월요일엔 38 이 끼어 37 은 후보 밖(○)
    assert ws.cell(r37, h[marks[0]]).value == "○"
    assert ws.cell(r0, h["지난주 순위"]).value == 3
    assert ws.cell(r0, h["신규/ 유지"]).value == "유지"
    r33 = rows[tick(33)]
    assert ws.cell(r33, h["신규/ 유지"]).value == "신규"
    assert ws.cell(r33, h["신규/ 유지"]).fill.fgColor.rgb.endswith("FFFF00")
    assert ws.cell(r0, h["v3 순위 (참고)"]).value == 40
    trend = wb["주간 추이"]
    ht = header(trend)
    assert trend.max_row == 7 + 32                                      # 후보 30 + 이탈 2
    wed = f"수 {HOLIDAY[5:]}"
    wed_rank, wed_comp = col_of(trend, wed, 0), col_of(trend, wed, 1)
    assert {trend.cell(r, wed_rank).value for r in range(8, trend.max_row + 1)} == {"✕"}
    assert {trend.cell(r, wed_comp).value for r in range(8, trend.max_row + 1)} == {None}
    tr = {trend.cell(r, 1).value: r for r in range(8, trend.max_row + 1)}
    assert trend.cell(tr[tick(34)], ht["구분"]).value == "이탈"
    assert trend.cell(tr[tick(0)], col_of(trend, f"화 {TUE[5:]}", 0)).value == 1
    assert trend.cell(tr[tick(0)], col_of(trend, f"월 {MON[5:]}", 1)).value == pytest.approx(96.0)


def test_weekly_moves_and_scorecard(weekly) -> None:
    _, wb = weekly
    ws = wb["이탈·진입"]
    h = header(ws)
    got = {(ws.cell(r, h["방향"]).value, ws.cell(r, h["코드"]).value)
           for r in range(8, ws.max_row + 1)}
    assert got == {("진입", tick(33)), ("진입", tick(37)), ("이탈", tick(34)), ("이탈", tick(39))}
    sc = wb["지난주 성적"]
    hs = header(sc)
    assert sc.max_row == 7 + 30
    rows = {sc.cell(r, hs["코드"]).value: r for r in range(8, sc.max_row + 1)}
    assert sc.cell(rows[tick(0)], hs["1주 수익률"]).value == pytest.approx(ret_of(0) * 100)
    assert sc.cell(rows[tick(35)], hs["1주 수익률"]).value == "결측(수정주가미해결)"
    last_ranked = [i for i in order_for(W38[1]) if i not in excluded_for(W38[1]) and i != 35]
    uni = statistics.fmean(ret_of(i) * 100 for i in last_ranked)
    assert sc.cell(rows[tick(0)], hs["유니버스 평균"]).value == pytest.approx(uni)
    assert "승률" in str(sc.cell(5, 1).value)


def test_weekly_factor_cards_and_meta(weekly) -> None:
    _, wb = weekly
    ws = wb["팩터 카드"]
    titles = [ws.cell(r, 1).value for r in range(6, ws.max_row + 1)
              if str(ws.cell(r, 1).value or "").split(".")[0].isdigit()]
    assert len(titles) == 30 and str(titles[0]).startswith(f"1. {tick(0)}")
    assert ws.freeze_panes == "A6"
    labels = {ws.cell(r, 1).value for r in range(6, ws.max_row + 1)}
    assert {"저위험", "리비전", "60일 변동성(%)", "최근 5일 순위", "기준일"} <= labels
    rank_row = next(r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == "최근 5일 순위")
    assert [ws.cell(rank_row, k).value for k in range(3, 8)] == ["월 2", "화 1", "수 ✕", "목 2",
                                                                "금 1"]
    meta = wb["메타"]
    pairs = {meta.cell(r, 1).value: meta.cell(r, 2).value for r in range(8, meta.max_row + 1)}
    assert pairs["주"] == "2026-W39"
    assert pairs[f"수 {HOLIDAY}"].startswith("판 없음")
    assert pairs["지난주 기준일"].startswith(W38[1])
    for ws2 in wb.worksheets:
        for row in ws2.iter_rows():
            assert all(c.data_type != "f" for c in row)


def test_weekly_without_last_week_or_runs(world: dict[str, Path], tmp_path: Path) -> None:
    with pytest.raises(DeliverError, match="하나도 없다"):
        build_weekly("2026-W41", model_root=world["model"], fi_root=world["fi"],
                     out_root=tmp_path)
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, FRI, write_fi_day(fi_root, FRI))
    res = build_weekly("2026-W39", model_root=model_root, fi_root=fi_root, out_root=tmp_path)
    wb = load_workbook(res.path)
    assert res.n_days == 1 and res.n_new == 0
    assert wb["이탈·진입"].max_row == 7 and wb["지난주 성적"].max_row == 7
    ws = wb["주간 후보"]
    h = header(ws)
    assert ws.cell(8, h["신규/ 유지"]).value is None


# ── CLI ───────────────────────────────────────────────────────────────────────
def _env(tmp_path: Path) -> Path:
    p = tmp_path / "fake.env"
    p.write_text(f"BOT_TOKEN={TOKEN}\nCHAT_ID_AIPLAYGROUND={CHAT}\nOTHER=x\n", encoding="utf-8")
    return p


def _roots(world: dict[str, Path], out: Path) -> list[str]:
    return ["--model-root", str(world["model"]), "--fi-root", str(world["fi"]),
            "--out-root", str(out)]


def test_cli_daily_and_weekly_dry_run(world: dict[str, Path], tmp_path: Path,
                                      capsys: pytest.CaptureFixture[str]) -> None:
    env = _env(tmp_path)
    rc = cli.main(["model-daily", "--date", "20260925", "--basis", "morning", "--dry-run",
                   "--env-file", str(env), *_roots(world, tmp_path)])
    out = capsys.readouterr()
    assert rc == 0
    assert (tmp_path / "daily" / "model_scores_20260925_morning.xlsx").exists()
    assert "dry-run" in out.out and "[모델 점수] 2026-09-25 morning" in out.out
    rc = cli.main(["model-weekly", "--week", "2026-W39", "--dry-run", "--send",
                   "--env-file", str(env), *_roots(world, tmp_path)])
    out2 = capsys.readouterr()
    assert rc == 0 and "[주간 후보] 2026-W39" in out2.out
    for text in (out.out, out.err, out2.out, out2.err):
        assert TOKEN not in text and CHAT not in text


def test_cli_send_uses_transport_and_errors(world: dict[str, Path], tmp_path: Path,
                                            monkeypatch: pytest.MonkeyPatch,
                                            capsys: pytest.CaptureFixture[str]) -> None:
    import deliver.telegram as tg
    calls: list[tuple[str, dict[str, str], str]] = []

    def fake(url: str, fields, name: str, data: bytes):
        calls.append((url, dict(fields), name))
        return {"ok": True}
    monkeypatch.setattr(tg, "urllib_transport", fake)
    rc = cli.main(["model-daily", "--date", "20260925", "--basis", "morning", "--send",
                   "--env-file", str(_env(tmp_path)), *_roots(world, tmp_path)])
    out = capsys.readouterr()
    assert rc == 0 and len(calls) == 1
    assert calls[0][0].endswith("/sendDocument") and calls[0][1]["chat_id"] == CHAT
    assert calls[0][2] == "model_scores_20260925_morning.xlsx"
    assert TOKEN not in out.out + out.err and CHAT not in out.out + out.err
    rc = cli.main(["model-daily", "--date", "20260926", "--basis", "morning",
                   *_roots(world, tmp_path)])
    assert rc == 2
    assert cli.main(["model-weekly", "--week", "2026W39", *_roots(world, tmp_path)]) == 2
    rc = cli.main(["model-daily", "--date", "20260925", "--basis", "morning", "--dry-run",
                   "--env-file", str(tmp_path / "none.env"), *_roots(world, tmp_path)])
    assert rc == 1                                                       # 비밀 키 없음


# ── 발송 장부·재발송 가드(N-25 Q9) · 생성 실패 rc 3(E-13) ──────────────────────────
LEDGER = "sent_model_daily.jsonl"
FRI_BUILD, FRI_GENERATED = "m_20260925T000000Z", "2026-09-25T00:00:00Z"


def _fake_send(monkeypatch: pytest.MonkeyPatch, *, ok: bool = True) -> list[dict[str, str]]:
    """가짜 전송(실제 텔레그램 없음) — 보낸 폼 필드(caption 포함)를 모은다."""
    import deliver.telegram as tg
    calls: list[dict[str, str]] = []

    def fake(url: str, fields, name: str, data: bytes):
        calls.append(dict(fields))
        return {"ok": True} if ok else {"ok": False, "description": "Bad Request: 가짜 실패"}
    monkeypatch.setattr(tg, "urllib_transport", fake)
    return calls


def _ledger(out: Path) -> list[dict[str, object]]:
    p = out / LEDGER
    if not p.exists():
        return []
    return [json.loads(ln) for ln in p.read_text(encoding="utf-8").splitlines() if ln.strip()]


def _send_args(world: dict[str, Path], tmp_path: Path) -> list[str]:
    return ["model-daily", "--date", "20260925", "--basis", "morning", "--send",
            "--env-file", str(_env(tmp_path)), *_roots(world, tmp_path)]


def test_cli_send_records_ledger_and_skips_the_same_day(world: dict[str, Path], tmp_path: Path,
                                                       monkeypatch: pytest.MonkeyPatch,
                                                       capsys: pytest.CaptureFixture[str]) -> None:
    """Q9 · K0-1 — 보낸 D·basis 는 발송 장부에 한 줄 남고, 같은 D·basis 의 두 번째 자동 발송은
    건너뛴다(rc 0, 로그만). 옛 코드는 장부가 없어 같은 D 를 다시 돌리면 다시 보냈다
    (E-08 — 10-02 5회)."""
    calls = _fake_send(monkeypatch)
    args = _send_args(world, tmp_path)
    assert cli.main(args) == 0
    xlsx = tmp_path / "daily" / "model_scores_20260925_morning.xlsx"
    sent = _ledger(tmp_path)
    assert len(calls) == 1 and len(sent) == 1
    e = sent[0]
    assert {k: e[k] for k in ("date", "basis", "build_id", "correction")} == {
        "date": FRI, "basis": "morning", "build_id": FRI_BUILD, "correction": 0}
    assert e["sha256"] == hashlib.sha256(xlsx.read_bytes()).hexdigest()
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", str(e["sent_utc"]))
    assert "정정" not in calls[0]["caption"]                 # 첫 발송 캡션은 종전 그대로
    mtime = xlsx.stat().st_mtime_ns
    capsys.readouterr()
    assert cli.main(args) == 0                               # 같은 D·basis 두 번째 실행
    out = capsys.readouterr().out
    assert len(calls) == 1 and len(_ledger(tmp_path)) == 1   # 발송 0회 · 장부 그대로
    assert "건너뜀" in out and "--resend" in out
    # 보낸 파일을 다시 만들지 않는다 — 디스크 파일이 장부 sha256 과 같게 남는다
    assert xlsx.stat().st_mtime_ns == mtime


def test_cli_resend_marks_correction_in_caption(world: dict[str, Path], tmp_path: Path,
                                                monkeypatch: pytest.MonkeyPatch) -> None:
    """Q9 — 같은 날 다시 보내려면 `--resend` 를 명시하고, 그때 캡션에 '정정 n'·판 id·생성
    시각을 단다."""
    calls = _fake_send(monkeypatch)
    args = _send_args(world, tmp_path)
    assert cli.main(args) == 0
    assert cli.main([*args, "--resend"]) == 0
    assert len(calls) == 2
    cap = calls[1]["caption"]
    assert "정정 1" in cap and FRI_BUILD in cap and FRI_GENERATED in cap
    assert [e["correction"] for e in _ledger(tmp_path)] == [0, 1]
    with pytest.raises(SystemExit) as ex:                    # --send 없는 --resend 는 인자 오류
        cli.main(["model-daily", "--date", "20260925", "--basis", "morning", "--resend",
                  *_roots(world, tmp_path)])
    assert ex.value.code == 2 and len(calls) == 2


def test_cli_failed_send_writes_no_ledger(world: dict[str, Path], tmp_path: Path,
                                          monkeypatch: pytest.MonkeyPatch) -> None:
    """장부는 발송이 성공한 뒤에만 쓴다 — 보내지 못한 날(rc 1)은 다음 자동 실행이 다시 보낸다."""
    _fake_send(monkeypatch, ok=False)
    args = _send_args(world, tmp_path)
    assert cli.main(args) == 1
    assert _ledger(tmp_path) == []
    calls = _fake_send(monkeypatch)
    assert cli.main(args) == 0
    assert len(calls) == 1 and len(_ledger(tmp_path)) == 1


def test_cli_unreadable_ledger_sends_nothing(world: dict[str, Path], tmp_path: Path,
                                             monkeypatch: pytest.MonkeyPatch,
                                             capsys: pytest.CaptureFixture[str]) -> None:
    """장부 줄을 읽지 못하면 보냈는지 알 수 없다 — 보내지 않고 rc 2(입력 오류, P1)."""
    calls = _fake_send(monkeypatch)
    (tmp_path / LEDGER).write_text('{"date": "2026-09-25", 깨진 줄\n', encoding="utf-8")
    assert cli.main(_send_args(world, tmp_path)) == 2
    assert calls == [] and f"{LEDGER}:1" in capsys.readouterr().err


@pytest.mark.parametrize("raw", [b'\xff\xfe{"date": "2026-09-25"}\n', b"[1, 2]\n"],
                         ids=["not_utf8", "not_object"])
def test_cli_malformed_ledger_is_an_input_error(world: dict[str, Path], tmp_path: Path,
                                                monkeypatch: pytest.MonkeyPatch,
                                                capsys: pytest.CaptureFixture[str],
                                                raw: bytes) -> None:
    """장부가 UTF-8 이 아니거나 줄이 JSON 객체가 아니어도 보냈는지 알 수 없다 — 옛 코드는
    UnicodeDecodeError·AttributeError 가 rc 3(생성 실패)으로 샜다. rc 2, 발송 0."""
    calls = _fake_send(monkeypatch)
    (tmp_path / LEDGER).write_bytes(raw)
    assert cli.main(_send_args(world, tmp_path)) == 2
    assert calls == [] and "보냈는지 알 수 없어 보내지 않는다" in capsys.readouterr().err


@pytest.mark.skipif(os.geteuid() == 0, reason="root 는 권한 없는 파일도 읽는다")
def test_cli_unreadable_ledger_permission_is_an_input_error(
        world: dict[str, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str]) -> None:
    """장부를 권한 때문에 못 읽어도(OSError) 보냈는지 알 수 없다 — rc 3(예상 밖)이 아니라
    rc 2, 발송 0."""
    calls = _fake_send(monkeypatch)
    ledger = tmp_path / LEDGER
    ledger.write_text("", encoding="utf-8")
    ledger.chmod(0o000)
    try:
        assert cli.main(_send_args(world, tmp_path)) == 2
    finally:
        ledger.chmod(0o644)
    assert calls == [] and "보냈는지 알 수 없어 보내지 않는다" in capsys.readouterr().err


def test_cli_ledger_append_failure_after_send_says_it_was_sent(
        world: dict[str, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str]) -> None:
    """B-58 — 발송이 성공한 뒤 장부 append 가 실패하면 rc 3 이고, 메시지가 '발송은 됐다'를 밝힌다
    (장부에 없으니 다음 자동 실행이 다시 보낼 수 있다 — 사람이 장부를 채워야 한다)."""
    calls = _fake_send(monkeypatch)
    (tmp_path / LEDGER).mkdir()           # 장부 자리에 폴더 — 읽기는 '없음', 발송 뒤 append 는 실패
    rc = cli.main(_send_args(world, tmp_path))
    err = capsys.readouterr().err
    assert rc == 3 and len(calls) == 1
    assert "발송은 됐다 — 장부 기록 실패" in err and "IsADirectoryError" in err


def test_cli_ledger_append_restores_a_missing_newline(world: dict[str, Path], tmp_path: Path,
                                                     monkeypatch: pytest.MonkeyPatch) -> None:
    """손 편집으로 장부 끝 줄바꿈이 빠졌어도 append 가 줄을 붙이지 않는다 — 붙으면 다음 실행부터
    매일 장부 손상 rc 2 로 멈춘다."""
    calls = _fake_send(monkeypatch)
    other = {"date": "2026-09-24", "basis": "morning", "build_id": "m_x", "sha256": "0" * 64,
             "sent_utc": "2026-09-24T01:00:00Z", "correction": 0}
    (tmp_path / LEDGER).write_text(json.dumps(other), encoding="utf-8")    # 끝 줄바꿈 없음
    args = _send_args(world, tmp_path)
    assert cli.main(args) == 0
    assert [e["date"] for e in _ledger(tmp_path)] == ["2026-09-24", FRI]
    assert cli.main(args) == 0 and len(calls) == 1     # 다음 실행도 장부를 읽고 건너뛴다


def test_cli_unexpected_error_is_rc_3(world: dict[str, Path], tmp_path: Path,
                                      monkeypatch: pytest.MonkeyPatch,
                                      capsys: pytest.CaptureFixture[str]) -> None:
    """E-13 — DeliverError 밖의 예외(엑셀 생성 실패)는 rc 3 이다. 옛 코드는 예외가 그대로 올라가
    파이썬 기본 rc 1(문서상 '발송 실패')과 섞였다. 아무것도 보내지 않고 장부도 쓰지 않는다."""
    calls = _fake_send(monkeypatch)

    def boom(*_a: object, **_k: object) -> None:
        raise ValueError("시트 '점수' 를 찾지 못했다(가짜)")
    monkeypatch.setattr(cli, "build_daily", boom)
    rc = cli.main(_send_args(world, tmp_path))
    err = capsys.readouterr().err
    assert rc == 3 and calls == [] and _ledger(tmp_path) == []
    assert "ValueError" in err and "가짜" in err and "--date 20260925" in err


# ── 실제 model 판 빌드 → 전달층 ────────────────────────────────────────────────
def test_real_model_build_feeds_deliver(tmp_path: Path) -> None:
    """`model.build` 가 쓴 판(40종목 합성 입력, 네 spec)을 그대로 읽는다 — 파일 계약 대조."""
    from model import build as mbuild
    from test_model_v4_rank import Board, _full

    b = Board()
    for i in range(40):
        _full(b, f"{100000 + i:06d}", i, sector=("G10", "G20", "G30")[i % 3])
    for r in b.tables["fi_prices"]:
        r["close"] = round(float(r["close"]))                  # type: ignore[arg-type]
    fi, fi_root, fi_bid = b.fi(), tmp_path / "fi", "m_20260929T000500_000000Z"
    for name, t in FI_TABLES.items():
        out = fi_root / name / f"v={fi_bid}" / "part0.parquet"
        mbuild.write_parquet(list(fi.tables.get(name, ())),
                             {c.name: c.dtype for c in t.columns}, out)
        (out.parent / "_meta.json").write_text(json.dumps(
            {"table": name, "build_id": fi_bid, "basis": "morning", "date": "2026-09-28"}))
    (fi_root / "latest_morning.json").write_text(json.dumps(
        {"layer": "factor_inputs", "status": "ok", "build_id": fi_bid, "date": "2026-09-28",
         "basis": "morning"}))
    res = mbuild.build("20260928", "morning", tmp_path / "model", fi_root,
                       min_prices_on_d=10, min_ranked=10)
    assert res.ok, res.specs

    d = build_daily("2026-09-28", "morning", model_root=tmp_path / "model", fi_root=fi_root,
                    out_root=tmp_path / "out")
    wb = load_workbook(d.path)
    primary = res.primary_spec                 # 기본 주 모델(scope@1.0)
    n_scores = res.specs[primary]["n_scores"]
    assert d.n_rows == n_scores and d.n_ranked == res.specs[primary]["n_ranked"]
    assert wb["점수"].max_row == 7 + n_scores
    pairs = {wb["메타"].cell(r, 1).value: wb["메타"].cell(r, 2).value
             for r in range(8, wb["메타"].max_row + 1)}
    assert pairs["model 판 id"] == res.build_id and "MG1 pass" in pairs[f"판 게이트 {primary}"]
    rule = pairs["유니버스 규칙"]
    assert "추정기관수 ≥ 1" in rule and "추정치 유예 없음" in rule
    # scope(v3 엔진) 점수 행엔 업종 열이 없다 — fi_universe 업종으로 채워 업종 시트가 3 대분류로 선다
    sector_rows = [wb["업종"].cell(r, 1).value for r in range(8, wb["업종"].max_row + 1)]
    assert len(sector_rows) >= 3, sector_rows
    # 신선도도 점수 행이 아니라 fi_universe 에서 — 빈칸이면 커버리지 열이 통째로 빈다
    ws = wb["점수"]
    cov = [ws.cell(r, header(ws)["커버리지"]).value for r in range(8, ws.max_row + 1)]
    assert cov and all(cov), cov[:5]
    # 종합 점수 열 이름은 엔진마다 다르다(v3 composite_score · v2 total_score) — 비면 안 된다
    comp = [ws.cell(r, header(ws)["종합 점수"]).value for r in range(8, ws.max_row + 1)]
    assert all(isinstance(c, float) for c in comp), comp[:5]
    sec = wb["업종"]
    med = [sec.cell(r, header(sec)["종합 점수 중앙값"]).value for r in range(8, sec.max_row + 1)]
    assert any(isinstance(m, float) for m in med), med
    with pytest.raises(DeliverError, match="업종 열"):
        load_day(tmp_path / "model", None, find_run(tmp_path / "model", "2026-09-28", "morning"),
                 with_fi=False)
    w = build_weekly("2026-W40", model_root=tmp_path / "model", fi_root=fi_root,
                     out_root=tmp_path / "out")
    assert w.base_date == "2026-09-28" and w.n_days == 1
    assert w.n_candidates == min(30, 3 * 9, res.specs[primary]["n_ranked"])   # 3 업종 × 상한 9
    for book in (wb, load_workbook(w.path)):
        for ws in book.worksheets:
            for row in ws.iter_rows():
                assert all(c.data_type != "f" for c in row)


def test_failed_comparison_model_is_left_out_of_the_excel(tmp_path: Path, monkeypatch) -> None:
    """비교 모델이 게이트에서 떨어지면 그 모델 열만 빠지고 메타에 사유가 남는다(N-11 격리)."""
    from model import build as mbuild
    from test_model_build import board_fi, edit_scores, patch_engine, write_fi_tree

    v2 = "v2_percentrank@1.0"
    fi_root = write_fi_tree(tmp_path / "fi", board_fi())
    patch_engine(monkeypatch, "v2_percentrank", edit_scores(lambda s: s[0].update(rank=None)))
    res = mbuild.build("20260928", "morning", tmp_path / "model", fi_root,
                       min_prices_on_d=10, min_ranked=10)
    assert res.ok and res.excluded == (v2,)
    d = build_daily("2026-09-28", "morning", model_root=tmp_path / "model", fi_root=fi_root,
                    out_root=tmp_path / "out")
    wb = load_workbook(d.path)
    pairs = {wb["메타"].cell(r, 1).value: wb["메타"].cell(r, 2).value
             for r in range(8, wb["메타"].max_row + 1)}
    assert "MG3 fail" in pairs[f"제외된 비교 모델 {v2}"]
    assert f"판 게이트 {v2}" not in pairs and v2 not in pairs["비교 모델"]
    assert model_label(v2) not in header(wb["모델 비교"])


def test_erroring_comparison_model_is_left_out_of_the_excel(tmp_path: Path, monkeypatch) -> None:
    """D-01 왕복(build ↔ 엑셀 계약): 비교 모델 엔진이 예외를 내면 `model.build` 가 그 모델을 빼고,
    일일 엑셀은 그 모델 열을 빼고 메타에 '실행 오류로 …' 와 사유를 적는다."""
    from model import build as mbuild
    from test_model_build import board_fi, boom, patch_engine, write_fi_tree

    v2 = "v2_percentrank@1.0"
    fi_root = write_fi_tree(tmp_path / "fi", board_fi())
    patch_engine(monkeypatch, "v2_percentrank", boom)
    res = mbuild.build("20260928", "morning", tmp_path / "model", fi_root,
                       min_prices_on_d=10, min_ranked=10)
    assert res.ok and res.excluded == (v2,)
    d = build_daily("2026-09-28", "morning", model_root=tmp_path / "model", fi_root=fi_root,
                    out_root=tmp_path / "out")
    wb = load_workbook(d.path)
    pairs = {wb["메타"].cell(r, 1).value: wb["메타"].cell(r, 2).value
             for r in range(8, wb["메타"].max_row + 1)}
    assert pairs[f"제외된 비교 모델 {v2}"] == (
        "실행 오류로 이번 판에서 뺐다 · ZeroDivisionError: 지표 분모 0")
    assert f"판 게이트 {v2}" not in pairs and v2 not in pairs["비교 모델"]
    assert model_label(v2) not in header(wb["모델 비교"])


def test_engine_error_comparison_model_reason_in_the_meta_sheet(tmp_path: Path) -> None:
    """D-01: 실행 예외로 뺀 비교 모델(`excluded_specs` 항목에 `error`)은 메타에 '실행 오류로 …'
    와 사유를 적는다. 게이트로 뺀 것은 '게이트 실패로 …' 문구 그대로다."""
    v2, v3 = "v2_percentrank@1.0", "v3_zscore@1.0"
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, FRI, write_fi_day(fi_root, FRI))
    path = model_root / "_runs" / "20260925_morning.json"
    meta = json.loads(path.read_text(encoding="utf-8"))
    gate_failed = meta["specs"].pop(v3)
    gate_failed["gates"] = {"MG3": {"status": "fail"}}
    meta["specs"].pop(v2)
    meta["excluded_specs"] = {v2: {"error": "ZeroDivisionError: x"}, v3: gate_failed}
    path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    res = build_daily(FRI, "morning", model_root=model_root, fi_root=fi_root, out_root=tmp_path)
    ws = load_workbook(res.path)["메타"]
    pairs = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(8, ws.max_row + 1)}
    assert pairs[f"제외된 비교 모델 {v2}"] == "실행 오류로 이번 판에서 뺐다 · ZeroDivisionError: x"
    assert pairs[f"제외된 비교 모델 {v3}"] == "게이트 실패로 이번 판에서 뺐다 · MG3 fail"


def test_display_names_strip_common_suffix_and_wics_prefix() -> None:
    from deliver.reader import display_names
    rows = display_names([
        {"name": "오리온홀딩스보통주", "sector_l1_name": "필수소비재",
         "sector_l2_name": "WICS 식품,음료,담배"},
        {"name": "현대차우", "sector_l1_name": "WICS 경기관련소비재", "sector_l2_name": None},
        {"name": "보통주", "sector_l1_name": None, "sector_l2_name": "WICS 자본재"},
    ])
    assert [r["name"] for r in rows] == ["오리온홀딩스", "현대차우", "보통주"]
    assert rows[0]["sector_l2_name"] == "식품,음료,담배"
    assert rows[1]["sector_l1_name"] == "경기관련소비재"
    assert rows[1]["sector_l2_name"] is None and rows[2]["sector_l2_name"] == "자본재"


# ── 배포 묶음 4-2a — 일간 엑셀 표시 바로잡기(E-09·E-02·E-08·E-06·E-10·Q10·N-27) ─────────────
def _meta_pairs(wb) -> list[tuple[object, object]]:
    """메타 시트 키·값 표(8행부터 첫 빈 행 전까지 — 그 뒤는 열 사전)."""
    ws, out = wb["메타"], []
    for r in range(8, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            break
        out.append((ws.cell(r, 1).value, ws.cell(r, 2).value))
    return out


def test_meta_lines_are_placed_by_key_name(daily) -> None:
    """E-09 — 비교 판 줄은 'model 생성 시각' 뒤에 이름으로 끼운다. 옛 코드는 insert(6)·[7:7]·
    insert(9) 고정 숫자라 1W 를 뺀 뒤 '순위 흐름 판'이 factor_inputs·equity 판 id 사이에 끼었다."""
    _, wb = daily
    keys = [k for k, _ in _meta_pairs(wb)]
    at = keys.index("model 생성 시각")
    assert keys[at + 1:at + 4] == ["전일 비교 판", "1M 비교 판", "순위 흐름 판"]
    assert keys[keys.index("factor_inputs 판 id") + 1] == "equity 판 id"


def test_exclude_reason_column_fits_every_v4_label(daily) -> None:
    """N-26 4.5 — '제외 사유' 열 상한 26. 상한 16 이면 v4 라벨 8개 중 6개가 잘려 보인다
    (가장 긴 '고점근접+반전 게이트(pull_gate)' qpack fit 25.05)."""
    from deliver.common import EXCLUDE_LABELS, exclude_label
    _, wb = daily
    ws = wb["점수"]
    col = header(ws)["제외 사유"]
    labels = [exclude_label(c) for c in (*EXCLUDE_LABELS, "pull_gate")]
    probe = Col("excl", "제외 사유", "txt", None, 99.0)
    need = max(fit_width(probe, [{"excl": s}]) for s in labels)
    assert ws.column_dimensions[ws.cell(7, col).column_letter].width >= need


def test_meta_records_excel_time_and_deployed_rev(world: dict[str, Path], tmp_path: Path,
                                                 monkeypatch: pytest.MonkeyPatch) -> None:
    """E-08 — 같은 model 판으로 코드만 바뀐 정정판을 파일로 구별하게 메타에 엑셀 생성 시각과 배포
    rev 를 싣는다. rev 는 **코드 루트**(`src/` 의 부모 — 서버 ~/quant-ledger)의 DEPLOYED.json
    (deploy.sh 가 쓴다)에서 읽는다 — QL_HOME 은 데이터 루트라 다른 곳을 가리킬 수 있다(검토 사소 2).
    못 읽으면 '알 수 없음'."""
    import deliver.reader as reader
    code, home = tmp_path / "code", tmp_path / "home"
    code.mkdir()
    home.mkdir()
    monkeypatch.setattr(reader, "CODE_ROOT", code)
    monkeypatch.setenv("QL_HOME", str(home))
    (code / "DEPLOYED.json").write_text(json.dumps(
        {"rev": "abc1234", "branch": "main", "at_utc": "2026-10-08T05:00:00Z"}), encoding="utf-8")
    (home / "DEPLOYED.json").write_text(json.dumps({"rev": "zzz9999"}), encoding="utf-8")
    res = build_daily(FRI, "morning", model_root=world["model"], fi_root=world["fi"],
                      out_root=tmp_path / "a")
    pairs = dict(_meta_pairs(load_workbook(res.path)))
    assert pairs["코드 rev"] == "abc1234 · 배포 2026-10-08T05:00:00Z"
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", str(pairs["엑셀 생성 시각"]))
    for broken in (None, "{깨진"):
        p = code / "DEPLOYED.json"
        if broken is None:
            p.unlink()
        else:
            p.write_text(broken, encoding="utf-8")
        res = build_daily(FRI, "morning", model_root=world["model"], fi_root=world["fi"],
                          out_root=tmp_path / "b")
        assert dict(_meta_pairs(load_workbook(res.path)))["코드 rev"] == "알 수 없음"


def _write_rank_run(model_root: Path, day: str, spec: str, ranks: dict[str, int]) -> None:
    """순위만 있는 최소 model 판(순위 흐름 재료)."""
    bid = f"m_{day.replace('-', '')}T000000Z"
    _write(model_root / spec / f"v={bid}" / "scores.parquet",
           pa.schema([("ticker", pa.string()), ("rank", pa.int64())]),
           [{"ticker": t, "rank": r} for t, r in ranks.items()])
    runs = model_root / "_runs"
    runs.mkdir(parents=True, exist_ok=True)
    meta = {"layer": "model", "status": "ok", "build_id": bid, "date": day, "basis": "morning",
            "fi_build_id": "x", "generated_at": f"{day}T00:00:00Z", "specs": {spec: {}},
            "primary_spec": spec}
    (runs / f"{day.replace('-', '')}_morning.json").write_text(json.dumps(meta), encoding="utf-8")


def test_trend_line_direction_matches_its_color_when_universe_grows(tmp_path: Path) -> None:
    """E-06 G1 — 유니버스가 100 → 110 으로 커지고 X 의 순위가 50 → 51(1칸 하락)이면 선 색은
    빨강(Δ −1)이다. 옛 코드는 선을 백분위(50.5 → 54.1, 위로)로 그려 선은 오르는데 색은 빨강이었다.
    선 = 원순위(위 = 상승)라 선의 처음 → 끝 방향과 색이 늘 같다."""
    from deliver.reader import load_run
    from deliver.trend import LINE_DOWN, load_trend, spark_groups
    from deliver.trend import write_sheet as write_trend
    from openpyxl import Workbook
    spec, d0, d1 = "v4_rank@0.1", "2026-09-01", "2026-10-01"
    others0 = {f"T{k:03d}": k + (1 if k >= 50 else 0) for k in range(1, 100)}
    others1 = {f"T{k:03d}": k + (1 if k >= 51 else 0) for k in range(1, 110)}
    _write_rank_run(tmp_path, d0, spec, {"X": 50, **others0})
    _write_rank_run(tmp_path, d1, spec, {"X": 51, **others1})
    run = load_run(tmp_path, d1, "morning")
    trend = load_trend(tmp_path, run, spec, "morning")
    assert trend.dates == (d0, d1) and trend.base["1M"] == d0
    delta = trend.delta("1M", "X", 51)
    assert delta == -1
    wb = Workbook()
    write_trend(wb, ["X"], trend)
    ws = wb["순위 흐름"]
    first, last = ws.cell(8, 2).value, ws.cell(8, 3).value
    assert isinstance(first, int | float) and isinstance(last, int | float)
    groups = spark_groups(["X"], trend, date.fromisoformat(d1), {"1M": "M"}, {"1M": {"X": delta}})
    assert [g.color for g in groups] == [LINE_DOWN]          # 선 색 = 하락
    assert last < first                                       # 선도 내려간다(색과 같은 방향)


def test_rank_delta_cell_color_is_centered_on_zero(daily) -> None:
    """E-06(Q8) — Δ순위 1M 칸 색의 가운데는 0 이다(음수 빨강 · 양수 초록, 끝점 ±M 대칭).
    옛 코드는 3색 백분위 10/50/90 이라 가운데가 중앙값(10-02 판 −21.5) — 소폭 하락이 연두로
    보였다."""
    _, wb = daily
    ws = wb["점수"]
    letter = ws.cell(7, header(ws)["Δ순위 1M"]).column_letter
    cs = next(r.colorScale for rng in ws.conditional_formatting for r in rng.rules
              if str(rng.sqref).startswith(letter))
    assert [v.type for v in cs.cfvo] == ["num", "num", "num"]
    lo, mid, hi = (float(v.val) for v in cs.cfvo)
    assert mid == 0 and lo == -hi and hi > 0
    assert [c.rgb[-6:] for c in cs.color] == ["F8696B", "FFEB84", "63BE7B"]


def test_scale_zero_mid_endpoints_are_symmetric_90th_magnitude() -> None:
    from deliver.qpack import scale_zero_mid
    vals: list[object] = [-192.0, -21.5, -1.0, 0.0, 5.0, 174.9, None, "x"]

    def points(values: list[object]) -> tuple[float, ...]:
        cs = scale_zero_mid(values).colorScale
        assert cs is not None and [v.type for v in cs.cfvo] == ["num"] * 3
        return tuple(float(v.val) for v in cs.cfvo)
    lo, mid, hi = points(vals)
    mags = sorted(abs(float(v)) for v in vals if isinstance(v, float))
    assert mid == 0 and hi == pytest.approx(quantile(mags, 0.9)) and lo == -hi
    assert points([None, 0]) == (-1.0, 0.0, 1.0)            # 값이 없거나 전부 0 이면 ±1


def test_daily_earnings_estimated_prior_year_yoy_is_grey_too(tmp_path: Path) -> None:
    """Q10 — Y−1 칸을 추정치로 채우면(회색 'E', N-19) 그 옆 y-y 칸도 회색 글자다.
    확정치 칸 옆 y-y 는 그대로."""
    day = "2027-01-04"
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, day, write_fi_day(fi_root, day, edit=_roll_to_2027))
    res = build_daily(day, "morning", model_root=model_root, fi_root=fi_root, out_root=tmp_path)
    ws = load_workbook(res.path)["실적"]
    rows = rows_by_code(ws)
    for occurrence in (0, 1, 2):                       # 매출 · 영업이익 · 순이익
        c = col_of(ws, "2026", occurrence)
        est_y = ws.cell(rows[tick(10)], c + 1)          # 2026 = 추정치 → y-y 회색
        assert est_y.value is not None and est_y.font.color.rgb.endswith("7F7F7F")
        act_y = ws.cell(rows[tick(0)], c + 1)           # 2026 = 확정치 → y-y 그대로
        assert act_y.font.color is None or not act_y.font.color.rgb.endswith("7F7F7F")
    grey = {code for code, r in rows.items()
            if (f := ws.cell(r, col_of(ws, "2026") + 1).font).color is not None
            and f.color.rgb.endswith("7F7F7F")}
    est = {code for code, r in rows.items()
           if '"E"' in ws.cell(r, col_of(ws, "2026")).number_format}
    assert grey == est                                   # y-y 회색 = 'E' 칸 행 그대로


def test_display_sheet_defines_roe_roa_debt_from_dart(daily) -> None:
    """지표 시트 ROE·ROA·부채비율은 fi 가 DART 사업보고서로 계산한다(factor_inputs queries 연간
    `d_net_income ÷ d_total_equity` 등). 옛 정의 문구는 'WISE'."""
    _, wb = daily
    ws = wb["지표(표시용)"]
    h = header(ws)
    for label in ("ROE(%)", "ROA(%)", "부채비율(%)"):
        text = str(ws.cell(5, h[label]).value)
        assert "DART" in text and "WISE" not in text, (label, text)


# ── scope 실물 모양 판(실제 model.build — 주 모델 scope@1.0) ─────────────────────
# 연간 확정 PER·PBR·배당수익률이 없는 종목 → scope 밸류 버킷 결측(279570 모양)
NO_VAL = "100007"
# 연간 행은 있는데(WISE 매출·PER) 손익·자산 재료가 다 빈 종목 → 퀄리티 = 변동성만(241560 모양)
VOL_ONLY = "100011"
SCOPE_FI_BID = "m_20260929T000500_000000Z"
LONE = "100040"            # 대분류 G50 의 유일한 종목


def _write_fi_tree(b, fi_root: Path, fi_bid: str, day: str) -> None:
    """test_model_v4_rank.Board → factor_inputs 고정 경로(8표 + latest)."""
    from model import build as mbuild
    fi = b.fi()
    for name, t in FI_TABLES.items():
        out = fi_root / name / f"v={fi_bid}" / "part0.parquet"
        mbuild.write_parquet(list(fi.tables.get(name, ())),
                             {c.name: c.dtype for c in t.columns}, out)
        (out.parent / "_meta.json").write_text(json.dumps(
            {"table": name, "build_id": fi_bid, "basis": "morning", "date": day}))
    (fi_root / "latest_morning.json").write_text(json.dumps(
        {"layer": "factor_inputs", "status": "ok", "build_id": fi_bid, "date": day,
         "basis": "morning"}))


@pytest.fixture(scope="module")
def scope_board(tmp_path_factory: pytest.TempPathFactory):
    """실제 `model.build` 판(40종목, 다섯 spec) — 주 모델 scope@1.0 은 원값을 점수 표 열로 싣고
    지표 긴 표는 0행이다(서버 10-02 판과 같은 모양). NO_VAL 만 밸류 재료가 없다."""
    from model import build as mbuild
    from test_model_v4_rank import Board, _full
    root = tmp_path_factory.mktemp("scope")
    b = Board()
    for i in range(40):
        _full(b, f"{100000 + i:06d}", i, sector=("G10", "G20", "G30")[i % 3])
    _full(b, LONE, 40, sector="G50")        # 혼자인 대분류 — 업종 z 는 유니버스 z 로 되돌린다
    for r in b.tables["fi_prices"]:
        r["close"] = round(float(r["close"]))                  # type: ignore[arg-type]
    for r in b.tables["fi_fin_summary"]:
        if r["period_type"] == "annual" and r["ticker"] != NO_VAL:
            k = int(str(r["ticker"])) % 11
            r.update(per=6.0 + k, pbr=0.6 + 0.1 * k, dividend_yield=0.5 + 0.2 * k, roa=2.0 + k,
                     debt_ratio=60.0 + 5 * k, gross_profit=100.0 + 3 * k)
        if r["period_type"] == "annual" and r["ticker"] == VOL_ONLY:
            r.update(roa=None, debt_ratio=None, gross_profit=None, fcf=None, total_assets=None)
    fi_root = root / "fi"
    _write_fi_tree(b, fi_root, SCOPE_FI_BID, "2026-09-28")
    res = mbuild.build("20260928", "morning", root / "model", fi_root,
                       min_prices_on_d=10, min_ranked=10)
    assert res.ok and res.primary_spec == "scope@1.0", res.specs
    d = build_daily("2026-09-28", "morning", model_root=root / "model", fi_root=fi_root,
                    out_root=root / "out")
    scores = {str(r["stock_code"]): r for r in duckdb_rows(
        root / "model" / "scope@1.0" / f"v={res.build_id}" / "scores.parquet")}
    return res, load_workbook(d.path), scores, root


def duckdb_rows(path: Path) -> list[dict[str, object]]:
    import duckdb
    con = duckdb.connect()
    try:
        cur = con.execute(f"SELECT * FROM read_parquet('{path}')")
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r, strict=True)) for r in cur.fetchall()]
    finally:
        con.close()


def test_scope_bucket_missing_is_marked_and_counted(scope_board) -> None:
    """E-02(N-26 4.5) — scope 는 지표 긴 표가 0행이라 옛 코드는 `t in view.ind` 로 '점수 대상'을
    가려 결측 버킷을 빈칸으로 두고 결측 축·메타 결측 수를 0 으로 냈다. 점수 행이 있으면 점수 대상:
    빈 버킷 = '결측(원천없음)', 결측 축 = 밸류, 메타 '축별 결측 수' 밸류 1."""
    _, wb, scores, _ = scope_board
    assert scores[NO_VAL]["valuation_score"] is None
    assert all(r["valuation_score"] is not None for t, r in scores.items() if t != NO_VAL)
    ws = wb["점수"]
    h, rows = header(ws), rows_by_code(ws)
    r = rows[NO_VAL]
    assert ws.cell(r, h["밸류 유니버스 z"]).value == "결측(원천없음)"     # scope 축 = z(10-08)
    assert ws.cell(r, h["밸류 업종 z"]).value == "결측(원천없음)"
    assert ws.cell(r, h["결측 축"]).value == "밸류"
    assert {ws.cell(rows[t], h["결측 축"]).value for t in scores if t != NO_VAL} <= {None, ""}
    meta = dict(_meta_pairs(wb))
    assert "밸류 1" in str(meta["축별 결측 수"]) and "모멘텀 0" in str(meta["축별 결측 수"])


# scope 점수 표 원값 27열 + 표식 6열(contracts.V3_SCORE_COLUMNS) — 엑셀 머리글(줄바꿈은 공백)
SCOPE_RAW_HEADERS = {
    "r1m": "1M 수익률 (비율)", "r3m": "3M 수익률 (비율)", "r6m": "6M 수익률 (비율)",
    "r9m": "9M 수익률 (비율)", "r12m": "12M 수익률 (비율)",
    "op_change_1w": "영업이익 추정 1W (비율)", "op_1w_flag": "영업이익 1W 표식",
    "ni_change_1w": "순이익 추정 1W (비율)", "ni_1w_flag": "순이익 1W 표식",
    "op_change_1m": "영업이익 추정 1M (비율)", "op_1m_flag": "영업이익 1M 표식",
    "ni_change_1m": "순이익 추정 1M (비율)", "ni_1m_flag": "순이익 1M 표식",
    "op_change_3m": "영업이익 추정 3M (비율)", "op_3m_flag": "영업이익 3M 표식",
    "ni_change_3m": "순이익 추정 3M (비율)", "ni_3m_flag": "순이익 3M 표식",
    "flow_inst_5d": "기관 5일 (순매수/시총)", "flow_inst_20d": "기관 20일 (순매수/시총)",
    "flow_for_5d": "외국인 5일 (순매수/시총)", "flow_for_20d": "외국인 20일 (순매수/시총)",
    "flow_pe_5d": "사모 5일 (순매수/시총)", "flow_pe_20d": "사모 20일 (순매수/시총)",
    "qual_gpa": "GP/A (비율)", "qual_roa": "ROA (%)", "qual_fcf_assets": "FCF/자산 (비율)",
    "qual_debt_ratio": "부채비율 (%)", "qual_gpa_change": "GP/A 변화 (비율)",
    "qual_std_20d": "20일 변동성 (비율)",
    "val_per": "PER (배)", "val_pbr": "PBR (배)", "val_ev_ebitda": "EV/EBITDA (배)",
    "val_dividend_yield": "배당수익률 (%)",
}


def test_scope_raw_sheet_carries_the_score_table_raw_columns(scope_board) -> None:
    """E-10(Q7) G1 — scope 엑셀 '점수 원자료'에 점수 표 원값 27열 + 표식 6열을 엔진 값 그대로
    싣는다. 옛 코드는 지표 긴 표(v4 전용)만 읽어 scope 원자료 시트가 코드·이름·기준 열뿐이었다
    (10-02 발송본)."""
    from model.contracts import V3_SCORE_COLUMNS
    _, wb, scores, _ = scope_board
    assert set(SCOPE_RAW_HEADERS) <= set(V3_SCORE_COLUMNS) and len(SCOPE_RAW_HEADERS) == 33
    ws = wb["점수 원자료"]
    h, rows = header(ws), rows_by_code(ws)
    assert set(SCOPE_RAW_HEADERS.values()) <= set(h), set(SCOPE_RAW_HEADERS.values()) - set(h)
    groups = [str(c.value) for c in ws[6] if c.value is not None]
    assert groups[:6] == ["종목", "모멘텀", "리비전", "수급", "퀄리티", "밸류"]
    t = "100005"
    for key, label in SCOPE_RAW_HEADERS.items():
        got, want = ws.cell(rows[t], h[label]).value, scores[t][key]
        if isinstance(want, float):
            assert got == pytest.approx(want), key                   # 배율 없이 엔진 값 그대로
        else:
            assert got == want, key                                  # 표식 글자 · 값 없음 = 빈칸
        assert str(ws.cell(5, h[label]).value).startswith(f"{key}:"), key   # 정의 = 열 이름부터
    assert isinstance(ws.cell(rows[t], h["1M 수익률 (비율)"]).value, float)
    defs = {key: str(ws.cell(5, h[label]).value) for key, label in SCOPE_RAW_HEADERS.items()}
    assert "비율" in defs["r1m"] and "%" in defs["val_dividend_yield"]
    assert "순매수" in defs["flow_for_5d"] and "시총" in defs["flow_for_5d"]
    unit = dict(_meta_pairs(wb))["점수 원자료 단위"]
    assert "r1m" in str(unit) and "val_dividend_yield" in str(unit) and "flow_" in str(unit)


def test_scope_sector_sheet_fills_per_revision_and_return(scope_board) -> None:
    """E-10(Q7) — 업종 시트 '업종 지표' 3열을 scope 점수 표 원값으로 채운다: PER 중앙값(val_per) ·
    리비전 상향 비율(op_change_1m > 0, 값 있는 종목 중) · 1M 수익률(r1m 시총가중 × 100).
    옛 코드는 v4 지표 긴 표만 봐서 scope 판에선 세 열이 전부 빈칸이었다."""
    res, wb, scores, root = scope_board
    ws = wb["업종"]
    h = header(ws)
    first = {ws.cell(r, h["코드"]).value: r for r in range(8, ws.max_row + 1)
             if ws.cell(r, h["구분"]).value == "대분류"}
    uni = duckdb_rows(root / "fi" / "fi_universe" / f"v={SCOPE_FI_BID}" / "part0.parquet")
    caps = {str(r["ticker"]): r["market_cap"] for r in uni if isinstance(r["market_cap"], float)}
    l1 = {str(r["ticker"]): r["sector_l1"] for r in uni}
    for code, r in first.items():
        members = [t for t in scores if l1[t] == code]
        per = [scores[t]["val_per"] for t in members if scores[t]["val_per"] is not None]
        rv = [scores[t]["op_change_1m"] for t in members if scores[t]["op_change_1m"] is not None]
        r1 = {t: scores[t]["r1m"] for t in members if scores[t]["r1m"] is not None}
        w = sum(caps[t] for t in r1)
        assert ws.cell(r, h["PER 중앙값 (배)"]).value == pytest.approx(statistics.median(per))
        assert ws.cell(r, h["리비전 상향 비율(%)"]).value == pytest.approx(
            sum(1 for v in rv if v > 0) / len(rv) * 100)
        assert ws.cell(r, h["1M 수익률 (%)"]).value == pytest.approx(
            sum(v * caps[t] for t, v in r1.items()) / w * 100)


def test_v4_comparison_columns_are_left_out_until_fixed(scope_board) -> None:
    """N-27 §8-17 — v4 비교 열(v4_rank@*)은 결함 수정 전까지 점수 시트 '다른 모델 순위'와
    '모델 비교' 시트에서 뺀다. v3 원본·v2 원본 비교와 '최대 차이'(남은 모델끼리)는 남고, 메타에
    한 줄 남긴다.
    v4 판 계산·저장은 그대로다(판 게이트 줄도 그대로)."""
    res, wb, _, _ = scope_board
    assert {"v4_rank@0.1", "v4_rank@0.2"} <= set(res.specs)
    sc, mc = wb["점수"], wb["모델 비교"]
    for ws in (sc, mc):
        h = header(ws)
        assert "v3 원본" in h and "v2 원본" in h, ws.title
        assert "v4 기본" not in h and "v4 동일가중" not in h, ws.title
    h, rows = header(sc), rows_by_code(sc)
    for t, r in rows.items():
        got = [sc.cell(r, h[k]).value for k in ("순위", "v3 원본", "v2 원본")]
        ranks = [x for x in got if x is not None]
        want = max(ranks) - min(ranks) if len(ranks) >= 2 else None
        assert sc.cell(r, h["최대 차이"]).value == want, t
    meta = dict(_meta_pairs(wb))
    left_out = str(meta["비교에서 뺀 모델"])
    assert "N-27" in left_out and "v4_rank@0.1" in left_out and "v4_rank@0.2" in left_out
    assert "v4_rank" not in str(meta["비교 모델"])
    assert "판 게이트 v4_rank@0.1" in meta                                   # 판은 그대로


# ── 4-2a 검토 후속(G-27 표기 · 순액 매출 y-y · 원값 열 부분 결측 · insert_after) ─────────────
VOL_ONLY_NOTE = "퀄리티 = 변동성만(손익 지표 없음)"
NET_NOTE = "매출 = 순영업이익(순액) — 총액 추정치와 y-y 비교 안 함"


def test_scope_volatility_only_quality_is_noted(scope_board) -> None:
    """G-27(컨트롤러 결정) — scope 점수 행의 퀄리티 손익 지표(qual_gpa·roa·fcf_assets·debt_ratio·
    gpa_change)가 다 비고 변동성(qual_std_20d)만 있으면 비고에 '퀄리티 = 변동성만(손익 지표 없음)'.
    원인(외화 재무 등)은 적지 않는다 — 다른 이유로 비는 종목도 있다(477850·0011T0)."""
    _, wb, scores, _ = scope_board
    row = scores[VOL_ONLY]
    assert row["qual_std_20d"] is not None and row["quality_score"] is not None
    assert all(row[k] is None for k in ("qual_gpa", "qual_roa", "qual_fcf_assets",
                                        "qual_debt_ratio", "qual_gpa_change"))
    ws = wb["점수"]
    h, rows = header(ws), rows_by_code(ws)
    notes = {t: ws.cell(r, h["비고"]).value for t, r in rows.items()}
    assert VOL_ONLY_NOTE in str(notes[VOL_ONLY]) and "외화" not in str(notes[VOL_ONLY])
    assert [t for t, n in notes.items() if VOL_ONLY_NOTE in str(n)] == [VOL_ONLY]
    need = fit_width(Col("note", "비고", "txt", None, 99.0), [{"note": notes[VOL_ONLY]}])
    assert ws.column_dimensions[ws.cell(7, h["비고"]).column_letter].width >= need   # 잘리지 않게


def _net_revenue(table: str, rows: list[dict[str, object]]) -> None:
    """12번 종목 = 증권사처럼 연간 확정 매출이 순액(순영업이익, fi1.2.0 `revenue_basis` 'net')
    두 해, 13번 = 2025 만 순액(2024 총액). 컨센서스 매출은 총액 그대로."""
    if table == "fi_fin_summary":
        for r in rows:
            net = r["ticker"] == tick(12) or (r["ticker"] == tick(13) and r["period"] == "2025/12")
            if r["period_type"] == "annual" and net:
                r["revenue_basis"] = "net"


def test_net_revenue_yoy_against_gross_estimate_is_left_blank(tmp_path: Path) -> None:
    """4-2b 명세 검토 #1 — 확정 매출이 순액(순영업이익)인데 추정치는 총액이라 2025 순액 대 2026E
    총액 y-y 가 가짜 급증(현장 증권·카드 6종목 +180~+1,100%)으로 찍혔다. 한쪽만 순액인 매출 y-y
    칸은 비우고 점수 시트 비고에 적는다. 둘 다 순액(12번 2025 대 2024)·영업이익·총액 종목은
    그대로."""
    model_root, fi_root = tmp_path / "model", tmp_path / "fi"
    write_model_day(model_root, FRI, write_fi_day(fi_root, FRI, edit=_net_revenue))
    wb = load_workbook(build_daily(FRI, "morning", model_root=model_root, fi_root=fi_root,
                                   out_root=tmp_path).path)
    ws = wb["실적"]
    rows = rows_by_code(ws)
    y25, y26, y27 = (col_of(ws, f"{y} y-y(%)", 0) for y in ("2025", "2026E", "2027E"))
    r12, r13, r10 = rows[tick(12)], rows[tick(13)], rows[tick(10)]
    assert ws.cell(r12, y26).value is None                        # 2026E 총액 vs 2025 순액
    num = (int, float)                                            # 엑셀 왕복에서 0.0 은 0
    assert isinstance(ws.cell(r12, y25).value, num)               # 2025 순액 vs 2024 순액
    assert isinstance(ws.cell(r12, y27).value, num)               # 추정 vs 추정
    assert ws.cell(r13, y25).value is None and ws.cell(r13, y26).value is None
    assert isinstance(ws.cell(r12, col_of(ws, "2026E y-y(%)", 1)).value, num)     # 영업이익 그대로
    assert ws.cell(r10, y26).value == pytest.approx((1300 / 1200 - 1) * 100)       # 총액 종목
    sc = wb["점수"]
    note, srows = header(sc)["비고"], rows_by_code(sc)
    assert sc.cell(srows[tick(12)], note).value == NET_NOTE
    assert sc.cell(srows[tick(13)], note).value == NET_NOTE
    assert sc.cell(srows[tick(10)], note).value is None


def test_scope_raw_columns_survive_a_missing_engine_column(scope_board, tmp_path: Path,
                                                          caplog: pytest.LogCaptureFixture) -> None:
    """검토 사소 1 — 원값 열은 '전부 아니면 전무'가 아니다. 엔진(v3_zscore)으로 판정하고, 점수 표에
    없는 열만 빈칸으로 두며 빠진 열 이름을 로그 한 줄로 남긴다. 옛 코드는 열 하나만 빠져도 원자료
    33열·업종 3열·메타 단위 줄이 통째로 사라졌다."""
    import shutil

    import duckdb
    res, _, scores, root = scope_board
    model = tmp_path / "model"
    shutil.copytree(root / "model", model)
    path = model / "scope@1.0" / f"v={res.build_id}" / "scores.parquet"
    tmp = path.with_name("cut.parquet")
    con = duckdb.connect()
    con.execute(f"COPY (SELECT * EXCLUDE (val_ev_ebitda) FROM read_parquet('{path}')) "
                f"TO '{tmp}' (FORMAT parquet)")
    con.close()
    tmp.replace(path)
    with caplog.at_level("WARNING"):
        d = build_daily("2026-09-28", "morning", model_root=model, fi_root=root / "fi",
                        out_root=tmp_path / "out")
    wb = load_workbook(d.path)
    ws = wb["점수 원자료"]
    h, rows = header(ws), rows_by_code(ws)
    assert set(SCOPE_RAW_HEADERS.values()) <= set(h)
    assert all(ws.cell(r, h["EV/EBITDA (배)"]).value is None for r in rows.values())
    t = "100005"
    for key, label in SCOPE_RAW_HEADERS.items():
        if key != "val_ev_ebitda" and isinstance(scores[t][key], float):
            assert ws.cell(rows[t], h[label]).value == pytest.approx(scores[t][key]), key
    sec = wb["업종"]
    hs = header(sec)
    for label in ("PER 중앙값 (배)", "리비전 상향 비율(%)", "1M 수익률 (%)"):
        assert isinstance(sec.cell(8, hs[label]).value, int | float), label
    assert "점수 원자료 단위" in dict(_meta_pairs(wb))
    assert any("val_ev_ebitda" in rec.getMessage() for rec in caplog.records)


def test_insert_after_appends_when_the_key_is_missing(caplog: pytest.LogCaptureFixture) -> None:
    """검토 사소 6 — 끼울 자리(키)가 없어도 그날 엑셀·발송을 막지 않는다: 끝에 덧붙이고
    경고 한 줄."""
    from deliver.excel_daily import insert_after
    pairs: list[tuple[str, object]] = [("기준일", "2026-10-02")]
    with caplog.at_level("WARNING"):
        insert_after(pairs, "없는 줄", [("엑셀 생성 시각", "x")])
    assert pairs == [("기준일", "2026-10-02"), ("엑셀 생성 시각", "x")]
    assert any("없는 줄" in rec.getMessage() for rec in caplog.records)


# ── 점수 시트 축 = z(사용자 결정 10-08 '둘 다 z') ────────────────────────────────────
SCOPE_BUCKETS = (("momentum", "모멘텀"), ("revision", "리비전"), ("flow", "수급"),
                 ("quality", "퀄리티"), ("valuation", "밸류"))


def _sector_z(values: list[float], v: float) -> float:
    """손 계산 — 대분류 평균·표본 표준편차, ±3σ 로 자른 뒤 표준화(엔진 z 와 같은 식)."""
    mean, sd = statistics.mean(values), statistics.stdev(values)
    return (max(mean - 3 * sd, min(mean + 3 * sd, v)) - mean) / sd


def test_scope_score_sheet_axes_are_engine_z(scope_board) -> None:
    """z 엔진(scope·v3_zscore) 주 모델이면 점수 시트 5팩터 칸이 z 다. 유니버스 z = 엔진 버킷 점수
    (종합 점수에 들어간 값) 그대로 — 다시 순위 매기지 않는다. 업종 z = 같은 z 를 WICS 대분류 안에서
    다시 z(표본 < 5 · 표준편차 0 이면 유니버스 z). 옛 코드는 버킷 z 를 0~100 백분위로 바꿔
    실었다."""
    _, wb, scores, root = scope_board
    ws = wb["점수"]
    h, rows = header(ws), rows_by_code(ws)
    uni = duckdb_rows(root / "fi" / "fi_universe" / f"v={SCOPE_FI_BID}" / "part0.parquet")
    l1 = {str(r["ticker"]): r["sector_l1"] for r in uni}
    for b, name in SCOPE_BUCKETS:
        u, s = h[f"{name} 유니버스 z"], h[f"{name} 업종 z"]
        assert f"{name} 유니버스" not in h
        for t, row in scores.items():
            got_u, got_s = ws.cell(rows[t], u).value, ws.cell(rows[t], s).value
            if row[f"{b}_score"] is None:
                assert got_u == got_s == "결측(원천없음)", (b, t)       # 결측 규칙 그대로
                continue
            assert got_u == pytest.approx(row[f"{b}_score"]), (b, t)
            assert -3.0 <= got_s <= 3.0, (b, t)
        # 손 계산 대조 — 100005(G30) 와 혼자인 G50 종목
        t = "100005"
        peers = [r[f"{b}_score"] for x, r in scores.items()
                 if l1[x] == l1[t] and r[f"{b}_score"] is not None]
        assert len(peers) >= 5
        assert ws.cell(rows[t], s).value == pytest.approx(_sector_z(peers, scores[t][f"{b}_score"]))
        assert ws.cell(rows[LONE], s).value == pytest.approx(scores[LONE][f"{b}_score"])
        assert ws.cell(rows[t], u).number_format == "#,##0.00"
    # 색 — z 칸은 고정 3색 −3·0·+3(초록 = 높음)
    for label in ("모멘텀 유니버스 z", "밸류 업종 z"):
        letter = ws.cell(7, h[label]).column_letter
        cs = next(r.colorScale for rng in ws.conditional_formatting for r in rng.rules
                  if str(rng.sqref).startswith(letter))
        assert [(v.type, float(v.val)) for v in cs.cfvo] == [("num", -3.0), ("num", 0.0),
                                                            ("num", 3.0)]
        assert [c.rgb[-6:] for c in cs.color] == ["F8696B", "FFEB84", "63BE7B"]
    assert "z" in str(ws.cell(5, 1).value)                               # 시트 각주


def test_scope_sector_sheet_axis_means_are_z(scope_board) -> None:
    """업종 시트 축별 평균도 z 엔진이면 유니버스 z(엔진 버킷 점수)의 업종 평균 — 머리글
    '축별 평균 z'."""
    _, wb, scores, root = scope_board
    ws = wb["업종"]
    assert "축별 평균 z" in [str(c.value) for c in ws[6] if c.value is not None]
    h = header(ws)
    uni = duckdb_rows(root / "fi" / "fi_universe" / f"v={SCOPE_FI_BID}" / "part0.parquet")
    l1 = {str(r["ticker"]): r["sector_l1"] for r in uni}
    first = {ws.cell(r, h["코드"]).value: r for r in range(8, ws.max_row + 1)
             if ws.cell(r, h["구분"]).value == "대분류"}
    for b, name in SCOPE_BUCKETS:
        for code, r in first.items():
            vals = [x[f"{b}_score"] for t, x in scores.items()
                    if l1[t] == code and x[f"{b}_score"] is not None]
            assert ws.cell(r, h[name]).value == pytest.approx(statistics.fmean(vals)), (b, code)


def test_v4_score_sheet_axes_stay_percentiles(daily) -> None:
    """회귀 가드 — 백분위 엔진(v4_rank) 주 모델은 지금처럼 0~100 백분위·백분위 색 10/50/90."""
    _, wb = daily
    ws = wb["점수"]
    h = header(ws)
    assert "저위험 유니버스" in h and "저위험 업종" in h
    assert not [k for k in h if k.endswith(" z")]
    vals = [ws.cell(r, h["저위험 유니버스"]).value for r in range(8, ws.max_row + 1)]
    assert all(0 <= v <= 100 for v in vals if isinstance(v, float))
    sec = wb["업종"]
    assert "축별 평균 백분위" in [str(c.value) for c in sec[6] if c.value is not None]


def test_group_z_falls_back_to_universe_value() -> None:
    from deliver.common import group_z
    vals = {"a": 1.0, "b": 2.0, "c": 3.0, "d": 4.0, "e": 5.0, "f": 0.7, "g": 0.7, "h": 9.0}
    grp = {"a": "X", "b": "X", "c": "X", "d": "X", "e": "X", "f": "Y", "g": "Y", "h": None}
    out = group_z(vals, grp, 5)
    assert out["c"] == pytest.approx(0.0) and out["e"] == pytest.approx(_sector_z(
        [1.0, 2.0, 3.0, 4.0, 5.0], 5.0))
    assert out["f"] == 0.7 and out["g"] == 0.7 and out["h"] == 9.0   # 표본 < 5 · 업종 없음
    same = group_z({k: 1.5 for k in "abcde"}, dict.fromkeys("abcde", "X"), 5)
    assert all(v == 1.5 for v in same.values())                        # 표준편차 0
